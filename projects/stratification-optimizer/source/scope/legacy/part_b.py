from __future__ import annotations

from datetime import datetime
from pathlib import Path
import math
import os
import shutil

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side

from test_config import load_test_config
from project_profile import load_project_profile, selected_demographics


# =====================================================
# Optimization Engine Part B
#
# Run order:
#   1. python part_a.py
#   2. python part_b.py
#
# Part B does not read an Excel config workbook.
# It consumes Part A's latest prep output and writes the
# three post-stratification deliverables requested.
#
# All test/clinical-form-specific values (target cell Ns,
# OSEP targets, template layout, output names, ...) are
# loaded from a YAML config selected via
# OPT_ENGINE_TEST_CONFIG (see test_config.py).
# =====================================================

TEST_CONFIG_ID = os.environ.get("OPT_ENGINE_TEST_CONFIG", "basc4_adhd")
CFG = load_test_config(TEST_CONFIG_ID)
PROJECT_PROFILE = load_project_profile()

PROJECT_DIR = Path(os.environ.get(
    "OPT_ENGINE_PROJECT_DIR",
    CFG.default_paths.default_project_dir or ".",
))
PART_A_FILE = Path(os.environ.get(
    "OPT_ENGINE_PART_A_FILE",
    str(PROJECT_DIR / CFG.default_paths.part_a_output_root / "Optimization_Engine_Part_A_prep_output_latest.xlsx"),
))

TARGET_TEMPLATE = Path(os.environ.get(
    "OPT_ENGINE_TARGET_WORKBOOK",
    str(PROJECT_DIR / CFG.default_paths.target_workbook),
))

OUTPUT_ROOT = Path(os.environ.get(
    "OPT_ENGINE_PART_B_OUTPUT_ROOT",
    str(PROJECT_DIR / CFG.default_paths.part_b_output_root),
))
STABLE_DELIVERABLE_DIR = OUTPUT_ROOT / "latest_deliverables"

UPDATED_TEID_NAME = CFG.output_names.updated_teid
DEMO_OSEP_NAME = CFG.output_names.demo_osep
SWAP_TABLE_NAME = CFG.output_names.swap_table

PROJECT = CFG.project
SEX_COL = CFG.sex_col
FEMALE_VALUE = CFG.female_value
MALE_VALUE = CFG.male_value
TARGET_FLAG = CFG.target_flag
MAX_CANDIDATES_PER_REMOVE = CFG.engine_params.max_candidates_per_remove
MAX_SOLVER_SECONDS = int(os.environ.get("OPT_ENGINE_MAX_SOLVER_SECONDS", str(CFG.engine_params.max_solver_seconds)))
CLINICAL_SEEDING = os.environ.get("OPT_ENGINE_CLINICAL_SEEDING", "").strip().upper() in {"1", "TRUE", "Y", "YES"}
FORCED_ONLY_SWAP = os.environ.get("OPT_ENGINE_FORCED_ONLY_SWAP", "").strip()
DISABLE_OPTIMIZATION = os.environ.get("OPT_ENGINE_DISABLE_OPTIMIZATION", "").strip().upper() in {
    "1", "TRUE", "Y", "YES",
}
NUM_WORKERS = CFG.engine_params.num_workers
SCALE = CFG.engine_params.scale
DEMO_WEIGHT = CFG.engine_params.demo_weight
SWAP_WEIGHT = CFG.engine_params.swap_weight
SWAP_LOCK_COLS = CFG.engine_params.swap_lock_cols
DEMO_TARGET_SHEETS = CFG.demo_target_sheets

TARGET_CELL_N = CFG.target_cell_n

INCLUDE_CENSUS = os.environ.get("OPT_ENGINE_INCLUDE_CENSUS", "1").strip().upper() in {"1", "TRUE", "Y", "YES"}
INCLUDE_OSEP = os.environ.get("OPT_ENGINE_INCLUDE_OSEP", "1").strip().upper() in {"1", "TRUE", "Y", "YES"}
INCLUDE_CLINICAL = os.environ.get("OPT_ENGINE_INCLUDE_CLINICAL", "1").strip().upper() in {"1", "TRUE", "Y", "YES"}
OSEP_TARGETS = {
    flag: target.as_dict()
    for flag, target in CFG.osep_targets.items()
    if (flag == "tmp_clin" and INCLUDE_CLINICAL) or (flag != "tmp_clin" and INCLUDE_OSEP)
}

VFORM_ORDER = CFG.vform_order
FORM_FAMILIES = CFG.form_families
DEMO_COLS = selected_demographics(CFG.demo_cols, PROJECT_PROFILE)
VALID_EDL_VALUES = CFG.valid_edl_values
FLAG_COLS = CFG.flag_cols
EXPORT_PRIORITY_COLS = CFG.export_priority_cols
FINAL_TOTAL_N = CFG.final_total_n
DEMO_WRITE_LAYOUT = CFG.demo_write_layout
PHASE2_MAX_SOLVER_SECONDS = CFG.engine_params.phase2_max_solver_seconds
FILTER_MASTER_EDL = CFG.filter_master_edl
SELECTED_TEID_NAME = CFG.output_names.selected_teid
AVAILABLE_POOL_NAME = CFG.output_names.available_pool


def clean_teid(value) -> str | pd.NA:
    if pd.isna(value):
        return pd.NA
    numeric = pd.to_numeric(pd.Series([str(value).strip()]), errors="coerce").iloc[0]
    if pd.notna(numeric):
        return str(int(numeric))
    text = str(value).strip()
    return text if text else pd.NA


def normalize(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(c).strip() for c in out.columns]
    if CFG.set_project_column:
        out["project"] = PROJECT
    out["teid"] = out["teid"].apply(clean_teid)
    out = out.dropna(subset=["teid"]).copy()
    out["vform"] = out["vform"].astype(str).str.strip().str.upper()
    out["agegrp"] = pd.to_numeric(out["agegrp"], errors="coerce").astype("Int64")
    out = out.dropna(subset=["agegrp"]).copy()
    out["agegrp"] = out["agegrp"].astype(int)
    for col in DEMO_COLS:
        if col in out.columns:
            if col == "edl":
                numeric = pd.to_numeric(out[col], errors="coerce")
                out[col] = numeric.astype("Int64").astype(str).replace("<NA>", "")
            else:
                out[col] = out[col].astype(str).str.strip().str.upper()
    for flag in FLAG_COLS:
        aliases = [flag, flag.lower(), flag.upper()]
        source = next((col for col in aliases if col in out.columns), None)
        if source is None:
            out[flag] = 0
        elif source != flag:
            out[flag] = out[source]
        out[flag] = pd.to_numeric(out[flag], errors="coerce").fillna(0).astype(int)
    out["tmp_clin"] = out[FLAG_COLS].astype(bool).any(axis=1).astype(int)
    return out


def load_part_a() -> tuple[pd.DataFrame, pd.DataFrame]:
    if not PART_A_FILE.exists():
        raise FileNotFoundError(
            f"Missing Part A output: {PART_A_FILE}. Run Optimization Engine Part A first."
        )
    if CFG.expected_scored_source:
        run_log = pd.read_excel(PART_A_FILE, sheet_name="run_log")
        source_rows = run_log.loc[
            run_log["item"].eq("scored_data_dir"), "value"
        ].dropna().astype(str)
        source_path = source_rows.iloc[-1].replace("\\", "/").lower() if not source_rows.empty else ""
        if not source_path.endswith(CFG.expected_scored_source.lower()):
            raise ValueError(
                "Part A input was not generated from the expected scored datasets "
                f"({CFG.expected_scored_source}). Re-run Part A for this test config first."
            )
    master = pd.read_excel(PART_A_FILE, sheet_name="master_prepped")
    available = pd.read_excel(PART_A_FILE, sheet_name="available_pool")
    master = normalize(master)
    available = normalize(available)
    # EDL=0 is not a reportable education category. Some forms also drop
    # invalid-EDL rows from the starting master sample; others (e.g. ASD)
    # only keep the available/replacement pool clean, per
    # CFG.filter_master_edl.
    if FILTER_MASTER_EDL:
        master = master.loc[master["edl"].isin(VALID_EDL_VALUES)].copy()
    available = available.loc[available["edl"].isin(VALID_EDL_VALUES)].copy()
    return master, available


def validate_target_capacity(master: pd.DataFrame, available: pd.DataFrame) -> None:
    """Confirm the combined universe can fill every documented target cell."""
    universe = pd.concat([master, available], ignore_index=True, sort=False)
    shortfalls = []
    for (vform, agegrp), target_n in TARGET_CELL_N.items():
        possible_n = len(
            universe.loc[universe["vform"].eq(vform) & universe["agegrp"].eq(agegrp)]
        )
        if possible_n < target_n:
            shortfalls.append(
                {
                    "vform": vform,
                    "agegrp": agegrp,
                    "target_n": target_n,
                    "available_universe_n": possible_n,
                    "shortfall": target_n - possible_n,
                }
            )
    if shortfalls:
        raise ValueError(
            "Part B cannot meet the documented target cell sizes with the "
            "current Part A universe. Add the missing scored cases, rerun "
            "Part A, and then rerun Part B:\n"
            + pd.DataFrame(shortfalls).to_string(index=False)
        )


def sex_counts(df: pd.DataFrame) -> dict[str, int]:
    values = df[SEX_COL].astype(str).str.strip().str.upper()
    return {
        FEMALE_VALUE: int((values == FEMALE_VALUE).sum()),
        MALE_VALUE: int((values == MALE_VALUE).sum()),
    }


def row_flag_value(row: pd.Series, flag: str) -> int:
    return int(pd.to_numeric(pd.Series([row.get(flag, 0)]), errors="coerce").fillna(0).iloc[0])


def pct_to_count(n: int, pct: float) -> int:
    return int(round(float(pct) * int(n) / 100.0 * SCALE))


def target_error_scaled(count: int, n: int, target_pct: float) -> float:
    return abs(int(count) * SCALE - pct_to_count(n, target_pct))


def range_error_scaled(count: int, n: int, min_pct: float, max_pct: float) -> float:
    lo = float(min_pct) * int(n) / 100.0 * SCALE
    hi = float(max_pct) * int(n) / 100.0 * SCALE
    val = int(count) * SCALE
    if val < lo:
        return lo - val
    if val > hi:
        return val - hi
    return 0.0


def parse_demo_targets(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    targets = {}

    def parse_demo_header(header):
        header = str(header).strip()
        if header.startswith("edl_"):
            return "edl", header.split("_", 1)[1], header
        if header.startswith("eth_"):
            suffix = header.split("_", 1)[1].lstrip("0123456789")
            return "eth", "BL" if suffix == "AA" else suffix, header
        if header.startswith("region_"):
            return "region", header.split("_", 1)[1].lstrip("0123456789"), header
        return None

    for sheet_name in DEMO_TARGET_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = [ws.cell(6, c).value for c in range(1, ws.max_column + 1)]
        current_vform = None
        for r in range(7, ws.max_row + 1):
            raw_vform = ws.cell(r, 1).value
            if raw_vform:
                current_vform = str(raw_vform).strip().upper()
            source = ws.cell(r, 4).value
            agegrp = ws.cell(r, 2).value
            if current_vform is None or source != "US population" or agegrp is None:
                continue
            key = (current_vform, int(agegrp))
            targets[key] = {}
            for c, header in enumerate(headers, start=1):
                if not header or c <= 4:
                    continue
                value = ws.cell(r, c).value
                if value is None:
                    continue
                parsed = parse_demo_header(header)
                if parsed:
                    targets[key][parsed] = float(value)
    return targets


def add_score(candidate: pd.Series, current_cell: pd.DataFrame) -> float:
    score = 0.0
    if TARGET_FLAG in current_cell.columns:
        target_rate = pd.to_numeric(current_cell[TARGET_FLAG], errors="coerce").fillna(0).mean()
        score += 1000.0 * abs(row_flag_value(candidate, TARGET_FLAG) - target_rate)
    for flag in ["tmp_EBD", "tmp_ASD", "tmp_DD", "tmp_SLD", "tmp_SLI", "tmp_OTHER"]:
        if flag in current_cell.columns:
            target_rate = pd.to_numeric(current_cell[flag], errors="coerce").fillna(0).mean()
            score += 100.0 * abs(row_flag_value(candidate, flag) - target_rate)
    for col in ["eth", "edl", "region"]:
        if col in current_cell.columns and col in candidate.index:
            mode = current_cell[col].mode(dropna=True)
            if not mode.empty:
                score += int(str(candidate.get(col)).upper() != str(mode.iloc[0]).upper())
    return score


def repair_target_n_and_sex_by_adding(
    master: pd.DataFrame,
    available: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Add cases until each cell reaches target N and F/M counts match exactly."""
    current = master.copy()
    pool = available.copy()
    additions = []

    for (vform, agegrp), target_n in TARGET_CELL_N.items():
        while True:
            cell = current.loc[(current["vform"] == vform) & (current["agegrp"] == agegrp)].copy()
            counts = sex_counts(cell)
            current_n = len(cell)
            if current_n >= target_n and counts[FEMALE_VALUE] == counts[MALE_VALUE]:
                break

            target_each = target_n // 2
            if counts[FEMALE_VALUE] < target_each:
                needed_sex = FEMALE_VALUE
            elif counts[MALE_VALUE] < target_each:
                needed_sex = MALE_VALUE
            else:
                break

            candidates = pool.loc[
                (pool["vform"] == vform)
                & (pool["agegrp"] == agegrp)
                & (pool[SEX_COL].astype(str).str.upper() == needed_sex)
            ].copy()
            if candidates.empty:
                additions.append({
                    "action": "unresolved_no_available_case",
                    "vform": vform,
                    "agegrp": agegrp,
                    "needed_sex": needed_sex,
                    "target_n": target_n,
                    "current_n": current_n,
                    "female_n": counts[FEMALE_VALUE],
                    "male_n": counts[MALE_VALUE],
                })
                break

            candidates["_score"] = candidates.apply(lambda row: add_score(row, cell), axis=1)
            add = candidates.sort_values(["_score", "teid"]).iloc[0].drop(labels=["_score"])
            add_row = add.copy()
            add_row["part_b_added_for_sex_n_balance"] = 1
            add_row["project"] = add_row.get("project", PROJECT)
            current = pd.concat([current, add_row.to_frame().T], ignore_index=True, sort=False)
            pool = pool.loc[pool["teid"] != add["teid"]].copy()

            additions.append({
                "action": "add_for_exact_sex_n_balance",
                "vform": vform,
                "agegrp": agegrp,
                "added_teid": add["teid"],
                "added_sex": needed_sex,
                "target_n": target_n,
                "n_before": current_n,
                "n_after": current_n + 1,
                "female_n_before": counts[FEMALE_VALUE],
                "male_n_before": counts[MALE_VALUE],
                "female_n_after": counts[FEMALE_VALUE] + int(needed_sex == FEMALE_VALUE),
                "male_n_after": counts[MALE_VALUE] + int(needed_sex == MALE_VALUE),
                "score": candidates["_score"].min(),
            })

    return normalize(current), normalize(pool), pd.DataFrame(additions)


def apply_forced_only_swap(
    master: pd.DataFrame,
    available: pd.DataFrame,
    specification: str,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Apply one explicitly requested swap without invoking repair or CP-SAT."""
    try:
        remove_teid, add_teid = [clean_teid(value) for value in specification.split(":", 1)]
    except ValueError as exc:
        raise ValueError("OPT_ENGINE_FORCED_ONLY_SWAP must be REMOVE_TEID:ADD_TEID") from exc

    remove_rows = master.loc[master["teid"].eq(remove_teid)]
    add_rows = available.loc[available["teid"].eq(add_teid)]
    if len(remove_rows) != 1 or len(add_rows) != 1:
        raise ValueError(
            f"Forced-only swap requires one master removal and one available addition; "
            f"found remove={len(remove_rows)}, add={len(add_rows)}."
        )
    remove_row = remove_rows.iloc[0]
    add_row = add_rows.iloc[0]
    lock_cols = ["vform", "agegrp", "sex"]
    mismatches = {
        col: (remove_row.get(col), add_row.get(col))
        for col in lock_cols
        if str(remove_row.get(col)).upper() != str(add_row.get(col)).upper()
    }
    if mismatches:
        raise ValueError(f"Forced-only swap violates form/age/sex locks: {mismatches}")

    final = pd.concat(
        [master.loc[~master["teid"].eq(remove_teid)].copy(), add_row.to_frame().T],
        ignore_index=True,
        sort=False,
    )
    remaining = pd.concat(
        [available.loc[~available["teid"].eq(add_teid)].copy(), remove_row.to_frame().T],
        ignore_index=True,
        sort=False,
    )
    swap = {
        "swap_id": 0,
        "vform": remove_row["vform"],
        "agegrp": int(remove_row["agegrp"]),
        "sex": remove_row["sex"],
        "remove_teid": remove_teid,
        "add_teid": add_teid,
        "remove_sex": remove_row["sex"],
        "add_sex": add_row["sex"],
        "solver_status": "BYPASSED_FORCED_ONLY_SWAP",
        "source": "OPT_ENGINE_FORCED_ONLY_SWAP",
    }
    for flag in set(FLAG_COLS + ["tmp_clin"]):
        swap[f"remove_{flag}"] = row_flag_value(remove_row, flag)
        swap[f"add_{flag}"] = row_flag_value(add_row, flag)
        swap[f"delta_{flag}"] = swap[f"add_{flag}"] - swap[f"remove_{flag}"]
    for col in DEMO_COLS:
        swap[f"remove_{col}"] = remove_row.get(col)
        swap[f"add_{col}"] = add_row.get(col)

    final_ids = set(final["teid"])
    expected_ids = (set(master["teid"]) - {remove_teid}) | {add_teid}
    if final_ids != expected_ids or len(final) != len(master):
        raise ValueError("Forced-only swap changed membership beyond the requested TEIDs.")
    return normalize(final), normalize(remaining), pd.DataFrame([swap])


def osep_loss(df: pd.DataFrame) -> float:
    total = 0.0
    for vform, grp in df.groupby("vform"):
        n = len(grp)
        if not n:
            continue
        for flag, rule in OSEP_TARGETS.items():
            if flag not in grp.columns:
                continue
            count = int(pd.to_numeric(grp[flag], errors="coerce").fillna(0).sum())
            weight = float(rule.get("weight", 1))
            if "target_pct" in rule:
                target = float(rule["target_pct"]) * n / 100.0
                total += weight * abs(count - target)
            else:
                lo = float(rule["min_pct"]) * n / 100.0
                hi = float(rule["max_pct"]) * n / 100.0
                total += weight * max(0.0, lo - count, count - hi)
    return total


def build_state(master: pd.DataFrame, demo_targets: dict) -> dict:
    state = {}
    for vform, grp in master.groupby("vform"):
        state[(vform, None)] = {
            "n": len(grp),
            "flags": {
                flag: int(pd.to_numeric(grp[flag], errors="coerce").fillna(0).sum())
                for flag in FLAG_COLS + ["tmp_clin"]
                if flag in grp.columns
            },
        }
    for (vform, agegrp), grp in master.groupby(["vform", "agegrp"]):
        state[(vform, agegrp)] = {
            "n": len(grp),
            "demo": {},
            "sex": sex_counts(grp),
        }
        for (target_vform, target_agegrp), target_map in demo_targets.items():
            if target_vform != vform or target_agegrp != agegrp:
                continue
            for (col, value, _header), _target_pct in target_map.items():
                if col in grp.columns:
                    state[(vform, agegrp)]["demo"][(col, value)] = int(
                        (grp[col].astype(str).str.upper() == str(value).upper()).sum()
                    )
    return state


def local_loss_delta(
    rem: pd.Series,
    add: pd.Series,
    state: dict,
    demo_targets: dict,
) -> float:
    vform = rem["vform"]
    agegrp = int(rem["agegrp"])
    delta = 0.0
    form_state = state[(vform, None)]
    n_form = form_state["n"]
    for flag, rule in OSEP_TARGETS.items():
        if flag not in form_state["flags"]:
            continue
        base_count = int(form_state["flags"][flag])
        new_count = base_count - row_flag_value(rem, flag) + row_flag_value(add, flag)
        weight = int(rule.get("weight", 1))
        if "target_pct" in rule:
            pre = target_error_scaled(base_count, n_form, rule["target_pct"])
            post = target_error_scaled(new_count, n_form, rule["target_pct"])
        else:
            pre = range_error_scaled(base_count, n_form, rule["min_pct"], rule["max_pct"])
            post = range_error_scaled(new_count, n_form, rule["min_pct"], rule["max_pct"])
        delta += weight * (pre - post)

    cell_targets = demo_targets.get((vform, agegrp), {})
    cell_state = state.get((vform, agegrp), {"n": 0, "demo": {}})
    n_cell = int(cell_state["n"])
    for (col, value, _header), target_pct in cell_targets.items():
        base_count = int(cell_state["demo"].get((col, value), 0))
        new_count = base_count - int(str(rem.get(col)).upper() == str(value).upper()) + int(
            str(add.get(col)).upper() == str(value).upper()
        )
        pre = target_error_scaled(base_count, n_cell, target_pct)
        post = target_error_scaled(new_count, n_cell, target_pct)
        delta += DEMO_WEIGHT * (pre - post)
    return delta


def add_exact_sex_constraints(model, y: dict, swaps: pd.DataFrame, state: dict) -> None:
    for (vform, agegrp), cell_state in state.items():
        if agegrp is None:
            continue
        n = int(cell_state["n"])
        if n % 2 != 0:
            raise ValueError(f"Cannot enforce exact sex balance for {vform} agegrp {agegrp}: odd N={n}.")
        target_each = n // 2
        cell_swaps = swaps.loc[(swaps["vform"] == vform) & (swaps["agegrp"] == agegrp)]
        female_expr = int(cell_state["sex"].get(FEMALE_VALUE, 0))
        male_expr = int(cell_state["sex"].get(MALE_VALUE, 0))
        for r in cell_swaps.itertuples():
            sid = int(r.swap_id)
            female_delta = int(str(r.add_sex).upper() == FEMALE_VALUE) - int(str(r.remove_sex).upper() == FEMALE_VALUE)
            male_delta = int(str(r.add_sex).upper() == MALE_VALUE) - int(str(r.remove_sex).upper() == MALE_VALUE)
            if female_delta:
                female_expr += female_delta * y[sid]
            if male_delta:
                male_expr += male_delta * y[sid]
        model.Add(female_expr == target_each)
        model.Add(male_expr == target_each)


def optimize_osep_swaps(
    master: pd.DataFrame,
    available: pd.DataFrame,
    demo_targets: dict,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    try:
        from ortools.sat.python import cp_model
    except ImportError as exc:
        raise ImportError("Install OR-Tools in your environment: python -m pip install ortools") from exc

    report_progress(40, "Building candidate OSEP swaps")
    state = build_state(master, demo_targets)
    swap_rows = []
    lock_cols = SWAP_LOCK_COLS
    for lock_vals, used_group in master.groupby(lock_cols, dropna=False):
        if not isinstance(lock_vals, tuple):
            lock_vals = (lock_vals,)
        mask = pd.Series(True, index=available.index)
        for col, val in zip(lock_cols, lock_vals):
            mask &= available[col].astype(str).eq(str(val))
        avail_group = available.loc[mask].copy()
        if avail_group.empty:
            continue
        for _, rem in used_group.iterrows():
            scored = []
            for _, add in avail_group.iterrows():
                score = local_loss_delta(rem, add, state, demo_targets)
                scored.append((score, add))
            for score, add in sorted(scored, key=lambda x: x[0], reverse=True)[:MAX_CANDIDATES_PER_REMOVE]:
                row = {
                    "swap_id": len(swap_rows),
                    "vform": rem["vform"],
                    "agegrp": int(rem["agegrp"]),
                    "sex": rem[SEX_COL],
                    "remove_teid": rem["teid"],
                    "add_teid": add["teid"],
                    "quick_loss_reduction": score,
                }
                for flag in set(FLAG_COLS + ["tmp_clin"]):
                    row[f"delta_{flag}"] = row_flag_value(add, flag) - row_flag_value(rem, flag)
                    row[f"remove_{flag}"] = row_flag_value(rem, flag)
                    row[f"add_{flag}"] = row_flag_value(add, flag)
                for col in DEMO_COLS:
                    row[f"remove_{col}"] = rem.get(col)
                    row[f"add_{col}"] = add.get(col)
                swap_rows.append(row)

    swaps = pd.DataFrame(swap_rows)
    if swaps.empty:
        return master, pd.DataFrame(), swaps

    model = cp_model.CpModel()
    y = {int(r.swap_id): model.NewBoolVar(f"swap_{int(r.swap_id)}") for r in swaps.itertuples()}

    if CLINICAL_SEEDING:
        # Optional project-level seeding: current-sample cases already flagged
        # for OSEP/clinical membership cannot be removed. With the option off,
        # the validated legacy BASC model is byte-for-byte equivalent here.
        seeded_flags = set(FLAG_COLS + ["tmp_clin"])
        for row in swaps.itertuples():
            if any(int(getattr(row, f"remove_{flag}", 0) or 0) == 1 for flag in seeded_flags):
                model.Add(y[int(row.swap_id)] == 0)

    for _, grp in swaps.groupby("remove_teid"):
        model.Add(sum(y[int(s)] for s in grp["swap_id"]) <= 1)
    for _, grp in swaps.groupby("add_teid"):
        model.Add(sum(y[int(s)] for s in grp["swap_id"]) <= 1)

    add_exact_sex_constraints(model, y, swaps, state)

    osep_terms = []
    demo_terms = []
    demo_dev_vars = []
    for vform in VFORM_ORDER:
        if (vform, None) not in state:
            continue
        n = int(state[(vform, None)]["n"])
        form_swaps = swaps.loc[swaps["vform"] == vform]
        for flag, rule in OSEP_TARGETS.items():
            base = int(state[(vform, None)]["flags"].get(flag, 0))
            expr = base
            for r in form_swaps.itertuples():
                delta = int(getattr(r, f"delta_{flag}"))
                if delta:
                    expr += delta * y[int(r.swap_id)]
            weight = int(rule.get("weight", 1))
            if "target_pct" in rule:
                target = int(round(float(rule["target_pct"]) * n / 100.0 * SCALE))
                dev = model.NewIntVar(0, n * SCALE, f"dev_{vform}_{flag}")
                model.Add(dev >= expr * SCALE - target)
                model.Add(dev >= target - expr * SCALE)
                osep_terms.append(weight * dev)
            else:
                lo = int(math.ceil(float(rule["min_pct"]) * n / 100.0 * SCALE))
                hi = int(math.floor(float(rule["max_pct"]) * n / 100.0 * SCALE))
                under = model.NewIntVar(0, n * SCALE, f"under_{vform}_{flag}")
                over = model.NewIntVar(0, n * SCALE, f"over_{vform}_{flag}")
                model.Add(expr * SCALE + under >= lo)
                model.Add(expr * SCALE - over <= hi)
                osep_terms.append(weight * (under + over))

    for (vform, agegrp), target_map in demo_targets.items():
        if (vform, agegrp) not in state:
            continue
        n = int(state[(vform, agegrp)]["n"])
        cell_swaps = swaps.loc[(swaps["vform"] == vform) & (swaps["agegrp"] == agegrp)]
        for (col, value, _header), target_pct in target_map.items():
            expr = int(state[(vform, agegrp)]["demo"].get((col, value), 0))
            for r in cell_swaps.itertuples():
                add_value = str(getattr(r, f"add_{col}")).upper()
                remove_value = str(getattr(r, f"remove_{col}")).upper()
                delta = int(add_value == str(value).upper()) - int(remove_value == str(value).upper())
                if delta:
                    expr += delta * y[int(r.swap_id)]
            target = pct_to_count(n, target_pct)
            dev = model.NewIntVar(0, n * SCALE, f"dev_{vform}_{agegrp}_{col}_{value}")
            model.Add(dev >= expr * SCALE - target)
            model.Add(dev >= target - expr * SCALE)
            demo_dev_vars.append(dev)
            demo_terms.append(DEMO_WEIGHT * dev)

    swap_term = SWAP_WEIGHT * sum(y.values())
    osep_objective = sum(osep_terms)
    demo_objective = sum(demo_terms)
    # Deviations are scaled case counts, not percentage points. A fixed
    # 100-case ceiling accidentally makes soft demographic targets hard.
    max_demo_bound = max((int(value["n"]) * SCALE for value in state.values()), default=0)
    max_demo_dev = model.NewIntVar(0, max_demo_bound, "max_demo_dev")
    for dev in demo_dev_vars:
        model.Add(max_demo_dev >= dev)
    model.Minimize(osep_objective)

    report_progress(52, f"Running OSEP solver phase 1 (up to {MAX_SOLVER_SECONDS} seconds)")
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = float(MAX_SOLVER_SECONDS)
    solver.parameters.num_search_workers = int(NUM_WORKERS)
    status = solver.Solve(model)
    if status not in [cp_model.OPTIMAL, cp_model.FEASIBLE]:
        raise ValueError(f"OSEP phase 1 optimizer failed: {solver.StatusName(status)}")

    best_osep_objective = int(round(solver.ObjectiveValue()))
    phase1_status = solver.StatusName(status)
    phase1_selected_ids = {
        int(swap_id)
        for swap_id in swaps["swap_id"]
        if solver.Value(y[int(swap_id)]) == 1
    }
    report_progress(70, "OSEP solver phase 1 complete; refining demographic fit")
    model.Add(osep_objective <= best_osep_objective)
    if TEST_CONFIG_ID == "basc4_mood":
        # At the phase-one OSEP fit, retention takes precedence over further
        # demographic refinement. This coefficient dominates its full range.
        demographic_bound = max_demo_bound * 1_000_000 + DEMO_WEIGHT * sum(int(value["n"]) * SCALE for value in state.values()) * max(1, len(demo_dev_vars))
        model.Minimize(sum(y.values()) * (demographic_bound + 1) + max_demo_dev * 1_000_000 + demo_objective)
        print("Mood refresh: minimize replacements at the best OSEP fit, then refine demographics")
    else:
        model.Minimize(max_demo_dev * 1_000_000 + demo_objective + swap_term)

    phase2_seconds = PHASE2_MAX_SOLVER_SECONDS if PHASE2_MAX_SOLVER_SECONDS is not None else MAX_SOLVER_SECONDS
    report_progress(76, f"Running demographic solver phase 2 (up to {phase2_seconds} seconds)")
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = float(phase2_seconds)
    solver.parameters.num_search_workers = int(NUM_WORKERS)
    status = solver.Solve(model)
    phase2_status = solver.StatusName(status)
    if status in [cp_model.OPTIMAL, cp_model.FEASIBLE]:
        selected_ids = {
            int(swap_id)
            for swap_id in swaps["swap_id"]
            if solver.Value(y[int(swap_id)]) == 1
        }
        final_solver_status = phase2_status
        final_solver_objective = solver.ObjectiveValue()
    elif PHASE2_MAX_SOLVER_SECONDS is not None:
        # Phase 1 already produced a feasible solution that satisfies the hard
        # sex-balance constraints and minimizes the OSEP objective. Retain it
        # when the optional demographic refinement exhausts its time budget
        # instead of failing the whole run.
        selected_ids = phase1_selected_ids
        final_solver_status = f"{phase2_status}_FALLBACK_TO_PHASE1"
        final_solver_objective = best_osep_objective
        print(
            "Demographic phase 2 did not return a solution within its time limit; "
            "using the feasible phase-1 OSEP solution."
        )
    else:
        raise ValueError(f"Demographic phase 2 optimizer failed: {phase2_status}")

    selected = swaps.loc[swaps["swap_id"].isin(selected_ids)].copy()
    remove_teids = set(selected["remove_teid"])
    add_teids = set(selected["add_teid"])
    final = pd.concat(
        [
            master.loc[~master["teid"].isin(remove_teids)].copy(),
            available.loc[available["teid"].isin(add_teids)].copy(),
        ],
        ignore_index=True,
        sort=False,
    )
    selected["solver_status"] = final_solver_status
    selected["solver_objective"] = final_solver_objective
    selected["phase1_solver_status"] = phase1_status
    selected["phase1_osep_objective"] = best_osep_objective
    selected["phase2_solver_status"] = phase2_status
    selected["pre_osep_loss"] = osep_loss(master)
    selected["post_osep_loss"] = osep_loss(final)
    return normalize(final), selected, swaps


def summarize_sex(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (vform, agegrp), grp in df.groupby(["vform", "agegrp"], sort=True):
        counts = sex_counts(grp)
        n = len(grp)
        rows.append({
            "vform": vform,
            "agegrp": agegrp,
            "n": n,
            "female_n": counts[FEMALE_VALUE],
            "male_n": counts[MALE_VALUE],
            "female_pct": round(counts[FEMALE_VALUE] / n * 100, 3) if n else 0,
            "male_pct": round(counts[MALE_VALUE] / n * 100, 3) if n else 0,
            "exact_n_match": counts[FEMALE_VALUE] == counts[MALE_VALUE],
            "exact_pct_match": counts[FEMALE_VALUE] == counts[MALE_VALUE],
        })
    return pd.DataFrame(rows)


def build_osep_target_qa(final_sample: pd.DataFrame, remaining_available: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for vform, grp in final_sample.groupby("vform", sort=True):
        n = len(grp)
        avail = remaining_available.loc[remaining_available["vform"].eq(vform)].copy()
        for flag, rule in OSEP_TARGETS.items():
            actual_count = int(pd.to_numeric(grp[flag], errors="coerce").fillna(0).sum())
            actual_pct = round(actual_count / n * 100, 3) if n else 0
            avail_positive = int(pd.to_numeric(avail.get(flag, 0), errors="coerce").fillna(0).sum()) if not avail.empty else 0
            upper_bound_count = actual_count + avail_positive
            row = {
                "vform": vform,
                "flag": flag,
                "n": n,
                "actual_count": actual_count,
                "actual_pct": actual_pct,
                "remaining_available_positive_count": avail_positive,
                "simple_first_admin_upper_bound_count": upper_bound_count,
            }
            if "target_pct" in rule:
                raw_target = float(rule["target_pct"]) * n / 100.0
                nearest_count = int(round(raw_target))
                row.update({
                    "target_pct": float(rule["target_pct"]),
                    "target_count_raw": raw_target,
                    "nearest_integer_target_count": nearest_count,
                    "nearest_integer_target_pct": round(nearest_count / n * 100, 3) if n else 0,
                    "target_min_count": "",
                    "target_max_count": "",
                })
                if actual_count == nearest_count:
                    status = "nearest_integer_target_met"
                elif upper_bound_count < nearest_count:
                    status = "not_achievable_with_remaining_first_admin_pool"
                elif actual_count < nearest_count:
                    status = "below_nearest_integer_target"
                else:
                    status = "above_nearest_integer_target"
            else:
                lo = int(math.ceil(float(rule["min_pct"]) * n / 100.0))
                hi = int(math.floor(float(rule["max_pct"]) * n / 100.0))
                row.update({
                    "target_pct": f"{float(rule['min_pct']):.2f}-{float(rule['max_pct']):.2f}",
                    "target_count_raw": "",
                    "nearest_integer_target_count": "",
                    "nearest_integer_target_pct": "",
                    "target_min_count": lo,
                    "target_max_count": hi,
                })
                if lo <= actual_count <= hi:
                    status = "target_range_met"
                elif upper_bound_count < lo:
                    status = "not_achievable_with_remaining_first_admin_pool"
                elif actual_count < lo:
                    status = "below_target_range"
                else:
                    status = "above_target_range"
            row["status"] = status
            rows.append(row)
    return pd.DataFrame(rows)


def export_sample_columns(df: pd.DataFrame) -> pd.DataFrame:
    cols = [c for c in EXPORT_PRIORITY_COLS if c in df.columns]
    rest = [c for c in df.columns if c not in cols]
    return df.loc[:, cols + rest].copy()


def pct_table(data: pd.DataFrame, var: str, values: list[str]) -> dict[str, float]:
    total = len(data)
    result = {}
    series = data[var].astype(str).str.upper() if var in data.columns else pd.Series([], dtype=str)
    for value in values:
        n = int((series == str(value).upper()).sum())
        result[value] = round(n / total * 100, 3) if total else 0
    return result


def write_demo_row(ws, row: int, data: pd.DataFrame) -> None:
    edl = pct_table(data, "edl", ["1", "2", "3", "4"])
    ws[f"E{row}"] = edl["1"]
    ws[f"F{row}"] = edl["2"]
    ws[f"G{row}"] = edl["3"]
    ws[f"H{row}"] = edl["4"]

    eth = pct_table(data, "eth", ["BL", "AS", "HI", "OT", "WH"])
    ws[f"I{row}"] = eth["BL"]
    ws[f"J{row}"] = eth["AS"]
    ws[f"K{row}"] = eth["HI"]
    ws[f"L{row}"] = eth["OT"]
    ws[f"M{row}"] = eth["WH"]

    region = pct_table(data, "region", ["MW", "NE", "SO", "WE"])
    ws[f"N{row}"] = region["MW"]
    ws[f"O{row}"] = region["NE"]
    ws[f"P{row}"] = region["SO"]
    ws[f"Q{row}"] = region["WE"]

    counts = sex_counts(data)
    ws[f"R{row}"] = counts[FEMALE_VALUE]
    ws[f"S{row}"] = counts[MALE_VALUE]


def write_demo_osep_workbook(final_sample: pd.DataFrame, output_path: Path) -> None:
    shutil.copyfile(TARGET_TEMPLATE, output_path)
    wb = load_workbook(output_path)
    df = final_sample.copy()
    df["form_family"] = df["vform"].str.split("_").str[0]

    layout = DEMO_WRITE_LAYOUT
    for sheet_name, family_vforms, row_map in [
        ("trs", ["TRS_C", "TRS_A"], layout.agegrp_rows),
        ("prs", ["PRS_C", "PRS_A"], layout.agegrp_rows),
        ("srp", ["SRP_C", "SRP_A"], layout.srp_agegrp_rows),
    ]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for agegrp, row in row_map.items():
            data = df.loc[df["vform"].isin(family_vforms) & df["agegrp"].eq(agegrp)].copy()
            if not data.empty:
                write_demo_row(ws, row, data)

    if "byage" in wb.sheetnames and "agey" in df.columns:
        ws = wb["byage"]
        age_series = pd.to_numeric(df["agey"], errors="coerce").round().astype("Int64")
        family_col_start = 8
        for idx, age in enumerate(range(layout.byage_min_age, layout.byage_max_age + 1), start=3):
            ws[f"A{idx}"] = age
            counts = {}
            for col_idx, vform in enumerate(VFORM_ORDER, start=2):
                count = int((df["vform"].eq(vform) & (age_series == age)).sum())
                counts[vform] = count
                ws.cell(idx, col_idx, count)
            for family_idx, family in enumerate(FORM_FAMILIES):
                family_total = sum(n for vform, n in counts.items() if vform.startswith(f"{family}_"))
                ws.cell(idx, family_col_start + family_idx, family_total)

    if "osep" in wb.sheetnames:
        ws = wb["osep"]
        form_to_col = layout.osep_form_to_col
        flag_to_row = layout.osep_flag_to_row
        for vform, col in form_to_col.items():
            data = df.loc[df["vform"].eq(vform)].copy()
            total = len(data)
            if not total:
                continue
            for flag, row in flag_to_row.items():
                if flag in data.columns:
                    n = pd.to_numeric(data[flag], errors="coerce").fillna(0).sum()
                    ws[f"{col}{row}"] = round(n / total * 100, 3)

    wb.save(output_path)


def write_updated_teid_workbook(
    final_sample: pd.DataFrame,
    output_path: Path,
    additions: pd.DataFrame,
    selected_swaps: pd.DataFrame,
    osep_target_qa: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        export_sample_columns(final_sample).to_excel(writer, sheet_name="final_all", index=False)
        summarize_sex(final_sample).to_excel(writer, sheet_name="sex_balance_qa", index=False)
        osep_target_qa.to_excel(writer, sheet_name="osep_target_qa", index=False)
        additions.to_excel(writer, sheet_name="part_b_additions", index=False)
        selected_swaps.to_excel(writer, sheet_name="part_b_osep_swaps", index=False)


def swap_row_from_addition(row: pd.Series) -> dict:
    return {
        "project": PROJECT,
        "vform": row.get("vform"),
        "agegrp": row.get("agegrp"),
        "step": row.name + 1,
        "action": row.get("action"),
        "removed_teid": "",
        "added_teid": row.get("added_teid"),
        "removed_sex": "",
        "added_sex": row.get("added_sex"),
        "female_n_before": row.get("female_n_before"),
        "male_n_before": row.get("male_n_before"),
        "female_n_after": row.get("female_n_after"),
        "male_n_after": row.get("male_n_after"),
        "source": "Optimization Engine Part B add-only sex N repair",
    }


def swap_row_from_osep_swap(row: pd.Series) -> dict:
    return {
        "project": PROJECT,
        "vform": row.get("vform"),
        "agegrp": row.get("agegrp"),
        "step": row.name + 1,
        "action": "osep_optimization_swap",
        "removed_teid": row.get("remove_teid"),
        "added_teid": row.get("add_teid"),
        "removed_sex": row.get("remove_sex"),
        "added_sex": row.get("add_sex"),
        "female_n_before": "",
        "male_n_before": "",
        "female_n_after": "",
        "male_n_after": "",
        "source": "Optimization Engine Part B OSEP swap optimizer",
    }


def write_swap_tables(additions: pd.DataFrame, selected_swaps: pd.DataFrame, output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for family in FORM_FAMILIES:
        ws = wb.create_sheet(family)
        ws["A1"] = f"{family}: Sample Changes during Optimization Engine Part B"
        add_rows = additions.loc[additions["vform"].astype(str).str.startswith(family)].copy() if not additions.empty else pd.DataFrame()
        swap_rows = selected_swaps.loc[selected_swaps["vform"].astype(str).str.startswith(family)].copy() if not selected_swaps.empty else pd.DataFrame()
        table = pd.concat(
            [
                pd.DataFrame([swap_row_from_addition(r) for _, r in add_rows.iterrows()]),
                pd.DataFrame([swap_row_from_osep_swap(r) for _, r in swap_rows.iterrows()]),
            ],
            ignore_index=True,
            sort=False,
        )
        if table.empty:
            table = pd.DataFrame(columns=[
                "project", "vform", "agegrp", "step", "action", "removed_teid", "added_teid",
                "removed_sex", "added_sex", "female_n_before", "male_n_before", "female_n_after",
                "male_n_after", "source",
            ])
        for col_idx, col in enumerate(table.columns, start=1):
            ws.cell(2, col_idx, col)
        for row_idx, record in enumerate(table.itertuples(index=False), start=3):
            for col_idx, value in enumerate(record, start=1):
                ws.cell(row_idx, col_idx, value)
    style_workbook(wb)
    wb.save(output_path)


def style_workbook(wb) -> None:
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    title_fill = PatternFill("solid", fgColor="1F4E78")
    border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = title_fill
        for cell in ws[2]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        ws.freeze_panes = "A3"
        for col_cells in ws.columns:
            letter = col_cells[0].column_letter
            width = min(40, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col_cells) + 2))
            ws.column_dimensions[letter].width = width


def copy_to_stable(run_file: Path, stable_file: Path) -> None:
    stable_file.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(run_file, stable_file)


def report_progress(percent: int, message: str) -> None:
    """Emit a console-friendly milestone that the GUI can display as progress."""
    print(f"[PROGRESS {percent:3d}%] {message}", flush=True)


def main() -> None:
    global TARGET_CELL_N, FINAL_TOTAL_N
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = OUTPUT_ROOT / f"run_{run_id}_part_b"
    run_dir.mkdir(parents=True, exist_ok=True)

    report_progress(5, "Loading the prepared Part A workbook")
    print("Optimization Engine Part B")
    print(f"Part A input: {PART_A_FILE}")

    master, available = load_part_a()
    profile_targets_by_age = PROJECT_PROFILE.get("target_n_by_agegrp") or {}
    profile_target_n = PROJECT_PROFILE.get("target_n_per_stratum")
    if profile_targets_by_age:
        try:
            normalized_targets = {int(agegrp): int(target) for agegrp, target in profile_targets_by_age.items()}
        except (TypeError, ValueError) as exc:
            raise ValueError("Project setup age-band targets must be whole numbers.") from exc
        if any(target < 1 for target in normalized_targets.values()):
            raise ValueError("Project setup age-band targets must be at least one.")
        observed = pd.concat([master[["vform", "agegrp"]], available[["vform", "agegrp"]]], ignore_index=True)
        observed = observed.drop_duplicates()
        missing = sorted(set(observed["agegrp"].astype(int)) - set(normalized_targets))
        if missing:
            raise ValueError(f"Project setup is missing target N values for age groups: {missing}")
        TARGET_CELL_N = {
            (str(row.vform), int(row.agegrp)): normalized_targets[int(row.agegrp)]
            for row in observed.itertuples()
        }
        FINAL_TOTAL_N = sum(TARGET_CELL_N.values())
        print("Project setup override: target N by age band " + ", ".join(
            f"{agegrp}={target:,}" for agegrp, target in sorted(normalized_targets.items())
        ))
    elif profile_target_n is not None:
        try:
            target_n = int(profile_target_n)
        except (TypeError, ValueError) as exc:
            raise ValueError("Project setup target N per stratum must be a whole number.") from exc
        if target_n < 1:
            raise ValueError("Project setup target N per stratum must be at least one.")
        observed = pd.concat([master[["vform", "agegrp"]], available[["vform", "agegrp"]]], ignore_index=True)
        TARGET_CELL_N = {(str(row.vform), int(row.agegrp)): target_n for row in observed.drop_duplicates().itertuples()}
        FINAL_TOTAL_N = sum(TARGET_CELL_N.values())
        print(f"Project setup override: target N={target_n:,} for each observed form/age group cell")
    report_progress(15, "Validating target capacity and reading target workbook")
    if CFG.validate_target_capacity:
        validate_target_capacity(master, available)
    demo_targets = parse_demo_targets(TARGET_TEMPLATE) if INCLUDE_CENSUS else {}
    report_progress(30, "Repairing target sample sizes and sex balance")
    if DISABLE_OPTIMIZATION:
        additions = pd.DataFrame()
        selected_swaps = pd.DataFrame()
        all_swaps = pd.DataFrame()
        final_sample = master.copy()
        remaining_available = available.copy()
        print("Pass-through mode: repair, balancing, and CP-SAT optimization bypassed")
    elif FORCED_ONLY_SWAP:
        additions = pd.DataFrame()
        final_sample, remaining_available, selected_swaps = apply_forced_only_swap(
            master, available, FORCED_ONLY_SWAP
        )
        all_swaps = selected_swaps.copy()
        print(f"Forced-only swap: {FORCED_ONLY_SWAP}; CP-SAT bypassed")
    else:
        repaired_sample, remaining_available, additions = repair_target_n_and_sex_by_adding(master, available)
        final_sample, selected_swaps, all_swaps = optimize_osep_swaps(
            repaired_sample, remaining_available, demo_targets
        )
    if CFG.rebuild_available_pool_from_universe:
        # The final available pool must reflect both add-only repairs and
        # selected OSEP swaps: it is every source TEID not in the final
        # sample. Rebuilding it from the original master + available
        # universe avoids leaving a swapped-out TEID unavailable.
        source_universe = pd.concat([master, available], ignore_index=True, sort=False)
        source_universe = source_universe.drop_duplicates(subset=["teid"], keep="first").copy()
        final_teids = set(final_sample["teid"])
        remaining_available = normalize(
            source_universe.loc[~source_universe["teid"].isin(final_teids)].copy()
        )
    report_progress(86, "Checking final sample and building quality assurance tables")
    sex_qa = summarize_sex(final_sample)
    osep_target_qa = build_osep_target_qa(final_sample, remaining_available)
    failures = sex_qa.loc[~sex_qa["exact_n_match"]]
    if not failures.empty:
        raise ValueError(
            "Part B could not make every vform+agegrp cell exact by sex N:\n"
            + failures.to_string(index=False)
        )
    n_failures = []
    for (vform, agegrp), target_n in TARGET_CELL_N.items():
        actual_n = len(final_sample.loc[final_sample["vform"].eq(vform) & final_sample["agegrp"].eq(agegrp)])
        if actual_n != target_n:
            n_failures.append({"vform": vform, "agegrp": agegrp, "target_n": target_n, "actual_n": actual_n})
    if n_failures:
        raise ValueError("Part B did not meet target cell Ns:\n" + pd.DataFrame(n_failures).to_string(index=False))
    if FINAL_TOTAL_N is not None and len(final_sample) != FINAL_TOTAL_N:
        raise ValueError(f"Part B final total N must be {FINAL_TOTAL_N}, got {len(final_sample)}")
    invalid_edl = final_sample.loc[~final_sample["edl"].isin(VALID_EDL_VALUES), ["vform", "agegrp", "teid", "edl"]]
    if not invalid_edl.empty:
        raise ValueError(
            "Part B final sample contains invalid EDL values; only EDL 1-4 are allowed:\n"
            + invalid_edl.to_string(index=False)
        )

    updated_teid_file = run_dir / UPDATED_TEID_NAME
    demo_osep_file = run_dir / DEMO_OSEP_NAME
    swap_table_file = run_dir / SWAP_TABLE_NAME

    report_progress(91, "Writing final Part B deliverables")
    write_updated_teid_workbook(final_sample, updated_teid_file, additions, selected_swaps, osep_target_qa)
    write_demo_osep_workbook(final_sample, demo_osep_file)
    write_swap_tables(additions, selected_swaps, swap_table_file)

    extra_deliverables: list[tuple[Path, Path]] = []
    if SELECTED_TEID_NAME:
        selected_teid_file = run_dir / SELECTED_TEID_NAME
        with pd.ExcelWriter(selected_teid_file, engine="openpyxl") as writer:
            export_sample_columns(final_sample).to_excel(writer, sheet_name="selected_teids", index=False)
        extra_deliverables.append((selected_teid_file, STABLE_DELIVERABLE_DIR / SELECTED_TEID_NAME))
    if AVAILABLE_POOL_NAME:
        available_pool_file = run_dir / AVAILABLE_POOL_NAME
        with pd.ExcelWriter(available_pool_file, engine="openpyxl") as writer:
            export_sample_columns(remaining_available).to_excel(writer, sheet_name="available_pool", index=False)
        extra_deliverables.append((available_pool_file, STABLE_DELIVERABLE_DIR / AVAILABLE_POOL_NAME))

    with pd.ExcelWriter(run_dir / "Optimization_Engine_Part_B_QA.xlsx", engine="openpyxl") as writer:
        sex_qa.to_excel(writer, sheet_name="sex_balance_qa", index=False)
        osep_target_qa.to_excel(writer, sheet_name="osep_target_qa", index=False)
        additions.to_excel(writer, sheet_name="part_b_additions", index=False)
        selected_swaps.to_excel(writer, sheet_name="part_b_osep_swaps", index=False)
        all_swaps.to_excel(writer, sheet_name="all_candidate_swaps", index=False)
        remaining_available.to_excel(writer, sheet_name="remaining_available_pool", index=False)

    report_progress(97, "Publishing latest deliverables")
    copy_to_stable(updated_teid_file, STABLE_DELIVERABLE_DIR / UPDATED_TEID_NAME)
    copy_to_stable(demo_osep_file, STABLE_DELIVERABLE_DIR / DEMO_OSEP_NAME)
    copy_to_stable(swap_table_file, STABLE_DELIVERABLE_DIR / SWAP_TABLE_NAME)
    for run_file, stable_file in extra_deliverables:
        copy_to_stable(run_file, stable_file)

    report_progress(100, "Part B is complete")
    print(f"Final sample N: {len(final_sample):,}")
    print(f"Added cases for exact N/sex balance: {len(additions.loc[additions['action'].eq('add_for_exact_sex_n_balance')]) if not additions.empty else 0}")
    print(f"OSEP optimization swaps selected: {len(selected_swaps)}")
    print(f"Updated TEID file: {updated_teid_file}")
    print(f"Demo/OSEP file: {demo_osep_file}")
    print(f"Swap table file: {swap_table_file}")
    print(f"Stable deliverables: {STABLE_DELIVERABLE_DIR}")


if __name__ == "__main__":
    main()
