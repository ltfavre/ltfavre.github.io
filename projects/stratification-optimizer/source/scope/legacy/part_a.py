from __future__ import annotations

from datetime import datetime
from pathlib import Path
import json
import math
import os
import shutil

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from test_config import load_test_config
from project_profile import (
    apply_age_bands, apply_clinical_groups, apply_field_map, apply_scope_filter,
    load_project_profile, selected_demographics,
)


# =====================================================
# Optimization Engine Part A
#
# Consolidated prep engine from the old 58 and 60-90
# workflow family. Part A cleans, organizes, visualizes,
# first-admin filters, builds used/available pools, and
# forces exact sex balance where the cell N allows it.
#
# All test/clinical-form-specific values (pool rules,
# clinical flags, file paths, output names, ...) are
# loaded from a YAML config selected via
# OPT_ENGINE_TEST_CONFIG (see test_config.py).
# =====================================================

TEST_CONFIG_ID = os.environ.get("OPT_ENGINE_TEST_CONFIG", "basc4_adhd")
CFG = load_test_config(TEST_CONFIG_ID)
PROJECT_PROFILE = load_project_profile()

PROJECT_DIR = Path(os.environ.get(
    "OPT_ENGINE_PROJECT_DIR",
    CFG.default_paths.default_project_dir or ".",
))

FR_TEID_FILE = Path(os.environ.get(
    "OPT_ENGINE_FR_TEID_FILE",
    str(PROJECT_DIR / CFG.default_paths.fr_teid_file),
))
DEFAULT_MASTER_INPUT = PROJECT_DIR / CFG.default_paths.default_master_input
DEFAULT_MASTER_SHEET = CFG.default_paths.default_master_sheet
FR_ADMIN_SHEETS = CFG.default_paths.fr_admin_sheets
CASE_SWAP_DIR = Path(os.environ.get(
    "OPT_ENGINE_CASE_SWAP_DIR",
    str(PROJECT_DIR / CFG.default_paths.case_swap_dir),
))
TARGET_WORKBOOK = Path(os.environ.get(
    "OPT_ENGINE_TARGET_WORKBOOK",
    str(PROJECT_DIR / CFG.default_paths.target_workbook),
))
SCORED_DATA_DIR = Path(os.environ.get(
    "OPT_ENGINE_SCORED_DATA_DIR",
    str(PROJECT_DIR / CFG.default_paths.scored_data_dir),
))
OUTPUT_ROOT = Path(os.environ.get(
    "OPT_ENGINE_PART_A_OUTPUT_ROOT",
    str(PROJECT_DIR / CFG.default_paths.part_a_output_root),
))
CANDIDATE_DATASETS = json.loads(os.environ.get("OPT_ENGINE_CANDIDATE_DATASETS_JSON", "[]") or "[]")

# Optional master-sample override. When supplied, Part A uses the indicated
# workbook/sheet as the current master and rebuilds the available pool from the
# scored SAS universe after excluding its TEIDs. With no override, the original
# FR-TEID workflow remains unchanged.
MASTER_INPUT_OVERRIDE = os.environ.get("OPT_ENGINE_MASTER_INPUT", "").strip()
MASTER_INPUT_SHEET_OVERRIDE = os.environ.get("OPT_ENGINE_MASTER_SHEET", "").strip()
DISABLE_OPTIMIZATION = os.environ.get("OPT_ENGINE_DISABLE_OPTIMIZATION", "").strip().upper() in {
    "1", "TRUE", "Y", "YES",
}

PROJECT = CFG.project
TARGET_FLAG = CFG.target_flag
TARGET_PCT = CFG.target_pct
SEX_COL = CFG.sex_col
FEMALE_VALUE = CFG.female_value
MALE_VALUE = CFG.male_value

CASE_SWAP_FILES = CFG.default_paths.case_swap_files

POOL_RULES = {
    vform: {
        "file": rule.file,
        "standard_testid": rule.standard_testid,
        "exclude_testid": rule.exclude_testid,
        "agegrp": rule.agegrp,
    }
    for vform, rule in CFG.pool_rules.items()
}

DEMO_COLS = selected_demographics(CFG.demo_cols, PROJECT_PROFILE)
FLAG_COLS = CFG.flag_cols
AGE_SOURCE_COLS = CFG.age_source_cols
KEY_COLS = CFG.key_cols
EXPORT_COLS = CFG.export_cols


def clean_teid(value) -> str | pd.NA:
    if pd.isna(value):
        return pd.NA
    text = str(value).strip()
    numeric = pd.to_numeric(pd.Series([text]), errors="coerce").iloc[0]
    if pd.notna(numeric):
        return str(int(numeric))
    return text if text else pd.NA


def clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [c.decode() if isinstance(c, bytes) else c for c in out.columns]
    out.columns = [str(c).strip() for c in out.columns]
    lower_map = {c.lower(): c for c in out.columns}
    for col in KEY_COLS + DEMO_COLS + FLAG_COLS + ["testid", "adminorder"]:
        if col not in out.columns and col.lower() in lower_map:
            out[col] = out[lower_map[col.lower()]]
    return out


def require_columns(df: pd.DataFrame, cols: list[str], label: str) -> None:
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise ValueError(f"{label} is missing required columns: {missing}")


def valid_cid(series: pd.Series) -> pd.Series:
    """True where a CID indicator column marks an actual diagnosis code."""
    text = series.fillna("").astype(str).str.strip()
    return text.ne("") & text.ne("0") & text.ne("0.0") & text.str.lower().ne("nan")


def derive_osep_flags(df: pd.DataFrame) -> pd.DataFrame:
    """Derive OSEP flags from a "study" text column and/or CID indicator
    columns, for source data without precomputed flag columns. Driven by
    CFG.flag_study_patterns (regex per flag) and CFG.cid_groups; both are
    empty for forms whose source data already carries flag columns, so this
    only changes behavior for configs that populate them (currently Mood)."""
    study = df.get("study", pd.Series("", index=df.index)).fillna("").astype(str)
    flags = pd.DataFrame(index=df.index)
    for flag in FLAG_COLS:
        pattern = CFG.flag_study_patterns.get(flag)
        flags[flag] = study.str.contains(pattern, case=False, regex=True, na=False) if pattern else False
    for group, cid_columns in CFG.cid_groups.items():
        flag = f"tmp_{group}"
        if flag not in flags.columns:
            continue
        cid_match = pd.Series(False, index=df.index)
        for cid in cid_columns:
            if cid in df.columns:
                cid_match |= valid_cid(df[cid])
        flags[flag] = flags[flag] | cid_match
    return flags


def derive_agegrp(age: pd.Series) -> pd.Series:
    """Map age-in-years to the standard 6-7 / 8-11 / 12-14 / 15-18 cells."""
    numeric_age = pd.to_numeric(age, errors="coerce")
    return (
        pd.Series(pd.NA, index=age.index, dtype="Int64")
        .mask(numeric_age.between(6, 7), 3)
        .mask(numeric_age.between(8, 11), 4)
        .mask(numeric_age.between(12, 14), 5)
        .mask(numeric_age.between(15, 18), 6)
    )


def normalize_sample(df: pd.DataFrame, label: str) -> pd.DataFrame:
    out = clean_columns(df)
    out = apply_field_map(out, PROJECT_PROFILE)
    require_columns(out, ["teid", "vform"], label)
    if CFG.set_project_column:
        out["project"] = PROJECT
    out["teid"] = out["teid"].apply(clean_teid)
    out = out.dropna(subset=["teid"]).copy()
    out["vform"] = out["vform"].astype(str).str.strip().str.upper()

    age_source = None
    for col in AGE_SOURCE_COLS:
        if col not in out.columns:
            continue
        values = pd.to_numeric(out[col], errors="coerce")
        age_source = values if age_source is None else age_source.fillna(values)
    if age_source is not None:
        out["agey"] = age_source

    out = apply_age_bands(out, age_source, PROJECT_PROFILE)

    if "agegrp" not in out.columns and CFG.derive_agegrp_if_missing and age_source is not None:
        out["agegrp"] = derive_agegrp(age_source)
    require_columns(out, ["agegrp"], label)
    out["agegrp"] = pd.to_numeric(out["agegrp"], errors="coerce").astype("Int64")
    out = out.dropna(subset=["agegrp"]).copy()
    out["agegrp"] = out["agegrp"].astype(int)

    for col in DEMO_COLS:
        if col in out.columns:
            if col == "edl":
                numeric = pd.to_numeric(out[col], errors="coerce")
                out[col] = numeric.astype("Int64").astype(str).replace("<NA>", "")
            else:
                out[col] = out[col].astype(str).str.strip().str.upper()
    derived_flags = derive_osep_flags(out) if (CFG.flag_study_patterns or CFG.cid_groups) else None
    for flag in FLAG_COLS:
        aliases = [flag, flag.lower(), flag.upper()]
        source = next((col for col in aliases if col in out.columns), None)
        if source is None:
            out[flag] = derived_flags[flag] if derived_flags is not None else 0
        elif source != flag:
            out[flag] = out[source]
        out[flag] = pd.to_numeric(out[flag], errors="coerce").fillna(0).astype(int)
    out = apply_clinical_groups(out, PROJECT_PROFILE)
    out["tmp_clin"] = out[FLAG_COLS].astype(bool).any(axis=1).astype(int)
    return out


def apply_first_admin_filter(df: pd.DataFrame, label: str) -> tuple[pd.DataFrame, dict]:
    require_columns(df, ["testid"], label)
    before = len(df)
    testid = pd.to_numeric(df["testid"], errors="coerce")
    keep = pd.Series(False, index=df.index)
    for vform, rule in POOL_RULES.items():
        standard_testid = rule.get("standard_testid", set())
        exclude_testid = rule.get("exclude_testid", set())
        age_eligible = pd.Series(True, index=df.index) if PROJECT_PROFILE.get("age_bands") else df["agegrp"].isin(rule["agegrp"])
        keep |= (
            df["vform"].eq(vform)
            & age_eligible
            & testid.isin(standard_testid)
            & ~testid.isin(exclude_testid)
        )
    out = df.loc[keep].copy()
    return out, {
        "stage": "first_admin_filter",
        "table": label,
        "rows_before": before,
        "rows_after": len(out),
        "rows_removed": before - len(out),
        "filter": "keep standardization English/Spanish testids; exclude retest/other-rater testids",
    }


def load_fr_teids() -> tuple[set[str], pd.DataFrame]:
    rows = []
    teids = set()
    for sheet in FR_ADMIN_SHEETS:
        raw = pd.read_excel(FR_TEID_FILE, sheet_name=sheet)
        if "teid" not in raw.columns:
            raise ValueError(f"{FR_TEID_FILE} sheet {sheet} is missing teid.")
        cleaned = raw["teid"].apply(clean_teid).dropna().astype(str)
        teids.update(cleaned)
        rows.append({
            "stage": "load_fr_teids",
            "sheet": sheet,
            "rows_read": len(raw),
            "unique_teids": cleaned.nunique(),
        })
    return teids, pd.DataFrame(rows)


def load_source_universe() -> tuple[pd.DataFrame, list[dict]]:
    parts = []
    logs = []
    for vform, rule in POOL_RULES.items():
        path = SCORED_DATA_DIR / rule["file"]
        raw = pd.read_sas(path, encoding="latin1")
        raw = clean_columns(raw)
        raw["vform"] = vform
        raw["source_file"] = rule["file"]
        if "agegrp" not in raw.columns and not CFG.derive_agegrp_if_missing:
            raise ValueError(f"{path} does not contain agegrp.")
        pool = apply_scope_filter(normalize_sample(raw, f"{vform} source pool"), PROJECT_PROFILE)
        pool, filter_log = apply_first_admin_filter(pool, f"{vform} source pool")
        parts.append(pool)
        logs.append({
            "stage": "load_source_pool",
            "vform": vform,
            "file": str(path),
            "rows_read": len(raw),
            "rows_after_first_admin_filter": len(pool),
        })
        logs.append(filter_log)
    return pd.concat(parts, ignore_index=True, sort=False), logs


def build_master_and_available_from_raw() -> tuple[pd.DataFrame, pd.DataFrame, list[dict]]:
    fr_teids, fr_log = load_fr_teids()
    universe, source_logs = load_source_universe()
    master = universe.loc[universe["teid"].isin(fr_teids)].copy()
    missing_teids = sorted(fr_teids - set(master["teid"]))
    before_dedup = len(master)
    master = master.drop_duplicates(subset=["teid"], keep="first").copy()
    available = universe.loc[~universe["teid"].isin(set(master["teid"]))].copy()
    available = available.drop_duplicates(subset=["teid"], keep="first").copy()
    logs = [
        {"stage": "config", "item": "fr_teid_file", "value": str(FR_TEID_FILE)},
        {"stage": "build_master_from_raw", "rows_before_dedup": before_dedup, "rows_after_dedup": len(master)},
        {"stage": "build_available_from_raw", "rows": len(available)},
        {"stage": "fr_teids_missing_after_first_admin_filter", "rows": len(missing_teids), "value": ", ".join(missing_teids[:50])},
        *fr_log.to_dict("records"),
        *source_logs,
    ]
    return master, available, logs


def build_master_and_available_from_workbook(path: Path, sheet_name: str) -> tuple[pd.DataFrame, pd.DataFrame, list[dict]]:
    if not path.exists():
        raise FileNotFoundError(f"Master input override does not exist: {path}")
    raw_master = pd.read_excel(path, sheet_name=sheet_name)
    if TEST_CONFIG_ID == "basc4_mood":
        # Membership comes from the previous selection; every case value
        # comes from the newly scored, eligible source universe.
        from membership_refresh import refresh_membership
        universe, source_logs = load_source_universe()
        master, available, audit = refresh_membership(raw_master, universe, clean_teid)
        audit_path = OUTPUT_ROOT / "Sample_Refresh_Audit.csv"
        audit_path.parent.mkdir(parents=True, exist_ok=True)
        audit.to_csv(audit_path, index=False)
        print(f"Refreshed {len(master):,} selected TEIDs from current scored records; audit: {audit_path}")
        return master, available, [
            {"stage": "refresh_membership", "item": status, "rows": int(count)}
            for status, count in audit.status.value_counts().items()
        ] + source_logs
    master = apply_scope_filter(normalize_sample(raw_master, f"master input override {path.name}:{sheet_name}"), PROJECT_PROFILE)
    before_dedup = len(master)
    master = master.drop_duplicates(subset=["teid"], keep="first").copy()
    if len(master) != before_dedup:
        raise ValueError(
            f"Master input override contains duplicate TEIDs: {before_dedup - len(master)} duplicate rows."
        )
    available, source_logs = load_source_available_pool(master)
    logs = [
        {"stage": "config", "item": "master_input_override", "value": str(path)},
        {"stage": "config", "item": "master_input_sheet", "value": sheet_name},
        {"stage": "build_master_from_workbook", "rows_before_dedup": before_dedup, "rows_after_dedup": len(master)},
        {"stage": "build_available_from_raw", "rows": len(available)},
        *source_logs,
    ]
    return master, available, logs


def infer_vform_agegrp_from_filename(path: Path) -> tuple[str | None, int | None]:
    parts = path.stem.upper().split("_")
    vform = None
    agegrp = None
    for i in range(len(parts) - 1):
        candidate = f"{parts[i]}_{parts[i + 1]}"
        if candidate in POOL_RULES:
            vform = candidate
            break
    for part in parts:
        if part.startswith("AGEGRP"):
            agegrp = int(part.replace("AGEGRP", ""))
    return vform, agegrp


def load_case_swap_pool_audit() -> pd.DataFrame:
    rows = []
    for file_name in CASE_SWAP_FILES:
        path = CASE_SWAP_DIR / file_name
        if not path.exists():
            rows.append({"file": file_name, "status": "missing", "rows": 0, "sheets": ""})
            continue
        xls = pd.ExcelFile(path)
        if "replacement_pool" not in xls.sheet_names:
            rows.append({
                "file": file_name,
                "status": "missing_replacement_pool_sheet",
                "rows": 0,
                "sheets": ", ".join(xls.sheet_names),
            })
            continue
        df = pd.read_excel(path, sheet_name="replacement_pool")
        rows.append({
            "file": file_name,
            "status": "ok",
            "rows": len(df),
            "sheets": ", ".join(xls.sheet_names),
        })
    return pd.DataFrame(rows)


def load_source_available_pool(master: pd.DataFrame) -> tuple[pd.DataFrame, list[dict]]:
    parts = []
    logs = []
    used_teids = set(master["teid"].dropna().astype(str))
    for vform, rule in POOL_RULES.items():
        path = SCORED_DATA_DIR / rule["file"]
        raw = pd.read_sas(path, encoding="latin1")
        raw = clean_columns(raw)
        raw["vform"] = vform
        if "agegrp" not in raw.columns and not CFG.derive_agegrp_if_missing:
            raise ValueError(f"{path} does not contain agegrp.")
        pool = apply_scope_filter(normalize_sample(raw, f"{vform} source pool"), PROJECT_PROFILE)
        pool, filter_log = apply_first_admin_filter(pool, f"{vform} source pool")
        before_used_drop = len(pool)
        pool = pool.loc[~pool["teid"].isin(used_teids)].copy()
        pool["source_file"] = rule["file"]
        parts.append(pool)
        logs.append({
            "stage": "load_source_pool",
            "vform": vform,
            "file": str(path),
            "rows_after_first_admin_filter": before_used_drop,
            "rows_available_after_used_drop": len(pool),
        })
        logs.append(filter_log)
    for item in CANDIDATE_DATASETS:
        path = Path(str(item.get("path", "")))
        if not path.is_file():
            raise FileNotFoundError(f"Supplemental candidate source does not exist: {path}")
        suffix = path.suffix.casefold()
        if suffix == ".csv": raw = pd.read_csv(path)
        elif suffix in {".sas7bdat", ".xpt"}: raw = pd.read_sas(path, format="xport" if suffix == ".xpt" else "sas7bdat", encoding="latin1")
        elif suffix in {".xlsx", ".xlsm", ".xls"}: raw = pd.read_excel(path)
        else: raise ValueError(f"Unsupported supplemental candidate source: {path.name}")
        raw = clean_columns(raw)
        pool = apply_scope_filter(normalize_sample(raw, f"supplemental source {path.name}"), PROJECT_PROFILE)
        pool, filter_log = apply_first_admin_filter(pool, f"supplemental source {path.name}")
        pool = pool.loc[~pool["teid"].isin(used_teids)].copy()
        pool["SCOPE_SOURCE_FILE"] = path.name
        pool["SCOPE_SOURCE_TYPE"] = str(item.get("source_type", "Unspecified"))
        pool["SCOPE_SOURCE_ROLE"] = str(item.get("source_role", "Supplemental"))
        pool["SCOPE_SOURCE_PRIORITY"] = int(item.get("priority", 999))
        parts.append(pool)
        logs.append({"stage": "load_supplemental_candidate", "file": str(path), "source_type": item.get("source_type"), "source_role": item.get("source_role"), "rows_available_after_used_drop": len(pool)})
        logs.append(filter_log)
    if not parts:
        return pd.DataFrame(columns=master.columns), logs
    available = pd.concat(parts, ignore_index=True, sort=False)
    if "SCOPE_SOURCE_PRIORITY" in available:
        available["SCOPE_SOURCE_PRIORITY"] = pd.to_numeric(available["SCOPE_SOURCE_PRIORITY"], errors="coerce").fillna(1000)
        available = available.sort_values("SCOPE_SOURCE_PRIORITY", kind="stable")
    available = available.drop_duplicates(subset=["teid"], keep="first").copy()
    return available, logs


def sex_counts(df: pd.DataFrame) -> dict[str, int]:
    if SEX_COL not in df.columns:
        return {FEMALE_VALUE: 0, MALE_VALUE: 0}
    values = df[SEX_COL].astype(str).str.strip().str.upper()
    return {
        FEMALE_VALUE: int((values == FEMALE_VALUE).sum()),
        MALE_VALUE: int((values == MALE_VALUE).sum()),
    }


def row_flag_value(row: pd.Series, flag: str) -> int:
    return int(pd.to_numeric(pd.Series([row.get(flag, 0)]), errors="coerce").fillna(0).iloc[0])


def swap_score(remove_row: pd.Series, add_row: pd.Series) -> float:
    score = 0.0
    for flag in FLAG_COLS:
        score += 1000.0 * abs(row_flag_value(remove_row, flag) - row_flag_value(add_row, flag))
    for col in ["eth", "edl", "region"]:
        score += 10.0 * int(str(remove_row.get(col, "")).upper() != str(add_row.get(col, "")).upper())
    remove_age = pd.to_numeric(pd.Series([remove_row.get("agey")]), errors="coerce").iloc[0]
    add_age = pd.to_numeric(pd.Series([add_row.get("agey")]), errors="coerce").iloc[0]
    if pd.notna(remove_age) and pd.notna(add_age):
        score += abs(float(remove_age) - float(add_age))
    return score


def force_exact_sex_balance(
    master: pd.DataFrame,
    available: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    current = master.copy()
    pool = available.copy()
    swaps = []
    unresolved = []

    for (vform, agegrp), cell in current.groupby(["vform", "agegrp"], sort=True):
        n = len(cell)
        counts = sex_counts(cell)
        if n % 2 != 0:
            unresolved.append({
                "vform": vform,
                "agegrp": agegrp,
                "reason": "odd_n_skip",
                "n": n,
                "female_n": counts[FEMALE_VALUE],
                "male_n": counts[MALE_VALUE],
            })
            continue
        target_each = n // 2
        female_surplus = counts[FEMALE_VALUE] - target_each
        if female_surplus == 0:
            continue
        remove_sex = FEMALE_VALUE if female_surplus > 0 else MALE_VALUE
        add_sex = MALE_VALUE if female_surplus > 0 else FEMALE_VALUE
        swaps_needed = abs(female_surplus)

        for _ in range(swaps_needed):
            cell = current.loc[(current["vform"] == vform) & (current["agegrp"] == agegrp)].copy()
            remove_candidates = cell.loc[cell[SEX_COL].astype(str).str.upper() == remove_sex].copy()
            add_candidates = pool.loc[
                (pool["vform"] == vform)
                & (pool["agegrp"] == agegrp)
                & (pool[SEX_COL].astype(str).str.upper() == add_sex)
            ].copy()
            if remove_candidates.empty or add_candidates.empty:
                unresolved.append({
                    "vform": vform,
                    "agegrp": agegrp,
                    "reason": "no_feasible_sex_balance_swap",
                    "remove_sex": remove_sex,
                    "add_sex": add_sex,
                    "remaining_needed": swaps_needed,
                })
                break

            best = None
            best_score = math.inf
            for _, rem in remove_candidates.iterrows():
                same_flag_adds = add_candidates.loc[
                    add_candidates[TARGET_FLAG] == row_flag_value(rem, TARGET_FLAG)
                ]
                candidate_adds = same_flag_adds if not same_flag_adds.empty else add_candidates
                for _, add in candidate_adds.iterrows():
                    score = swap_score(rem, add)
                    if score < best_score:
                        best_score = score
                        best = (rem, add)
            if best is None:
                unresolved.append({
                    "vform": vform,
                    "agegrp": agegrp,
                    "reason": "no_scored_candidate",
                    "remove_sex": remove_sex,
                    "add_sex": add_sex,
                })
                break

            rem, add = best
            current = current.loc[current["teid"] != rem["teid"]].copy()
            add_row = add.copy()
            add_row["part_a_added_for_exact_sex_balance"] = 1
            current = pd.concat([current, add_row.to_frame().T], ignore_index=True, sort=False)
            pool = pool.loc[pool["teid"] != add["teid"]].copy()
            pool = pd.concat([pool, rem.to_frame().T], ignore_index=True, sort=False)
            swaps.append({
                "vform": vform,
                "agegrp": agegrp,
                "remove_teid": rem["teid"],
                "remove_sex": remove_sex,
                "add_teid": add["teid"],
                "add_sex": add_sex,
                "score": best_score,
                "preserved_target_flag": row_flag_value(rem, TARGET_FLAG) == row_flag_value(add, TARGET_FLAG),
            })

    current = normalize_sample(current, "sex-balanced master")
    pool = normalize_sample(pool, "sex-balanced available pool")
    return current, pool, pd.DataFrame(swaps + unresolved)


def summarize_cells(df: pd.DataFrame, label: str) -> pd.DataFrame:
    rows = []
    for (vform, agegrp), grp in df.groupby(["vform", "agegrp"], sort=True):
        counts = sex_counts(grp)
        row = {
            "table": label,
            "vform": vform,
            "agegrp": agegrp,
            "n": len(grp),
            "female_n": counts[FEMALE_VALUE],
            "male_n": counts[MALE_VALUE],
            "exact_sex_balance": counts[FEMALE_VALUE] == counts[MALE_VALUE],
            "odd_n": len(grp) % 2 != 0,
        }
        for flag in FLAG_COLS + ["tmp_clin"]:
            if flag in grp.columns:
                flag_n = int(pd.to_numeric(grp[flag], errors="coerce").fillna(0).sum())
                row[f"{flag}_n"] = flag_n
                row[f"{flag}_pct"] = round(flag_n / len(grp) * 100, 3) if len(grp) else 0
        rows.append(row)
    return pd.DataFrame(rows)


def summarize_overall(df: pd.DataFrame, label: str) -> pd.DataFrame:
    rows = [{"metric": "rows", "table": label, "value": len(df)}]
    rows.append({"metric": "unique_teids", "table": label, "value": df["teid"].nunique()})
    for flag in FLAG_COLS + ["tmp_clin"]:
        if flag in df.columns and len(df):
            n = int(pd.to_numeric(df[flag], errors="coerce").fillna(0).sum())
            rows.append({"metric": f"{flag}_n", "table": label, "value": n})
            rows.append({"metric": f"{flag}_pct", "table": label, "value": round(n / len(df) * 100, 3)})
    return pd.DataFrame(rows)


def parse_target_workbook_audit() -> pd.DataFrame:
    wb = load_workbook(TARGET_WORKBOOK, read_only=True, data_only=True)
    rows = []
    for sheet in ["trs", "prs", "srp", "osep", "byage"]:
        rows.append({
            "target_sheet": sheet,
            "exists": sheet in wb.sheetnames,
            "max_row": wb[sheet].max_row if sheet in wb.sheetnames else None,
            "max_column": wb[sheet].max_column if sheet in wb.sheetnames else None,
        })
    return pd.DataFrame(rows)


def prep_export(df: pd.DataFrame) -> pd.DataFrame:
    cols = [col for col in EXPORT_COLS if col in df.columns]
    extra = [col for col in KEY_COLS + DEMO_COLS + FLAG_COLS if col in df.columns and col not in cols]
    return df.loc[:, cols + extra].copy()


def autosize(ws) -> None:
    for col_cells in ws.columns:
        letter = col_cells[0].column_letter
        width = min(45, max(10, max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells) + 2))
        ws.column_dimensions[letter].width = width


def write_workbook(
    path: Path,
    master_pre: pd.DataFrame,
    master_final: pd.DataFrame,
    available: pd.DataFrame,
    run_log: pd.DataFrame,
    case_swap_audit: pd.DataFrame,
    sex_balance_log: pd.DataFrame,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.concat(
            [
                summarize_overall(master_pre, "master_before_sex_balance"),
                summarize_overall(master_final, "master_after_sex_balance"),
                summarize_overall(available, "available_after_sex_balance"),
            ],
            ignore_index=True,
        ).to_excel(writer, sheet_name="summary", index=False)
        pd.concat(
            [
                summarize_cells(master_pre, "master_before_sex_balance"),
                summarize_cells(master_final, "master_after_sex_balance"),
                summarize_cells(available, "available_after_sex_balance"),
            ],
            ignore_index=True,
        ).to_excel(writer, sheet_name="cell_summary", index=False)
        prep_export(master_final).to_excel(writer, sheet_name="master_prepped", index=False)
        prep_export(available).to_excel(writer, sheet_name="available_pool", index=False)
        sex_balance_log.to_excel(writer, sheet_name="sex_balance_log", index=False)
        case_swap_audit.to_excel(writer, sheet_name="case_swap_audit", index=False)
        parse_target_workbook_audit().to_excel(writer, sheet_name="target_workbook_audit", index=False)
        run_log.to_excel(writer, sheet_name="run_log", index=False)

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F81BD")
        autosize(ws)

    if "cell_summary" in wb.sheetnames:
        ws = wb["cell_summary"]
        chart = BarChart()
        chart.title = "Cell Counts by Sex"
        chart.y_axis.title = "N"
        chart.x_axis.title = "Cell"
        data = Reference(ws, min_col=5, max_col=6, min_row=1, max_row=min(ws.max_row, 40))
        cats = Reference(ws, min_col=3, min_row=2, max_row=min(ws.max_row, 40))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 9
        chart.width = 16
        ws.add_chart(chart, "T2")

    wb.save(path)


def report_progress(percent: int, message: str) -> None:
    """Emit a console-friendly milestone that the GUI can display as progress."""
    print(f"[PROGRESS {percent:3d}%] {message}", flush=True)


def main() -> None:
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = OUTPUT_ROOT / f"run_{run_id}_part_a_prep"
    output_file = run_dir / "Optimization_Engine_Part_A_prep_output.xlsx"

    master_input = Path(MASTER_INPUT_OVERRIDE) if MASTER_INPUT_OVERRIDE else None
    master_sheet = MASTER_INPUT_SHEET_OVERRIDE or "final_all"
    if master_input is None and not FR_TEID_FILE.exists():
        if DEFAULT_MASTER_INPUT.exists():
            master_input = DEFAULT_MASTER_INPUT
            master_sheet = MASTER_INPUT_SHEET_OVERRIDE or DEFAULT_MASTER_SHEET
        else:
            raise FileNotFoundError(
                "No Part A master input was found. The legacy FR TEID workbook is missing at "
                f"{FR_TEID_FILE}, and the default master workbook is missing at "
                f"{DEFAULT_MASTER_INPUT}. Set OPT_ENGINE_MASTER_INPUT and, if needed, "
                "OPT_ENGINE_MASTER_SHEET."
            )

    report_progress(5, "Checking Part A inputs")
    if master_input is not None:
        with pd.ExcelFile(master_input) as book:
            master_sheet = next((name for name in (master_sheet, "final_all", "master_prepped", "selected_teids") if name in book.sheet_names), book.sheet_names[0])
    print("Optimization Engine Part A")
    report_progress(15, "Loading master sample and source pool")
    if master_input is not None:
        print(f"Master workbook input: {master_input}")
        print(f"Master workbook sheet: {master_sheet}")
    else:
        print(f"FR TEID input: {FR_TEID_FILE}")
    print(f"SAS scored folder: {SCORED_DATA_DIR}")
    print(f"Case swap audit folder: {CASE_SWAP_DIR}")
    print(f"Target workbook: {TARGET_WORKBOOK}")
    if PROJECT_PROFILE:
        print(f"Project setup profile: {os.environ.get('OPT_ENGINE_PROJECT_PROFILE')}")

    if master_input is not None:
        master, available, raw_logs = build_master_and_available_from_workbook(
            master_input,
            master_sheet,
        )
    else:
        master, available, raw_logs = build_master_and_available_from_raw()
    report_progress(45, "Auditing case-swap inputs")
    case_swap_audit = load_case_swap_pool_audit()

    source_label = "master workbook" if master_input is not None else "FR TEIDs after first-admin filter"
    print(f"Master from {source_label}: {len(master):,}")
    print(f"Available source pool after first-admin filter and used drop: {len(available):,}")

    report_progress(55, "Preparing exact sex-balance inputs")
    if DISABLE_OPTIMIZATION:
        master_balanced = master.copy()
        available_balanced = available.copy()
        sex_balance_log = pd.DataFrame()
        print("Pass-through mode: exact-sex-balance adjustments bypassed")
    else:
        master_balanced, available_balanced, sex_balance_log = force_exact_sex_balance(master, available)
    print(f"Master after exact-sex-balance stage: {len(master_balanced):,}")

    run_log = pd.DataFrame(
        [
            {"stage": "config", "item": "project_dir", "value": str(PROJECT_DIR)},
            {"stage": "config", "item": "fr_teid_file", "value": str(FR_TEID_FILE)},
            {"stage": "config", "item": "master_input_override", "value": str(master_input or "")},
            {"stage": "config", "item": "master_input_sheet", "value": master_sheet if master_input else ""},
            {"stage": "config", "item": "scored_data_dir", "value": str(SCORED_DATA_DIR)},
            {"stage": "config", "item": "case_swap_dir", "value": str(CASE_SWAP_DIR)},
            {"stage": "config", "item": "target_workbook", "value": str(TARGET_WORKBOOK)},
            {"stage": "config", "item": "first_admin_filter", "value": "keep standardization English/Spanish testids by vform/agegrp; exclude retest/other-rater testids"},
            {"stage": "config", "item": "optimization_disabled", "value": DISABLE_OPTIMIZATION},
            {"stage": "config", "item": "exact_sex_balance", "value": "bypassed; supplied master preserved exactly" if DISABLE_OPTIMIZATION else "forced for even-N cells when feasible; odd-N cells logged"},
            *raw_logs,
        ]
    )

    report_progress(75, "Writing Part A workbook")
    write_workbook(
        output_file,
        master,
        master_balanced,
        available_balanced,
        run_log,
        case_swap_audit,
        sex_balance_log,
    )

    report_progress(92, "Publishing the latest Part A workbook")
    latest_file = OUTPUT_ROOT / "Optimization_Engine_Part_A_prep_output_latest.xlsx"
    shutil.copy2(output_file, latest_file)

    report_progress(100, "Part A is complete")
    print(f"Done: {output_file}")
    print(f"Latest copy: {latest_file}")


if __name__ == "__main__":
    main()
