"""Profile-driven demographic reports rendered from the approved SCOPE template."""
from __future__ import annotations

from copy import copy
from dataclasses import dataclass
import json
from pathlib import Path
import re
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .profiles import TestProfile


TEMPLATE_SHEETS = {1: "TEMPLATE_1WAY", 2: "TEMPLATE_2WAY", 3: "TEMPLATE_3WAY"}
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().casefold())


def _as_number(value: object) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _display_count(value: float | None) -> float | int | None:
    if value is None:
        return None
    rounded = round(value)
    return int(rounded) if abs(value - rounded) < 1e-6 else round(value, 2)


def _semantic_name(value: str) -> str:
    key = _key(value)
    aliases = {
        "age": "age_group", "ageband": "age_group", "agegroup": "age_group", "targetageband": "age_group",
        "gender": "sex", "eth": "ethnicity", "raceethnicity": "ethnicity",
        "edl": "education", "ped": "education", "sed": "education", "pedsed": "education",
        "reg": "region",
    }
    return aliases.get(key, key)


def _normalized_level(value: object, semantic: str) -> str:
    text = str(value).strip()
    if semantic == "age_group":
        numbers = [int(number) for number in re.findall(r"\d+", text)]
        if len(numbers) >= 6:
            return f"{numbers[0]}:{numbers[1]}-{numbers[3]}:{numbers[4]}"
        if len(numbers) >= 4:
            return f"{numbers[0]}:{numbers[1]}-{numbers[2]}:{numbers[3]}"
    if re.fullmatch(r"-?\d+(?:\.0+)?", text):
        return str(int(float(text)))
    aliases = {"female": "f", "woman": "f", "girl": "f", "male": "m", "man": "m", "boy": "m"}
    normalized = _key(text)
    return aliases.get(normalized, normalized)


@dataclass(frozen=True)
class CategoryLevel:
    value: object
    label: str


@dataclass(frozen=True)
class ReportingVariable:
    name: str
    canonical: str
    field: str
    display: str
    levels: tuple[CategoryLevel, ...]


@dataclass(frozen=True)
class TargetValue:
    pct: float | None = None
    n: float | None = None


def _style(cell):
    return copy(cell._style)


def _row_style(sheet, row: int, max_column: int) -> tuple[list[object], float | None]:
    return ([_style(sheet.cell(row, column)) for column in range(1, max_column + 1)], sheet.row_dimensions[row].height)


def _apply_row_style(sheet, row: int, styles: list[object], height: float | None, max_column: int | None = None) -> None:
    for column in range(1, (max_column or len(styles)) + 1):
        source_index = min(column, len(styles)) - 1
        sheet.cell(row, column)._style = copy(styles[source_index])
    if height is not None:
        sheet.row_dimensions[row].height = height


def _clear_merges(sheet) -> None:
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))


class DemographicReportRenderer:
    """Render one-, two-, and three-way tables without assessment constants."""

    def __init__(self, workbook, sample: pd.DataFrame, profile: TestProfile, targets: Iterable[dict] | None = None):
        self.workbook = workbook
        self.sample = sample.copy()
        self.profile = profile
        self.targets = list(targets or [])
        self._target_records: list[tuple[str, str, TargetValue]] = []
        self._criteria_target_records: list[tuple[dict[str, str], TargetValue]] = []
        self._target_index = self._build_target_index()

    def _definitions(self) -> dict[str, dict]:
        data = dict(self.profile.reporting_variables or {})
        if not data:
            data = dict((self.profile.data_specification or {}).get("reporting_variables") or {})
        if not data:
            data = dict((self.profile.reporting or {}).get("variables") or {})
        return data

    def _variable(self, requested: str) -> ReportingVariable:
        definitions = self._definitions()
        match_name, specification = "", None
        for name, value in definitions.items():
            candidate = value if isinstance(value, dict) else {"canonical": str(value)}
            aliases = {name, candidate.get("canonical", ""), candidate.get("field", ""), candidate.get("display", "")}
            if _key(requested) in {_key(alias) for alias in aliases if alias}:
                match_name, specification = str(name), candidate
                break
        if specification is None:
            mapping = self.profile.schema_mapping or {}
            canonical = next((name for name in mapping if _key(name) == _key(requested)), requested)
            field = str(mapping.get(canonical, canonical))
            display = requested
            observed = list(pd.unique(self.sample[field].dropna())) if field in self.sample else []
            levels = tuple(CategoryLevel(value, str(value)) for value in observed)
        else:
            canonical = str(specification.get("canonical") or match_name or requested)
            field = str(specification.get("field") or (self.profile.schema_mapping or {}).get(canonical) or canonical)
            display = str(specification.get("display") or match_name or requested)
            levels_list = []
            for item in specification.get("levels") or []:
                if isinstance(item, dict):
                    value = item.get("value", item.get("label")); label = str(item.get("label", value))
                else:
                    value, label = item, str(item)
                levels_list.append(CategoryLevel(value, label))
            if not levels_list and field in self.sample:
                levels_list = [CategoryLevel(value, str(value)) for value in pd.unique(self.sample[field].dropna())]
            levels = tuple(levels_list)
        if field not in self.sample.columns:
            raise ValueError(f"{self.profile.display_name} reporting variable '{requested}' maps to missing field '{field}'.")
        if not levels:
            raise ValueError(f"{self.profile.display_name} reporting variable '{requested}' has no ordered category levels.")
        return ReportingVariable(requested, canonical, field, display, levels)

    def _build_target_index(self) -> dict[tuple[str, str], TargetValue]:
        index: dict[tuple[str, str], TargetValue] = {}
        for row in self.targets:
            category = row.get("category", row.get("Category", "")); group = row.get("group", row.get("Group", ""))
            number = _as_number(row.get("target", row.get("Target")))
            if number is None: continue
            unit = str(row.get("unit", row.get("Unit", "%"))).strip().casefold()
            value = TargetValue(n=number) if unit in {"n", "count"} else TargetValue(pct=number / 100 if number > 1 else number)
            raw_criteria = row.get("criteria", row.get("Criteria", ""))
            if raw_criteria:
                try:
                    parsed = json.loads(raw_criteria) if isinstance(raw_criteria, str) else raw_criteria
                except (TypeError, ValueError, json.JSONDecodeError):
                    parsed = None
                if isinstance(parsed, dict) and parsed:
                    criteria = {_semantic_name(str(name)): str(raw) for name, raw in parsed.items()}
                    self._criteria_target_records.append((criteria, value))
            index[(_key(category), _key(group))] = value
            self._target_records.append((_key(category), _key(group), value))
        declared = (self.profile.reporting or {}).get("targets") or {}
        if isinstance(declared, dict):
            for category, values in declared.items():
                if not isinstance(values, dict): continue
                for group, raw in values.items():
                    if isinstance(raw, dict):
                        pct, n = _as_number(raw.get("pct", raw.get("target_pct"))), _as_number(raw.get("n", raw.get("target_n")))
                    else:
                        pct, n = _as_number(raw), None
                    if pct is not None and pct > 1: pct /= 100
                    value = TargetValue(pct=pct, n=n)
                    index.setdefault((_key(category), _key(group)), value)
                    self._target_records.append((_key(category), _key(group), value))
        return index

    def _criteria_target(self, variables: tuple[ReportingVariable, ...], levels: tuple[CategoryLevel, ...]) -> TargetValue | None:
        """Aggregate a fully crossed target workbook to the requested marginal/cross.

        Matrix percentages are conditional across the column variable, matching
        the approved template's 100% total for every row block.
        """
        if not self._criteria_target_records:
            return None
        filters = [(_semantic_name(variable.canonical or variable.name), level) for variable, level in zip(variables, levels)]

        def aggregate(wanted: list[tuple[str, CategoryLevel]]) -> TargetValue:
            counts, percentages = [], []
            for criteria, target in self._criteria_target_records:
                if all(
                    semantic in criteria
                    and _normalized_level(criteria[semantic], semantic) in {
                        _normalized_level(level.value, semantic), _normalized_level(level.label, semantic)
                    }
                    for semantic, level in wanted
                ):
                    if target.n is not None: counts.append(target.n)
                    if target.pct is not None: percentages.append(target.pct)
            return TargetValue(n=sum(counts) if counts else None, pct=sum(percentages) if percentages else None)

        numerator = aggregate(filters)
        denominator = aggregate(filters[:-1])
        if numerator.n is not None and denominator.n:
            return TargetValue(pct=numerator.n / denominator.n, n=numerator.n)
        if numerator.pct is not None and denominator.pct:
            return TargetValue(pct=numerator.pct / denominator.pct)
        return None

    @staticmethod
    def _combine_targets(values: Iterable[TargetValue]) -> TargetValue:
        values = list(values); percentages = [value.pct for value in values if value.pct is not None]; counts = [value.n for value in values if value.n is not None]
        return TargetValue(
            pct=sum(percentages) / len(percentages) if percentages else None,
            n=sum(counts) if counts else None,
        )

    def _target(self, variables: tuple[ReportingVariable, ...], levels: tuple[CategoryLevel, ...]) -> TargetValue:
        criteria_target = self._criteria_target(variables, levels)
        if criteria_target is not None:
            return criteria_target
        category_candidates = []
        for names in ((v.name for v in variables), (v.canonical for v in variables), (v.display for v in variables)):
            category_candidates.append(" × ".join(names))
        group_candidates = []
        for names in ((level.value for level in levels), (level.label for level in levels)):
            values = [str(value) for value in names]
            group_candidates.extend((" | ".join(values), " / ".join(values), " × ".join(values)))
        for category in category_candidates:
            for group in group_candidates:
                value = self._target_index.get((_key(category), _key(group)))
                if value is not None: return value
        # A marginal target for the column variable is a valid common target
        # when a profile does not declare intersection-specific percentages.
        last_variable, last_level = variables[-1], levels[-1]
        for category in (last_variable.name, last_variable.canonical, last_variable.display, last_variable.field, ""):
            for group in (last_level.value, last_level.label):
                value = self._target_index.get((_key(category), _key(group)))
                if value is not None: return value
        # Canonical target files may qualify a category with form/age-cell
        # context (for example, ``Form / age group / category``).  Match the
        # profile-owned category aliases and aggregate those applicable rows;
        # the rendering code remains independent of assessment/category names.
        category_keys = {_key(value) for value in (last_variable.name, last_variable.canonical, last_variable.display, last_variable.field) if value}
        level_keys = {_key(value) for value in (last_level.value, last_level.label) if str(value).strip()}
        qualified = [
            value for category, group, value in self._target_records
            if category in category_keys and any(group == level or group.endswith(level) for level in level_keys)
        ]
        if qualified:
            return self._combine_targets(qualified)
        # Count targets for a one-way variable are commonly expressed as
        # qualified sample cells.  Sum cells whose group text carries both the
        # variable alias and ordered level rather than relying on a fixed label.
        if len(variables) == 1:
            qualified_counts = [
                value for _category, group, value in self._target_records
                if value.n is not None
                and any(len(alias) > 1 and alias in group for alias in category_keys)
                and any(group.endswith(level) for level in level_keys)
            ]
            if qualified_counts:
                return self._combine_targets(qualified_counts)
        return TargetValue()

    def _mask(self, variable: ReportingVariable, level: CategoryLevel) -> pd.Series:
        semantic = _semantic_name(variable.canonical or variable.name)
        values = self.sample[variable.field].map(lambda value: _normalized_level(value, semantic))
        return values.eq(_normalized_level(level.value, semantic)) | values.eq(_normalized_level(level.label, semantic))

    def _sheet_name(self, displays: Iterable[str]) -> str:
        base = INVALID_SHEET_CHARS.sub("-", " × ".join(displays)).strip()[:31] or "Demographic Table"
        name, counter = base, 2
        while name in self.workbook.sheetnames:
            suffix = f" ({counter})"; name = base[:31-len(suffix)] + suffix; counter += 1
        return name

    def render_one_way(self, variable: str):
        definition = self._variable(variable)
        sheet = self.workbook.copy_worksheet(self.workbook[TEMPLATE_SHEETS[1]])
        sheet.title = self._sheet_name((definition.display,))
        data_style = _row_style(sheet, 4, 5); total_style = _row_style(sheet, 9, 5)
        if sheet.max_column > 5:
            _clear_merges(sheet)
            sheet.delete_cols(6, sheet.max_column - 5)
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        sheet.delete_rows(4, max(0, sheet.max_row - 3)); sheet["A1"] = definition.display
        total = len(self.sample); target_total_n = 0.0; target_pct_total = 0.0; has_target_n = has_target_pct = False
        for offset, level in enumerate(definition.levels):
            row = 4 + offset; _apply_row_style(sheet, row, *data_style)
            sample_n = int(self._mask(definition, level).sum()); sample_pct = sample_n / total if total else 0.0
            target = self._target((definition,), (level,)); target_pct = target.pct
            target_n = _display_count(target.n) if target.n is not None else (round(target_pct * total) if target_pct is not None else None)
            values = (level.label, target_n, target_pct, sample_n, sample_pct)
            for column, value in enumerate(values, 1): sheet.cell(row, column, value)
            for column in (3, 5): sheet.cell(row, column).number_format = "0.0%"
            if target_n is not None: target_total_n += target_n; has_target_n = True
            if target_pct is not None: target_pct_total += target_pct; has_target_pct = True
        total_row = 4 + len(definition.levels); _apply_row_style(sheet, total_row, *total_style)
        totals = ("Total", _display_count(target_total_n) if has_target_n else None, target_pct_total if has_target_pct else None, total, 1.0 if total else 0.0)
        for column, value in enumerate(totals, 1): sheet.cell(total_row, column, value)
        for column in (3, 5): sheet.cell(total_row, column).number_format = "0.0%"
        sheet.conditional_formatting._cf_rules.clear()
        sheet.freeze_panes = "A4"
        return sheet

    def _matrix_sheet(self, dimensions: int, variables: tuple[ReportingVariable, ...]):
        template = self.workbook[TEMPLATE_SHEETS[dimensions]]
        sheet = self.workbook.copy_worksheet(template); sheet.title = self._sheet_name(v.display for v in variables)
        max_column = template.max_column
        title_style = _row_style(template, 1, max_column)
        section_style = _row_style(template, 3, max_column) if dimensions == 3 else None
        header_row = 4 if dimensions == 3 else 3; block_start = 5 if dimensions == 3 else 4
        header_style = _row_style(template, header_row, max_column)
        block_styles = [_row_style(template, block_start + index, max_column) for index in range(3)]
        widths = {column: template.column_dimensions[get_column_letter(column)].width for column in range(1, max_column + 1)}
        _clear_merges(sheet); sheet.delete_rows(1, sheet.max_row); sheet.delete_cols(1, sheet.max_column)
        column_variable = variables[-1]; last_column = 2 + len(column_variable.levels) + 1
        _apply_row_style(sheet, 1, *title_style, max_column=last_column); sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
        sheet.cell(1, 1, " × ".join(variable.display for variable in variables))
        for column in range(1, last_column + 1):
            source_column = 1 if column == 1 else 2 if column == 2 else max_column if column == last_column else 3
            sheet.column_dimensions[get_column_letter(column)].width = widths.get(source_column, 13)
        sheet.conditional_formatting._cf_rules.clear()
        return sheet, section_style, header_style, block_styles, last_column

    def _write_matrix(self, sheet, row: int, row_variable: ReportingVariable, column_variable: ReportingVariable, row_filter: pd.Series, header_style, block_styles, last_column: int) -> int:
        _apply_row_style(sheet, row, *header_style, max_column=last_column)
        for column in range(3, last_column): sheet.cell(row, column)._style = copy(header_style[0][2])
        sheet.cell(row, last_column)._style = copy(header_style[0][-1])
        sheet.cell(row, 1, row_variable.display); sheet.cell(row, 2, "Metric")
        for offset, level in enumerate(column_variable.levels, 3): sheet.cell(row, offset, level.label)
        sheet.cell(row, last_column, "Total")
        row += 1
        for row_level in row_variable.levels:
            block_start = row
            row_mask = row_filter & self._mask(row_variable, row_level); denominator = int(row_mask.sum())
            sheet.merge_cells(start_row=block_start, start_column=1, end_row=block_start + 2, end_column=1)
            sheet.cell(block_start, 1, row_level.label)
            metrics = ("Target %", "Sample %", "Sample N")
            for metric_index, metric in enumerate(metrics):
                current_row = block_start + metric_index; styles, height = block_styles[metric_index]; _apply_row_style(sheet, current_row, styles, height, max_column=last_column)
                for column in range(3, last_column): sheet.cell(current_row, column)._style = copy(styles[2])
                sheet.cell(current_row, last_column)._style = copy(styles[-1]); sheet.cell(current_row, 2, metric)
            sample_total = 0; target_total = 0.0; has_target = False
            for column, column_level in enumerate(column_variable.levels, 3):
                count = int((row_mask & self._mask(column_variable, column_level)).sum()); sample_total += count
                sample_pct = count / denominator if denominator else 0.0
                target = self._target((row_variable, column_variable), (row_level, column_level)); target_pct = target.pct
                sheet.cell(block_start, column, target_pct); sheet.cell(block_start + 1, column, sample_pct); sheet.cell(block_start + 2, column, count)
                for current_row in (block_start, block_start + 1): sheet.cell(current_row, column).number_format = "0.0%"
                if target_pct is not None: target_total += target_pct; has_target = True
            sheet.cell(block_start, last_column, target_total if has_target else None); sheet.cell(block_start + 1, last_column, sample_total / denominator if denominator else 0.0); sheet.cell(block_start + 2, last_column, sample_total)
            sheet.cell(block_start, last_column).number_format = "0.0%"; sheet.cell(block_start + 1, last_column).number_format = "0.0%"
            row += 3
        return row

    def render_two_way(self, var_a: str, var_b: str):
        first, second = self._variable(var_a), self._variable(var_b)
        sheet, _section_style, header_style, block_styles, last_column = self._matrix_sheet(2, (first, second))
        self._write_matrix(sheet, 3, first, second, pd.Series(True, index=self.sample.index), header_style, block_styles, last_column)
        sheet.freeze_panes = "C4"
        return sheet

    def render_three_way(self, var_a: str, var_b: str, var_c: str):
        first, second, third = self._variable(var_a), self._variable(var_b), self._variable(var_c)
        sheet, section_style, header_style, block_styles, last_column = self._matrix_sheet(3, (first, second, third))
        row = 3
        for index, first_level in enumerate(first.levels):
            if index: row += 2
            _apply_row_style(sheet, row, *section_style, max_column=last_column); sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
            sheet.cell(row, 1, f"{first.display}: {first_level.label}"); row += 1
            row = self._write_matrix(sheet, row, second, third, self._mask(first, first_level), header_style, block_styles, last_column)
        sheet.freeze_panes = "C5"
        return sheet


def render_one_way(variable: str, *, workbook, sample: pd.DataFrame, profile: TestProfile, targets: Iterable[dict] | None = None):
    return DemographicReportRenderer(workbook, sample, profile, targets).render_one_way(variable)


def render_two_way(var_a: str, var_b: str, *, workbook, sample: pd.DataFrame, profile: TestProfile, targets: Iterable[dict] | None = None):
    return DemographicReportRenderer(workbook, sample, profile, targets).render_two_way(var_a, var_b)


def render_three_way(var_a: str, var_b: str, var_c: str, *, workbook, sample: pd.DataFrame, profile: TestProfile, targets: Iterable[dict] | None = None):
    return DemographicReportRenderer(workbook, sample, profile, targets).render_three_way(var_a, var_b, var_c)


def render_demographic_workbook(
    path: str | Path,
    sample: pd.DataFrame,
    profile: TestProfile,
    targets: Iterable[dict] | None = None,
    template: str | Path | None = None,
    selected_tables: dict[str, list[str]] | None = None,
) -> Path:
    """Create selected demographic worksheets from a clean approved template copy."""
    template_path = Path(template) if template else profile.resource_path("demographic_template")
    if not template_path or not template_path.is_file():
        raise FileNotFoundError("The approved SCOPE demographic reporting template is unavailable.")
    workbook = load_workbook(template_path)
    specification = selected_tables or profile.required_demographic_tables()
    used_sheets = {sheet for key, sheet in (("one_way", "TEMPLATE_1WAY"), ("two_way", "TEMPLATE_2WAY"), ("three_way", "TEMPLATE_3WAY")) if specification.get(key)}
    missing = [name for name in used_sheets if name not in workbook.sheetnames]
    if missing: raise ValueError("Demographic template is missing the sheet(s) your profile uses: " + ", ".join(missing))
    renderer = DemographicReportRenderer(workbook, sample, profile, targets)
    rendered = []
    for variable in specification["one_way"]: rendered.append(renderer.render_one_way(variable).title)
    for table in specification["two_way"]:
        parts = profile.parse_report_table(table)
        if len(parts) != 2: raise ValueError(f"Two-way reporting table must name two variables: {table}")
        rendered.append(renderer.render_two_way(*parts).title)
    for table in specification["three_way"]:
        parts = profile.parse_report_table(table)
        if len(parts) != 3: raise ValueError(f"Three-way reporting table must name three variables: {table}")
        rendered.append(renderer.render_three_way(*parts).title)
    if not rendered: raise ValueError(f"{profile.display_name} does not declare demographic reporting tables.")
    for name in ("README", *TEMPLATE_SHEETS.values()):
        if name in workbook.sheetnames: workbook.remove(workbook[name])
    workbook.active = 0
    output = Path(path); output.parent.mkdir(parents=True, exist_ok=True); workbook.save(output)
    return output
