"""Strict, non-GUI SCOPE runner for the validated BASC-4 NB legacy workflow.

This module intentionally orchestrates the bundled, known-working Part A and
Part B programs.  It does not duplicate or replace their optimization logic.
"""
from __future__ import annotations

from contextlib import contextmanager, redirect_stderr, redirect_stdout
from dataclasses import dataclass
from datetime import datetime
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import sys
from typing import Callable, Mapping

from openpyxl import load_workbook
import pandas as pd
import yaml

from .profiles import load_profile


ROOT = Path(__file__).resolve().parents[1]
LEGACY_DIR = ROOT / "scope" / "legacy"
CONFIG_FILE = ROOT / "scope" / "configs" / "basc4_adhd.yaml"
SUPPORTED_PROFILE = "BASC4_NB"
SUPPORTED_WORKFLOW = "sample_optimization"
STAGES = ("PROFILE", "INPUT PREP", "PART A", "TARGETS", "PART B/CP-SAT", "OUTPUT", "PASS")


@dataclass(frozen=True)
class ScopeRun:
    """Files and baseline metrics from a completed strict SCOPE run."""

    output_dir: Path
    outputs: dict[str, Path]
    legacy_outputs: dict[str, Path]
    metrics: dict[str, object]


class _ProgressStream:
    """Forward complete legacy log lines without swallowing exceptions."""

    def __init__(self, callback: Callable[[str], None] | None):
        self.callback = callback
        self.buffer = ""

    def write(self, value: str) -> int:
        value = str(value)
        self.buffer += value
        while "\n" in self.buffer:
            line, self.buffer = self.buffer.split("\n", 1)
            if line.strip() and self.callback:
                self._emit(line)
        return len(value)

    def flush(self) -> None:
        if self.buffer.strip() and self.callback:
            self._emit(self.buffer.strip())
        self.buffer = ""

    def _emit(self, line: str) -> None:
        with redirect_stdout(sys.__stdout__), redirect_stderr(sys.__stderr__):
            self.callback(line)


def _path(value: object, *, base: Path | None = None) -> Path | None:
    if value is None or not str(value).strip():
        return None
    result = Path(os.path.expandvars(str(value))).expanduser()
    if not result.is_absolute() and base:
        result = base / result
    return result.resolve()


def _project_inputs(value: Mapping[str, object] | str | Path) -> tuple[dict[str, object], Path | None]:
    if isinstance(value, Mapping):
        return dict(value), None
    source = Path(value).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Project-input JSON was not found: {source}")
    try:
        loaded = json.loads(source.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Project-input JSON is invalid: {source}: {exc}") from exc
    if not isinstance(loaded, dict):
        raise ValueError(f"Project-input JSON must contain one object: {source}")
    return loaded, source.parent


def _first(inputs: Mapping[str, object], *names: str) -> object | None:
    for name in names:
        if name in inputs and inputs[name] not in (None, ""):
            return inputs[name]
    return None


def _workbook_sheets(path: Path) -> list[str]:
    try:
        return load_workbook(path, read_only=True, data_only=True).sheetnames
    except Exception as exc:
        raise ValueError(f"Workbook could not be opened: {path}: {exc}") from exc


def _load_config() -> dict:
    if not CONFIG_FILE.is_file():
        raise FileNotFoundError(f"BASC legacy engine configuration is missing: {CONFIG_FILE}")
    raw = yaml.safe_load(CONFIG_FILE.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ValueError(f"BASC legacy engine configuration is empty or invalid: {CONFIG_FILE}")
    required = ("pool_rules", "target_cell_n", "osep_targets", "demo_target_sheets", "engine_params", "output_names")
    missing = [key for key in required if not raw.get(key)]
    if missing:
        raise ValueError(f"BASC legacy engine configuration is missing required keys: {', '.join(missing)} ({CONFIG_FILE})")
    return raw


def _validate_target_workbook(path: Path, config: dict) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"BASC target workbook is missing: {path}")
    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"BASC Part B requires an .xlsx or .xlsm target workbook, not {path.suffix}: {path}")
    sheets = set(_workbook_sheets(path))
    required = set(config["demo_target_sheets"].keys()) | {"osep", "byage"}
    missing = sorted(required - sheets)
    if missing:
        raise ValueError(f"BASC target workbook is missing required sheet(s) {missing}: {path}")
    parsed = _demographic_target_rows(path, config)
    if not parsed:
        raise ValueError(f"BASC target workbook contains no usable US population demographic targets: {path}")
    covered = {(row["vform"], int(row["agegrp"])) for row in parsed}
    expected = {(str(row["vform"]), int(row["agegrp"])) for row in config["target_cell_n"]}
    missing_cells = sorted(expected - covered)
    if missing_cells:
        formatted = ", ".join(f"{form}/agegrp {age}" for form, age in missing_cells)
        raise ValueError(f"BASC target workbook has no demographic targets for required cell(s): {formatted}. Source: {path}")


def _validate_inputs(inputs: dict[str, object], base: Path | None, profile_target: Path, config: dict) -> dict[str, object]:
    scored = _path(_first(inputs, "scored_data", "scored_data_dir", "scored_dir"), base=base)
    if not scored or not scored.is_dir():
        raise FileNotFoundError(f"BASC scored-data folder is required and was not found: {scored or '<not supplied>'}")
    allow_parquet = bool(_first(inputs, "allow_parquet_fixtures", "allow_parquet_fixture"))
    missing_scored = []
    for rule in config["pool_rules"].values():
        sas = scored / str(rule["file"])
        if sas.is_file():
            continue
        parquet = Path(str(sas) + ".parquet")
        if allow_parquet and parquet.is_file():
            continue
        suffix = " (or its .parquet test fixture)" if allow_parquet else ""
        missing_scored.append(str(sas) + suffix)
    if missing_scored:
        raise FileNotFoundError("BASC scored-data source(s) required by the legacy pool rules are missing:\n- " + "\n- ".join(missing_scored))

    current = _path(_first(inputs, "current_sample"), base=base)
    master = current or _path(_first(inputs, "master_dataset", "master_file"), base=base)
    initial = _path(_first(inputs, "initial_teid_source", "fr_teid_file", "all_cases_teid_source"), base=base)
    if master:
        if not master.is_file():
            raise FileNotFoundError(f"BASC master/current-sample workbook was not found: {master}")
        if master.suffix.casefold() not in {".xlsx", ".xlsm"}:
            raise ValueError(f"The known-working BASC Part A master/current-sample path requires an .xlsx or .xlsm workbook: {master}")
        master_sheet = str(_first(inputs, "master_sheet", "current_sample_sheet") or ("final_all" if current else config["default_paths"]["default_master_sheet"]))
        if master_sheet not in _workbook_sheets(master):
            raise ValueError(f"BASC master/current-sample workbook does not contain sheet '{master_sheet}': {master}")
    else:
        master_sheet = ""
        if not initial:
            raise ValueError(
                "BASC first-run optimization needs initial_teid_source (the FR/All Cases TEID workbook) "
                "or master_dataset. current_sample is optional and is not required."
            )
        if not initial.is_file():
            raise FileNotFoundError(f"BASC initial TEID source was not found: {initial}")
        sheets = set(_workbook_sheets(initial))
        required = set(config["default_paths"]["fr_admin_sheets"])
        missing = sorted(required - sheets)
        if missing:
            raise ValueError(f"BASC initial TEID source is missing required administration sheet(s) {missing}: {initial}")

    target = _path(_first(inputs, "target_specification", "target_workbook"), base=base) or profile_target.resolve()
    _validate_target_workbook(target, config)
    case_review = _path(_first(inputs, "case_review_data", "case_swap_dir"), base=base)
    if case_review and not case_review.is_dir():
        raise FileNotFoundError(f"BASC case-review folder was supplied but was not found: {case_review}")
    max_seconds = int(_first(inputs, "max_solver_seconds") or config["engine_params"].get("max_solver_seconds", 300))
    if max_seconds < 1:
        raise ValueError("max_solver_seconds must be at least 1.")
    return {
        "scored": scored, "master": master, "master_sheet": master_sheet,
        "initial": initial, "target": target, "case_review": case_review,
        "allow_parquet": allow_parquet, "max_solver_seconds": max_seconds,
    }


@contextmanager
def _legacy_environment(values: Mapping[str, object], allow_parquet: bool):
    old_env = {key: os.environ.get(key) for key in values}
    old_path = list(sys.path)
    old_read_sas = pd.read_sas
    old_modules = {name: sys.modules.get(name) for name in ("test_config", "project_profile")}

    if allow_parquet:
        def fixture_read_sas(path, *args, **kwargs):
            source = Path(path)
            fixture = Path(str(source) + ".parquet")
            if fixture.is_file():
                return pd.read_parquet(fixture)
            return old_read_sas(path, *args, **kwargs)
        pd.read_sas = fixture_read_sas

    sys.path.insert(0, str(LEGACY_DIR))
    os.environ.update({key: str(value) for key, value in values.items() if value is not None and str(value) != ""})
    for name in old_modules:
        sys.modules.pop(name, None)
    try:
        yield
    finally:
        pd.read_sas = old_read_sas
        sys.path[:] = old_path
        for key, value in old_env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        for name, module in old_modules.items():
            sys.modules.pop(name, None)
            if module is not None:
                sys.modules[name] = module


def _invoke_legacy(script_name: str, log: Callable[[str], None] | None) -> None:
    script = LEGACY_DIR / script_name
    if not script.is_file():
        raise FileNotFoundError(f"Bundled known-working SOE program is missing: {script}")
    name = f"scope_strict_{script.stem}_{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
    spec = importlib.util.spec_from_file_location(name, script)
    if not spec or not spec.loader:
        raise RuntimeError(f"Could not load bundled SOE program: {script}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    stream = _ProgressStream(log)
    try:
        spec.loader.exec_module(module)
        with redirect_stdout(stream), redirect_stderr(stream):
            module.main()
        stream.flush()
    finally:
        sys.modules.pop(name, None)


def _demographic_target_rows(path: Path, config: dict) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []
    for sheet_name in config["demo_target_sheets"]:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        headers = [sheet.cell(6, column).value for column in range(1, sheet.max_column + 1)]
        current_vform = None
        for row_number in range(7, sheet.max_row + 1):
            raw_vform = sheet.cell(row_number, 1).value
            if raw_vform:
                current_vform = str(raw_vform).strip().upper()
            agegrp = sheet.cell(row_number, 2).value
            if current_vform is None or sheet.cell(row_number, 4).value != "US population" or agegrp is None:
                continue
            for column, header in enumerate(headers, start=1):
                if column <= 4 or not header:
                    continue
                value = sheet.cell(row_number, column).value
                text = str(header).strip()
                if value is None or not text.startswith(("edl_", "eth_", "region_")):
                    continue
                if text.startswith("edl_"):
                    variable, level = "edl", text.split("_", 1)[1]
                elif text.startswith("eth_"):
                    suffix = text.split("_", 1)[1].lstrip("0123456789")
                    variable, level = "eth", "BL" if suffix == "AA" else suffix
                else:
                    variable, level = "region", text.split("_", 1)[1].lstrip("0123456789")
                rows.append({
                    "vform": current_vform, "agegrp": int(agegrp), "variable": variable,
                    "level": level, "value": float(value), "sheet": sheet_name,
                    "cell": sheet.cell(row_number, column).coordinate,
                })
    return rows


def _derived_repair_rows(part_a_file: Path, config: dict) -> list[dict[str, object]]:
    master = pd.read_excel(part_a_file, sheet_name="master_prepped")
    rows: list[dict[str, object]] = []
    target_flag = str(config["target_flag"])
    flags = [target_flag, "tmp_EBD", "tmp_ASD", "tmp_DD", "tmp_SLD", "tmp_SLI", "tmp_OTHER"]
    for (vform, agegrp), cell in master.groupby(["vform", "agegrp"], dropna=False):
        for flag in flags:
            if flag not in cell:
                continue
            rows.append({
                "target": f"repair preserve {flag} rate in {vform}/agegrp {int(agegrp)}",
                "category": "derived repair preference", "value": float(pd.to_numeric(cell[flag], errors="coerce").fillna(0).mean()),
                "source_file": str(part_a_file), "source_sheet_config_key": "master_prepped (derived current-cell rate)",
                "hard_soft": "soft", "tolerance": "", "priority": 1000 if flag == target_flag else 100,
            })
        for variable in ("eth", "edl", "region"):
            if variable in cell and not cell[variable].mode(dropna=True).empty:
                rows.append({
                    "target": f"repair preserve modal {variable} in {vform}/agegrp {int(agegrp)}",
                    "category": "derived repair preference", "value": str(cell[variable].mode(dropna=True).iloc[0]),
                    "source_file": str(part_a_file), "source_sheet_config_key": "master_prepped (derived current-cell mode)",
                    "hard_soft": "soft", "tolerance": "", "priority": 1,
                })
    return rows


def _write_target_lineage(path: Path, target_workbook: Path, part_a_file: Path, config: dict) -> None:
    rows: list[dict[str, object]] = []
    config_source = str(CONFIG_FILE)
    for index, cell in enumerate(config["target_cell_n"]):
        form, age, n = str(cell["vform"]), int(cell["agegrp"]), int(cell["n"])
        rows.extend((
            {"target": f"sample N {form}/agegrp {age}", "category": "sample size", "value": n,
             "source_file": config_source, "source_sheet_config_key": f"target_cell_n[{index}].n", "hard_soft": "hard", "tolerance": 0, "priority": "required"},
            {"target": f"sex balance {form}/agegrp {age}", "category": "sex balance", "value": f"F={n // 2}; M={n // 2}",
             "source_file": config_source, "source_sheet_config_key": f"target_cell_n[{index}] + female_value/male_value", "hard_soft": "hard", "tolerance": 0, "priority": "required"},
        ))
    rows.append({"target": "final total N", "category": "sample size", "value": int(config["final_total_n"]),
                 "source_file": config_source, "source_sheet_config_key": "final_total_n", "hard_soft": "hard", "tolerance": 0, "priority": "required"})
    rows.append({"target": "swap locks", "category": "structural constraint", "value": ", ".join(config["engine_params"]["swap_lock_cols"]),
                 "source_file": config_source, "source_sheet_config_key": "engine_params.swap_lock_cols", "hard_soft": "hard", "tolerance": 0, "priority": "required"})
    for flag, rule in config["osep_targets"].items():
        value = rule.get("target_pct")
        tolerance: object = 0
        if value is None:
            value = f"{rule['min_pct']}–{rule['max_pct']}%"
            tolerance = f"range {rule['min_pct']}–{rule['max_pct']}%"
        rows.append({"target": flag, "category": "clinical" if flag == "tmp_clin" else "OSEP", "value": value,
                     "source_file": config_source, "source_sheet_config_key": f"osep_targets.{flag}", "hard_soft": "soft",
                     "tolerance": tolerance, "priority": rule["weight"]})
    for target in _demographic_target_rows(target_workbook, config):
        rows.append({"target": f"{target['vform']}/agegrp {target['agegrp']} {target['variable']}={target['level']}",
                     "category": "demographic", "value": target["value"], "source_file": str(target_workbook),
                     "source_sheet_config_key": f"{target['sheet']}!{target['cell']}", "hard_soft": "soft", "tolerance": 0,
                     "priority": config["engine_params"].get("demo_weight", 1)})
    rows.extend(_derived_repair_rows(part_a_file, config))
    for row in rows:
        locator = str(row["source_sheet_config_key"])
        source = Path(str(row["source_file"]))
        if source.suffix.casefold() in {".yaml", ".yml"}:
            row["source_sheet"] = ""
            row["source_config_key"] = locator
        elif "!" in locator:
            row["source_sheet"], row["source_config_key"] = locator.split("!", 1)
        elif locator.startswith("master_prepped"):
            row["source_sheet"] = "master_prepped"
            row["source_config_key"] = locator.removeprefix("master_prepped").strip(" ()")
        else:
            row["source_sheet"] = ""
            row["source_config_key"] = locator
    pd.DataFrame(rows, columns=[
        "target", "category", "value", "source_file", "source_sheet", "source_config_key", "source_sheet_config_key",
        "hard_soft", "tolerance", "priority",
    ]).to_csv(path, index=False)


def _latest_run(root: Path, suffix: str) -> Path:
    runs = sorted((path for path in root.glob(f"run_*_{suffix}") if path.is_dir()), key=lambda path: path.stat().st_mtime)
    if not runs:
        raise RuntimeError(f"Legacy {suffix} completed without creating a run directory under {root}")
    return runs[-1]


def _hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _ids(path: Path, sheet: str) -> list[str]:
    frame = pd.read_excel(path, sheet_name=sheet, dtype={"teid": str})
    if "teid" not in frame:
        raise ValueError(f"Expected a teid column in {path} [{sheet}]")
    return sorted(frame["teid"].dropna().astype(str).str.strip().tolist())


def _metrics(part_a_file: Path, selected_file: Path, qa_file: Path) -> dict[str, object]:
    master_ids = _ids(part_a_file, "master_prepped")
    candidate_ids = sorted(set(master_ids) | set(_ids(part_a_file, "available_pool")))
    selected_ids = _ids(selected_file, "final_all")
    final_available = _ids(qa_file, "remaining_available_pool")
    sex = pd.read_excel(qa_file, sheet_name="sex_balance_qa")
    osep = pd.read_excel(qa_file, sheet_name="osep_target_qa")
    swaps = pd.read_excel(qa_file, sheet_name="part_b_osep_swaps")
    result: dict[str, object] = {
        "part_a_master_n": len(master_ids), "candidate_pool_n": len(candidate_ids),
        "part_a_available_n": len(_ids(part_a_file, "available_pool")),
        "final_n": len(selected_ids), "final_available_n": len(final_available),
        "selected_swap_n": len(swaps), "all_exact_n_constraints_met": bool(sex["exact_n_match"].fillna(False).all()),
        "all_exact_sex_constraints_met": bool(sex["exact_pct_match"].fillna(False).all()),
        "osep_status_counts": {str(key): int(value) for key, value in osep["status"].value_counts().to_dict().items()},
    }
    if not swaps.empty:
        for column in ("solver_status", "solver_objective", "phase1_solver_status", "phase1_osep_objective", "pre_osep_loss", "post_osep_loss"):
            if column in swaps:
                value = swaps[column].iloc[0]
                result[column] = value.item() if hasattr(value, "item") else value
    return result


def run_scope(
    profile: str = SUPPORTED_PROFILE,
    workflow: str = SUPPORTED_WORKFLOW,
    project_inputs: Mapping[str, object] | str | Path | None = None,
    output_dir: str | Path | None = None,
    *,
    progress: Callable[[str], None] | None = None,
    log: Callable[[str], None] | None = None,
) -> ScopeRun:
    """Run BASC4_NB Sample Optimization through legacy SOE Part A and Part B.

    ``project_inputs`` may be a mapping or a JSON file. For a first run,
    provide ``initial_teid_source`` and ``scored_data``; ``current_sample`` is
    intentionally optional. A ``master_dataset`` may replace the initial TEID
    source for an intentional workbook-based rerun.
    """
    emit = progress or (lambda _stage: None)
    emit("PROFILE")
    if str(profile).strip().upper().replace("-", "_") != SUPPORTED_PROFILE:
        raise ValueError(f"This strict runner supports profile='{SUPPORTED_PROFILE}' only; received {profile!r}.")
    if str(workflow).strip().casefold().replace(" ", "_") != SUPPORTED_WORKFLOW:
        raise ValueError(f"This strict runner supports workflow='{SUPPORTED_WORKFLOW}' only; received {workflow!r}.")
    installed = load_profile(SUPPORTED_PROFILE)
    if not installed.supports(SUPPORTED_WORKFLOW) or installed.source_config != "basc4_adhd":
        raise ValueError(f"Installed {SUPPORTED_PROFILE} profile is not connected to validated basc4_adhd Sample Optimization.")
    profile_target = installed.resource_path("approved_targets")
    if not profile_target:
        raise FileNotFoundError(f"Installed {SUPPORTED_PROFILE} profile does not declare approved_targets.")
    config = _load_config()

    emit("INPUT PREP")
    if project_inputs is None:
        raise ValueError("project_inputs is required. Supply initial_teid_source + scored_data for a first run, or master_dataset + scored_data.")
    raw_inputs, base = _project_inputs(project_inputs)
    resolved = _validate_inputs(raw_inputs, base, profile_target, config)
    destination = _path(output_dir)
    if not destination:
        raise ValueError("output_dir is required for the strict BASC run.")
    destination.mkdir(parents=True, exist_ok=True)
    legacy_root = destination / "_legacy"
    part_a_root, part_b_root = legacy_root / "part_a", legacy_root / "part_b"
    part_a_root.mkdir(parents=True, exist_ok=True)
    part_b_root.mkdir(parents=True, exist_ok=True)
    project_dir = _path(_first(raw_inputs, "project_dir"), base=base) or destination
    case_review = resolved["case_review"] or (project_dir / "case_review_not_supplied")
    env = {
        "OPT_ENGINE_TEST_CONFIG": "basc4_adhd", "OPT_ENGINE_PROJECT_DIR": project_dir,
        "OPT_ENGINE_SCORED_DATA_DIR": resolved["scored"], "OPT_ENGINE_CASE_SWAP_DIR": case_review,
        "OPT_ENGINE_TARGET_WORKBOOK": resolved["target"], "OPT_ENGINE_PART_A_OUTPUT_ROOT": part_a_root,
        "OPT_ENGINE_PART_B_OUTPUT_ROOT": part_b_root,
        "OPT_ENGINE_PART_A_FILE": part_a_root / "Optimization_Engine_Part_A_prep_output_latest.xlsx",
        "OPT_ENGINE_MAX_SOLVER_SECONDS": resolved["max_solver_seconds"],
        "OPT_ENGINE_INCLUDE_CENSUS": "1", "OPT_ENGINE_INCLUDE_OSEP": "1", "OPT_ENGINE_INCLUDE_CLINICAL": "1",
        "OPT_ENGINE_CLINICAL_SEEDING": "0",
    }
    if resolved["master"]:
        env["OPT_ENGINE_MASTER_INPUT"] = resolved["master"]
        env["OPT_ENGINE_MASTER_SHEET"] = resolved["master_sheet"]
    else:
        env["OPT_ENGINE_FR_TEID_FILE"] = resolved["initial"]

    with _legacy_environment(env, bool(resolved["allow_parquet"])):
        emit("PART A")
        _invoke_legacy("part_a.py", log)
        part_a_file = Path(env["OPT_ENGINE_PART_A_FILE"])
        if not part_a_file.is_file():
            raise RuntimeError(f"Legacy Part A did not create its required prepared workbook: {part_a_file}")
        emit("TARGETS")
        lineage = destination / "target_lineage.csv"
        _write_target_lineage(lineage, Path(resolved["target"]), part_a_file, config)
        emit("PART B/CP-SAT")
        _invoke_legacy("part_b.py", log)

    part_b_run = _latest_run(part_b_root, "part_b")
    names = config["output_names"]
    raw_selected = part_b_run / names["updated_teid"]
    raw_report = part_b_run / names["demo_osep"]
    raw_swap = part_b_run / names["swap_table"]
    raw_qa = part_b_run / "Optimization_Engine_Part_B_QA.xlsx"
    required_raw = (raw_selected, raw_report, raw_swap, raw_qa)
    missing_raw = [str(path) for path in required_raw if not path.is_file()]
    if missing_raw:
        raise RuntimeError("Legacy Part B completed without required output(s): " + "; ".join(missing_raw))

    emit("OUTPUT")
    selected = destination / "Selected_TEIDs.xlsx"
    report = destination / names["demo_osep"]
    constraint = destination / "Constraint_Summary.xlsx"
    available = destination / "Available_Pool.xlsx"
    shutil.copy2(raw_selected, selected)
    shutil.copy2(raw_report, report)
    shutil.copy2(raw_qa, constraint)
    available_frame = pd.read_excel(raw_qa, sheet_name="remaining_available_pool")
    with pd.ExcelWriter(available, engine="openpyxl") as writer:
        available_frame.to_excel(writer, sheet_name="available_pool", index=False)
    metrics = _metrics(part_a_file, raw_selected, raw_qa)
    if metrics["final_n"] != int(config["final_total_n"]):
        raise RuntimeError(f"Legacy BASC run returned final N={metrics['final_n']}; required N={config['final_total_n']}.")
    if not metrics["all_exact_n_constraints_met"] or not metrics["all_exact_sex_constraints_met"]:
        raise RuntimeError("Legacy BASC run returned a failed exact-N or exact-sex constraint.")

    outputs = {
        "selected_teids": selected, "available_pool": available, "stratification": report,
        "constraint_summary": constraint, "target_lineage": lineage,
    }
    legacy_outputs = {
        "part_a": part_a_file, "selected_teids": raw_selected, "available_pool_and_qa": raw_qa,
        "stratification": raw_report, "swap_table": raw_swap,
    }
    manifest = destination / "run_manifest.json"
    payload = {
        "status": "PASS", "profile": SUPPORTED_PROFILE, "profile_version": installed.version,
        "workflow": SUPPORTED_WORKFLOW, "completed": datetime.now().isoformat(timespec="seconds"),
        "implementation": {"part_a": str(LEGACY_DIR / "part_a.py"), "part_b": str(LEGACY_DIR / "part_b.py"), "engine_config": str(CONFIG_FILE)},
        "inputs": {key: str(value) for key, value in resolved.items() if isinstance(value, Path)},
        "outputs": {key: str(value.resolve()) for key, value in outputs.items()},
        "output_sha256": {key: _hash(value) for key, value in outputs.items()},
        "legacy_outputs": {key: str(value.resolve()) for key, value in legacy_outputs.items()},
        "metrics": metrics,
    }
    manifest.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    outputs["run_manifest"] = manifest
    emit("PASS")
    return ScopeRun(destination, outputs, legacy_outputs, metrics)
