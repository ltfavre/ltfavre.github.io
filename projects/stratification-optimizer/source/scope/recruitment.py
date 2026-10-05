"""Count-level CP-SAT planning for Field Research recruitment targets."""
from __future__ import annotations

from itertools import product
from pathlib import Path
import re

import pandas as pd
from ortools.sat.python import cp_model
from openpyxl import load_workbook


TARGET_SHEET = "Recruitment Targets"


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def _column(frame: pd.DataFrame, *names: str) -> str | None:
    indexed = {_key(name): name for name in frame.columns}
    return next((indexed[_key(name)] for name in names if _key(name) in indexed), None)


def read_targets(path: str | Path | pd.DataFrame, total_n: int | None) -> tuple[pd.DataFrame, int]:
    """Read the documented SCOPE Recruitment Targets table and normalize target Ns."""
    if isinstance(path, pd.DataFrame):
        raw = path.copy()
    else:
        source = Path(path)
        if source.suffix.casefold() == ".csv":
            raw = pd.read_csv(source)
        else:
            with pd.ExcelFile(source) as book:
                sheet = TARGET_SHEET if TARGET_SHEET in book.sheet_names else book.sheet_names[0]
                raw = pd.read_excel(book, sheet_name=sheet)
    dimension, category = _column(raw, "dimension"), _column(raw, "category")
    pct, target_n = _column(raw, "target_pct", "target percent", "target %"), _column(raw, "target_n", "target n", "n")
    tolerance = _column(raw, "tolerance_n", "tolerance", "tolerance n")
    if not dimension or not category or (not pct and not target_n):
        raise ValueError("Recruitment plan needs Dimension, Category, and Target_N or Target_Pct columns. Use the SCOPE template as a starting point.")
    frame = raw[[dimension, category] + [name for name in (pct, target_n, tolerance) if name]].copy()
    frame.columns = ["dimension", "category"] + (["target_pct"] if pct else []) + (["target_n"] if target_n else []) + (["tolerance_n"] if tolerance else [])
    frame["dimension"] = frame["dimension"].astype(str).str.strip()
    frame["category"] = frame["category"].astype(str).str.strip()
    frame = frame.loc[(frame.dimension != "") & (frame.category != "")].copy()
    if frame.empty or frame.duplicated(["dimension", "category"]).any():
        raise ValueError("Recruitment targets must contain unique, nonblank Dimension + Category rows.")
    if "target_n" in frame:
        frame["target_n"] = pd.to_numeric(frame.target_n, errors="coerce")
    if "target_pct" in frame:
        frame["target_pct"] = pd.to_numeric(frame.target_pct, errors="coerce")
        frame.loc[frame.target_pct > 1, "target_pct"] /= 100
    inferred = None
    if total_n is None and "target_n" in frame and frame.target_n.notna().any():
        totals = frame.groupby("dimension").target_n.sum().dropna()
        if not totals.empty: inferred = int(totals.iloc[0])
    total_n = total_n or inferred
    if not total_n or total_n < 1:
        raise ValueError("Enter Requested final N, or provide Target_N values that total to N for at least one dimension.")
    frame["target_n"] = frame.get("target_n", pd.Series(index=frame.index, dtype=float)).fillna(
        (frame.get("target_pct", pd.Series(index=frame.index, dtype=float)) * total_n).round()
    )
    if frame.target_n.isna().any() or (frame.target_n < 0).any():
        raise ValueError("Each recruitment target must have a non-negative Target_N or Target_Pct.")
    frame["target_n"] = frame.target_n.astype(int)
    tolerance_values = frame["tolerance_n"] if "tolerance_n" in frame else pd.Series(0, index=frame.index)
    frame["tolerance_n"] = pd.to_numeric(tolerance_values, errors="coerce").fillna(0).astype(int)
    return frame[["dimension", "category", "target_n", "tolerance_n"]], int(total_n)


def validate_plan(path: str | Path, total_n: int | None) -> list[tuple[str, bool, str]]:
    targets, total_n = read_targets(path, total_n)
    checks: list[tuple[str, bool, str]] = [("Recruitment plan", True, f"{len(targets)} target rows; total N = {total_n}")]
    for dimension, rows in targets.groupby("dimension"):
        target_sum = int(rows.target_n.sum())
        checks.append((f"{dimension} targets", True, f"{target_sum} planned across {len(rows)} categories (overall N {total_n})"))
    return checks


def validate_profile_age_targets(path: str | Path | pd.DataFrame, total_n: int | None, profile, study_type: str) -> None:
    """Require an exact, ordered Age margin for approved native reports."""
    targets, resolved_n = read_targets(path, total_n)
    expected = profile_age_bands(profile, study_type)
    if not expected:
        return
    age_dimensions = targets.loc[targets.dimension.map(_key).isin({"age", "ageband", "agegroup", "agegrp"})]
    if age_dimensions.empty:
        raise ValueError("The recruitment target amendment must include the approved Age dimension: " + ", ".join(expected))
    actual = [_template_age_label(value) for value in age_dimensions.category]
    if actual != expected:
        raise ValueError("Recruitment age bands must exactly match the approved profile order. Expected: " + ", ".join(expected) + "; found: " + ", ".join(actual))
    if int(age_dimensions.target_n.sum()) != int(resolved_n):
        raise ValueError(f"Recruitment age-band Target Ns must reconcile to total N {resolved_n}.")


def _actuals(path: str | Path | None, targets: pd.DataFrame) -> pd.DataFrame:
    if not path:
        return pd.DataFrame(columns=["dimension", "category", "actual_n"])
    source = Path(path)
    if source.suffix.casefold() == ".csv":
        raw = pd.read_csv(source)
    else:
        with pd.ExcelFile(source) as book:
            raw = pd.read_excel(book, sheet_name=book.sheet_names[0])
    dim, cat, actual = _column(raw, "dimension"), _column(raw, "category"), _column(raw, "actual_n", "usable_n", "n")
    if dim and cat and actual:
        data = raw[[dim, cat, actual]].copy(); data.columns = ["dimension", "category", "actual_n"]
        data.actual_n = pd.to_numeric(data.actual_n, errors="coerce").fillna(0)
        return data.groupby(["dimension", "category"], as_index=False).actual_n.sum()
    results = []
    by_key = {_key(column): column for column in raw.columns}
    for dimension, rows in targets.groupby("dimension"):
        source = by_key.get(_key(dimension))
        if not source: continue
        counts = raw[source].astype(str).str.strip().value_counts()
        for category in rows.category:
            results.append({"dimension": dimension, "category": category, "actual_n": int(counts.get(category, 0))})
    return pd.DataFrame(results, columns=["dimension", "category", "actual_n"])


def build_recruitment_plan(plan_file: str | Path | pd.DataFrame, total_n: int | None, actuals_file: str | Path | None = None, max_seconds: int = 300) -> dict[str, pd.DataFrame]:
    """Generate integer intersection targets that reconcile all marginal targets."""
    targets, total_n = read_targets(plan_file, total_n)
    groups = [(dimension, rows.category.tolist()) for dimension, rows in targets.groupby("dimension", sort=False)]
    dimensions = [name for name, _ in groups]
    combinations = list(product(*(categories for _, categories in groups)))
    if len(combinations) > 50000:
        raise ValueError(f"This plan creates {len(combinations):,} intersections. Consolidate categories or separate the planning problem.")
    model = cp_model.CpModel()
    cells = [model.NewIntVar(0, total_n, f"cell_{index}") for index in range(len(combinations))]
    model.Add(sum(cells) == total_n)
    errors = []
    for row in targets.itertuples(index=False):
        position = dimensions.index(row.dimension)
        margin = sum(cells[index] for index, combo in enumerate(combinations) if combo[position] == row.category)
        error = model.NewIntVar(0, total_n, f"error_{len(errors)}")
        model.AddAbsEquality(error, margin - int(row.target_n))
        if int(row.tolerance_n) > 0:
            model.Add(error <= int(row.tolerance_n))
        errors.append(error)
    model.Minimize(sum(errors))
    solver = cp_model.CpSolver(); solver.parameters.max_time_in_seconds = max(1, int(max_seconds)); solver.parameters.num_search_workers = 8
    status = solver.Solve(model)
    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        raise ValueError("The recruitment targets are infeasible within the stated tolerances. Increase a tolerance or revise the target table.")
    cell_rows = [{**dict(zip(dimensions, combo)), "target_n": solver.Value(cells[index])} for index, combo in enumerate(combinations) if solver.Value(cells[index])]
    cells_frame = pd.DataFrame(cell_rows, columns=dimensions + ["target_n"])
    margin_rows = []
    for row in targets.itertuples(index=False):
        actual_target = int(cells_frame.loc[cells_frame[row.dimension].eq(row.category), "target_n"].sum()) if not cells_frame.empty else 0
        margin_rows.append({"dimension": row.dimension, "category": row.category, "requested_n": int(row.target_n), "planned_n": actual_target, "difference": actual_target - int(row.target_n), "tolerance_n": int(row.tolerance_n)})
    margins = pd.DataFrame(margin_rows)
    actuals = _actuals(actuals_file, targets)
    needs = margins.merge(actuals, on=["dimension", "category"], how="left")
    needs["actual_n"] = pd.to_numeric(needs["actual_n"], errors="coerce").fillna(0).astype(int)
    needs["remaining_n"] = (needs.planned_n - needs.actual_n).clip(lower=0)
    needs["target_pct"] = needs["planned_n"] / total_n
    actual_totals = needs.groupby("dimension")["actual_n"].transform("sum")
    needs["current_pct"] = needs["actual_n"].div(actual_totals.where(actual_totals.ne(0)))
    needs["diff"] = needs["current_pct"] - needs["target_pct"]
    field_research = needs.rename(columns={
        "dimension": "Demographic / Group", "category": "Category",
        "planned_n": "Target N", "actual_n": "Current / Recruited N",
        "remaining_n": "Remaining N", "target_pct": "Target %",
        "current_pct": "Current %", "diff": "Diff",
    })[["Demographic / Group", "Category", "Target N", "Current / Recruited N", "Remaining N", "Target %", "Current %", "Diff"]]
    age_plan = needs.loc[needs.dimension.str.casefold().isin(("age", "age band", "age group", "agegrp")), ["category", "planned_n", "actual_n", "remaining_n"]].rename(columns={"category": "age_band"})
    canonical_targets = margins.rename(columns={"dimension": "Category", "category": "Group", "planned_n": "Target", "tolerance_n": "Tolerance_N"})
    canonical_targets["Rule"] = "Target"; canonical_targets["Unit"] = "N"; canonical_targets["Priority"] = "Required"; canonical_targets["Required"] = True
    canonical_targets = canonical_targets[["Category", "Group", "Rule", "Target", "Unit", "Tolerance_N", "Priority", "Required"]]
    lowered = field_research["Demographic / Group"].astype(str).str.casefold()
    demographics = field_research.loc[~lowered.str.contains("osep|clinical", regex=True)].copy()
    special = field_research.loc[lowered.str.contains("osep|clinical", regex=True)].copy()
    outputs = {
        "Field Research Report": field_research,
        "Recruitment Targets": field_research[["Demographic / Group", "Category", "Target N", "Current / Recruited N", "Remaining N"]],
        "Demographic Targets": demographics,
        "Recruitment Plan": age_plan,
        "Recruitment Cell Targets": cells_frame,
        "SCOPE Targets": canonical_targets,
        "Marginal Targets": margins,
        "Remaining Recruitment Need": needs,
    }
    if not special.empty:
        outputs["OSEP Clinical Targets"] = special
    return outputs


def _level_values(profile, variable: str) -> list[str]:
    raw = (profile.reporting_variables or {}).get(variable) or {}
    return [str(item.get("value", item.get("label", ""))) if isinstance(item, dict) else str(item) for item in raw.get("levels", [])]


def _template_age_label(value: object) -> str:
    """Normalize Aurora and display age ranges to an inclusive year:month label."""
    text = str(value or "").strip().replace("–", "-")
    match = re.match(r"^(\d{1,2})[.:](\d{1,2})(?:[.:]\d{1,2})?\s*-\s*(\d{1,2})[.:](\d{1,2})", text)
    if not match:
        return text
    y1, m1, y2, m2 = map(int, match.groups())
    return f"{y1}:{m1}-{y2}:{m2}"


def profile_age_bands(profile, study_type: str = "") -> list[str]:
    specification = profile.recruitment_spec(study_type)
    declared = specification.get("age_bands") or profile.age_structure
    return [_template_age_label(value.get("label", "") if isinstance(value, dict) else value) for value in declared]


def targets_from_approved_template(template: str | Path, profile, study_type: str, total_n: int | None) -> tuple[pd.DataFrame, int]:
    """Read approved target proportions from the template's first native Aurora block.

    This keeps targets grounded in the approved workbook.  If project N differs
    from the workbook example, the approved proportions are converted to integer
    margins that reconcile exactly to project N.
    """
    path = Path(template)
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "aurora_data" not in workbook.sheetnames:
        raise ValueError(f"Approved recruitment template has no aurora_data sheet: {path.name}")
    sheet = workbook["aurora_data"]
    data = list(sheet.iter_rows(values_only=True))
    header_index = next((index for index, row in enumerate(data[:80]) if len(row) >= 3 and _key(row[0]) == "ageband" and _key(row[2]) == "report"), None)
    if header_index is None:
        raise ValueError(f"Approved recruitment template has no recognizable native target table: {path.name}")
    headers = list(data[header_index])
    total_col = next((index + 1 for index, value in enumerate(headers) if _key(value) == "total"), None)
    if not total_col:
        raise ValueError(f"Approved recruitment template has no Total column in its native target table: {path.name}")
    approved_ages = profile_age_bands(profile, study_type)
    target_rows: list[tuple[str, int, list[object]]] = []
    seen_ages: set[str] = set()
    for row_number, values in enumerate(data[header_index + 1:], start=header_index + 2):
        if len(values) < 3 or _key(values[2]) != "targetn":
            continue
        age = _template_age_label(values[0])
        if age in approved_ages and age not in seen_ages:
            target_rows.append((age, row_number, list(values[:total_col])))
            seen_ages.add(age)
            if len(seen_ages) == len(approved_ages):
                break
    if not target_rows:
        raise ValueError(f"No approved {study_type or 'default'} age bands were found in {path.name}.")
    workbook_n = sum(float(values[total_col - 1] or 0) for _age, _row, values in target_rows)
    final_n = int(total_n or round(workbook_n))
    if final_n < 1 or workbook_n <= 0:
        raise ValueError("Approved template targets do not contain a usable total N.")

    def reconcile(weights: list[float]) -> list[int]:
        total = sum(weights)
        exact = [final_n * value / total for value in weights]
        result = [int(value) for value in exact]
        for index in sorted(range(len(result)), key=lambda i: exact[i] - result[i], reverse=True)[:final_n - sum(result)]:
            result[index] += 1
        return result

    rows: list[dict] = []
    age_counts = reconcile([float(values[total_col - 1] or 0) for _age, _row, values in target_rows])
    rows.extend({"Dimension": "Age", "Category": age, "Target_N": count, "Tolerance_N": 0} for (age, _row, _values), count in zip(target_rows, age_counts))

    # Read the first complete native column segment for each profile variable.
    for variable in profile.reporting_variables or {}:
        if _key(variable) == "age":
            continue
        levels = _level_values(profile, variable)
        if not levels:
            continue
        level_keys = [_key(value) for value in levels]
        candidate_columns: list[int] = []
        for column in range(4, total_col):
            header = _key(headers[column - 1])
            if any(header == level or header.startswith(level) or header.endswith(level) for level in level_keys):
                candidate_columns.append(column)
        if not candidate_columns:
            continue
        weights = []
        for level_key in level_keys:
            matched = [column for column in candidate_columns if (_key(headers[column - 1]) == level_key or _key(headers[column - 1]).startswith(level_key) or _key(headers[column - 1]).endswith(level_key))]
            weights.append(sum(float(values[column - 1] or 0) for _age, _row, values in target_rows for column in matched))
        if not all(value >= 0 for value in weights) or sum(weights) <= 0:
            continue
        counts = reconcile(weights)
        rows.extend({"Dimension": variable, "Category": level, "Target_N": count, "Tolerance_N": 1} for level, count in zip(levels, counts))
    frame = pd.DataFrame(rows)
    for dimension, part in frame.groupby("Dimension", sort=False):
        if int(part.Target_N.sum()) != final_n:
            raise ValueError(f"Approved {dimension} targets do not reconcile to total N {final_n}.")
    return frame, final_n


def validate_actual_age_mapping(path: str | Path | None, profile, study_type: str) -> None:
    """Block raw progress data containing ages outside the exact approved bands."""
    if not path:
        return
    source = Path(path)
    raw = pd.read_csv(source) if source.suffix.casefold() == ".csv" else pd.read_excel(source)
    age_group = _column(raw, "age_band", "age group", "agegrp")
    age_years = _column(raw, "age", "agey", "age_years")
    if not age_group and not age_years:
        return  # A documented long count table need not contain case ages.
    approved = profile_age_bands(profile, study_type)
    intervals = []
    for label in approved:
        match = re.fullmatch(r"(\d+):(\d+)-(\d+):(\d+)", label)
        if match:
            y1, m1, y2, m2 = map(int, match.groups()); intervals.append((label, y1 * 12 + m1, y2 * 12 + m2))
    unmapped: list[str] = []
    if age_group:
        for value in raw[age_group].dropna().astype(str).unique():
            if _template_age_label(value) not in approved:
                unmapped.append(value)
    else:
        for value in pd.to_numeric(raw[age_years], errors="coerce").dropna():
            months = int(round(float(value) * 12))
            matches = [label for label, low, high in intervals if low <= months <= high]
            if len(matches) != 1:
                unmapped.append(str(value))
    if unmapped:
        raise ValueError("Recruitment progress contains ages that do not map to exactly one approved age band: " + ", ".join(unmapped[:12]))


def create_template(path: str | Path) -> Path:
    """Create a short, editable starting template for a Field Research plan."""
    output = Path(path); output.parent.mkdir(parents=True, exist_ok=True)
    example = pd.DataFrame([
        ("Age Band", "6", 25, "", 1), ("Age Band", "7", 25, "", 1), ("Age Band", "8", 25, "", 1), ("Age Band", "9", 25, "", 1),
        ("Sex", "Female", 50, "", 1), ("Sex", "Male", 50, "", 1),
    ], columns=["Dimension", "Category", "Target_Pct", "Target_N", "Tolerance_N"])
    instructions = pd.DataFrame({"How to use this file": [
        "One row is one marginal recruitment target. Use Target_Pct (0-100) or Target_N; do not need both.",
        "Use the same total N within every complete dimension. Tolerance_N allows a target to vary when exact integer reconciliation is impossible.",
        "Add dimensions such as Ethnicity, Education, Region, OSEP, Clinical Group, Language, Form, or Paper/Digital.",
        "SCOPE jointly optimizes all marginal targets and outputs integer intersection cells and remaining recruitment need.",
    ]})
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        example.to_excel(writer, sheet_name=TARGET_SHEET, index=False)
        instructions.to_excel(writer, sheet_name="Instructions", index=False)
    return output
