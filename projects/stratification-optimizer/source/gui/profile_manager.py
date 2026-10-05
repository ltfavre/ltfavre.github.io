"""Information-first authoring and management for SCOPE assessment profiles.

The editor is organized as a guided workflow — Sources → Sample Rules →
Demographics → Targets → Clinical/OSEP → Optimization Rules → Output Tables →
Validate — so international adaptations (for example WAIS-5 Canada) can be
configured from a dataset, a census/target workbook, a reference profile, and
direct RD decisions without inventing placeholder specifications.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk

import yaml

ROOT = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scope.profiles import (
    DIMENSION_DISPLAY, ORIGIN_LABELS, PROFILE_KINDS, SAMPLE_FILTER_FIELDS,
    TARGET_CELLS_AS, TARGET_DIMENSION_ALIASES, TARGET_SOURCE_FORMATS,
    USER_PROFILES_DIR, WORKFLOW_LABELS, derive_marginal_target_rows,
    dimension_semantic, discover_profile_catalog, format_sample_filter,
    install_profile_package, normalize_optimization_rule, normalize_sample_filter,
    profile_directories, profile_validation_summary, validate_profile_data,
    workflow_readiness,
)

NAVY, BACKGROUND, SURFACE = "#123B63", "#F6F3EC", "#FFFFFF"
TEXT, MUTED, BORDER = "#233240", "#6E7B86", "#D8D5CD"
SUCCESS, WARNING, ERROR = "#3B8D68", "#B77A16", "#B84444"
RESOURCE_ROLES = (
    ("research_specification", "Research Plan / approved specification (optional)"),
    ("approved_targets", "Default CP-SAT target workbook"),
    ("census_targets", "Census / demographic target source"),
    ("osep_specification", "OSEP / clinical definition source"),
    ("recruitment_targets", "Recruitment target specification"),
    ("recruitment_template", "Default recruitment template"),
    ("demographic_template", "Demographic reporting template"),
    ("osep_template", "OSEP reporting template"),
    ("clinical_template", "Clinical reporting template"),
    ("targeted_swap_template", "Targeted case/group table template"),
    ("profile_validation", "Profile validation specification"),
)
SOURCE_TYPES = (
    "Assessment Dataset", "Census / Target Workbook", "Demographic Targets",
    "Existing Profile", "Research Plan", "Sampling Plan / Sampling Specification",
    "Data Analysis Plan", "Scoring Specification", "Technical Specification",
    "OSEP / Clinical Definitions", "Data Dictionary", "Case / Variable Mapping",
    "Recruitment Specification", "Other authoritative project document",
)
CAPABILITY_LABELS = (
    ("recruitment_planning", "Recruitment Planning"),
    ("targeted_swap", "Targeted Case Swap"),
    ("targeted_adjustment", "OSEP / Clinical Adjustments"),
    ("sample_optimization", "Sample Optimization (CP-SAT)"),
)
WORKFLOW_INPUT_DEFAULTS = {
    "recruitment_planning": [
        {"role": "recruitment_targets", "group": "TARGETS", "label": "Recruitment target source", "status": "required", "kind": "file", "accepted_types": [".xlsx", ".csv", ".json"], "purpose": "Authoritative age-band and demographic targets; a Research Plan is not required.", "adapter": "scope_recruitment_targets"},
        {"role": "recruitment_progress", "group": "PROGRESS", "label": "Current recruitment / progress workbook", "status": "optional", "kind": "file", "accepted_types": [".xlsx", ".csv"], "purpose": "Current recruited or usable counts.", "adapter": "scope_recruitment_progress"},
    ],
    "sample_optimization": [
        {"role": "master_dataset", "group": "CASE DATA", "label": "Candidate case data / TEID list", "status": "required", "kind": "file", "accepted_types": [".xlsx", ".sas7bdat", ".csv"], "purpose": "Complete candidate universe. It must include TEID plus every demographic field used by the target workbook.", "adapter": "canonical_table"},
        {"role": "current_sample", "group": "SAMPLE", "label": "Current selected sample", "status": "optional", "kind": "file", "accepted_types": [".xlsx"], "purpose": "Optional starting membership for re-optimization.", "adapter": "canonical_sample"},
        {"role": "target_specification", "group": "TARGETS", "label": "Demographic target workbook", "status": "required", "kind": "file", "accepted_types": [".xlsx", ".csv", ".json"], "purpose": "Authoritative target stratification. A Research Plan is not required.", "adapter": "canonical_table"},
    ],
    "targeted_swap": [
        {"role": "swap_sample", "group": "MEMBERSHIP", "label": "Current selected TEID list", "status": "required", "kind": "file", "accepted_types": [".xlsx", ".csv"], "purpose": "Defines which TEIDs are currently selected. A combined SCOPE membership workbook is also supported.", "adapter": "scope_membership_workbook"},
        {"role": "swap_available_pool", "group": "MEMBERSHIP", "label": "Available / candidate TEID list", "status": "required", "kind": "file", "accepted_types": [".xlsx", ".csv"], "purpose": "Defines possible additions. TEIDs that overlap the selected list remain selected and are automatically excluded from the add pool.", "adapter": "canonical_pool"},
        {"role": "swap_sample_data", "group": "CASE DATA", "label": "Original dataset for selected TEIDs", "status": "optional", "kind": "file", "accepted_types": [".xlsx", ".sas7bdat", ".csv"], "purpose": "Hydrates TEID-only selected membership with original participant attributes.", "adapter": "canonical_case_universe"},
        {"role": "swap_available_data", "group": "CASE DATA", "label": "Updated dataset for available TEIDs", "status": "optional", "kind": "file", "accepted_types": [".xlsx", ".sas7bdat", ".csv"], "purpose": "Hydrates TEID-only candidate membership with refreshed participant attributes.", "adapter": "canonical_case_universe"},
        {"role": "swap_case_universe", "group": "CASE DATA", "label": "Shared participant dataset", "status": "optional", "kind": "file", "accepted_types": [".xlsx", ".sas7bdat", ".csv"], "purpose": "Fallback when one participant dataset can hydrate both membership lists.", "adapter": "canonical_case_universe"},
    ],
}

STEPS = ("overview", "sources", "sample_rules", "demographics", "targets", "clinical_osep", "optimization", "outputs", "validate")
STEP_TITLES = {
    "overview": "Overview",
    "sources": "Sources",
    "sample_rules": "Sample Rules",
    "demographics": "Demographics",
    "targets": "Targets",
    "clinical_osep": "Clinical / OSEP",
    "optimization": "Optimization Rules",
    "outputs": "Output Tables",
    "validate": "Validate",
}
STEP_GUIDANCE = {
    "overview": "Identity, market, and assessment type. Unknown values may remain blank and appear as Needs Review.",
    "sources": "Record what this profile is built from. Documents are optional evidence — a Research Plan is never required. International adaptations are typically built from the dataset, a census/target workbook, a reference profile, and RD decisions.",
    "sample_rules": "Decide exactly which cases are eligible for optimization. These rules filter the candidate universe directly (for example Study = STDZ).",
    "demographics": "Demographic categories, dataset column mappings, and structure used by targets and reports.",
    "targets": "Declare target sources and how SCOPE should interpret them. A fully crossed census workbook is a source format; the CP-SAT objective uses the derived marginal/group targets with your tolerances.",
    "clinical_osep": "OSEP / clinical applicability for this market. Mark Not Applicable when the adaptation has no OSEP/clinical program.",
    "optimization": "Tolerances, priorities, and engine settings used to build the CP-SAT objective.",
    "outputs": "Reporting tables this profile must produce after a run.",
    "validate": "Core coherence and workflow readiness are evaluated independently.",
}
STEP_NEXT_HINTS = {
    "overview": "Next: record the dataset, census/target workbook, and reference profile this is built from",
    "sources": "Next: decide exactly which cases are eligible (for example Study = STDZ)",
    "sample_rules": "Next: map dataset columns to canonical demographic fields",
    "demographics": "Next: declare the census/target workbook and how SCOPE should interpret it",
    "targets": "Next: confirm whether OSEP / clinical applies in this market",
    "clinical_osep": "Next: set optimization tolerances such as Region ±10%",
    "optimization": "Next: choose the reporting tables SCOPE must produce",
    "outputs": "Next: validate the profile and review workflow readiness",
    "validate": "Validate the profile, then Save Draft or Finish / Install below",
}
STEP_REVIEW_KEYS = {
    "overview": ("identity", "locale"),
    "sources": ("sources",),
    "sample_rules": ("eligibility",),
    "demographics": ("demographics", "structure", "age_structure", "mappings"),
    "targets": ("targets",),
    "clinical_osep": ("clinical_osep",),
    "optimization": ("workflows",),
    "outputs": ("reporting",),
    "validate": (),
}
SECTION_STEP = {
    "identity": "overview", "locale": "overview", "sources": "sources",
    "eligibility": "sample_rules", "demographics": "demographics",
    "structure": "demographics", "age_structure": "demographics", "mappings": "demographics",
    "targets": "targets", "clinical_osep": "clinical_osep",
    "reporting": "outputs", "workflows": "optimization", "validation": "validate",
}
PROVENANCE_KEY_MAP = {
    "identity": "overview", "locale": "overview", "sources": "sources",
    "structure": "demographics", "age_structure": "demographics", "demographics": "demographics",
    "mappings": "demographics", "eligibility": "sample_rules", "targets": "targets",
    "clinical_osep": "clinical_osep", "workflows": "optimization", "reporting": "outputs",
}
MAPPING_FIELDS = (
    ("teid", "TEID / participant ID"), ("form", "Form"), ("age", "Age"),
    ("age_group", "Age group / band"), ("sex", "Sex"), ("ethnicity", "Ethnicity"),
    ("education", "Education (PED/SED)"), ("region", "Region"),
    ("language", "Language"), ("study_type", "Study / study type"),
)
RULE_CHOICES = ("Target", "Range", "Minimum", "Maximum", "Exact")
PRIORITY_CHOICES = ("Required", "High", "Demographic", "Low")
TARGET_FAMILIES = ("Census", "OSEP", "Clinical")
ENGINE_GUIDANCE = "Use generic_cp_sat when the candidate TEID workbook and demographic target workbook are the complete optimization specification; no Research Plan, forms, age structure, or exclusions are required."


def normalized_workflow_inputs(raw: dict | None, capabilities: dict[str, bool]) -> dict[str, list[dict]]:
    """Migrate older inputs while retaining assessment-specific entries."""
    result = {str(workflow): [dict(item) for item in items] for workflow, items in (raw or {}).items()}
    for workflow in ("recruitment_planning", "sample_optimization"):
        if capabilities.get(workflow) and not result.get(workflow):
            result[workflow] = [dict(item) for item in WORKFLOW_INPUT_DEFAULTS[workflow]]
    if capabilities.get("targeted_swap") or capabilities.get("targeted_adjustment"):
        existing = {str(item.get("role")): dict(item) for item in result.get("targeted_swap", [])}
        migrated = [{**existing.get(default["role"], {}), **default} for default in WORKFLOW_INPUT_DEFAULTS["targeted_swap"]]
        known = {item["role"] for item in migrated}
        migrated.extend(item for role, item in existing.items() if role not in known)
        result["targeted_swap"] = migrated
    return result


def _lines(widget: tk.Text) -> list[str]:
    return [value.strip() for value in widget.get("1.0", "end").splitlines() if value.strip()]


def _set_text(widget: tk.Text, values) -> None:
    widget.delete("1.0", "end")
    if values:
        widget.insert("1.0", "\n".join(str(value) for value in values))


def _slug(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "", value).upper()


def _locale_code(value: str) -> str:
    known = {"canada": "CA", "united states": "US", "usa": "US", "us": "US", "united kingdom": "UK", "uk": "UK"}
    return known.get(value.strip().casefold(), _slug(value)[:3] or "INTL")


def _norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().casefold())


def _default_inherited_origins(data: dict) -> dict[str, str]:
    """Mark sections populated by a parent profile as inherited."""
    sampling = data.get("sampling_specification") or {}
    reporting = data.get("reporting_specification") or data.get("reporting") or {}
    origins: dict[str, str] = {}
    if any(data.get(key) for key in ("forms", "age_structure", "languages", "study_types", "demographics", "schema_mapping", "reporting_variables")):
        origins["demographics"] = "inherited"
    if data.get("eligibility_rules") or sampling.get("sample_filters"):
        origins["sample_rules"] = "inherited"
    if data.get("default_target_n") or sampling.get("target_tolerances") or sampling.get("target_interpretation") or sampling.get("target_sources"):
        origins["targets"] = "inherited"
    if data.get("target_families") or data.get("target_flags") or data.get("osep_mappings"):
        origins["clinical_osep"] = "inherited"
    if data.get("source_config") or sampling.get("engine_config") or data.get("workflow_inputs") or any((data.get("capabilities") or {}).values()):
        origins["optimization"] = "inherited"
    if (reporting.get("required_tables") or reporting.get("recruitment")):
        origins["outputs"] = "inherited"
    return origins


def build_profile_draft(
    assessment: str,
    edition: str,
    profile_version: str,
    market: str,
    languages: str,
    method: str,
    parent_id: str = "",
    parent_payload: dict | None = None,
    parent_path: str = "",
    profile_kind: str = "",
) -> dict:
    """Create a new manual/document/adaptation/international draft without inventing content."""
    if method in {"adapt", "international"} and parent_payload:
        data = deepcopy(parent_payload)
        for key in ("last_modified", "workflow_readiness", "missing"):
            data.pop(key, None)
        data["source_documents"] = [{"type": "Existing Profile", "path": parent_path, "notes": f"Adapted from {parent_id}; inherited values require review."}]
        review_keys = ("locale", "structure", "age_structure", "demographics", "clinical_osep", "eligibility", "targets", "mappings", "reporting")
        data["adaptation"] = {"parent_profile_id": parent_id, "review": {key: "review" for key in review_keys}}
        data["status"] = "Needs Review"
    elif method == "adapt":
        raise ValueError("Choose an existing profile to adapt.")
    else:
        data = {"forms": [], "demographics": [], "target_flags": [], "schema_mapping": {}, "age_structure": [], "eligibility_rules": [], "languages": [], "study_types": [], "capabilities": {key: False for key, _ in CAPABILITY_LABELS}, "status": "Draft", "source_documents": []}
    display = assessment.strip()
    if market.strip() and market.strip().casefold() not in display.casefold():
        display += f" {market.strip()}"
    data.update({"profile_schema_version": 3, "id": f"{_slug(assessment)}_{_locale_code(market)}" if market.strip() else _slug(assessment), "display_name": display, "edition": edition.strip(), "version": profile_version.strip() or "0.1", "market": market.strip(), "description": "", "creation_method": method, "catalog_visibility": "user"})
    if languages.strip():
        data["languages"] = [value.strip() for value in languages.split(",") if value.strip()]
    if method == "international" or profile_kind:
        data["profile_kind"] = profile_kind or "international_adaptation"
    provenance = dict(data.get("provenance") or {})
    provenance["identity"] = {"source": "Manual configuration", "confirmed": True, "confirmed_at": datetime.now(timezone.utc).isoformat()}
    if data.get("adaptation"):
        provenance["field_origins"] = _default_inherited_origins(data)
    data["provenance"] = provenance
    return data


def _read_crossed_target_cells(path: Path) -> tuple[list[dict], list[str]]:
    """Read one-row-per-cell target workbooks without pandas (lazy openpyxl)."""
    try:
        from openpyxl import load_workbook  # noqa: F401
        from openpyxl.workbook import Workbook  # noqa: F401  (ensures PyInstaller collects openpyxl)
    except ImportError as exc:
        raise ValueError("Previewing workbook targets requires openpyxl in this Python environment.") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    sheets: list[str] = []
    try:
        for sheet in workbook.worksheets:
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                continue
            indexed = {_norm(header): index for index, header in enumerate(headers) if header is not None}
            dimensions: list[tuple[str, int]] = []
            for semantic, aliases in TARGET_DIMENSION_ALIASES.items():
                column = next((indexed[_norm(alias)] for alias in (semantic, *aliases) if _norm(alias) in indexed), None)
                if column is not None:
                    dimensions.append((semantic, column))
            count = next((indexed[_norm(name)] for name in ("count", "target count", "target_n", "target n") if _norm(name) in indexed), None)
            percent = next((indexed[_norm(name)] for name in ("pcnt", "percent", "target_pct", "target percent", "target %") if _norm(name) in indexed), None)
            if len(dimensions) < 2 or (count is None and percent is None):
                continue
            sheets.append(sheet.title)
            for record in values:
                criteria = {semantic: str(record[column]).strip() for semantic, column in dimensions if column < len(record) and record[column] not in (None, "")}
                if len(criteria) != len(dimensions):
                    continue
                if count is not None and count < len(record) and record[count] not in (None, ""):
                    target, unit = record[count], "N"
                elif percent is not None and percent < len(record) and record[percent] not in (None, ""):
                    target, unit = record[percent], "%"
                else:
                    continue
                rows.append({"category": " × ".join(DIMENSION_DISPLAY.get(name, name) for name, _column in dimensions), "group": " | ".join(f"{name}={value}" for name, value in criteria.items()), "rule": "Target", "target": target, "unit": unit, "criteria": dict(criteria), "source": f"{path.name} · {sheet.title}"})
    finally:
        workbook.close()
    return rows, sheets


class ScrollFrame(ttk.Frame):
    """A vertically scrollable container so page content is never clipped."""

    def __init__(self, parent, **kwargs):
        super().__init__(parent, **kwargs)
        self.canvas = tk.Canvas(self, bg=BACKGROUND, highlightthickness=0)
        self.scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=self.scrollbar.set)
        self.scrollbar.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.inner = ttk.Frame(self.canvas)
        self._window = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", self._on_frame_configure)
        self.canvas.bind("<Configure>", self._on_canvas_configure)
        for widget in (self.canvas, self.inner):
            widget.bind("<Enter>", lambda _event: self.canvas.bind_all("<MouseWheel>", self._on_mousewheel))
            widget.bind("<Leave>", lambda _event: self.canvas.unbind_all("<MouseWheel>"))

    def _on_frame_configure(self, _event=None) -> None:
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def _on_canvas_configure(self, event) -> None:
        self.canvas.itemconfigure(self._window, width=event.width)

    def _on_mousewheel(self, event) -> None:
        self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")


class ProfileManager:
    """Step-based profile editor with progressive disclosure and rule provenance."""

    def __init__(self) -> None:
        self.root = tk.Tk(); self.root.title("SCOPE Profile Manager"); self.root.geometry("1480x920"); self.root.minsize(1180, 780); self.root.configure(bg=BACKGROUND)
        self.base_payload: dict = {}; self.current_descriptor: Path | None = None; self.catalog_profiles = []
        self.source_documents: list[dict] = []; self.target_sources: list[dict] = []
        self.sample_filter_rules: list[dict] = []; self.tolerance_rule_rows: list[dict] = []
        self.field_origins: dict[str, str] = {}; self.adaptation_review: dict[str, str] = {}; self.last_summary: dict = {}
        self.current_step = "overview"
        self.name = tk.StringVar(); self.edition = tk.StringVar(); self.version = tk.StringVar(value="0.1"); self.profile_id = tk.StringVar(); self.market = tk.StringVar(); self.resource_folder = tk.StringVar()
        self.profile_kind = tk.StringVar(value="standard"); self.engine_config = tk.StringVar(); self.default_target_n = tk.StringVar(); self.creation_method = tk.StringVar(value="manual")
        self.capabilities = {key: tk.BooleanVar() for key, _label in CAPABILITY_LABELS}
        self.family_vars = {family: tk.BooleanVar(value=family == "Census") for family in TARGET_FAMILIES}
        self.provenance_vars = {key: tk.StringVar() for key in STEPS if key != "validate"}
        self.origin_vars = {key: tk.StringVar(value="") for key in STEPS if key != "validate"}
        self.mapping_vars = {semantic: tk.StringVar() for semantic, _label in MAPPING_FIELDS}
        self.interpretation_format = tk.StringVar(value=""); self.interpretation_cells = tk.StringVar(value="report_only"); self.interpretation_notes = tk.StringVar()
        self.derive_vars = {semantic: tk.BooleanVar() for semantic in TARGET_DIMENSION_ALIASES}
        self.glance_cases = tk.StringVar(); self.glance_targets = tk.StringVar(); self.glance_tolerances = tk.StringVar(); self.glance_origins = tk.StringVar()
        self.status = tk.StringVar(value="Choose an existing assessment or add a new one."); self.section_review_status = tk.StringVar()
        self._styles(); self._build(); self._refresh_catalog(); self.root.mainloop()

    # ------------------------------------------------------------------ chrome
    def _styles(self) -> None:
        style = ttk.Style(self.root)
        if "clam" in style.theme_names(): style.theme_use("clam")
        style.configure("TFrame", background=BACKGROUND); style.configure("TLabel", background=BACKGROUND, foreground=TEXT, font=("Segoe UI", 9))
        style.configure("TLabelframe", background=BACKGROUND, bordercolor=BORDER); style.configure("TLabelframe.Label", background=BACKGROUND, foreground=NAVY, font=("Segoe UI", 10, "bold"))
        style.configure("Primary.TButton", background=NAVY, foreground="white", padding=(13, 8), font=("Segoe UI", 9, "bold")); style.configure("Secondary.TButton", background=SURFACE, foreground=NAVY, padding=(10, 7))
        style.configure("Danger.TButton", background="#F9EFEF", foreground=ERROR, padding=(10, 7), font=("Segoe UI", 9, "bold")); style.map("Danger.TButton", background=[("active", "#F1E2E2")])
        style.configure("Rail.TButton", background=BACKGROUND, foreground=TEXT, padding=(10, 7), anchor="w"); style.configure("RailActive.TButton", background=NAVY, foreground="white", padding=(10, 7), anchor="w", font=("Segoe UI", 9, "bold"))
        style.configure("Treeview", rowheight=26, background=SURFACE, fieldbackground=SURFACE); style.configure("Treeview.Heading", background="#E8EDF1", foreground=NAVY, font=("Segoe UI", 9, "bold"))

    def _build(self) -> None:
        header = tk.Frame(self.root, bg=NAVY, height=78); header.pack(fill="x"); header.pack_propagate(False)
        tk.Label(header, text="SCOPE PROFILE MANAGER", bg=NAVY, fg="white", font=("Segoe UI", 19, "bold")).pack(anchor="w", padx=18, pady=(11, 0))
        tk.Label(header, text="Sources → Sample Rules → Demographics → Targets → Clinical/OSEP → Optimization Rules → Output Tables → Validate", bg=NAVY, fg="#D8E5EF", font=("Segoe UI", 9)).pack(anchor="w", padx=18)
        body = tk.PanedWindow(self.root, orient="horizontal", bg=BACKGROUND, sashwidth=6, bd=0); body.pack(fill="both", expand=True, padx=14, pady=12)
        left = tk.Frame(body, bg=BACKGROUND, width=440); body.add(left, minsize=360)
        tk.Label(left, text="ASSESSMENTS", bg=BACKGROUND, fg=NAVY, font=("Segoe UI", 11, "bold")).pack(anchor="w", pady=(0, 5))
        columns = ("assessment", "edition", "market", "status", "modified", "readiness")
        self.catalog = ttk.Treeview(left, columns=columns, show="headings", height=14, selectmode="browse")
        for column, title, width in (("assessment", "Assessment", 140), ("edition", "Edition", 55), ("market", "Market", 80), ("status", "Status", 85), ("modified", "Modified", 80), ("readiness", "Workflow readiness", 220)):
            self.catalog.heading(column, text=title); self.catalog.column(column, width=width, minwidth=50)
        self.catalog.pack(fill="both", expand=True); self.catalog.bind("<Double-1>", self._open_selected)
        actions = ttk.Frame(left); actions.pack(fill="x", pady=(6, 0))
        first = ttk.Frame(actions); first.pack(fill="x")
        ttk.Button(first, text="+ Add Assessment", style="Primary.TButton", command=self._add_assessment).pack(side="left", fill="x", expand=True); ttk.Button(first, text="Import Profile", style="Secondary.TButton", command=self._browse_descriptor).pack(side="left", padx=(5, 0))
        second = ttk.Frame(actions); second.pack(fill="x", pady=(4, 0))
        ttk.Button(second, text="Open Existing Profile", style="Secondary.TButton", command=self._open_selected).pack(side="left", fill="x", expand=True); ttk.Button(second, text="Duplicate / Adapt", style="Secondary.TButton", command=self._duplicate_selected).pack(side="left", padx=(5, 0))
        third = ttk.Frame(actions); third.pack(fill="x", pady=(4, 0)); ttk.Button(third, text="Export…", style="Secondary.TButton", command=self._export_profile).pack(side="left"); ttk.Button(third, text="Refresh", style="Secondary.TButton", command=self._refresh_catalog).pack(side="left", padx=(5, 0)); ttk.Button(third, text="Delete…", style="Danger.TButton", command=self._delete_selected).pack(side="right")
        tk.Label(left, text="Bundled and user profiles are shown together. Adapt an existing profile for international markets; missing specifications are normal there.", bg=BACKGROUND, fg=MUTED, justify="left", wraplength=400, font=("Segoe UI", 8)).pack(anchor="w", pady=(6, 0))
        right = tk.Frame(body, bg=BACKGROUND); body.add(right, minsize=760)
        self.editor_title = tk.Label(right, text="Select an assessment", bg=BACKGROUND, fg=NAVY, font=("Segoe UI", 16, "bold"), anchor="w"); self.editor_title.pack(fill="x")
        self.editor_subtitle = tk.Label(right, text="Open a profile or use Add Assessment to begin.", bg=BACKGROUND, fg=MUTED, font=("Segoe UI", 9), anchor="w"); self.editor_subtitle.pack(fill="x", pady=(0, 6))
        glance = tk.Frame(right, bg=SURFACE, highlightbackground=BORDER, highlightthickness=1); glance.pack(fill="x", pady=(0, 7))
        tk.Label(glance, text="AT A GLANCE", bg=SURFACE, fg=NAVY, font=("Segoe UI", 8, "bold")).grid(row=0, column=0, sticky="nw", padx=(10, 8), pady=(7, 7))
        for row, variable in enumerate((self.glance_cases, self.glance_targets, self.glance_tolerances, self.glance_origins)):
            tk.Label(glance, textvariable=variable, bg=SURFACE, fg=TEXT, font=("Segoe UI", 8), anchor="w", justify="left", wraplength=680).grid(row=row, column=1, sticky="ew", padx=(0, 10))
        glance.grid_columnconfigure(1, weight=1)
        review = ttk.Frame(right); review.pack(fill="x", pady=(0, 5)); ttk.Label(review, textvariable=self.section_review_status, foreground=WARNING).pack(side="left")
        ttk.Button(review, text="Confirm Current Section", style="Secondary.TButton", command=self._confirm_current_section).pack(side="right"); ttk.Button(review, text="Mark Not Applicable", style="Secondary.TButton", command=self._not_applicable_current_section).pack(side="right", padx=(0, 6))
        workspace = ttk.Frame(right); workspace.pack(fill="both", expand=True)
        rail = tk.Frame(workspace, bg=BACKGROUND, width=180); rail.pack(side="left", fill="y", padx=(0, 8)); rail.pack_propagate(False)
        self.step_buttons = {}
        for key in STEPS:
            button = ttk.Button(rail, text=STEP_TITLES[key], style="Rail.TButton", command=lambda value=key: self._show_step(value)); button.pack(fill="x", pady=1); self.step_buttons[key] = button
        self.content = ttk.Frame(workspace); self.content.pack(side="left", fill="both", expand=True)
        self.pages = {}
        for key in STEPS:
            scroll = ScrollFrame(self.content)
            page = scroll.inner
            page._scroll_host = scroll
            self.pages[key] = page
        self._build_overview_page(); self._build_sources_page(); self._build_sample_rules_page(); self._build_demographics_page(); self._build_targets_page(); self._build_clinical_page(); self._build_optimization_page(); self._build_outputs_page(); self._build_validate_page()
        step_nav = tk.Frame(right, bg=BACKGROUND); step_nav.pack(fill="x", pady=(7, 0))
        self.step_counter = tk.Label(step_nav, text="", bg=BACKGROUND, fg=MUTED, font=("Segoe UI", 9, "bold")); self.step_counter.pack(side="left")
        self.step_hint = tk.Label(step_nav, text="", bg=BACKGROUND, fg=MUTED, font=("Segoe UI", 8)); self.step_hint.pack(side="left", padx=(12, 0))
        self.next_step_button = ttk.Button(step_nav, text="Next →", style="Primary.TButton", command=self._next_step); self.next_step_button.pack(side="right")
        self.back_step_button = ttk.Button(step_nav, text="← Back", style="Secondary.TButton", command=self._prev_step); self.back_step_button.pack(side="right", padx=(0, 6))
        footer = ttk.Frame(right); footer.pack(fill="x", pady=(5, 0)); ttk.Button(footer, text="Save Draft", style="Secondary.TButton", command=self._save_draft).pack(side="left"); ttk.Button(footer, text="Validate", style="Primary.TButton", command=self._validate).pack(side="left", padx=(7, 0))
        self.mark_button = ttk.Button(footer, text="Mark as Validated", style="Secondary.TButton", command=self._mark_validated, state="disabled"); self.mark_button.pack(side="left", padx=(7, 0)); self.install_button = ttk.Button(footer, text="Finish / Install", style="Primary.TButton", command=self._install); self.install_button.pack(side="right")
        tk.Label(self.root, textvariable=self.status, bg="#EAE7E0", fg=NAVY, anchor="w", padx=14, pady=8, font=("Segoe UI", 9, "bold")).pack(fill="x")
        self._show_step("overview"); self._update_glance()

    # ------------------------------------------------------------------ helpers
    def _show_step(self, key: str) -> None:
        self.current_step = key
        for name, frame in self.pages.items():
            frame._scroll_host.pack_forget()
        self.pages[key]._scroll_host.pack(fill="both", expand=True)
        self.pages[key]._scroll_host.canvas.yview_moveto(0)
        for name, button in self.step_buttons.items():
            button.configure(style="RailActive.TButton" if name == key else "Rail.TButton")
        index = STEPS.index(key)
        self.step_counter.configure(text=f"Step {index + 1} of {len(STEPS)}")
        self.step_hint.configure(text=STEP_NEXT_HINTS.get(key, ""))
        self.back_step_button.configure(state="disabled" if index == 0 else "normal")
        self.next_step_button.configure(text="Validate →" if key == "outputs" else "Next →", state="disabled" if key == "validate" else "normal")
        self._tab_changed(); self._update_glance()

    def _next_step(self) -> None:
        index = STEPS.index(self.current_step)
        if index < len(STEPS) - 1:
            self._show_step(STEPS[index + 1])

    def _prev_step(self) -> None:
        index = STEPS.index(self.current_step)
        if index > 0:
            self._show_step(STEPS[index - 1])

    def _page_header(self, key: str) -> None:
        page = self.pages[key]
        title = ttk.Frame(page); title.pack(fill="x", pady=(0, 4))
        ttk.Label(title, text=STEP_TITLES[key], foreground=NAVY, font=("Segoe UI", 12, "bold")).pack(side="left")
        ttk.Label(page, text=STEP_GUIDANCE[key], foreground=MUTED, wraplength=840, justify="left").pack(anchor="w", pady=(0, 4))
        if key != "validate":
            row = ttk.Frame(page); row.pack(fill="x", pady=(0, 4))
            ttk.Label(row, text="Origin of these settings:").pack(side="left")
            combo = ttk.Combobox(row, textvariable=self.origin_vars[key], values=[f"{value} — {label}" for value, label in ORIGIN_LABELS.items()], state="readonly", width=34)
            combo.pack(side="left", padx=(6, 0)); combo.bind("<<ComboboxSelected>>", lambda _event, step=key: self._origin_chosen(step))

    def _origin_chosen(self, step: str, value: str = "") -> None:
        raw = value or self.origin_vars[step].get()
        origin = raw.split(" — ")[0].strip().casefold()
        if origin in ORIGIN_LABELS:
            self.field_origins[step] = origin
        self._update_glance()

    def _mark_manual_edit(self, step: str) -> None:
        if self.field_origins.get(step) == "inherited":
            self.field_origins[step] = "manual"
            self.origin_vars[step].set(f"manual — {ORIGIN_LABELS['manual']}")
        self._update_glance()

    def _text_field(self, parent, title: str, guidance: str, height: int = 4) -> scrolledtext.ScrolledText:
        frame = ttk.LabelFrame(parent, text=title, padding=7); frame.pack(fill="both", expand=True, pady=4); ttk.Label(frame, text=guidance, foreground=MUTED, wraplength=760).pack(anchor="w")
        widget = scrolledtext.ScrolledText(frame, height=height, wrap="word", font=("Consolas", 9)); widget.pack(fill="both", expand=True, pady=(4, 0)); return widget

    def _provenance_row(self, parent, key: str) -> None:
        row = ttk.Frame(parent); row.pack(fill="x", pady=(2, 6)); ttk.Label(row, text="Source / provenance:", width=21).pack(side="left"); ttk.Entry(row, textvariable=self.provenance_vars[key]).pack(side="left", fill="x", expand=True); ttk.Label(row, text="Optional document name or 'Manual configuration'.", foreground=MUTED).pack(side="left", padx=(7, 0))

    def _two_columns(self, parent):
        row = ttk.Frame(parent); row.pack(fill="both", expand=True); left, right = ttk.Frame(row), ttk.Frame(row); left.pack(side="left", fill="both", expand=True, padx=(0, 4)); right.pack(side="left", fill="both", expand=True, padx=(4, 0)); return left, right

    def _collapsible(self, parent, title: str):
        frame = ttk.LabelFrame(parent, text=f"▸ {title} (advanced — usually not needed)", padding=7)
        button_host = ttk.Frame(parent); button_host.pack(fill="x", pady=(6, 0))
        def toggle() -> None:
            if frame.winfo_manager(): frame.pack_forget(); button.configure(text=f"Show {title}")
            else: frame.pack(fill="both", expand=True, pady=4, after=button_host); button.configure(text=f"Hide {title}")
        button = ttk.Button(button_host, text=f"Show {title}", style="Secondary.TButton", command=toggle); button.pack(side="left")
        return frame

    # ------------------------------------------------------------------ pages
    def _build_overview_page(self) -> None:
        page = self.pages["overview"]; self._page_header("overview")
        identity = ttk.LabelFrame(page, text="Assessment identity", padding=10); identity.pack(fill="x")
        for index, (label, variable) in enumerate((("Assessment name", self.name), ("Edition", self.edition), ("Profile version", self.version), ("Profile ID", self.profile_id), ("Market / locale / country", self.market))):
            row, column = divmod(index, 3); ttk.Label(identity, text=label).grid(row=row * 2, column=column, sticky="w", padx=(0, 10)); ttk.Entry(identity, textvariable=variable).grid(row=row * 2 + 1, column=column, sticky="ew", padx=(0, 10), pady=(1, 7)); identity.grid_columnconfigure(column, weight=1)
        kind = ttk.LabelFrame(page, text="Assessment type", padding=8); kind.pack(fill="x", pady=5)
        ttk.Radiobutton(kind, text="Standard assessment", variable=self.profile_kind, value="standard").pack(anchor="w")
        ttk.Radiobutton(kind, text="International / adapted assessment — missing Research Plan, scoring, or data-analysis specifications are normal; configure from the dataset, census/target workbook, a reference profile, and RD decisions.", variable=self.profile_kind, value="international_adaptation").pack(anchor="w")
        self.description = self._text_field(page, "Short description", "Optional description of the assessment, edition, or adaptation.", 3)
        self.adaptation_label = ttk.Label(page, text="", foreground=MUTED, wraplength=840, justify="left"); self.adaptation_label.pack(anchor="w", pady=(3, 0))
        self._provenance_row(page, "overview")

    def _build_sources_page(self) -> None:
        page = self.pages["sources"]; self._page_header("sources")
        self.source_tree = ttk.Treeview(page, columns=("type", "file", "notes"), show="headings", height=7)
        for column, title, width in (("type", "Source type", 210), ("file", "File", 300), ("notes", "Notes / role in this profile", 370)): self.source_tree.heading(column, text=title); self.source_tree.column(column, width=width)
        self.source_tree.pack(fill="x")
        actions = ttk.Frame(page); actions.pack(fill="x", pady=5)
        ttk.Button(actions, text="Add Source…", style="Primary.TButton", command=self._add_source_document).pack(side="left"); ttk.Button(actions, text="Add Multiple Files…", style="Secondary.TButton", command=self._add_multiple_sources).pack(side="left", padx=(5, 0)); ttk.Button(actions, text="Edit…", style="Secondary.TButton", command=self._edit_source_document).pack(side="left", padx=(5, 0)); ttk.Button(actions, text="Remove Selected", style="Danger.TButton", command=self._remove_source_document).pack(side="right")
        ttk.Label(page, text="Typical international adaptation: the assessment dataset, the census/target workbook (for example WAIS 5 EN CDN Targets.xlsx), and the reference profile it was adapted from.", foreground=MUTED, wraplength=860, justify="left").pack(anchor="w")
        resources = self._collapsible(page, "reusable templates and resources")
        ttk.Label(resources, text="Adding a resource does not make that document type required.", foreground=MUTED).grid(row=0, column=0, columnspan=3, sticky="w")
        self.resource_vars = {}
        for row, (role, label) in enumerate(RESOURCE_ROLES, start=1):
            variable = tk.StringVar(); self.resource_vars[role] = variable; ttk.Label(resources, text=label, width=35).grid(row=row, column=0, sticky="w", pady=1); ttk.Entry(resources, textvariable=variable).grid(row=row, column=1, sticky="ew", pady=1); ttk.Button(resources, text="Browse…", style="Secondary.TButton", command=lambda value=role: self._choose_resource(value)).grid(row=row, column=2, padx=(7, 0), pady=1)
        resources.grid_columnconfigure(1, weight=1)
        self._provenance_row(page, "sources")

    def _build_sample_rules_page(self) -> None:
        page = self.pages["sample_rules"]; self._page_header("sample_rules")
        editor = ttk.LabelFrame(page, text="Sample filter / eligibility rules", padding=8); editor.pack(fill="both", expand=True)
        self.filter_tree = ttk.Treeview(editor, columns=("field", "operator", "values", "origin", "note"), show="headings", height=7)
        for column, title, width in (("field", "Field", 150), ("operator", "Rule", 90), ("values", "Value(s)", 190), ("origin", "Origin", 110), ("note", "Note / RD decision", 330)):
            self.filter_tree.heading(column, text=title); self.filter_tree.column(column, width=width)
        self.filter_tree.pack(fill="both", expand=True); self.filter_tree.bind("<Double-1>", lambda _event: self._edit_filter())
        actions = ttk.Frame(editor); actions.pack(fill="x", pady=5)
        ttk.Button(actions, text="Add Rule…", style="Primary.TButton", command=self._add_filter).pack(side="left"); ttk.Button(actions, text="Edit…", style="Secondary.TButton", command=self._edit_filter).pack(side="left", padx=(5, 0)); ttk.Button(actions, text="Delete Selected", style="Danger.TButton", command=self._remove_filter).pack(side="right")
        presets = ttk.Frame(editor); presets.pack(fill="x")
        ttk.Label(presets, text="Quick add:").pack(side="left")
        ttk.Button(presets, text="Study = STDZ", style="Secondary.TButton", command=lambda: self._quick_filter("study_type", ["STDZ"])).pack(side="left", padx=(6, 0))
        ttk.Button(presets, text="Language = English", style="Secondary.TButton", command=lambda: self._quick_filter("language", ["English"])).pack(side="left", padx=(5, 0))
        self.eligibility_extra = self._text_field(page, "Additional eligibility notes", "Any further approved rules, one per line (for example 'First administration only').", 4)
        self._provenance_row(page, "sample_rules")

    def _build_demographics_page(self) -> None:
        page = self.pages["demographics"]; self._page_header("demographics")
        left, right = self._two_columns(page)
        self.demographics = self._text_field(left, "Demographic categories", "Canonical demographic fields, one per line (for example sex, eth, edl, region).", 4)
        mapping = ttk.LabelFrame(left, text="Dataset column mapping", padding=7); mapping.pack(fill="x", pady=4)
        ttk.Label(mapping, text="Map each canonical field to the column in the assessment dataset. TEID is needed for case-level workflows.", foreground=MUTED, wraplength=430).grid(row=0, column=0, columnspan=2, sticky="w")
        for index, (semantic, label) in enumerate(MAPPING_FIELDS, start=1):
            ttk.Label(mapping, text=label).grid(row=index, column=0, sticky="w", pady=1); ttk.Entry(mapping, textvariable=self.mapping_vars[semantic], width=26).grid(row=index, column=1, sticky="ew", pady=1)
        mapping.grid_columnconfigure(1, weight=1)
        self.forms = self._text_field(right, "Forms", "One form per line. Unknown values may remain blank.", 3)
        self.languages = self._text_field(right, "Languages", "One supported language per line.", 3)
        self.study_types = self._text_field(right, "Study types", "One study type per line (for example STDZ).", 3)
        self.age_structure = self._text_field(right, "Age structure", "One age band per line, such as 16:0–17:11.", 4)
        advanced = self._collapsible(page, "reporting variables, missing values, and test IDs")
        adv_left, adv_right = self._two_columns(advanced)
        self.reporting_variables_text = self._text_field(adv_left, "Reporting variables", "Name | canonical | field | display label | value=label, value=label.", 5)
        self.missing_values = self._text_field(adv_right, "Missing-value coding", "One code per line.", 2)
        self.test_ids = self._text_field(adv_right, "Test IDs", "One approved test ID per line.", 2)
        self._provenance_row(page, "demographics")

    def _build_targets_page(self) -> None:
        page = self.pages["targets"]; self._page_header("targets")
        settings = ttk.Frame(page); settings.pack(fill="x", pady=(0, 4))
        ttk.Label(settings, text="Default target N:").pack(side="left"); ttk.Entry(settings, textvariable=self.default_target_n, width=12).pack(side="left", padx=(5, 10)); ttk.Label(settings, text="Leave blank when unknown.", foreground=MUTED).pack(side="left")
        sources = ttk.LabelFrame(page, text="Target sources", padding=8); sources.pack(fill="both", expand=True)
        self.target_source_tree = ttk.Treeview(sources, columns=("path", "format", "family", "notes"), show="headings", height=5)
        for column, title, width in (("path", "Workbook / file", 300), ("format", "Source format", 170), ("family", "Family", 90), ("notes", "Notes", 300)):
            self.target_source_tree.heading(column, text=title); self.target_source_tree.column(column, width=width)
        self.target_source_tree.pack(fill="both", expand=True); self.target_source_tree.bind("<Double-1>", lambda _event: self._edit_target_source())
        actions = ttk.Frame(sources); actions.pack(fill="x", pady=5)
        ttk.Button(actions, text="Add Target Source…", style="Primary.TButton", command=self._add_target_source).pack(side="left"); ttk.Button(actions, text="Add Multiple Files…", style="Secondary.TButton", command=self._add_multiple_target_sources).pack(side="left", padx=(5, 0)); ttk.Button(actions, text="Edit…", style="Secondary.TButton", command=self._edit_target_source).pack(side="left", padx=(5, 0)); ttk.Button(actions, text="Remove Selected", style="Danger.TButton", command=self._remove_target_source).pack(side="right")
        interpretation = ttk.LabelFrame(page, text="Target interpretation — source format vs. optimization format", padding=8); interpretation.pack(fill="x", pady=5)
        row = ttk.Frame(interpretation); row.pack(fill="x")
        ttk.Label(row, text="Source target format:", width=20).pack(side="left")
        ttk.Combobox(row, textvariable=self.interpretation_format, values=[f"{value} — {label}" for value, label in TARGET_SOURCE_FORMATS], state="readonly", width=56).pack(side="left", padx=(6, 0))
        crossed = ttk.LabelFrame(interpretation, text="For fully crossed cell-by-cell sources", padding=6); crossed.pack(fill="x", pady=5)
        ttk.Label(crossed, text="Derive marginal/group targets for:", foreground=MUTED).pack(anchor="w")
        boxes = ttk.Frame(crossed); boxes.pack(fill="x", pady=2)
        for semantic, variable in self.derive_vars.items():
            ttk.Checkbutton(boxes, text=DIMENSION_DISPLAY.get(semantic, semantic), variable=variable, command=self._update_glance).pack(side="left", padx=(0, 12))
        cells_row = ttk.Frame(crossed); cells_row.pack(fill="x", pady=(3, 0))
        ttk.Label(cells_row, text="Raw cells:").pack(side="left")
        for value, label in TARGET_CELLS_AS:
            ttk.Radiobutton(cells_row, text=label, variable=self.interpretation_cells, value=value, command=self._update_glance).pack(side="left", padx=(8, 0))
        note_row = ttk.Frame(interpretation); note_row.pack(fill="x", pady=(3, 0))
        ttk.Label(note_row, text="Interpretation notes:", width=20).pack(side="left"); ttk.Entry(note_row, textvariable=self.interpretation_notes).pack(side="left", fill="x", expand=True, padx=(6, 0))
        ttk.Button(interpretation, text="Preview Derived Optimization Targets", style="Secondary.TButton", command=self._preview_derived_targets).pack(anchor="w", pady=(6, 0))
        self._provenance_row(page, "targets")

    def _build_clinical_page(self) -> None:
        page = self.pages["clinical_osep"]; self._page_header("clinical_osep")
        families = ttk.LabelFrame(page, text="Target families that apply in this market", padding=8); families.pack(fill="x")
        ttk.Label(families, text="OSEP is a US program; many international adaptations apply Census only.", foreground=MUTED).pack(anchor="w")
        row = ttk.Frame(families); row.pack(fill="x", pady=3)
        for family, variable in self.family_vars.items():
            ttk.Checkbutton(row, text=family, variable=variable, command=self._update_glance).pack(side="left", padx=(0, 16))
        left, right = self._two_columns(page)
        self.osep_mappings = self._text_field(left, "OSEP / clinical category definitions", "group: value, value — for example 'osep: ADHD, ASD' or 'clinical: GT, IDMI'.", 5)
        self.rules = self._text_field(left, "Clinical / OSEP rules", "Only enter documented rules.", 5)
        self.osep_tables = self._text_field(right, "OSEP reporting definitions", "One report per line.", 5)
        self.clinical_tables = self._text_field(right, "Clinical reporting definitions", "One report per line.", 5)
        self._provenance_row(page, "clinical_osep")

    def _build_optimization_page(self) -> None:
        page = self.pages["optimization"]; self._page_header("optimization")
        capabilities = ttk.LabelFrame(page, text="Workflow capabilities to configure", padding=8); capabilities.pack(fill="x")
        row = ttk.Frame(capabilities); row.pack(fill="x", pady=3)
        for key, label in CAPABILITY_LABELS: ttk.Checkbutton(row, text=label, variable=self.capabilities[key]).pack(side="left", padx=(0, 14))
        engine = ttk.Frame(capabilities); engine.pack(fill="x", pady=(3, 0)); ttk.Label(engine, text="Optimizer / engine configuration:", width=31).pack(side="left"); ttk.Entry(engine, textvariable=self.engine_config).pack(side="left", fill="x", expand=True)
        ttk.Label(capabilities, text=ENGINE_GUIDANCE, foreground=MUTED, wraplength=880, justify="left").pack(anchor="w", pady=(4, 0))
        rules_frame = ttk.LabelFrame(page, text="Optimization tolerances and priorities", padding=8); rules_frame.pack(fill="both", expand=True, pady=4)
        ttk.Label(rules_frame, text="These rules shape the CP-SAT objective for the optimization target format — for example Region ±10% becomes a Range rule around each derived Region marginal.", foreground=MUTED, wraplength=880, justify="left").pack(anchor="w")
        self.rules_tree = ttk.Treeview(rules_frame, columns=("category", "applies", "rule", "target_n", "tolerance_percent", "tolerance_n", "priority", "origin", "note"), show="headings", height=6)
        for column, title, width in (("category", "Category", 115), ("applies", "Applies to", 130), ("rule", "Rule", 70), ("target_n", "Target N", 70), ("tolerance_percent", "Tolerance %", 85), ("tolerance_n", "Tolerance N", 85), ("priority", "Priority", 80), ("origin", "Origin", 95), ("note", "Note / RD decision", 200)):
            self.rules_tree.heading(column, text=title); self.rules_tree.column(column, width=width)
        self.rules_tree.pack(fill="both", expand=True, pady=3); self.rules_tree.bind("<Double-1>", lambda _event: self._edit_rule())
        actions = ttk.Frame(rules_frame); actions.pack(fill="x")
        ttk.Button(actions, text="Add Rule…", style="Primary.TButton", command=self._add_rule).pack(side="left"); ttk.Button(actions, text="Edit…", style="Secondary.TButton", command=self._edit_rule).pack(side="left", padx=(5, 0)); ttk.Button(actions, text="Delete Selected", style="Danger.TButton", command=self._remove_rule).pack(side="right")
        presets = ttk.Frame(rules_frame); presets.pack(fill="x", pady=(4, 0))
        ttk.Label(presets, text="Quick add:").pack(side="left")
        for label, category, tolerance in (("Region ±10%", "Region", 10), ("Sex ±5%", "Sex", 5), ("Ethnicity ±5%", "Ethnicity", 5), ("PED/SED ±5%", "Education", 5)):
            ttk.Button(presets, text=label, style="Secondary.TButton", command=lambda cat=category, tol=tolerance: self._quick_rule(cat, tol)).pack(side="left", padx=(6, 0))
        self.default_locks = self._text_field(page, "Default targeted-change locks", "Canonical fields that must match during swaps.", 3)
        advanced = self._collapsible(page, "task input definitions")
        self.workflow_inputs_text = self._text_field(advanced, "Advanced task input definitions", "task | role | group | required/optional | file/folder | types | adapter | label | purpose. Defaults are supplied for enabled workflows.", 8)
        self._provenance_row(page, "optimization")

    def _build_outputs_page(self) -> None:
        page = self.pages["outputs"]; self._page_header("outputs")
        left, right = self._two_columns(page)
        self.demographic_tables = self._text_field(left, "Selectable one-way tables", "One table per line. SCOPE users can choose any of these for a run.", 5)
        self.cross_tabs = self._text_field(left, "Selectable two-way crosses", "One table per line, such as Age × Sex.", 5)
        self.three_way_tables = self._text_field(right, "Selectable three-way crosses", "One table per line, such as Age × Sex × Region.", 5)
        self.recruitment_specs = self._text_field(right, "Recruitment reports by study type", "Study Type | template path | Age, Sex, Ethnicity, ...", 5)
        self._provenance_row(page, "outputs")

    def _build_validate_page(self) -> None:
        page = self.pages["validate"]; self._page_header("validate")
        self.validation_tree = ttk.Treeview(page, columns=("status", "section", "detail", "action"), show="headings", height=10)
        for column, title, width in (("status", "Status", 120), ("section", "Core profile section", 220), ("detail", "Result", 450), ("action", "Action", 150)): self.validation_tree.heading(column, text=title); self.validation_tree.column(column, width=width)
        self.validation_tree.pack(fill="both", expand=True); self.validation_tree.bind("<Double-1>", self._open_validation_section)
        ttk.Label(page, text="Workflow readiness", foreground=NAVY, font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(8, 3))
        self.workflow_tree = ttk.Treeview(page, columns=("workflow", "status", "detail"), show="headings", height=6)
        for column, title, width in (("workflow", "Workflow", 240), ("status", "Readiness", 150), ("detail", "Information needed", 550)): self.workflow_tree.heading(column, text=title); self.workflow_tree.column(column, width=width)
        self.workflow_tree.pack(fill="both", expand=True)

    # ------------------------------------------------------------------ at a glance
    def _update_glance(self) -> None:
        if not hasattr(self, "pages"):
            return
        included = [format_sample_filter(item) for item in self.sample_filter_rules if normalize_sample_filter(item)["operator"] == "include"]
        excluded = [format_sample_filter(item) for item in self.sample_filter_rules if normalize_sample_filter(item)["operator"] == "exclude"]
        parts = [f"include {value}" for value in included] + [f"exclude {value}" for value in excluded]
        self.glance_cases.set("Cases included: " + ("; ".join(parts) if parts else "all cases in the candidate dataset (no sample rules defined yet)"))
        derived = [DIMENSION_DISPLAY.get(semantic, semantic) for semantic, variable in self.derive_vars.items() if variable.get()]
        source_format = self.interpretation_format.get().split(" — ")[0].strip()
        if source_format == "crossed_cells" and derived:
            cells = dict(TARGET_CELLS_AS).get(self.interpretation_cells.get(), self.interpretation_cells.get())
            self.glance_targets.set(f"Targets optimized: derived marginal/group targets for {', '.join(derived)} — crossed cells are {cells.casefold() if isinstance(cells, str) else cells}")
        elif self.target_sources:
            self.glance_targets.set(f"Targets optimized: as provided by {len(self.target_sources)} target source(s)")
        else:
            self.glance_targets.set("Targets optimized: no target source declared yet")
        rules = []
        for row in self.tolerance_rule_rows:
            normalized = normalize_optimization_rule(row.get("category", ""), row)
            tolerance = f" ±{normalized['tolerance_percent']}%" if normalized["tolerance_percent"] not in (None, "") else f" ±{normalized['tolerance_n']} N" if normalized["tolerance_n"] not in (None, "") else ""
            rules.append(f"{normalized['category']}{tolerance}" if tolerance else f"{normalized['category']} ({normalized['rule']})")
        self.glance_tolerances.set("Tolerances: " + ("; ".join(rules) if rules else "none configured yet"))
        counts: dict[str, int] = {}
        for origin in self.field_origins.values():
            counts[origin] = counts.get(origin, 0) + 1
        for item in [*self.sample_filter_rules, *self.tolerance_rule_rows, *self.target_sources]:
            origin = str(item.get("origin", "")).casefold()
            if origin:
                counts[origin] = counts.get(origin, 0) + 1
        self.glance_origins.set("Rule origins: " + ("; ".join(f"{count} {ORIGIN_LABELS.get(origin, origin).casefold()}" for origin, count in sorted(counts.items())) if counts else "everything is new for this profile"))

    # ------------------------------------------------------------------ catalog
    def _read_descriptor(self, path: Path) -> dict:
        data = yaml.safe_load(path.read_text(encoding="utf-8")) if path.suffix.lower() in {".yaml", ".yml"} else json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict): raise ValueError("Descriptor must contain an object.")
        return data

    def _locate_descriptor(self, profile_id: str) -> Path | None:
        for root in profile_directories():
            for path in (*root.glob("*/profile.yaml"), *root.glob("*/profile.yml"), *root.glob("*/profile.json"), *root.glob("*.json")):
                try:
                    if str(self._read_descriptor(path).get("id")) == profile_id: return path
                except Exception: continue
        return None

    def _refresh_catalog(self) -> None:
        catalog = discover_profile_catalog(); self.catalog.delete(*self.catalog.get_children()); self.catalog_profiles = catalog.profiles
        for profile in self.catalog_profiles:
            readiness = workflow_readiness(profile); short = []
            for key, label in (("targeted_swap", "Swap"), ("recruitment_planning", "Recruitment"), ("sample_optimization", "Optimization")):
                marker = "✓" if readiness[key]["status"] == "READY" else "⚠" if readiness[key]["status"] == "NEEDS REVIEW" else "—"; short.append(f"{label} {marker}")
            self.catalog.insert("", "end", iid=profile.id, values=(profile.display_name, profile.edition or "—", profile.market or "—", profile.status, profile.last_modified[:10] if profile.last_modified else "—", " · ".join(short)))
        self.status.set(f"Discovered {len(catalog.profiles)} profile(s). User profiles are stored in {USER_PROFILES_DIR}.")

    def _selected_profile(self):
        selection = self.catalog.selection(); return next((profile for profile in self.catalog_profiles if profile.id == selection[0]), None) if selection else None

    def _open_selected(self, _event=None) -> None:
        profile = self._selected_profile()
        if profile is None: messagebox.showinfo("Open Existing Profile", "Select an assessment first."); return
        path = self._locate_descriptor(profile.id)
        if path: self._load_descriptor(path)

    def _browse_descriptor(self) -> None:
        selected = filedialog.askopenfilename(title="Import or open SCOPE profile", filetypes=[("Profile descriptors", "*.yaml *.yml *.json"), ("All files", "*.*")])
        if selected: self._load_descriptor(Path(selected))

    def _delete_selected(self) -> None:
        profile = self._selected_profile()
        if profile is None: messagebox.showinfo("Delete Profile", "Select an assessment profile first."); return
        bundled = Path(profile.descriptor_path).resolve().is_relative_to((ROOT / "profiles").resolve()) if profile.descriptor_path else False
        if bundled: messagebox.showinfo("Delete Profile", f"{profile.display_name} is a bundled read-only profile. Delete its user override instead from {USER_PROFILES_DIR}, or Duplicate / Adapt it."); return
        if not messagebox.askyesno("Delete Profile", f"Delete the user profile '{profile.display_name}' ({profile.id})?\n\nThis removes {profile.descriptor_path} and cannot be undone."): return
        import shutil
        shutil.rmtree(Path(profile.descriptor_path).parent, ignore_errors=True)
        self.status.set(f"Deleted user profile {profile.id}."); self._refresh_catalog()

    def _load_descriptor(self, path: Path) -> None:
        try: data = self._read_descriptor(path)
        except Exception as exc: messagebox.showerror("Profile Manager", f"Could not open {path.name}: {exc}"); return
        self._populate(data, path)

    # ------------------------------------------------------------------ populate
    def _populate(self, data: dict, path: Path | None = None) -> None:
        self.base_payload = deepcopy(data); self.current_descriptor = path
        self.name.set(str(data.get("display_name", ""))); self.edition.set(str(data.get("edition", ""))); self.version.set(str(data.get("version", "0.1"))); self.profile_id.set(str(data.get("id", ""))); self.market.set(str(data.get("market", data.get("locale", data.get("country", "")))))
        self.profile_kind.set(str(data.get("profile_kind", "standard")) or "standard")
        _set_text(self.description, [data.get("description", "")])
        stored = data.get("capabilities") or {}
        for key, variable in self.capabilities.items(): variable.set(bool(stored.get(key, stored.get("targeted_swap", False) if key == "targeted_adjustment" else False)))
        families = {str(value).strip().casefold() for value in (data.get("target_families") or (data.get("sampling_specification") or {}).get("target_families") or TARGET_FAMILIES)}
        for family, variable in self.family_vars.items(): variable.set(family.casefold() in families)
        ds, ss, rs = data.get("data_specification") or {}, data.get("sampling_specification") or {}, data.get("reporting_specification") or data.get("reporting") or {}
        mapping = data.get("schema_mapping") or ds.get("variable_map") or {}
        for semantic, variable in self.mapping_vars.items(): variable.set(str(mapping.get(semantic, "")))
        reporting_variables = data.get("reporting_variables") or ds.get("reporting_variables") or rs.get("variables") or {}; reporting_lines = []
        for name, spec in reporting_variables.items():
            if isinstance(spec, dict):
                levels = ", ".join(f"{item.get('value')}={item.get('label')}" if isinstance(item, dict) else str(item) for item in spec.get("levels", [])); reporting_lines.append(" | ".join((str(name), str(spec.get("canonical", "")), str(spec.get("field", "")), str(spec.get("display", name)), levels)))
        _set_text(self.reporting_variables_text, reporting_lines)
        widget_values = ((self.missing_values, ds.get("missing_values", [])), (self.test_ids, data.get("test_ids") or ds.get("test_ids", [])), (self.demographics, data.get("demographics") or ds.get("demographics", [])), (self.languages, data.get("languages") or ds.get("languages", [])), (self.study_types, data.get("study_types") or ds.get("study_types", [])), (self.age_structure, data.get("age_structure") or ds.get("age_structure", [])), (self.forms, data.get("forms") or ds.get("forms", [])), (self.default_locks, ss.get("default_locks") or []), (self.rules, ss.get("seeding_rules", [])), (self.demographic_tables, (rs.get("required_tables") or {}).get("demographics", [])), (self.cross_tabs, (rs.get("required_tables") or {}).get("demo_cross_tabs", [])), (self.three_way_tables, (rs.get("required_tables") or {}).get("demo_three_way", [])), (self.osep_tables, (rs.get("required_tables") or {}).get("osep", [])), (self.clinical_tables, (rs.get("required_tables") or {}).get("clinical", [])))
        for widget, values in widget_values: _set_text(widget, values)
        _set_text(self.osep_mappings, [f"{group}: {', '.join(str(value) for value in values)}" if isinstance(values, list) else str(values) for group, values in (data.get("osep_mappings") or ds.get("osep_mappings") or {}).items()])
        default_origin = "inherited" if (data.get("adaptation") or {}).get("parent_profile_id") else "manual"
        self.sample_filter_rules = []
        for item in ss.get("sample_filters") or data.get("sample_filters") or []:
            normalized = normalize_sample_filter(item)
            if not (isinstance(item, dict) and item.get("origin")): normalized["origin"] = default_origin
            self.sample_filter_rules.append(normalized)
        mirrors = {format_sample_filter(item) for item in self.sample_filter_rules}
        extras = ss.get("eligibility_notes")
        if extras is None:
            extras = [line for line in (data.get("eligibility_rules") or ss.get("eligibility_rules", [])) if line not in mirrors]
        _set_text(self.eligibility_extra, extras)
        self.tolerance_rule_rows = []
        for category, spec in (ss.get("target_tolerances") or data.get("target_tolerances") or {}).items():
            normalized = normalize_optimization_rule(category, spec)
            if not (isinstance(spec, dict) and spec.get("origin")): normalized["origin"] = default_origin
            self.tolerance_rule_rows.append(normalized)
        self.target_sources = []
        for item in ss.get("target_sources") or data.get("target_sources") or []:
            if isinstance(item, dict):
                source = dict(item); source.setdefault("origin", default_origin); self.target_sources.append(source)
        if not self.target_sources:
            resources = data.get("resources") or {}
            for role, family in (("census_targets", "Census"), ("approved_targets", "Census")):
                if resources.get(role):
                    self.target_sources.append({"path": str(resources[role]), "format": "", "family": family, "notes": f"Declared as reusable resource '{role}'.", "origin": default_origin}); break
        interpretation = ss.get("target_interpretation") or data.get("target_interpretation") or {}
        if not interpretation and self.profile_kind.get() == "international_adaptation":
            interpretation = {"source_format": "crossed_cells", "cells_as": "report_only", "derive_marginals": self._demographic_semantics()}
        self.interpretation_format.set(next((f"{value} — {label}" for value, label in TARGET_SOURCE_FORMATS if value == str(interpretation.get("source_format", ""))), str(interpretation.get("source_format", ""))))
        self.interpretation_cells.set(str(interpretation.get("cells_as", "report_only")) or "report_only"); self.interpretation_notes.set(str(interpretation.get("notes", "")))
        derived = {dimension_semantic(value) for value in interpretation.get("derive_marginals", [])}
        for semantic, variable in self.derive_vars.items(): variable.set(semantic in derived)
        capabilities = {key: variable.get() for key, variable in self.capabilities.items()}; _set_text(self.workflow_inputs_text, self._workflow_input_lines(normalized_workflow_inputs(data.get("workflow_inputs"), capabilities)))
        recruitment = rs.get("recruitment") or {}; _set_text(self.recruitment_specs, [f"{study} | {spec.get('template', '')} | {', '.join(spec.get('required_tables', []))}" for study, spec in recruitment.items() if isinstance(spec, dict)])
        self.engine_config.set(str(data.get("source_config") or ss.get("engine_config") or "")); target_n = data.get("default_target_n", ss.get("default_target_n", "")); self.default_target_n.set("" if target_n in (None, "") else str(target_n))
        for role, variable in self.resource_vars.items(): variable.set(str((data.get("resources") or {}).get(role, "")))
        package = path.parent if path else Path(); self.source_documents = []
        for item in data.get("source_documents") or []:
            if isinstance(item, dict):
                document = dict(item); source = Path(str(document.get("path", "")))
                if source and not source.is_absolute() and path: document["path"] = str(package / source)
                self.source_documents.append(document)
        self._refresh_source_tree(); self._refresh_filter_tree(); self._refresh_rules_tree(); self._refresh_target_source_tree()
        provenance = data.get("provenance") or {}
        for key, variable in self.provenance_vars.items(): variable.set("")
        for raw_key, value in provenance.items():
            if raw_key == "field_origins": continue
            step = PROVENANCE_KEY_MAP.get(raw_key, raw_key if raw_key in self.provenance_vars else "")
            if step and not self.provenance_vars[step].get():
                self.provenance_vars[step].set(str(value.get("source", "")) if isinstance(value, dict) else str(value))
        self.field_origins = {str(key): str(value) for key, value in (provenance.get("field_origins") or {}).items()}
        if not self.field_origins and data.get("adaptation"):
            self.field_origins = _default_inherited_origins(data)
        for step, variable in self.origin_vars.items():
            origin = self.field_origins.get(step, "")
            variable.set(f"{origin} — {ORIGIN_LABELS[origin]}" if origin in ORIGIN_LABELS else "")
        self.adaptation_review = dict((data.get("adaptation") or {}).get("review") or {}); self.creation_method.set(str(data.get("creation_method", "manual"))); self.resource_folder.set(str(path.parent) if path else "")
        parent = (data.get("adaptation") or {}).get("parent_profile_id", "")
        self.adaptation_label.configure(text=f"Adapted from {parent}. Values marked 'inherited' came from the parent profile until you confirm or change them." if parent else "")
        self.editor_title.configure(text=self.name.get() or "Unnamed assessment")
        kind_label = "International / adapted assessment" if self.profile_kind.get() == "international_adaptation" else "Standard assessment"
        self.editor_subtitle.configure(text=f"{self.market.get() or 'Market not specified'} · {kind_label} · {data.get('status', 'Draft')} · {self.profile_id.get() or 'ID pending'}")
        self.validation_tree.delete(*self.validation_tree.get_children()); self.workflow_tree.delete(*self.workflow_tree.get_children()); self.mark_button.configure(state="disabled"); self.install_button.configure(state="normal"); self.status.set(f"Editing {path}" if path else "New assessment profile ready for configuration.")
        self._update_step_icons(); self._show_step("overview")

    def _demographic_semantics(self) -> list[str]:
        aliases = {"sex": "sex", "eth": "ethnicity", "ethnicity": "ethnicity", "edl": "education", "education": "education", "region": "region", "age": "age_group", "age_group": "age_group", "agegrp": "age_group"}
        semantics: list[str] = []
        for value in _lines(self.demographics):
            semantic = aliases.get(_norm(value), dimension_semantic(value))
            if semantic and semantic not in semantics: semantics.append(semantic)
        return semantics

    # ------------------------------------------------------------------ wizard
    def _add_assessment(self) -> None: self._assessment_wizard()
    def _duplicate_selected(self) -> None:
        profile = self._selected_profile()
        if profile is None: messagebox.showinfo("Duplicate / Adapt", "Select an existing assessment profile first."); return
        self._assessment_wizard(profile.id)

    def _assessment_wizard(self, parent_id: str = "") -> None:
        window = tk.Toplevel(self.root); window.title("Add Assessment"); window.geometry("720x560"); window.configure(bg=BACKGROUND); window.transient(self.root); window.grab_set()
        step = tk.IntVar(value=1); assessment = tk.StringVar(); edition = tk.StringVar(); profile_version = tk.StringVar(value="0.1"); market = tk.StringVar(); languages = tk.StringVar(); method = tk.StringVar(value="international" if parent_id else "documents"); parent = tk.StringVar()
        labels = {profile.id: f"{profile.display_name} · {profile.market or 'market not recorded'} · {profile.status}" for profile in self.catalog_profiles}; reverse = {label: key for key, label in labels.items()}; parent.set(labels.get(parent_id, ""))
        title = tk.Label(window, text="STEP 1 OF 2 — ASSESSMENT IDENTITY", bg=BACKGROUND, fg=NAVY, font=("Segoe UI", 15, "bold")); title.pack(anchor="w", padx=22, pady=(20, 6)); host = ttk.Frame(window); host.pack(fill="both", expand=True, padx=22); first, second = ttk.Frame(host), ttk.Frame(host)
        for row, (label, variable) in enumerate((("Assessment name:", assessment), ("Edition / version:", edition), ("Market / locale / country:", market), ("Language(s), if known:", languages), ("Profile version:", profile_version))): ttk.Label(first, text=label).grid(row=row, column=0, sticky="w", pady=5); ttk.Entry(first, textvariable=variable).grid(row=row, column=1, sticky="ew", pady=5)
        first.grid_columnconfigure(1, weight=1); ttk.Label(first, text="Unknown information may remain blank and will appear as Needs Review.", foreground=MUTED).grid(row=5, column=0, columnspan=2, sticky="w", pady=(10, 0))
        options = (
            ("international", "International / adapted assessment", "For a new market or language (for example WAIS-5 Canada). No Research Plan, scoring spec, or data-analysis spec is expected — configure from the dataset, census/target workbook, a reference profile, and RD decisions."),
            ("adapt", "Adapt an existing profile", "Review inherited information for this market/version."),
            ("documents", "Build from source documents", "Add any authoritative documents you have; none is mandatory."),
            ("manual", "Configure manually", "Create without requiring source documents."),
        )
        for row, (value, label, help_text) in enumerate(options): ttk.Radiobutton(second, text=label, variable=method, value=value).grid(row=row * 2, column=0, sticky="w", pady=(8, 0)); ttk.Label(second, text=help_text, foreground=MUTED, wraplength=600, justify="left").grid(row=row * 2 + 1, column=0, sticky="w", padx=(22, 0))
        ttk.Label(second, text="Reference profile to inherit from (optional for international):").grid(row=8, column=0, sticky="w", pady=(16, 3)); ttk.Combobox(second, textvariable=parent, values=list(reverse), state="readonly").grid(row=9, column=0, sticky="ew"); second.grid_columnconfigure(0, weight=1)
        if parent_id:
            source = next((p for p in self.catalog_profiles if p.id == parent_id), None)
            if source: assessment.set(source.display_name); edition.set(source.edition)
        first.pack(fill="both", expand=True); buttons = ttk.Frame(window); buttons.pack(fill="x", padx=22, pady=18); back = ttk.Button(buttons, text="Back", style="Secondary.TButton", state="disabled"); back.pack(side="left"); next_button = ttk.Button(buttons, text="Next", style="Primary.TButton"); next_button.pack(side="right")
        def show(which: int) -> None:
            step.set(which); first.pack_forget(); second.pack_forget(); (first if which == 1 else second).pack(fill="both", expand=True); title.configure(text="STEP 1 OF 2 — ASSESSMENT IDENTITY" if which == 1 else "STEP 2 OF 2 — HOW TO CREATE THE PROFILE"); back.configure(state="disabled" if which == 1 else "normal"); next_button.configure(text="Next" if which == 1 else "Create Profile")
        back.configure(command=lambda: show(1))
        def advance() -> None:
            if step.get() == 1:
                if not assessment.get().strip(): messagebox.showerror("Add Assessment", "Enter an assessment name.", parent=window); return
                show(2); return
            if method.get() == "adapt" and parent.get() not in reverse: messagebox.showerror("Add Assessment", "Choose an existing profile to adapt.", parent=window); return
            self._create_from_wizard(assessment.get(), edition.get(), profile_version.get(), market.get(), languages.get(), method.get(), reverse.get(parent.get(), "")); window.destroy()
        next_button.configure(command=advance)

    def _create_from_wizard(self, assessment: str, edition: str, profile_version: str, market: str, languages: str, method: str, parent_id: str) -> None:
        path = self._locate_descriptor(parent_id) if method in {"adapt", "international"} and parent_id else None
        parent_payload = self._read_descriptor(path) if path else None
        try: data = build_profile_draft(assessment, edition, profile_version, market, languages, method, parent_id, parent_payload, str(path or ""))
        except Exception as exc: messagebox.showerror("Add Assessment", str(exc)); return
        self._populate(data); self._show_step("sources" if method == "documents" else "demographics" if method == "manual" else "overview"); self.status.set("Assessment created as a draft. Review the indicated sections and Save Draft at any time.")

    # ------------------------------------------------------------------ sources step
    def _source_dialog(self, index: int | None = None) -> None:
        existing = self.source_documents[index] if index is not None else {}; path = filedialog.askopenfilename(title="Add source", initialdir=self.resource_folder.get() or None)
        if not path and index is None: return
        window = tk.Toplevel(self.root); window.title("Source"); window.geometry("640x260"); window.configure(bg=BACKGROUND); window.transient(self.root); window.grab_set(); source_type = tk.StringVar(value=str(existing.get("type", SOURCE_TYPES[0]))); source_path = tk.StringVar(value=path or str(existing.get("path", ""))); notes = tk.StringVar(value=str(existing.get("notes", ""))); form = ttk.Frame(window, padding=16); form.pack(fill="both", expand=True)
        for row, (label, widget) in enumerate((("Source type:", ttk.Combobox(form, textvariable=source_type, values=SOURCE_TYPES, state="readonly")), ("File:", ttk.Entry(form, textvariable=source_path)), ("Notes / role in this profile:", ttk.Entry(form, textvariable=notes)))): ttk.Label(form, text=label).grid(row=row, column=0, sticky="w", pady=5); widget.grid(row=row, column=1, sticky="ew", pady=5)
        form.grid_columnconfigure(1, weight=1)
        def save() -> None:
            item = {"type": source_type.get(), "path": source_path.get().strip(), "notes": notes.get().strip()}; self.source_documents.append(item) if index is None else self.source_documents.__setitem__(index, item); self._refresh_source_tree(); self._mark_manual_edit("sources"); window.destroy()
        actions = ttk.Frame(form); actions.grid(row=3, column=1, sticky="e", pady=(18, 0))
        ttk.Button(actions, text="Cancel", style="Secondary.TButton", command=window.destroy).pack(side="right", padx=(7, 0)); ttk.Button(actions, text="Save Source", style="Primary.TButton", command=save).pack(side="right")

    def _add_multiple_sources(self) -> None:
        paths = filedialog.askopenfilenames(title="Add source files (multi-select)", initialdir=self.resource_folder.get() or None)
        if not paths: return
        for path in paths:
            self.source_documents.append({"type": "Other authoritative project document", "path": path, "notes": "Added via multi-select — set the source type with Edit."})
        self._refresh_source_tree(); self._mark_manual_edit("sources"); self.status.set(f"Added {len(paths)} source file(s). Use Edit to set each source type.")

    def _add_source_document(self) -> None: self._source_dialog()
    def _edit_source_document(self) -> None:
        selected = self.source_tree.selection()
        if selected: self._source_dialog(int(selected[0]))
    def _remove_source_document(self) -> None:
        selected = self.source_tree.selection()
        if selected: self.source_documents.pop(int(selected[0])); self._refresh_source_tree(); self._mark_manual_edit("sources")
    def _refresh_source_tree(self) -> None:
        if not hasattr(self, "source_tree"): return
        self.source_tree.delete(*self.source_tree.get_children())
        for index, item in enumerate(self.source_documents): self.source_tree.insert("", "end", iid=str(index), values=(item.get("type", "Other"), Path(str(item.get("path", ""))).name, item.get("notes", "")))
    def _choose_resource(self, role: str) -> None:
        selected = filedialog.askopenfilename(title=dict(RESOURCE_ROLES)[role], initialdir=self.resource_folder.get() or None)
        if selected: self.resource_vars[role].set(selected)

    # ------------------------------------------------------------------ sample rules step
    def _filter_dialog(self, index: int | None = None) -> None:
        existing = self.sample_filter_rules[index] if index is not None else {}
        window = tk.Toplevel(self.root); window.title("Sample Filter Rule"); window.geometry("560x330"); window.configure(bg=BACKGROUND); window.transient(self.root); window.grab_set()
        field = tk.StringVar(value=str(existing.get("field", "study_type"))); operator = tk.StringVar(value="Include only" if existing.get("operator", "include") == "include" else "Exclude"); values = tk.StringVar(value=", ".join(str(value) for value in existing.get("values", []))); origin = tk.StringVar(value=str(existing.get("origin", "manual"))); note = tk.StringVar(value=str(existing.get("note", "")))
        form = ttk.Frame(window, padding=16); form.pack(fill="both", expand=True)
        ttk.Label(form, text="Field:").grid(row=0, column=0, sticky="w", pady=4); ttk.Combobox(form, textvariable=field, values=[f"{semantic} — {label}" for semantic, label in SAMPLE_FILTER_FIELDS], state="readonly").grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Rule:").grid(row=1, column=0, sticky="w", pady=4); ttk.Combobox(form, textvariable=operator, values=("Include only", "Exclude"), state="readonly").grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Value(s), comma-separated:").grid(row=2, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=values).grid(row=2, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Origin:").grid(row=3, column=0, sticky="w", pady=4); ttk.Combobox(form, textvariable=origin, values=list(ORIGIN_LABELS), state="readonly").grid(row=3, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Note / RD decision:").grid(row=4, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=note).grid(row=4, column=1, sticky="ew", pady=4)
        form.grid_columnconfigure(1, weight=1)
        ttk.Label(form, text="Example: field study_type, Include only STDZ keeps only standardization cases eligible for optimization.", foreground=MUTED, wraplength=500, justify="left").grid(row=5, column=0, columnspan=2, sticky="w", pady=(6, 0))
        def save() -> None:
            if not values.get().strip(): messagebox.showerror("Sample Filter Rule", "Enter at least one value.", parent=window); return
            item = normalize_sample_filter({"field": field.get().split(" — ")[0], "operator": "exclude" if operator.get() == "Exclude" else "include", "values": values.get(), "origin": origin.get(), "note": note.get()})
            self.sample_filter_rules.append(item) if index is None else self.sample_filter_rules.__setitem__(index, item)
            self._refresh_filter_tree(); self._mark_manual_edit("sample_rules"); window.destroy()
        actions = ttk.Frame(form); actions.grid(row=6, column=1, sticky="e", pady=(14, 0))
        ttk.Button(actions, text="Cancel", style="Secondary.TButton", command=window.destroy).pack(side="right", padx=(7, 0)); ttk.Button(actions, text="Save Rule", style="Primary.TButton", command=save).pack(side="right")

    def _add_filter(self) -> None: self._filter_dialog()
    def _edit_filter(self) -> None:
        selected = self.filter_tree.selection()
        if selected: self._filter_dialog(int(selected[0]))
    def _remove_filter(self) -> None:
        selected = self.filter_tree.selection()
        if selected: self.sample_filter_rules.pop(int(selected[0])); self._refresh_filter_tree(); self._mark_manual_edit("sample_rules")
    def _quick_filter(self, field: str, values: list[str]) -> None:
        if any(normalize_sample_filter(item)["field"] == field for item in self.sample_filter_rules): self.status.set(f"A {field} sample rule already exists — edit it instead."); return
        self.sample_filter_rules.append(normalize_sample_filter({"field": field, "values": values, "origin": "manual", "note": "Quick-added preset"})); self._refresh_filter_tree(); self._mark_manual_edit("sample_rules")
    def _refresh_filter_tree(self) -> None:
        if not hasattr(self, "filter_tree"): return
        self.filter_tree.delete(*self.filter_tree.get_children())
        for index, item in enumerate(self.sample_filter_rules):
            normalized = normalize_sample_filter(item); self.filter_tree.insert("", "end", iid=str(index), values=(normalized["label"], "Include only" if normalized["operator"] == "include" else "Exclude", ", ".join(normalized["values"]), ORIGIN_LABELS.get(normalized["origin"], normalized["origin"]), normalized["note"]))
        self._update_glance()

    # ------------------------------------------------------------------ targets step
    def _target_source_dialog(self, index: int | None = None) -> None:
        existing = self.target_sources[index] if index is not None else {}
        path = filedialog.askopenfilename(title="Add target source", initialdir=self.resource_folder.get() or None, filetypes=[("Target workbooks", "*.xlsx *.xlsm *.xls *.csv *.json *.docx"), ("All files", "*.*")])
        if not path and index is None: return
        window = tk.Toplevel(self.root); window.title("Target Source"); window.geometry("660x300"); window.configure(bg=BACKGROUND); window.transient(self.root); window.grab_set()
        source_path = tk.StringVar(value=path or str(existing.get("path", ""))); source_format = tk.StringVar(value=str(existing.get("format", ""))); family = tk.StringVar(value=str(existing.get("family", "Census"))); notes = tk.StringVar(value=str(existing.get("notes", "")))
        form = ttk.Frame(window, padding=16); form.pack(fill="both", expand=True)
        ttk.Label(form, text="File:").grid(row=0, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=source_path).grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Source format:").grid(row=1, column=0, sticky="w", pady=4); ttk.Combobox(form, textvariable=source_format, values=[f"{value} — {label}" for value, label in TARGET_SOURCE_FORMATS], state="readonly").grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Target family:").grid(row=2, column=0, sticky="w", pady=4); ttk.Combobox(form, textvariable=family, values=TARGET_FAMILIES + ("Mixed",), state="readonly").grid(row=2, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Notes:").grid(row=3, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=notes).grid(row=3, column=1, sticky="ew", pady=4)
        form.grid_columnconfigure(1, weight=1)
        def save() -> None:
            item = {"path": source_path.get().strip(), "format": source_format.get().split(" — ")[0], "family": family.get(), "notes": notes.get().strip(), "origin": str(existing.get("origin", "manual"))}
            self.target_sources.append(item) if index is None else self.target_sources.__setitem__(index, item)
            if item["format"] and not self.interpretation_format.get(): self.interpretation_format.set(next(f"{value} — {label}" for value, label in TARGET_SOURCE_FORMATS if value == item["format"]))
            self._refresh_target_source_tree(); self._mark_manual_edit("targets"); window.destroy()
        actions = ttk.Frame(form); actions.grid(row=4, column=1, sticky="e", pady=(14, 0))
        ttk.Button(actions, text="Cancel", style="Secondary.TButton", command=window.destroy).pack(side="right", padx=(7, 0)); ttk.Button(actions, text="Save Source", style="Primary.TButton", command=save).pack(side="right")

    def _add_multiple_target_sources(self) -> None:
        paths = filedialog.askopenfilenames(title="Add target sources (multi-select)", initialdir=self.resource_folder.get() or None, filetypes=[("Target workbooks", "*.xlsx *.xlsm *.xls *.csv *.json *.docx"), ("All files", "*.*")])
        if not paths: return
        for path in paths:
            self.target_sources.append({"path": path, "format": "", "family": "Census", "notes": "Added via multi-select — set the format with Edit.", "origin": "assessment"})
        self._refresh_target_source_tree(); self._mark_manual_edit("targets"); self.status.set(f"Added {len(paths)} target source(s). Use Edit to set each source format.")

    def _add_target_source(self) -> None: self._target_source_dialog()
    def _edit_target_source(self) -> None:
        selected = self.target_source_tree.selection()
        if selected: self._target_source_dialog(int(selected[0]))
    def _remove_target_source(self) -> None:
        selected = self.target_source_tree.selection()
        if selected: self.target_sources.pop(int(selected[0])); self._refresh_target_source_tree(); self._mark_manual_edit("targets")
    def _refresh_target_source_tree(self) -> None:
        if not hasattr(self, "target_source_tree"): return
        self.target_source_tree.delete(*self.target_source_tree.get_children())
        for index, item in enumerate(self.target_sources):
            label = dict(TARGET_SOURCE_FORMATS).get(str(item.get("format", "")), str(item.get("format", "")) or "—"); self.target_source_tree.insert("", "end", iid=str(index), values=(Path(str(item.get("path", ""))).name, label, item.get("family", ""), item.get("notes", "")))
        self._update_glance()

    def _tolerances_dict(self) -> dict[str, dict]:
        tolerances: dict[str, dict] = {}
        for row in self.tolerance_rule_rows:
            normalized = normalize_optimization_rule(row.get("category", ""), row); spec: dict = {}
            if normalized["rule"] != "Target": spec["rule"] = normalized["rule"]
            for key in ("target_n", "tolerance_percent", "tolerance_n", "priority", "origin", "note"):
                if normalized[key] not in (None, ""): spec[key] = normalized[key]
            if normalized["applies_to"] == "group" and normalized["group"]:
                spec["applies_to"] = "group"; spec["group"] = normalized["group"]
            key = normalized["category"] if normalized["applies_to"] == "all" else f"{normalized['category']} | {normalized['group']}"
            tolerances[key] = spec
        return tolerances
    def _interpretation(self) -> dict:
        source_format = self.interpretation_format.get().split(" — ")[0].strip()
        derived = [semantic for semantic, variable in self.derive_vars.items() if variable.get()]
        interpretation: dict = {}
        if source_format or derived:
            interpretation = {"source_format": source_format or "mixed", "cells_as": self.interpretation_cells.get() or "report_only"}
            if derived: interpretation["derive_marginals"] = derived
            if self.interpretation_notes.get().strip(): interpretation["notes"] = self.interpretation_notes.get().strip()
        return interpretation

    def _preview_derived_targets(self) -> None:
        try:
            import pandas  # noqa: F401
        except ImportError:
            messagebox.showerror("Preview Derived Targets", "The standalone Profile Manager build cannot read target workbooks because pandas is not bundled. Rebuild with the latest misc/Build_Scope_Profile_Manager_EXE.ps1."); return
        sources = [item for item in self.target_sources if str(item.get("path", "")).strip()]
        if not sources: messagebox.showinfo("Preview Derived Targets", "Add a target source first — for example the census workbook."); return
        dimensions = [semantic for semantic, variable in self.derive_vars.items() if variable.get()]
        if not dimensions: messagebox.showinfo("Preview Derived Targets", "Choose at least one marginal dimension to derive (for example Region, Sex)."); return
        base = Path(self.resource_folder.get() or Path.cwd()); rows: list[dict] = []; read_notes: list[str] = []
        for item in sources:
            candidate = Path(str(item["path"])); candidate = candidate if candidate.is_absolute() else base / candidate
            if not candidate.is_file(): read_notes.append(f"{candidate.name}: file not found"); continue
            try: cells, sheets = _read_crossed_target_cells(candidate)
            except Exception as exc: read_notes.append(f"{candidate.name}: {exc}"); continue
            if cells: rows.extend(cells); read_notes.append(f"{candidate.name}: {len(cells):,} crossed cell(s) from {', '.join(sheets)}")
            else: read_notes.append(f"{candidate.name}: no fully crossed sheets detected")
        derived = derive_marginal_target_rows(rows, dimensions, self._tolerances_dict(), source_label="Derived from target source")
        window = tk.Toplevel(self.root); window.title("Derived Optimization Targets"); window.geometry("900x480"); window.configure(bg=BACKGROUND); window.transient(self.root)
        tk.Label(window, text="DERIVED OPTIMIZATION TARGETS", bg=BACKGROUND, fg=NAVY, font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=18, pady=(14, 2))
        behavior = dict(TARGET_CELLS_AS).get(self.interpretation_cells.get(), "")
        tk.Label(window, text=f"These marginal/group targets become the CP-SAT objective. Raw crossed cells: {behavior}.\n" + ("\n".join(read_notes) if read_notes else "No sources read."), bg=BACKGROUND, fg=MUTED, justify="left", wraplength=860).pack(anchor="w", padx=18)
        tree = ttk.Treeview(window, columns=("category", "group", "rule", "target", "unit", "priority", "origin"), show="headings")
        for column, title, width in (("category", "Category", 110), ("group", "Group", 150), ("rule", "Rule", 80), ("target", "Target", 120), ("unit", "Unit", 60), ("priority", "Priority", 100), ("origin", "Origin", 120)): tree.heading(column, text=title); tree.column(column, width=width)
        tree.pack(fill="both", expand=True, padx=18, pady=10)
        for row in derived: tree.insert("", "end", values=(row.get("category", ""), row.get("group", ""), row.get("rule", ""), row.get("target", ""), row.get("unit", ""), row.get("priority", ""), ORIGIN_LABELS.get(str(row.get("origin", "")), row.get("origin", ""))))
        if not derived: tree.insert("", "end", values=("—", "No derived targets", "", "", "", "", ""))

    # ------------------------------------------------------------------ optimization step
    def _rule_dialog(self, index: int | None = None) -> None:
        existing = self.tolerance_rule_rows[index] if index is not None else {}
        window = tk.Toplevel(self.root); window.title("Optimization Rule"); window.geometry("620x470"); window.configure(bg=BACKGROUND); window.transient(self.root); window.grab_set()
        category = tk.StringVar(value=str(existing.get("category", ""))); applies_to = tk.StringVar(value=str(existing.get("applies_to", "all"))); group = tk.StringVar(value=str(existing.get("group", "")))
        rule = tk.StringVar(value=str(existing.get("rule", "Target"))); target_n = tk.StringVar(value=str(existing.get("target_n", "")))
        tolerance_percent = tk.StringVar(value=str(existing.get("tolerance_percent", ""))); tolerance_n = tk.StringVar(value=str(existing.get("tolerance_n", ""))); priority = tk.StringVar(value=str(existing.get("priority", "Demographic"))); origin = tk.StringVar(value=str(existing.get("origin", "manual"))); note = tk.StringVar(value=str(existing.get("note", "")))
        form = ttk.Frame(window, padding=16); form.pack(fill="both", expand=True)
        example = tk.Label(form, text="", bg=BACKGROUND, fg=NAVY, font=("Segoe UI", 9, "bold"), anchor="w", justify="left", wraplength=560)
        def refresh_example(*_args) -> None:
            cat = category.get().strip() or "Category"
            scope = "all groups" if applies_to.get() == "all" else f"group '{group.get().strip() or '?'}'"
            if target_n.get().strip():
                example.configure(text=f"→ {cat} ({scope}): Exact N = {target_n.get().strip()}")
            elif tolerance_percent.get().strip():
                example.configure(text=f"→ {cat} ({scope}): Range around the target, ±{tolerance_percent.get().strip()}%")
            elif tolerance_n.get().strip():
                example.configure(text=f"→ {cat} ({scope}): Range around the target, ±{tolerance_n.get().strip()} N")
            else:
                example.configure(text=f"→ {cat} ({scope}): {rule.get()} as provided by the target source")
        ttk.Label(form, text="Category:").grid(row=0, column=0, sticky="w", pady=4); category_box = ttk.Combobox(form, textvariable=category, values=list(DIMENSION_DISPLAY.values()) + ["Age Group", "Sample Cell", "OSEP / Clinical"], state="normal"); category_box.grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Applies to:").grid(row=1, column=0, sticky="w", pady=4); scope_row = ttk.Frame(form); scope_row.grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Radiobutton(scope_row, text="All groups in this category", variable=applies_to, value="all", command=refresh_example).pack(side="left")
        ttk.Radiobutton(scope_row, text="Only one group:", variable=applies_to, value="group", command=refresh_example).pack(side="left", padx=(12, 4)); ttk.Entry(scope_row, textvariable=group, width=22).pack(side="left")
        ttk.Label(form, text="Rule:").grid(row=2, column=0, sticky="w", pady=4); ttk.Combobox(form, textvariable=rule, values=RULE_CHOICES, state="readonly").grid(row=2, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Target N (exact count):").grid(row=3, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=target_n).grid(row=3, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Tolerance ± %:").grid(row=4, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=tolerance_percent).grid(row=4, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Tolerance ± N:").grid(row=5, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=tolerance_n).grid(row=5, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Priority:").grid(row=6, column=0, sticky="w", pady=4); ttk.Combobox(form, textvariable=priority, values=PRIORITY_CHOICES, state="readonly").grid(row=6, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Origin:").grid(row=7, column=0, sticky="w", pady=4); ttk.Combobox(form, textvariable=origin, values=list(ORIGIN_LABELS), state="readonly").grid(row=7, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Note / RD decision:").grid(row=8, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=note).grid(row=8, column=1, sticky="ew", pady=4)
        example.grid(row=9, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        form.grid_columnconfigure(1, weight=1)
        ttk.Label(form, text="Set Target N to require an exact count (for example Age Group = 100 for every band, or Only one group 16:0–17:11 = 100). A ± % tolerance turns a percentage target into a Range rule instead (Region 23.4% ± 10% → 13.4–33.4%).", foreground=MUTED, wraplength=560, justify="left").grid(row=10, column=0, columnspan=2, sticky="w", pady=(6, 0))
        for variable in (category, applies_to, group, rule, target_n, tolerance_percent, tolerance_n): variable.trace_add("write", refresh_example)
        refresh_example()
        def save() -> None:
            if not category.get().strip(): messagebox.showerror("Optimization Rule", "Enter a category.", parent=window); return
            if applies_to.get() == "group" and not group.get().strip(): messagebox.showerror("Optimization Rule", "Enter the group this rule applies to (for example 16:0–17:11).", parent=window); return
            item = normalize_optimization_rule(category.get(), {"rule": rule.get(), "target_n": target_n.get(), "tolerance_percent": tolerance_percent.get(), "tolerance_n": tolerance_n.get(), "priority": priority.get(), "origin": origin.get(), "note": note.get(), "applies_to": applies_to.get(), "group": group.get()})
            self.tolerance_rule_rows.append(item) if index is None else self.tolerance_rule_rows.__setitem__(index, item)
            self._refresh_rules_tree(); self._mark_manual_edit("optimization"); window.destroy()
        actions = ttk.Frame(form); actions.grid(row=11, column=1, sticky="e", pady=(12, 0))
        ttk.Button(actions, text="Cancel", style="Secondary.TButton", command=window.destroy).pack(side="right", padx=(7, 0)); ttk.Button(actions, text="Save Rule", style="Primary.TButton", command=save).pack(side="right")

    def _add_rule(self) -> None: self._rule_dialog()
    def _edit_rule(self) -> None:
        selected = self.rules_tree.selection()
        if selected: self._rule_dialog(int(selected[0]))
    def _remove_rule(self) -> None:
        selected = self.rules_tree.selection()
        if selected: self.tolerance_rule_rows.pop(int(selected[0])); self._refresh_rules_tree(); self._mark_manual_edit("optimization")
    def _quick_rule(self, category: str, tolerance_percent: float) -> None:
        if any(_norm(row.get("category", "")) == _norm(category) for row in self.tolerance_rule_rows): self.status.set(f"A {category} rule already exists — edit it instead."); return
        self.tolerance_rule_rows.append(normalize_optimization_rule(category, {"tolerance_percent": tolerance_percent, "priority": "Demographic", "origin": "manual", "note": "Quick-added preset"})); self._refresh_rules_tree(); self._mark_manual_edit("optimization")
    def _refresh_rules_tree(self) -> None:
        if not hasattr(self, "rules_tree"): return
        self.rules_tree.delete(*self.rules_tree.get_children())
        for index, row in enumerate(self.tolerance_rule_rows):
            normalized = normalize_optimization_rule(row.get("category", ""), row)
            applies = "All groups" if normalized["applies_to"] == "all" else f"Only {normalized['group']}"
            self.rules_tree.insert("", "end", iid=str(index), values=(normalized["category"], applies, normalized["rule"], normalized["target_n"], normalized["tolerance_percent"], normalized["tolerance_n"], normalized["priority"], ORIGIN_LABELS.get(normalized["origin"], normalized["origin"]), normalized["note"]))
        self._update_glance()

    # ------------------------------------------------------------------ payload
    def _reporting_variables(self) -> dict[str, dict]:
        result = {}
        for line in _lines(self.reporting_variables_text):
            parts = [part.strip() for part in line.split("|", 4)]
            if len(parts) < 4: raise ValueError(f"Reporting variable needs 'Name | canonical | field | display | levels': {line}")
            levels = []
            for item in parts[4].split(",") if len(parts) > 4 else []:
                item = item.strip()
                if "=" in item: value, label = (piece.strip() for piece in item.split("=", 1)); levels.append({"value": value, "label": label})
                elif item: levels.append(item)
            result[parts[0]] = {"canonical": parts[1], "field": parts[2], "display": parts[3], "levels": levels}
        return result

    def _osep_mappings(self) -> dict[str, list[str]]:
        result: dict[str, list[str]] = {}
        for line in _lines(self.osep_mappings):
            group, _, values = line.partition(":")
            items = [value.strip() for value in re.split(r"[,;]", values or group) if value.strip()]
            if items: result.setdefault((group.strip().casefold() or "osep") if values else "osep", []).extend(items)
        return result

    @staticmethod
    def _workflow_input_lines(workflow_inputs: dict[str, list[dict]]) -> list[str]:
        return [" | ".join((workflow, str(item.get("role", "")), str(item.get("group", "INPUTS")), str(item.get("status", "required")), str(item.get("kind", "file")), ",".join(str(v) for v in item.get("accepted_types", [])), str(item.get("adapter", "canonical_table")), str(item.get("label", "")), str(item.get("purpose", "")))) for workflow, items in workflow_inputs.items() for item in items]

    def _workflow_inputs(self) -> dict[str, list[dict]]:
        capabilities = {key: variable.get() for key, variable in self.capabilities.items()}; parsed = {}
        for line in _lines(self.workflow_inputs_text):
            parts = [part.strip() for part in line.split("|", 8)]
            if len(parts) != 9: raise ValueError(f"Task input needs 9 pipe-separated values: {line}")
            workflow, role, group, status, kind, accepted, adapter, label, purpose = parts
            if status.casefold() not in {"required", "optional", "conditional"}: raise ValueError(f"Task input status is invalid: {line}")
            if kind.casefold() not in {"file", "folder"}: raise ValueError(f"Task input kind is invalid: {line}")
            parsed.setdefault(workflow, []).append({"role": role, "group": group or "INPUTS", "label": label, "status": status.casefold(), "kind": kind.casefold(), "accepted_types": [v.strip() for v in accepted.split(",") if v.strip()], "purpose": purpose, "adapter": adapter or "canonical_table"})
        return normalized_workflow_inputs(parsed, capabilities)

    def _payload(self, status: str | None = None) -> dict:
        profile_id = re.sub(r"[^A-Za-z0-9_]+", "_", self.profile_id.get().strip() or f"{self.name.get()}_{self.market.get()}").strip("_").upper()
        if not profile_id or not self.name.get().strip(): raise ValueError("Assessment name and Profile ID are required.")
        try: target_n = int(self.default_target_n.get().strip()) if self.default_target_n.get().strip() else None
        except ValueError as exc: raise ValueError("Default target N must be a whole number or blank.") from exc
        recruitment = {}
        for line in _lines(self.recruitment_specs):
            parts = [part.strip() for part in line.split("|", 2)]
            if len(parts) < 2: raise ValueError(f"Recruitment report needs 'Study Type | template path | tables': {line}")
            recruitment[parts[0]] = {"template": parts[1], "required_tables": [v.strip() for v in parts[2].split(",") if v.strip()] if len(parts) > 2 else []}
        resources = {role: var.get().strip() for role, var in self.resource_vars.items() if var.get().strip()}
        census_sources = [item for item in self.target_sources if str(item.get("family", "")).casefold() == "census" and str(item.get("path", "")).strip()]
        if census_sources and "census_targets" not in resources: resources["census_targets"] = str(census_sources[0]["path"])
        reporting_variables = self._reporting_variables(); mapping = {semantic: variable.get().strip() for semantic, variable in self.mapping_vars.items() if variable.get().strip()}
        filters = [normalize_sample_filter(item) for item in self.sample_filter_rules]; extras = _lines(self.eligibility_extra)
        eligibility_rules = [format_sample_filter(item) for item in filters] + extras
        data_spec = dict(self.base_payload.get("data_specification") or {}); data_spec.update({"variable_map": mapping, "id_variable": mapping.get("teid") or mapping.get("id"), "forms": _lines(self.forms), "age_structure": _lines(self.age_structure), "languages": _lines(self.languages), "study_types": _lines(self.study_types), "demographics": _lines(self.demographics), "test_ids": _lines(self.test_ids), "missing_values": _lines(self.missing_values), "reporting_variables": reporting_variables, "accepted_case_data_types": [".xlsx", ".sas7bdat"]})
        capabilities = {key: variable.get() for key, variable in self.capabilities.items()}
        engine_config = self.engine_config.get().strip() or ("generic_cp_sat" if capabilities.get("sample_optimization") else "")
        families = [family for family, variable in self.family_vars.items() if variable.get()]
        sampling = dict(self.base_payload.get("sampling_specification") or {}); sampling.update({"engine_config": engine_config, "default_target_n": target_n, "target_families": families, "eligibility_rules": eligibility_rules, "eligibility_notes": extras, "default_locks": _lines(self.default_locks), "seeding_rules": _lines(self.rules)})
        for key, value in (("sample_filters", filters), ("target_tolerances", self._tolerances_dict()), ("target_interpretation", self._interpretation()), ("target_sources", [dict(item) for item in self.target_sources])):
            if value: sampling[key] = value
            else: sampling.pop(key, None)
        reporting = dict(self.base_payload.get("reporting_specification") or self.base_payload.get("reporting") or {}); required = dict(reporting.get("required_tables") or {}); required.update({"demographics": _lines(self.demographic_tables), "demo_cross_tabs": _lines(self.cross_tabs), "demo_three_way": _lines(self.three_way_tables), "osep": _lines(self.osep_tables), "clinical": _lines(self.clinical_tables)}); reporting.update({"required_tables": required, "comparison_columns": ["Target N", "Sample N", "Target %", "Sample %"], "recruitment": recruitment, "variables": reporting_variables}); reporting.pop("diff_formula", None)
        family_keys = {value.casefold() for value in families}
        optimization_outputs = ["Final Selected TEIDs", "Available Case Pool", "Demographic / Target Adherence Tables"]
        if "osep" in family_keys: optimization_outputs.append("OSEP Tables")
        if "clinical" in family_keys: optimization_outputs.append("Clinical Tables")
        optimization_outputs.append("QA / Constraint Summary")
        output = dict(self.base_payload.get("output_specification") or {}); deliverables = dict(output.get("workflow_deliverables") or {}); defaults = {"recruitment_planning": [{"title": "Field Research Report", "kind": "workbook"}], "sample_optimization": [{"title": v, "kind": "workbook"} for v in optimization_outputs], "targeted_swap": [{"title": v, "kind": "workbook", "description": "CP-SAT NOT RERUN"} for v in ("Updated Selected Cases TEID List", "Updated Available Case Pool", "Swap Log", "Refreshed Demographic / OSEP / Clinical Tables")]}
        for workflow, items in defaults.items():
            if workflow == "sample_optimization" and engine_config.casefold() in {"generic_cp_sat", "generic-cp-sat", "target_workbook_cp_sat"}:
                deliverables[workflow] = items
            else:
                deliverables.setdefault(workflow, items)
        output.update({"required_resources": list(output.get("required_resources") or []), "workflow_deliverables": deliverables, "project_output_choices": ["teid_lists", "tables", "both"]})
        provenance = dict(self.base_payload.get("provenance") or {})
        for key, variable in self.provenance_vars.items():
            if variable.get().strip(): old = provenance.get(key, {}); provenance[key] = {**(old if isinstance(old, dict) else {}), "source": variable.get().strip()}
        if self.provenance_vars["overview"].get().strip(): old = provenance.get("identity", {}); provenance["identity"] = {**(old if isinstance(old, dict) else {}), "source": self.provenance_vars["overview"].get().strip()}
        provenance["field_origins"] = {key: value for key, value in self.field_origins.items() if value in ORIGIN_LABELS}
        adaptation = dict(self.base_payload.get("adaptation") or {}); adaptation["review"] = dict(self.adaptation_review)
        osep_mappings = self._osep_mappings()
        payload = deepcopy(self.base_payload); payload.update({"profile_schema_version": 3, "id": profile_id, "display_name": self.name.get().strip(), "edition": self.edition.get().strip(), "version": self.version.get().strip() or "0.1", "market": self.market.get().strip(), "profile_kind": self.profile_kind.get() or "standard", "description": self.description.get("1.0", "end").strip(), "status": status or str(self.base_payload.get("status", "Draft")), "catalog_visibility": "user", "creation_method": self.creation_method.get(), "forms": _lines(self.forms), "demographics": _lines(self.demographics), "test_ids": _lines(self.test_ids), "languages": _lines(self.languages), "study_types": _lines(self.study_types), "age_structure": _lines(self.age_structure), "eligibility_rules": eligibility_rules, "target_families": families, "default_target_n": target_n, "source_config": engine_config, "schema_mapping": mapping, "reporting_variables": reporting_variables, "capabilities": capabilities, "workflow_inputs": self._workflow_inputs(), "data_specification": data_spec, "sampling_specification": sampling, "reporting_specification": reporting, "output_specification": output, "resources": resources, "source_documents": deepcopy(self.source_documents), "provenance": provenance, "adaptation": adaptation})
        if osep_mappings: payload["osep_mappings"] = osep_mappings
        profile, checks = validate_profile_data(payload, self.current_descriptor.parent if self.current_descriptor else ROOT)
        if profile is None: raise ValueError(checks[0][2])
        readiness = workflow_readiness(profile); payload["workflow_readiness"] = readiness; payload["missing"] = {key: list(value.get("missing", [])) for key, value in readiness.items() if value.get("missing")}; payload["runnable"] = any(value["status"] == "READY" for key, value in readiness.items() if key != "reporting")
        return payload

    # ------------------------------------------------------------------ validate
    def _validate(self) -> bool:
        try: payload = self._payload()
        except Exception as exc: messagebox.showerror("Validate Profile", str(exc)); return False
        profile, checks = validate_profile_data(payload, self.current_descriptor.parent if self.current_descriptor else ROOT)
        if profile is None: messagebox.showerror("Validate Profile", checks[0][2]); return False
        summary = profile_validation_summary(profile); self.last_summary = summary; self._render_validation(summary); core_valid = bool(summary["core_valid"]); reviews = sum(item["status"] == "NEEDS REVIEW" for item in summary["sections"]); workflow_reviews = sum(item["status"] == "NEEDS REVIEW" for item in summary["workflows"].values()); self.mark_button.configure(state="normal" if core_valid else "disabled")
        stored = str(self.base_payload.get("status", payload.get("status", "Draft"))); prefix = "PROFILE VALIDATED ✓" if core_valid and stored.casefold() == "validated" else "PROFILE VALID ✓" if core_valid else "PROFILE INVALID ✕"; self.status.set(f"{prefix} — {reviews} section(s) and {workflow_reviews} workflow(s) need review. Finish / Install remains available for drafts and Needs Review profiles." if core_valid else "Profile identity or schema is corrupted. Open the indicated section."); self._show_step("validate"); self._update_step_icons(summary); return core_valid

    def _render_validation(self, summary: dict) -> None:
        self.validation_tree.delete(*self.validation_tree.get_children()); symbols = {"COMPLETE": "✓ Complete", "NEEDS REVIEW": "⚠ Needs Review", "NOT CONFIGURED": "— Not Configured", "NOT APPLICABLE": "— Not Applicable", "ERROR": "✕ Error"}; actions = {"locale": "Review Overview", "structure": "Configure Demographics", "age_structure": "Configure Demographics", "mappings": "Configure Demographics", "demographics": "Configure Demographics", "clinical_osep": "Configure Clinical/OSEP", "eligibility": "Edit Sample Rules", "targets": "Configure Targets", "reporting": "Configure Output Tables", "sources": "Add Source"}
        for item in summary["sections"]:
            state = item["status"]; self.validation_tree.insert("", "end", iid=item["key"], values=(symbols.get(state, state), item["label"], item["detail"], actions.get(item["key"], "Review")), tags=(state.casefold().replace(" ", "_"),))
        for tag, color in (("complete", SUCCESS), ("needs_review", WARNING), ("error", ERROR), ("not_configured", MUTED), ("not_applicable", MUTED)): self.validation_tree.tag_configure(tag, foreground=color)
        self.workflow_tree.delete(*self.workflow_tree.get_children())
        for key, item in summary["workflows"].items(): self.workflow_tree.insert("", "end", values=(WORKFLOW_LABELS.get(key, key), item["status"], item["detail"]), tags=(item["status"].casefold().replace(" ", "_"),))
        for tag, color in (("ready", SUCCESS), ("needs_review", WARNING), ("not_configured", MUTED)): self.workflow_tree.tag_configure(tag, foreground=color)

    def _open_validation_section(self, _event=None) -> None:
        selected = self.validation_tree.selection()
        if selected: self._show_step(SECTION_STEP.get(selected[0], "overview"))
    def _tab_changed(self, _event=None) -> None:
        states = [self.adaptation_review.get(key, "") for key in STEP_REVIEW_KEYS.get(self.current_step, ())]
        self.section_review_status.set("⚠ Inherited information requires confirmation for this adaptation." if "review" in states else "✓ This section has been confirmed." if states and all(state == "confirmed" for state in states) else "— This section is marked not applicable." if states and all(state == "not_applicable" for state in states) else "")
    def _confirm_current_section(self) -> None:
        key = self.current_step
        if key == "validate": return
        for value in STEP_REVIEW_KEYS.get(key, ()): self.adaptation_review[value] = "confirmed"
        variable = self.provenance_vars.get(key)
        if variable and not variable.get().strip(): variable.set("Manual confirmation")
        if key in self.origin_vars and not self.origin_vars[key].get(): self._origin_chosen(key, "manual")
        self._tab_changed(); self._update_step_icons(); self.status.set(f"{STEP_TITLES[key]} confirmed.")
    def _not_applicable_current_section(self) -> None:
        key = self.current_step
        if key in {"overview", "validate"}: messagebox.showinfo("Not Applicable", "Overview and validation cannot be marked Not Applicable."); return
        for value in STEP_REVIEW_KEYS.get(key, ()): self.adaptation_review[value] = "not_applicable"
        self._tab_changed(); self._update_step_icons()
    def _update_step_icons(self, summary: dict | None = None) -> None:
        if summary is None:
            try: profile, _ = validate_profile_data(self._payload(), self.current_descriptor.parent if self.current_descriptor else ROOT); summary = profile_validation_summary(profile) if profile else {}
            except Exception: summary = {}
        states = {item["key"]: item["status"] for item in summary.get("sections", [])}
        for key in STEPS:
            related = STEP_REVIEW_KEYS.get(key, ()); values = [states.get(item) for item in related if states.get(item)]; icon = "✕" if "ERROR" in values else "⚠" if "NEEDS REVIEW" in values or any(self.adaptation_review.get(item) == "review" for item in related) else "✓" if values and all(v in {"COMPLETE", "NOT APPLICABLE"} for v in values) else "—"; self.step_buttons[key].configure(text=f"{icon}  {STEP_TITLES[key]}")

    # ------------------------------------------------------------------ save/install/export
    def _mark_validated(self) -> None:
        if not self._validate(): return
        self.base_payload = self._payload("Validated"); self.status.set("Profile marked Validated. Finish / Install to save it to the catalog.")
    def _resource_sources(self) -> dict[str, Path]:
        sources = {}; base = self.current_descriptor.parent if self.current_descriptor else Path(self.resource_folder.get() or Path.cwd())
        for role, variable in self.resource_vars.items():
            value = variable.get().strip()
            if not value: continue
            candidate = Path(value)
            if not candidate.is_absolute(): local, root_candidate = base / candidate, ROOT / candidate; candidate = local if local.exists() or not root_candidate.exists() else root_candidate
            if candidate.is_file(): sources[role] = candidate
        return sources
    def _save_draft(self) -> None:
        try: payload = self._payload("Draft"); profile = install_profile_package(payload, self._resource_sources())
        except Exception as exc: messagebox.showerror("Save Draft", str(exc)); return
        self.current_descriptor = USER_PROFILES_DIR / profile.id / "profile.yaml"; self.base_payload = payload; self.status.set(f"Draft saved. It will reopen from {self.current_descriptor}"); self._refresh_catalog()
    def _install(self) -> None:
        if not self._validate(): return
        try:
            payload = self._payload(); profile, _ = validate_profile_data(payload, self.current_descriptor.parent if self.current_descriptor else ROOT); summary = profile_validation_summary(profile); existing = str(self.base_payload.get("status", "")); status = "Validated" if existing.casefold() == "validated" or not summary["needs_review"] else "Needs Review"; payload = self._payload(status); installed = install_profile_package(payload, self._resource_sources())
        except Exception as exc: messagebox.showerror("Finish / Install", str(exc)); return
        self.current_descriptor = USER_PROFILES_DIR / installed.id / "profile.yaml"; self.base_payload = payload; self._refresh_catalog(); self.status.set(f"Installed {installed.display_name} as {status}. Stored in {self.current_descriptor.parent}."); messagebox.showinfo("Profile Installed", f"{installed.display_name}\n\nStatus: {status}\nThe profile is now available in SCOPE.")
    def _export_profile(self) -> None:
        if not self.name.get().strip(): messagebox.showinfo("Export", "Open or create a profile first."); return
        try: payload = self._payload()
        except Exception as exc: messagebox.showerror("Export", str(exc)); return
        selected = filedialog.asksaveasfilename(title="Export SCOPE Profile", defaultextension=".yaml", initialfile=f"{payload['id']}_profile.yaml", filetypes=[("YAML", "*.yaml"), ("JSON", "*.json")])
        if selected:
            path = Path(selected); path.write_text(json.dumps(payload, indent=2) if path.suffix.lower() == ".json" else yaml.safe_dump(payload, sort_keys=False, allow_unicode=True), encoding="utf-8"); self.status.set(f"Profile exported: {path}")


if __name__ == "__main__":
    ProfileManager()
