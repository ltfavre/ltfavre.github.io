"""Polished SCOPE desktop shell. All calculations are delegated to scope.engine."""
from __future__ import annotations

import json
import os
from pathlib import Path
import queue
import re
import sys
import threading
import time
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk

import pandas as pd
from PIL import Image, ImageTk

ROOT = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scope.configure import ProjectConfig
from scope.canonical import build_target_specification, canonical_target_records, resolved_profile_input
from scope.profiles import describe_profile_optimization_rules, describe_target_interpretation, profile_sample_filters, profile_study_type_defaults
from scope.engine import run, validate
from scope.profiles import discover_profile_catalog
from scope.targeted_swap import default_lock_fields, normalize_lock_fields, stage_swap_selection


COLORS = {
    "background": "#F6F3EC", "surface": "#FFFFFF", "navy": "#123B63", "navy_hover": "#1B4D78",
    "text": "#233240", "muted": "#6E7B86", "border": "#D8D5CD", "success": "#3B8D68",
    "warning": "#B77A16", "error": "#B84444", "recruitment_planning": "#178E86",
    "sample_optimization": "#6950A1", "targeted_swap": "#2D73B9", "targeted_adjustment": "#B06C2F",
}

APP_MODE = os.environ.get("SCOPE_APP_MODE", "scope").strip().casefold()
ALL_WORKFLOWS = {
    "recruitment_planning": ("Sample Planning", "For Field Research", "Generate targeted recruitment Ns from a Research Plan."),
    "targeted_swap": ("Targeted Case Swaps", "For Psychometrics", "Replace specific selected cases and refresh selected/available membership."),
    "targeted_adjustment": ("Targeted OSEP / Clinical Adjustments", "For Psychometrics", "Make specific OSEP or clinical-group membership adjustments."),
    "sample_optimization": ("Run CP-SAT for the Sample", "For Psychometrics", "Optimize the selected sample against profile-defined targets."),
}
WORKFLOWS = ({"recruitment_planning": ALL_WORKFLOWS["recruitment_planning"]}
             if APP_MODE == "sample_planning"
             else {key: ALL_WORKFLOWS[key] for key in ("targeted_swap", "targeted_adjustment", "sample_optimization")})

NAV_ITEMS = (("Test Profile", "Planning Setup", "Review", "Run", "Results")
             if APP_MODE == "sample_planning"
             else ("Sample Data & Task", "Operation Setup", "Review", "Output", "Results"))
ASSET_DIR = ROOT / "gui" / "assets"


class ScrollableFrame(tk.Frame):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, bg=COLORS["background"], **kwargs)
        self.canvas = tk.Canvas(self, bg=COLORS["background"], highlightthickness=0)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.inner = tk.Frame(self.canvas, bg=COLORS["background"])
        self.window = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda _event: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda event: self.canvas.itemconfigure(self.window, width=event.width))
        self.canvas.configure(yscrollcommand=scrollbar.set)
        self.canvas.pack(side="left", fill="both", expand=True); scrollbar.pack(side="right", fill="y")


class RoundedCard(tk.Canvas):
    """Small canvas-backed card so Tkinter cards have genuinely rounded edges."""
    def __init__(self, parent, width: int, height: int, radius: int = 16):
        super().__init__(parent, width=width, height=height, bg=COLORS["background"], highlightthickness=0, bd=0)
        self.radius = radius; self.border = COLORS["border"]
        self.content = tk.Frame(self, bg=COLORS["surface"])
        self.content_window = self.create_window(9, 9, window=self.content, anchor="nw")
        self.bind("<Configure>", self._redraw)

    def _redraw(self, event=None) -> None:
        width = max(24, event.width if event else self.winfo_width()); height = max(24, event.height if event else self.winfo_height()); r = min(self.radius, width // 3, height // 3)
        self.delete("card_shape")
        points = [r, 1, width-r, 1, width-1, 1, width-1, r, width-1, height-r, width-1, height-1, width-r, height-1, r, height-1, 1, height-1, 1, height-r, 1, r, 1, 1]
        shape = self.create_polygon(points, smooth=True, splinesteps=24, fill=COLORS["surface"], outline=self.border, width=3 if self.border != COLORS["border"] else 1, tags="card_shape")
        self.tag_lower(shape); self.coords(self.content_window, 9, 9); self.itemconfigure(self.content_window, width=max(1, width-18), height=max(1, height-18))

    def set_selected(self, color: str | None) -> None:
        self.border = color or COLORS["border"]; self._redraw()


class ToolTip:
    """Tiny accessible tooltip for compact toolbar actions."""
    def __init__(self, widget, text: str):
        self.widget, self.text, self.window = widget, text, None
        widget.bind("<Enter>", self._show, add="+"); widget.bind("<Leave>", self._hide, add="+")

    def _show(self, _event=None) -> None:
        if self.window or not self.text: return
        self.window = tk.Toplevel(self.widget); self.window.wm_overrideredirect(True)
        self.window.wm_geometry(f"+{self.widget.winfo_rootx()}+{self.widget.winfo_rooty() + self.widget.winfo_height() + 5}")
        tk.Label(self.window, text=self.text, bg="#FFF8D8", fg=COLORS["text"], relief="solid", bd=1, padx=7, pady=4, font=("Segoe UI", 8)).pack()

    def _hide(self, _event=None) -> None:
        if self.window: self.window.destroy(); self.window = None


class SCOPEApp:
    def __init__(self) -> None:
        self.root = tk.Tk(); self.root.title("SCOPE Sample Planning" if APP_MODE == "sample_planning" else "SCOPE"); self.root.geometry("1180x860"); self.root.minsize(1000, 740); self.root.configure(bg=COLORS["background"])
        self._suspend_dirty = True; self.project_dirty = False; self.current_project_file: Path | None = None
        self.events: queue.Queue = queue.Queue(); self.busy = False; self.validation_passed = False; self.last_result = None
        self.run_history: list[dict] = []
        self.workflow = tk.StringVar(); self.profile_id = tk.StringVar(); self.completed: set[int] = set(); self.unlocked_step = 0; self.active_step = 0
        self.swap_rows: list[dict] = []; self.form_vars: dict[str, tk.BooleanVar] = {}; self.demographic_vars: dict[str, tk.BooleanVar] = {}; self.lock_vars: dict[str, tk.BooleanVar] = {}
        self.report_table_lists: dict[str, tk.Listbox] = {}; self.report_table_values: dict[str, list[str]] = {}
        self.profile_catalog = discover_profile_catalog(); self.profiles = self.profile_catalog.user_profiles; self.profile_errors = self.profile_catalog.errors
        self.profile_cards = {}; self.workflow_cards = {}; self.run_started = 0.0; self.progress_value = 0
        self.input_status_labels = {}; self.canonical_target_rows: list[dict] = []; self.show_additional_sources = False
        self.target_refresh_pending = False
        self.candidate_datasets: list[dict] = []; self.target_overrides: list[dict] = []
        self.technical_messages: list[str] = []; self.user_messages: list[str] = []; self.activity_expanded = False
        self.companion_name = ""; self.companion_frame = 0
        if not self.profiles and not self.profile_errors: raise RuntimeError("No SCOPE Test Profile descriptors were found.")
        self.animal_images = self._load_animal_images()
        self._vars(); self._styles(); self._build(); self._log_profile_catalog(); self._track_project_changes(); self._suspend_dirty = False; self._update_project_banner(); self.root.after(100, self._drain); self.root.after(600, self._animate_companion); self.root.mainloop()

    def _vars(self) -> None:
        default_project_dir = Path.home() / "Documents" / "SCOPE Projects"
        try:
            default_project_dir.mkdir(parents=True, exist_ok=True)
        except OSError:
            default_project_dir = Path.cwd()
        self.project_name = tk.StringVar(); self.project_dir = tk.StringVar(value=str(default_project_dir)); self.output_dir = tk.StringVar(value=str(default_project_dir / "scope_outputs")); self.sample_size = tk.StringVar()
        self.language = tk.StringVar(value="English"); self.study_type = tk.StringVar(value="STDZ")
        self.include_census = tk.BooleanVar(value=True); self.include_osep = tk.BooleanVar(value=True); self.include_clinical = tk.BooleanVar(value=True); self.clinical_seeding = tk.BooleanVar(value=False)
        self.dataset = tk.StringVar(); self.master_file = tk.StringVar(); self.scored_dir = tk.StringVar(); self.case_swap_dir = tk.StringVar(); self.target_file = tk.StringVar(); self.fr_report_file = tk.StringVar()
        self.recruitment_plan_file = tk.StringVar(); self.recruitment_actuals_file = tk.StringVar()
        self.swap_sample = tk.StringVar(); self.swap_available = tk.StringVar(); self.swap_universe = tk.StringVar()
        self.swap_sample_data = tk.StringVar(); self.swap_available_data = tk.StringVar()
        self.swap_remove = tk.StringVar(); self.swap_add = tk.StringVar(); self.swap_reason = tk.StringVar(); self.swap_selection_summary = tk.StringVar(value="0 swap selections added")
        self.status = tk.StringVar(value="Choose what SCOPE should do."); self.companion = tk.StringVar(value="")
        self.project_banner = tk.StringVar(value="Project: Untitled  •  Unsaved changes")
        self.elapsed = tk.StringVar(value="Elapsed 00:00"); self.result_summary = tk.StringVar(value="No completed run yet.")
        self.targets_confirmed = tk.BooleanVar(value=False); self.target_review_summary = tk.StringVar(value="Effective targets are built automatically after validation.")
        self.use_existing_sample = tk.BooleanVar(value=False); self.override_age_defaults = tk.BooleanVar(value=False)
        self.output_selection = tk.StringVar(value="both")

    def _load_animal_images(self) -> dict[str, ImageTk.PhotoImage]:
        images = {}
        for name in (
            "kitty", "kitty_motion", "bunny", "bunny_motion", "puppy", "puppy_motion",
            "guinea_pig", "guinea_pig_motion", "success_party", "attention_party",
        ):
            path = ASSET_DIR / f"{name}.png"
            if not path.is_file():
                continue
            image = Image.open(path).convert("RGBA")
            image.thumbnail((68 if "party" in name else 42, 34), Image.Resampling.LANCZOS)
            images[name] = ImageTk.PhotoImage(image)
        return images

    def _styles(self) -> None:
        style = ttk.Style(self.root)
        if "clam" in style.theme_names(): style.theme_use("clam")
        style.configure("TFrame", background=COLORS["background"]); style.configure("TLabel", background=COLORS["background"], foreground=COLORS["text"], font=("Segoe UI", 9))
        style.configure("TLabelframe", background=COLORS["background"], bordercolor=COLORS["border"]); style.configure("TLabelframe.Label", background=COLORS["background"], foreground=COLORS["navy"], font=("Segoe UI", 10, "bold"))
        style.configure("Primary.TButton", background=COLORS["navy"], foreground="white", padding=(15, 9), font=("Segoe UI", 10, "bold"), borderwidth=0)
        style.map("Primary.TButton", background=[("active", COLORS["navy_hover"]), ("disabled", "#AAB4BC")])
        style.configure("Secondary.TButton", background=COLORS["surface"], foreground=COLORS["navy"], padding=(11, 7), bordercolor=COLORS["border"])
        style.configure("Run.Horizontal.TProgressbar", troughcolor="#E7E3DC", background=COLORS["navy"], bordercolor="#E7E3DC", lightcolor=COLORS["navy"], darkcolor=COLORS["navy"])
        style.configure("Treeview", rowheight=28, background=COLORS["surface"], fieldbackground=COLORS["surface"], bordercolor=COLORS["border"])
        style.configure("Treeview.Heading", background="#E8EDF1", foreground=COLORS["navy"], font=("Segoe UI", 9, "bold"))

    def _build(self) -> None:
        header = tk.Frame(self.root, bg=COLORS["navy"], height=88); header.pack(fill="x"); header.pack_propagate(False)
        title = tk.Frame(header, bg=COLORS["navy"]); title.pack(side="left", padx=20, pady=(9, 8))
        tk.Label(title, text="SCOPE SAMPLE PLANNING" if APP_MODE == "sample_planning" else "SCOPE", bg=COLORS["navy"], fg="white", font=("Segoe UI", 21, "bold")).pack(anchor="w")
        subtitle = "Recruitment and sample planning" if APP_MODE == "sample_planning" else "Sample Configuration, Optimization, and Profiling Engine"
        tk.Label(title, text=subtitle, bg=COLORS["navy"], fg="#D8E5EF", font=("Segoe UI", 9)).pack(anchor="w")
        toolbar = tk.Frame(header, bg=COLORS["navy"]); toolbar.pack(side="right", padx=18)
        self.new_project_button = self._header_button(toolbar, "New Project", self._new_project); self.new_project_button.pack(side="left", padx=3)
        self.save_project_button = self._header_button(toolbar, "Save Project", self._save_project); self.save_project_button.pack(side="left", padx=3)
        self.open_project_button = self._header_button(toolbar, "Open Project", self._open_project); self.open_project_button.pack(side="left", padx=3)
        self.recent_button = tk.Menubutton(toolbar, text="Recent Projects ▾", bg=COLORS["navy_hover"], fg="white", activebackground="#285D86", activeforeground="white", relief="flat", padx=10, pady=6, font=("Segoe UI", 9))
        self.recent_menu = tk.Menu(self.recent_button, tearoff=False); self.recent_button.configure(menu=self.recent_menu); self.recent_button.pack(side="left", padx=3); self._refresh_recent_menu()
        ToolTip(self.new_project_button, "Start a new SCOPE project configuration. Test Profiles are not changed.")
        ToolTip(self.save_project_button, "Save workflow choices, paths, targets, and settings to a .scope project file.")
        ToolTip(self.open_project_button, "Open a saved .scope project without running it.")
        ToolTip(self.recent_button, "Open a recently used SCOPE project configuration.")
        tk.Label(header, textvariable=self.project_banner, bg=COLORS["navy"], fg="#D8E5EF", font=("Segoe UI", 8), anchor="e").place(relx=1.0, x=-20, y=62, anchor="ne")

        body = tk.Frame(self.root, bg=COLORS["background"]); body.pack(fill="both", expand=True)
        nav = tk.Frame(body, bg=COLORS["navy"], width=205); nav.pack(side="left", fill="y"); nav.pack_propagate(False)
        tk.Label(nav, text="PROJECT WORKFLOW", bg=COLORS["navy"], fg="#AFC5D6", font=("Segoe UI", 8, "bold")).pack(anchor="w", padx=16, pady=(18, 8))
        self.nav_buttons = []
        for index, name in enumerate(NAV_ITEMS):
            button = tk.Button(nav, text=f"{index + 1}.  {name}", anchor="w", bg=COLORS["navy"], fg="white", activebackground=COLORS["navy_hover"], activeforeground="white", disabledforeground="#7890A3", relief="flat", bd=0, padx=16, pady=12, font=("Segoe UI", 10), command=lambda value=index: self._show_step(value))
            button.pack(fill="x"); self.nav_buttons.append(button)
        tk.Label(nav, text=" → ".join(NAV_ITEMS[:3]) + "\n" + " → ".join(NAV_ITEMS[3:]), bg=COLORS["navy"], fg="#89A7BD", wraplength=170, justify="left", font=("Segoe UI", 8)).pack(side="bottom", anchor="w", padx=16, pady=18)

        self.page_host = tk.Frame(body, bg=COLORS["background"]); self.page_host.pack(side="left", fill="both", expand=True, padx=22, pady=15)
        self.pages = [tk.Frame(self.page_host, bg=COLORS["background"]) for _ in NAV_ITEMS]
        for page in self.pages: page.grid(row=0, column=0, sticky="nsew")
        self.page_host.grid_rowconfigure(0, weight=1); self.page_host.grid_columnconfigure(0, weight=1)
        self._profile_page(self.pages[0]); self._project_page(self.pages[1]); self._validate_page(self.pages[2]); self._run_page(self.pages[3]); self._results_page(self.pages[4])

        self.activity = tk.Frame(self.root, bg=COLORS["background"])
        activity_header = tk.Frame(self.activity, bg=COLORS["background"]); activity_header.pack(fill="x")
        self.activity_toggle = tk.Button(activity_header, text="▸ Activity", command=self._toggle_activity, bg=COLORS["background"], fg=COLORS["navy"], activebackground=COLORS["background"], relief="flat", bd=0, font=("Segoe UI", 9, "bold")); self.activity_toggle.pack(side="left")
        ttk.Button(activity_header, text="Open Technical Log", style="Secondary.TButton", command=self._open_technical_log).pack(side="right")
        self.activity_content = tk.Frame(self.activity, bg=COLORS["background"])
        self.log = scrolledtext.ScrolledText(self.activity_content, height=4, state="disabled", wrap="word", bg="#FBFAF7", fg=COLORS["text"], relief="solid", bd=1, font=("Consolas", 8)); self.log.pack(fill="x", pady=(2, 0))
        self.statusbar = tk.Frame(self.root, bg="#EAE7E0", height=50); self.statusbar.pack_propagate(False)
        self.animal_label = tk.Label(self.statusbar, text="", bg="#EAE7E0", width=72); self.animal_label.pack(side="left", padx=(10, 6))
        status_text = tk.Frame(self.statusbar, bg="#EAE7E0"); status_text.pack(side="left", fill="both", expand=True, pady=5)
        tk.Label(status_text, textvariable=self.status, bg="#EAE7E0", fg=COLORS["navy"], font=("Segoe UI", 10, "bold"), anchor="w").pack(fill="x")
        tk.Label(status_text, textvariable=self.companion, bg="#EAE7E0", fg=COLORS["muted"], font=("Segoe UI", 9), anchor="w").pack(fill="x")
        progress_side = tk.Frame(self.statusbar, bg="#EAE7E0"); progress_side.pack(side="right", padx=14, pady=6)
        self.progress = ttk.Progressbar(progress_side, style="Run.Horizontal.TProgressbar", maximum=100, length=260); self.progress.pack()
        tk.Label(progress_side, textvariable=self.elapsed, bg="#EAE7E0", fg=COLORS["muted"], font=("Segoe UI", 8)).pack(anchor="e", pady=(3, 0))
        self._show_step(0, force=True)

    def _header_button(self, parent, text, command):
        return tk.Button(parent, text=text, command=command, bg=COLORS["navy_hover"], fg="white", activebackground="#285D86", activeforeground="white", relief="flat", bd=0, padx=10, pady=6, font=("Segoe UI", 9))

    def _page_title(self, page, step, title, subtitle):
        tk.Label(page, text=f"STEP {step}", bg=COLORS["background"], fg=COLORS["muted"], font=("Segoe UI", 8, "bold")).pack(anchor="w")
        tk.Label(page, text=title, bg=COLORS["background"], fg=COLORS["navy"], font=("Segoe UI", 19, "bold")).pack(anchor="w", pady=(2, 2))
        tk.Label(page, text=subtitle, bg=COLORS["background"], fg=COLORS["muted"], wraplength=850, justify="left", font=("Segoe UI", 10)).pack(anchor="w", pady=(0, 14))

    def _bottom_navigation(self, page, back_index: int | None, back_label: str, continue_label: str = "", continue_command=None):
        bar = tk.Frame(page, bg=COLORS["background"]); bar.pack(side="bottom", fill="x", pady=(10, 0))
        back = ttk.Button(bar, text=f"←  {back_label}", style="Secondary.TButton", command=(lambda: self._show_step(back_index)) if back_index is not None else None)
        back.pack(side="left"); back.configure(state="normal" if back_index is not None else "disabled")
        forward = None
        if continue_label:
            forward = ttk.Button(bar, text=f"{continue_label}  →", style="Primary.TButton", command=continue_command)
            forward.pack(side="right")
        return back, forward

    def _workflow_page(self, page) -> None:
        self._bottom_navigation(page, None, "Back", "Continue to Test Profile", self._workflow_continue)
        self._page_title(page, 1, "Choose what SCOPE should do", "Start with the work you need to accomplish. SCOPE will show only the relevant project controls.")
        cards = tk.Frame(page, bg=COLORS["background"]); cards.pack(fill="x", pady=8)
        for column, (key, (name, _audience, description)) in enumerate(WORKFLOWS.items()):
            accent = COLORS[key]; card = RoundedCard(cards, width=270, height=220); card.grid(row=0, column=column, sticky="nsew", padx=(0 if column == 0 else 8, 0)); cards.grid_columnconfigure(column, weight=1); panel = card.content
            tk.Frame(panel, bg=accent, height=7).pack(fill="x")
            icon = {"recruitment_planning": "◎", "sample_optimization": "◇", "targeted_swap": "⇄", "targeted_adjustment": "◫"}[key]
            tk.Label(panel, text=icon, bg=COLORS["surface"], fg=accent, font=("Segoe UI Symbol", 26)).pack(anchor="w", padx=16, pady=(12, 3))
            tk.Label(panel, text=name.upper(), bg=COLORS["surface"], fg=COLORS["navy"], font=("Segoe UI", 11, "bold"), wraplength=235, justify="left").pack(anchor="w", padx=16)
            tk.Label(panel, text=description, bg=COLORS["surface"], fg=COLORS["text"], font=("Segoe UI", 9), wraplength=235, justify="left").pack(anchor="w", padx=16, pady=(7, 0))
            if key == "sample_optimization":
                tk.Label(panel, text="Optional clinical seeding available.", bg=COLORS["surface"], fg=accent, font=("Segoe UI", 8, "bold")).pack(anchor="w", padx=16, pady=(5, 0))
            choose = tk.Button(panel, text="Choose", command=lambda value=key: self._choose_workflow(value), bg=accent, fg="white", activebackground=accent, activeforeground="white", relief="flat", padx=13, pady=5, font=("Segoe UI", 9, "bold")); choose.pack(anchor="w", padx=16, pady=10)
            self.workflow_cards[key] = card

    def _profile_page(self, page) -> None:
        scroll = ScrollableFrame(page); scroll.pack(fill="both", expand=True); content = scroll.inner
        if APP_MODE == "sample_planning":
            title, subtitle = "Select Test Profile", "Choose the assessment for this sample-planning project."
        else:
            title, subtitle = "Enter sample data and choose a task", "Start with the participant-level data to update, then choose the operation SCOPE should perform."
        self._page_title(content, 1, title, subtitle)
        self.profile_grid = tk.Frame(content, bg=COLORS["background"]); self.profile_grid.pack(fill="x")
        self._build_profile_cards()
        catalog_actions = tk.Frame(content, bg=COLORS["background"]); catalog_actions.pack(fill="x", pady=(7, 0))
        tk.Label(catalog_actions, text="Need another assessment? Validated profiles are installed with the separate SCOPE Profile Manager.", bg=COLORS["background"], fg=COLORS["muted"], font=("Segoe UI", 8), anchor="w").pack(side="left", fill="x", expand=True)
        ttk.Button(catalog_actions, text="Refresh Installed Profiles", style="Secondary.TButton", command=self._refresh_installed_profiles).pack(side="right")
        self.profile_detail = tk.Label(content, text="Select a Test Profile to see its capabilities.", bg=COLORS["surface"], fg=COLORS["text"], justify="left", anchor="nw", wraplength=840, highlightbackground=COLORS["border"], highlightthickness=1, padx=14, pady=12)
        self.profile_detail.pack(fill="x", pady=(12, 8))
        if APP_MODE != "sample_planning":
            sample_data = ttk.LabelFrame(content, text="Sample data", padding=11); sample_data.pack(fill="x", pady=(5, 8))
            self._path_row(sample_data, "Participant-level case data", self.dataset, file=True)
            ttk.Label(sample_data, text="Accepted: .xlsx or .sas7bdat. This dataset supplies participant attributes used to validate changes and refresh tables.", foreground=COLORS["muted"]).pack(anchor="w", pady=(4, 0))
        tk.Label(content, text="WHAT WOULD YOU LIKE SCOPE TO DO?", bg=COLORS["background"], fg=COLORS["navy"], font=("Segoe UI", 11, "bold")).pack(anchor="w", pady=(8, 5))
        tk.Label(content, text="Choose an operation to continue. Operations become available after you select a Test Profile.", bg=COLORS["background"], fg=COLORS["muted"], font=("Segoe UI", 9)).pack(anchor="w", pady=(0, 6))
        cards = tk.Frame(content, bg=COLORS["background"]); cards.pack(fill="x", pady=(0, 12))
        self.workflow_buttons = {}
        for column, (key, (name, _audience, description)) in enumerate(WORKFLOWS.items()):
            accent = COLORS[key]; card = RoundedCard(cards, width=270, height=166); card.grid(row=0, column=column, sticky="nsew", padx=(0 if column == 0 else 8, 0)); cards.grid_columnconfigure(column, weight=1); panel = card.content
            tk.Frame(panel, bg=accent, height=6).pack(fill="x")
            tk.Label(panel, text=name.upper(), bg=COLORS["surface"], fg=COLORS["navy"], font=("Segoe UI", 10, "bold"), wraplength=235, justify="left").pack(anchor="w", padx=14, pady=(10, 3))
            tk.Label(panel, text=description, bg=COLORS["surface"], fg=COLORS["muted"], font=("Segoe UI", 8), wraplength=235, justify="left").pack(anchor="w", padx=14)
            button = tk.Button(panel, text=f"Choose {name}", command=lambda value=key: self._choose_workflow(value, navigate=True), bg=accent, fg="white", activebackground=accent, activeforeground="white", disabledforeground="#E5E5E5", relief="flat", padx=10, pady=5, font=("Segoe UI", 8, "bold")); button.pack(anchor="w", padx=14, pady=10)
            self.workflow_cards[key] = card; self.workflow_buttons[key] = button
        self._update_workflow_buttons()

    def _build_profile_cards(self) -> None:
        for child in self.profile_grid.winfo_children(): child.destroy()
        self.profile_cards = {}
        workflow = self.workflow.get()
        for index, profile in enumerate(self.profiles):
            row, column = divmod(index, 3); card = RoundedCard(self.profile_grid, width=245, height=174, radius=13); card.grid(row=row, column=column, sticky="nsew", padx=(0 if column == 0 else 8, 0), pady=(0 if row == 0 else 8, 0)); self.profile_grid.grid_columnconfigure(column, weight=1); panel = card.content
            muted = False
            normalized_status = profile.status.casefold(); status_icon = "✓" if normalized_status == "validated" else "⚠" if normalized_status == "draft" else "○"
            status_color = COLORS["success"] if normalized_status == "validated" else COLORS["warning"] if normalized_status == "draft" else COLORS["muted"]
            name = tk.Label(panel, text=profile.display_name, bg=COLORS["surface"], fg=COLORS["muted"] if muted else COLORS["navy"], font=("Segoe UI", 11, "bold"), anchor="w")
            name.pack(fill="x", padx=12, pady=(12, 5))
            status = tk.Label(panel, text=f"{status_icon} {profile.status}  ·  v{profile.version}", bg=COLORS["surface"], fg=status_color, font=("Segoe UI", 9), anchor="w")
            status.pack(fill="x", padx=12)
            modes = "\n".join(f"{'✓' if profile.supports(key) else '○'} {WORKFLOWS[key][0]}" for key in WORKFLOWS)
            mode_label = tk.Label(panel, text=modes, bg=COLORS["surface"], fg=COLORS["muted"] if muted else COLORS["text"], font=("Segoe UI", 8), anchor="w", justify="left"); mode_label.pack(fill="x", padx=12, pady=(5, 0))
            hint = tk.Label(panel, text="Select profile", bg=COLORS["surface"], fg=COLORS["muted"], font=("Segoe UI", 8), anchor="w"); hint.pack(fill="x", padx=12, pady=(5, 0))
            for widget in (card, panel, name, status, mode_label, hint): widget.bind("<Button-1>", lambda _event, value=profile.id: self._select_profile(value))
            if self.profile_id.get() == profile.id: card.set_selected(COLORS.get(workflow, COLORS["navy"]))
            self.profile_cards[profile.id] = card
        offset = len(self.profiles)
        for error_index, error in enumerate(self.profile_errors, start=offset):
            row, column = divmod(error_index, 3); card = RoundedCard(self.profile_grid, width=245, height=174, radius=13); card.grid(row=row, column=column, sticky="nsew", padx=(0 if column == 0 else 8, 0), pady=(0 if row == 0 else 8, 0)); panel = card.content
            tk.Label(panel, text=error.assessment_name, bg=COLORS["surface"], fg=COLORS["muted"], font=("Segoe UI", 11, "bold"), anchor="w").pack(fill="x", padx=12, pady=(12, 5))
            tk.Label(panel, text="⚠ Profile Error", bg=COLORS["surface"], fg=COLORS["error"], font=("Segoe UI", 9, "bold"), anchor="w").pack(fill="x", padx=12)
            detail = tk.Label(panel, text="This profile could not be loaded. Open the Technical Log for QA details.", bg=COLORS["surface"], fg=COLORS["muted"], wraplength=205, justify="left", anchor="w", font=("Segoe UI", 8)); detail.pack(fill="x", padx=12, pady=(8, 0)); ToolTip(detail, error.message)

    def _refresh_installed_profiles(self) -> None:
        selected = self.profile_id.get()
        self.profile_catalog = discover_profile_catalog(); self.profiles = self.profile_catalog.user_profiles; self.profile_errors = self.profile_catalog.errors
        if selected not in {profile.id for profile in self.profiles}: self.profile_id.set(""); self.workflow.set("")
        self._build_profile_cards(); self._refresh_profile_detail(); self._update_workflow_buttons()
        self.status.set(f"Profile catalog refreshed: {len(self.profiles)} user profile(s) available.")

    def _project_page(self, page) -> None:
        self._bottom_navigation(page, 0, NAV_ITEMS[0], "Continue to Review", self._project_continue)
        scroll = ScrollableFrame(page); scroll.pack(fill="both", expand=True); content = scroll.inner
        self._page_title(content, 2, "Set up this operation", "Provide the sources and choices needed for the operation selected on the first page.")
        self.project_frame = ttk.LabelFrame(content, text="Project", padding=11); self.project_frame.pack(fill="x", pady=4)
        name_row = ttk.Frame(self.project_frame); name_row.pack(fill="x", pady=2)
        ttk.Label(name_row, text="Project name:", width=28).pack(side="left"); ttk.Entry(name_row, textvariable=self.project_name).pack(side="left", fill="x", expand=True)
        self._path_row(self.project_frame, "Project folder", self.project_dir, folder=True)

        self.profile_resources_frame = ttk.LabelFrame(content, text="Profile resources", padding=11); self.profile_resources_frame.pack(fill="x", pady=4)
        self._render_profile_resources()

        self.scope_frame = ttk.LabelFrame(content, text="Project scope", padding=11); self.scope_frame.pack(fill="x", pady=4)
        scope_actions = ttk.Frame(self.scope_frame); scope_actions.pack(fill="x", pady=(0, 7))
        ttk.Label(scope_actions, text="Choose only what applies to this project.", foreground=COLORS["muted"]).pack(side="left")
        ttk.Button(scope_actions, text="Reset selections", style="Secondary.TButton", command=self._reset_project_scope).pack(side="right")
        top = ttk.Frame(self.scope_frame); top.pack(fill="x")
        ttk.Label(top, text="Target N:").pack(side="left"); ttk.Entry(top, textvariable=self.sample_size, width=10).pack(side="left", padx=(5, 16))
        ttk.Label(top, text="Language:").pack(side="left"); self.language_picker = ttk.Combobox(top, textvariable=self.language, state="readonly", width=14); self.language_picker.pack(side="left", padx=(5, 16))
        ttk.Label(top, text="Study type:").pack(side="left"); self.study_picker = ttk.Combobox(top, textvariable=self.study_type, state="readonly", width=18); self.study_picker.pack(side="left", padx=(5, 0))
        self.profile_rules_label = ttk.Label(self.scope_frame, text="", foreground=COLORS["muted"], wraplength=930, justify="left"); self.profile_rules_label.pack(fill="x", pady=(2, 0))
        self.forms_frame = ttk.Frame(self.scope_frame); self.forms_frame.pack(fill="x", pady=(8, 2))
        self.demographics_frame = ttk.Frame(self.scope_frame)
        population = ttk.Frame(self.scope_frame); population.pack(fill="x", pady=4); ttk.Label(population, text="Target families:").pack(side="left")
        for label, variable in (("Census demographics", self.include_census), ("OSEP", self.include_osep), ("Clinical", self.include_clinical)):
            ttk.Checkbutton(population, text=label, variable=variable).pack(side="left", padx=(8, 0))
        self.seeding_check = ttk.Checkbutton(population, text="Allow clinical/OSEP seeding", variable=self.clinical_seeding); self.seeding_check.pack(side="left", padx=(18, 0))
        self.report_tables_frame = ttk.LabelFrame(self.scope_frame, text="Demographic table output", padding=8)
        self.report_tables_frame.pack(fill="x", pady=(6, 2))
        override_row = ttk.Frame(self.scope_frame); override_row.pack(fill="x", pady=(4, 0))
        ttk.Checkbutton(override_row, text="Override Test Profile age-band defaults", variable=self.override_age_defaults, command=self._toggle_age_override).pack(anchor="w")
        self.age_frame = ttk.LabelFrame(self.scope_frame, text="Age bands and target Ns", padding=7); self.age_frame.pack(fill="x", pady=(5, 0))
        ttk.Label(self.age_frame, text="One per line, for example: 5:0-5:5 = 100. Leave blank to use the Test Profile.").pack(anchor="w")
        self.age_text = scrolledtext.ScrolledText(self.age_frame, height=3, wrap="word"); self.age_text.pack(fill="x", pady=(3, 0)); self.age_text.bind("<<Modified>>", self._age_text_changed); self.age_text.edit_modified(False)
        self.age_frame.pack_forget()

        self.workflow_inputs = ttk.LabelFrame(content, text="Project sources", padding=11); self.workflow_inputs.pack(fill="x", pady=4, before=self.scope_frame)
        self.output_frame = ttk.LabelFrame(content, text="Output location", padding=11); self.output_frame.pack(fill="x", pady=4)
        self._path_row(self.output_frame, "SCOPE run outputs", self.output_dir, folder=True)

    def _validate_page(self, page) -> None:
        _back, self.continue_run_button = self._bottom_navigation(page, 1, NAV_ITEMS[1], "Continue to Output", lambda: self._show_step(3))
        self.continue_run_button.configure(state="disabled")
        self._page_title(page, 3, "Review the operation", "Confirm that the sample data, membership sources, and requested changes are ready.")
        validation_actions = ttk.Frame(page); validation_actions.pack(fill="x")
        ttk.Button(validation_actions, text="VALIDATE PROJECT", style="Primary.TButton", command=self._validate).pack(side="left")
        self.fix_project_button = ttk.Button(validation_actions, text="Fix in Operation Setup", style="Secondary.TButton", command=lambda: self._show_step(1)); self.fix_project_button.pack(side="left", padx=(8, 0))
        self.validation_tree = ttk.Treeview(page, columns=("status", "check", "detail"), show="headings", height=8)
        for column, title, width in (("status", "Status", 110), ("check", "Check", 230), ("detail", "Result / correction needed", 600)):
            self.validation_tree.heading(column, text=title); self.validation_tree.column(column, width=width)
        self.validation_tree.pack(fill="both", expand=True, pady=12)
        # Review canonical targets through the user-facing Review Targets step.
        self.target_review_frame = ttk.LabelFrame(page, text="Review Targets", padding=9)
        ttk.Label(self.target_review_frame, textvariable=self.target_review_summary, foreground=COLORS["muted"]).pack(anchor="w", fill="x", pady=(0, 5))
        self.target_review_tree = ttk.Treeview(self.target_review_frame, columns=("category", "group", "rule", "target", "tolerance", "priority", "source"), show="headings", height=7)
        for column, title, width in (("category", "Category", 105), ("group", "Group", 190), ("rule", "Rule", 75), ("target", "Target", 85), ("tolerance", "Tolerance", 75), ("priority", "Priority", 80), ("source", "Source", 170)):
            self.target_review_tree.heading(column, text=title); self.target_review_tree.column(column, width=width, stretch=column in {"group", "source"})
        self.target_review_tree.tag_configure("override", foreground=COLORS["warning"])
        self.target_review_tree.pack(fill="x")
        target_actions = ttk.Frame(self.target_review_frame); target_actions.pack(fill="x", pady=(8, 0))
        ttk.Button(target_actions, text="← Back", style="Secondary.TButton", command=lambda: self._show_step(1)).pack(side="left")
        ttk.Button(target_actions, text="Edit Targets", style="Secondary.TButton", command=self._edit_targets).pack(side="right", padx=(7, 0))
        self.accept_targets_button = ttk.Button(target_actions, text="Accept Targets & Run →", style="Primary.TButton", command=self._accept_targets_and_run); self.accept_targets_button.pack(side="right")

    def _run_page(self, page) -> None:
        self._bottom_navigation(page, 2, "Review")
        title = "Run Sample Planning" if APP_MODE == "sample_planning" else "Choose outputs and run"
        subtitle = "Generate the approved planning workbook." if APP_MODE == "sample_planning" else "Choose the deliverables for this run, then generate them after review has passed."
        self._page_title(page, 4, title, subtitle)
        self.run_unavailable = tk.Frame(page, bg=COLORS["background"]); self.run_unavailable.pack(fill="x", pady=(0, 12))
        self.run_unavailable_message = tk.Label(self.run_unavailable, text="Run is not available yet. Complete Project setup and Validation first.", bg=COLORS["background"], fg=COLORS["warning"], font=("Segoe UI", 10, "bold")); self.run_unavailable_message.pack(anchor="w", pady=(0, 6))
        ttk.Button(self.run_unavailable, text="Go to Project", style="Secondary.TButton", command=lambda: self._show_step(1)).pack(side="left")
        ttk.Button(self.run_unavailable, text="Go to Validate", style="Secondary.TButton", command=lambda: self._show_step(2)).pack(side="left", padx=(7, 0))
        self.run_heading = tk.Label(page, text="", bg=COLORS["background"], fg=COLORS["navy"], font=("Segoe UI", 15, "bold")); self.run_heading.pack(anchor="w")
        self.run_blurb = tk.Label(page, text="", bg=COLORS["background"], fg=COLORS["muted"], font=("Segoe UI", 10), wraplength=820, justify="left"); self.run_blurb.pack(anchor="w", pady=(3, 14))
        self.output_choices_frame = ttk.LabelFrame(page, text="Output package", padding=10)
        if APP_MODE != "sample_planning":
            self.output_choices_frame.pack(fill="x", pady=(0, 12))
            ttk.Label(self.output_choices_frame, text="What should SCOPE create for this run?", foreground=COLORS["muted"]).pack(anchor="w", pady=(0, 4))
            for label, value in (("TEID lists", "teid_lists"), ("Demo / OSEP / clinical tables", "tables"), ("Both", "both")):
                ttk.Radiobutton(self.output_choices_frame, text=label, value=value, variable=self.output_selection).pack(side="left", padx=(0, 18))
        run_progress_row = tk.Frame(page, bg=COLORS["background"]); run_progress_row.pack(fill="x", pady=(0, 12))
        self.run_screen_progress = ttk.Progressbar(run_progress_row, style="Run.Horizontal.TProgressbar", maximum=100); self.run_screen_progress.pack(side="left", fill="x", expand=True)
        tk.Label(run_progress_row, textvariable=self.elapsed, bg=COLORS["background"], fg=COLORS["muted"], width=15, anchor="e").pack(side="right", padx=(10, 0))
        stages_card = tk.Frame(page, bg=COLORS["surface"], highlightbackground=COLORS["border"], highlightthickness=1); stages_card.pack(fill="x", pady=(0, 14))
        self.stage_labels = []
        for _index in range(6):
            label = tk.Label(stages_card, text="", bg=COLORS["surface"], fg=COLORS["muted"], anchor="w", font=("Segoe UI", 10), padx=15, pady=7)
            label.pack(fill="x"); self.stage_labels.append(label)
        self.run_readiness = tk.Label(page, text="Validate the project first.", bg=COLORS["background"], fg=COLORS["warning"], font=("Segoe UI", 10, "bold")); self.run_readiness.pack(anchor="w", pady=(0, 10))
        self.run_button = ttk.Button(page, text="RUN SCOPE", style="Primary.TButton", state="disabled", command=self._run_selected); self.run_button.pack(anchor="w")

    def _results_page(self, page) -> None:
        self._bottom_navigation(page, 3, "Run")
        self._page_title(page, 5, "Review outputs", "Open completed deliverables or preview the profile-defined outputs before a run.")
        tk.Label(page, textvariable=self.result_summary, bg=COLORS["background"], fg=COLORS["navy"], font=("Segoe UI", 11, "bold"), wraplength=850, justify="left").pack(anchor="w", pady=(0, 10))
        self.result_metrics = tk.Frame(page, bg=COLORS["background"]); self.result_metrics.pack(fill="x", pady=(0, 10))
        self.result_details = tk.Label(page, text="", bg=COLORS["surface"], fg=COLORS["text"], anchor="nw", justify="left", wraplength=850, padx=14, pady=12, highlightbackground=COLORS["border"], highlightthickness=1)
        self.result_details.pack(fill="x", pady=(0, 10))
        self.results_buttons = tk.Frame(page, bg=COLORS["background"]); self.results_buttons.pack(fill="x")
        self.preview_cards = tk.Frame(page, bg=COLORS["background"]); self.preview_cards.pack(fill="both", expand=True, pady=(4, 0))

    def _update_run_availability(self) -> None:
        if not hasattr(self, "run_unavailable"): return
        if self.validation_passed:
            self.run_unavailable.pack_forget()
            self.run_readiness.configure(text="✓ Validation passed. This project is ready to run.", fg=COLORS["success"])
        else:
            if not self.run_unavailable.winfo_manager(): self.run_unavailable.pack(fill="x", pady=(0, 12), before=self.run_heading)
            self.run_unavailable_message.configure(text="Run is not available yet. Complete Project setup and Validation first.")
            self.run_readiness.configure(text="Validate the project before running.", fg=COLORS["warning"])

    def _show_results_preview(self) -> None:
        for frame in (self.result_metrics, self.results_buttons, self.preview_cards):
            for child in frame.winfo_children(): child.destroy()
        if not self.preview_cards.winfo_manager(): self.preview_cards.pack(fill="both", expand=True, pady=(4, 0))
        if not self.profile_id.get() or not self.workflow.get():
            self.result_summary.set("RESULTS PREVIEW\nSelect a Test Profile and workflow to see its expected outputs.")
            self.result_details.configure(text="Results are profile-driven. Return to Test Profile and choose an available workflow.")
            return
        profile = self._profile(); workflow = self.workflow.get(); previews = profile.output_previews(workflow, self.study_type.get())
        self.result_summary.set("RESULTS PREVIEW\nSCOPE will generate the following outputs for this workflow.")
        if workflow == "recruitment_planning":
            self.result_details.configure(text=f"SCOPE will populate the approved Field Research report format for: {profile.display_name} • {self.study_type.get() or 'Profile default'}")
        elif workflow in {"targeted_swap", "targeted_adjustment"}:
            self.result_details.configure(text="Targeted Case Swap refreshes the selected cases, available pool, swap log, and stratification tables. CP-SAT NOT RERUN.")
        else:
            self.result_details.configure(text=f"The {profile.display_name} Reporting / Output Specification determines this package.")
        for index, item in enumerate(previews):
            card = RoundedCard(self.preview_cards, width=255, height=130, radius=12); card.grid(row=index // 3, column=index % 3, sticky="nsew", padx=(0 if index % 3 == 0 else 8, 0), pady=(0, 8)); self.preview_cards.grid_columnconfigure(index % 3, weight=1)
            panel = card.content; title = str(item.get("title", "Output")); tables = [str(value) for value in item.get("tables", [])]
            tk.Label(panel, text=title.upper(), bg=COLORS["surface"], fg=COLORS["navy"], font=("Segoe UI", 9, "bold"), wraplength=220, justify="left").pack(anchor="w", padx=12, pady=(10, 3))
            detail = str(item.get("description") or (", ".join(tables[:3]) + ("…" if len(tables) > 3 else "")) or "Profile-defined deliverable")
            tk.Label(panel, text=detail, bg=COLORS["surface"], fg=COLORS["muted"], font=("Segoe UI", 8), wraplength=220, justify="left").pack(anchor="w", padx=12)
            ttk.Button(panel, text="Preview", style="Secondary.TButton", command=lambda value=item: self._open_output_preview(value)).pack(anchor="w", padx=12, pady=8)

    def _open_output_preview(self, item: dict[str, object]) -> None:
        window = tk.Toplevel(self.root); window.title(f"SCOPE Preview — {item.get('title', 'Output')}"); window.geometry("760x430"); window.configure(bg=COLORS["background"])
        tk.Label(window, text=str(item.get("title", "Output")), bg=COLORS["background"], fg=COLORS["navy"], font=("Segoe UI", 16, "bold")).pack(anchor="w", padx=18, pady=(16, 3))
        tables = [str(value) for value in item.get("tables", [])]
        template = str(item.get("template", ""))
        detail = str(item.get("description", ""))
        if template and Path(template).is_file():
            try:
                from openpyxl import load_workbook
                sheet_names = load_workbook(template, read_only=True).sheetnames
                detail += "\nTemplate sheets: " + ", ".join(sheet_names[:12]) + ("…" if len(sheet_names) > 12 else "")
            except Exception: pass
        detail_label = tk.Label(window, text=detail or "Example structure; no participant data are required for this preview.", bg=COLORS["background"], fg=COLORS["muted"], wraplength=715, justify="left"); detail_label.pack(anchor="w", padx=18, pady=(0, 12))
        template_sheet = str(item.get("template_sheet", ""))
        if template_sheet == "TEMPLATE_1WAY":
            columns = ("category", "target_n", "target_pct", "sample_n", "sample_pct")
            headings = ("Category", "Target N", "Target %", "Sample N", "Sample %"); widths = (190, 90, 90, 90, 90)
            rows = [("First profile level", 100, "25.0%", 88, "22.0%"), ("Second profile level", 100, "25.0%", 104, "26.0%"), ("Total", 400, "100.0%", 400, "100.0%")]
        elif template_sheet in {"TEMPLATE_2WAY", "TEMPLATE_3WAY"}:
            first_table = tables[0] if tables else "Variable A × Variable B"; parts = self._profile().parse_report_table(first_table)
            column_labels = ["Level 1", "Level 2", "Level N"]
            definitions = self._profile().reporting_variables or {}
            if parts:
                requested = parts[-1]
                for name, raw in definitions.items():
                    specification = raw if isinstance(raw, dict) else {}; aliases = (name, specification.get("canonical", ""), specification.get("field", ""), specification.get("display", ""))
                    if re.sub(r"[^a-z0-9]+", "", requested.casefold()) in {re.sub(r"[^a-z0-9]+", "", str(alias).casefold()) for alias in aliases if alias}:
                        labels = [str(value.get("label", value.get("value", ""))) if isinstance(value, dict) else str(value) for value in specification.get("levels", [])]
                        if labels: column_labels = labels[:4]
                        break
            columns = tuple(["group", "metric"] + [f"level_{index}" for index in range(len(column_labels))] + ["total"])
            headings = tuple([parts[1] if len(parts) > 1 else "Group", "Metric"] + column_labels + ["Total"]); widths = tuple([150, 100] + [95] * len(column_labels) + [80])
            rows = []
            if template_sheet == "TEMPLATE_3WAY": rows.append((f"{parts[0] if parts else 'Variable A'}: first level", "", *("" for _ in column_labels), ""))
            rows.extend([
                ("First row level", "Target %", *("25.0%" for _ in column_labels), "100.0%"),
                ("", "Sample %", *(f"{22 + index}.0%" for index in range(len(column_labels))), "100.0%"),
                ("", "Sample N", *(str(22 + index) for index in range(len(column_labels))), str(sum(22 + index for index in range(len(column_labels))))),
            ])
            detail += f"\nApproved structure: {template_sheet}. Rows and columns expand from the selected Test Profile."
        else:
            columns = ("category", "group", "target_n", "current_n", "remaining", "target_pct", "current_pct", "diff")
            headings = ("Category", "Group", "Target N", "Current N", "Remaining N", "Target %", "Current %", "Diff"); widths = (120, 150, 75, 75, 85, 75, 75, 70)
            examples = tables[:3] or ["Age", "Sex", "Ethnicity"]
            rows = [(name, "Example category", 100, 72 + index, 28 - index, "25.0%", f"{22 + index}.0%", f"{-3 + index}.0%") for index, name in enumerate(examples)]
        tree = ttk.Treeview(window, columns=columns, show="headings", height=10)
        detail_label.configure(text=detail or "Example structure; no participant data are required for this preview.")
        for column, heading, width in zip(columns, headings, widths): tree.heading(column, text=heading); tree.column(column, width=width, anchor="center" if column not in {"category", "group"} else "w")
        for row in rows: tree.insert("", "end", values=row)
        tree.pack(fill="both", expand=True, padx=18, pady=(0, 12))
        ttk.Button(window, text="Close", style="Primary.TButton", command=window.destroy).pack(anchor="e", padx=18, pady=(0, 16))

    def _show_step(self, index: int, force: bool = False) -> None:
        self.active_step = index; self.pages[index].tkraise(); self._update_navigation()
        if index == 3: self._update_run_availability()
        if index == 4 and self.last_result is None: self._show_results_preview()

    def _go_back(self) -> None:
        if self.active_step > 0 and not self.busy:
            self._show_step(self.active_step - 1)

    def _update_navigation(self) -> None:
        accent = COLORS.get(self.workflow.get(), COLORS["navy_hover"])
        for index, button in enumerate(self.nav_buttons):
            complete = index in self.completed; unavailable = index > self.unlocked_step
            prefix = "✓" if complete else "○" if unavailable else str(index + 1)
            button.configure(text=f"{prefix}   {NAV_ITEMS[index]}", state="normal", bg=accent if index == self.active_step else COLORS["navy"], fg="white")

    def _choose_workflow(self, workflow: str, navigate: bool = False) -> None:
        self._hide_status_bar()
        if not self.profile_id.get():
            messagebox.showinfo("Test Profile", "Select a Test Profile first. Its capabilities determine which workflows are available."); return
        profile = self._profile()
        if not profile.supports(workflow):
            messagebox.showinfo("Workflow connection pending", self._profile_unavailable_message(profile, workflow)); return
        self.workflow.set(workflow); self.completed = {0}; self.unlocked_step = 1; self.validation_passed = False; self.last_result = None
        if workflow == "sample_optimization" and self.dataset.get().strip() and not self.candidate_datasets:
            self.candidate_datasets = [{"path": self.dataset.get().strip(), "source_type": self.study_type.get() or "STDZ", "source_role": "Primary", "priority": 1}]
        for key, card in self.workflow_cards.items(): card.set_selected(COLORS[workflow] if key == workflow else None)
        self._render_workflow_inputs(); self._update_run_copy(); self._refresh_profile_detail(); self._update_workflow_buttons(); self.status.set(f"{WORKFLOWS[workflow][0]} selected. Define this project."); self.companion.set("")
        self._append_log(f"Workflow selected: {WORKFLOWS[workflow][0]}")
        if navigate: self._show_step(1)
        else: self._update_navigation()

    def _workflow_continue(self) -> None:
        if not self.workflow.get(): messagebox.showinfo("Workflow", "Choose what SCOPE should do before continuing."); return
        self.completed.add(0); self.unlocked_step = max(self.unlocked_step, 1); self._show_step(1)

    def _select_profile(self, profile_id: str, navigate: bool = False) -> None:
        self._hide_status_bar()
        candidate = next(profile for profile in self.profiles if profile.id == profile_id)
        if self.profile_id.get() and self.profile_id.get() != profile_id:
            self._reset_activity_context()
            self.canonical_target_rows = []; self.targets_confirmed.set(False)
        self.profile_id.set(profile_id); self.validation_passed = False; self.last_result = None
        families = {str(value).strip().casefold() for value in (candidate.target_families or [])}
        self.include_census.set(bool(families & {"census", "demographic", "demographics"}) or not families)
        self.include_osep.set("osep" in families)
        self.include_clinical.set("clinical" in families)
        if self.workflow.get() and not candidate.supports(self.workflow.get()): self.workflow.set("")
        self.completed.discard(0); self.unlocked_step = 0
        for key, card in self.profile_cards.items(): card.set_selected(COLORS.get(self.workflow.get(), COLORS["navy"]) if key == profile_id else None)
        self._refresh_profile_detail(); self._refresh_profile_controls(); self._render_workflow_inputs(); self._update_workflow_buttons()
        self.status.set(f"{self._profile().display_name} selected. Choose what SCOPE should do.")
        self._append_log(f"{candidate.display_name} profile loaded.")
        if self.workflow.get() == "sample_optimization": self._refresh_target_review(show_errors=False)
        self._update_navigation()

    def _profile_continue(self) -> None:
        if not self.profile_id.get(): messagebox.showinfo("Test Profile", "Choose a Test Profile before continuing."); return
        profile = self._profile()
        if not profile.supports(self.workflow.get()):
            messagebox.showinfo("Workflow connection pending", self._profile_unavailable_message(profile, self.workflow.get())); return
        self.completed.add(1); self.unlocked_step = max(self.unlocked_step, 2); self._show_step(2)

    def _profile_unavailable_message(self, profile, workflow: str) -> str:
        workflow_name = WORKFLOWS.get(workflow, (workflow.replace("_", " ").title(), "", ""))[0]
        missing = ", ".join(profile.missing_for(workflow)) or "a validated SCOPE workflow connection"
        guidance = "Open the SCOPE Profile Manager to enable and confirm it, then Refresh Installed Profiles." if profile.status.casefold() != "validated" else "Enable this workflow in the SCOPE Profile Manager, then Refresh Installed Profiles."
        return f"{workflow_name} is not yet available for {profile.display_name}. Pending: {missing}. {guidance}"

    def _update_profile_continue_state(self) -> None:
        self._update_workflow_buttons()

    def _update_workflow_buttons(self) -> None:
        if not hasattr(self, "workflow_buttons"): return
        profile = self._profile() if self.profile_id.get() else None
        for workflow, button in self.workflow_buttons.items():
            enabled = bool(profile and profile.supports(workflow))
            button.configure(state="normal" if enabled else "disabled")
            detail = "Select a Test Profile first." if profile is None else self._profile_unavailable_message(profile, workflow) if not enabled else f"Open {WORKFLOWS[workflow][0]} project setup."
            ToolTip(button, detail)

    def _project_continue(self) -> None:
        try: self._config()
        except Exception as exc: self._show_correction("Project", exc); return
        self.completed.add(1); self.completed.discard(2); self.completed.discard(3); self.completed.discard(4); self.unlocked_step = max(self.unlocked_step, 2); self.validation_passed = False
        self.status.set("Project scope saved. Validate before running."); self._show_step(2)

    def _profile(self):
        return next(profile for profile in self.profiles if profile.id == self.profile_id.get())

    def _refresh_profile_detail(self) -> None:
        if not self.profile_id.get():
            if hasattr(self, "profile_detail"): self.profile_detail.configure(text="Select a Test Profile to see its status and workflow capabilities.")
            return
        profile = self._profile(); workflow = self.workflow.get(); supported = profile.supports(workflow) if workflow else False
        capabilities = []
        for key, (name, _audience, _description) in WORKFLOWS.items(): capabilities.append(f"{'✓' if profile.supports(key) else '○'} {name}")
        missing = ", ".join(profile.missing_for(workflow)) if workflow and not supported else ""
        text = f"{profile.display_name} · Profile v{profile.version} · {profile.status}\n" + "   ".join(capabilities)
        if workflow and not supported: text += f"\n\n{self._profile_unavailable_message(profile, workflow)}"
        else:
            tables = profile.required_demographic_tables()
            text += f"\n\nForms: {', '.join(profile.forms) or 'Profile-defined'}\nAge structure: {', '.join(profile.age_structure or []) or 'Profile-defined'}\nDemographics: {', '.join(profile.demographics) or 'None configured'}\nSelectable demographic tables: {len(tables['one_way'])} one-way, {len(tables['two_way'])} two-way, {len(tables['three_way'])} three-way"
        self.profile_detail.configure(text=text)

    def _refresh_profile_controls(self) -> None:
        if not self.profile_id.get(): return
        profile = self._profile()
        for child in self.forms_frame.winfo_children(): child.destroy()
        ttk.Label(self.forms_frame, text="Forms:").pack(side="left")
        self.form_vars = {form: tk.BooleanVar(value=True) for form in profile.forms}
        for form, variable in self.form_vars.items(): variable.trace_add("write", self._target_inputs_changed); ttk.Checkbutton(self.forms_frame, text=form, variable=variable).pack(side="left", padx=(8, 0))
        if not profile.forms: ttk.Label(self.forms_frame, text="Loaded when the profile is configured", foreground=COLORS["muted"]).pack(side="left", padx=(8, 0))
        for child in self.demographics_frame.winfo_children(): child.destroy()
        ttk.Label(self.demographics_frame, text="Demographic targets:").pack(side="left")
        labels = {"sex": "Sex", "eth": "Ethnicity", "edl": "Education", "region": "Region"}
        self.demographic_vars = {name: tk.BooleanVar(value=True) for name in profile.demographics}
        for name, variable in self.demographic_vars.items(): variable.trace_add("write", self._target_inputs_changed); ttk.Checkbutton(self.demographics_frame, text=labels.get(name, name.replace("_", " ").title()), variable=variable).pack(side="left", padx=(8, 0))
        if not profile.demographics: ttk.Label(self.demographics_frame, text="Profile-defined", foreground=COLORS["muted"]).pack(side="left", padx=(8, 0))
        languages = profile.languages or ["English", "Spanish"]; self.language_picker.configure(values=languages)
        if self.language.get() not in languages: self.language.set(languages[0])
        studies = list(profile.study_types or [])
        defaults = profile_study_type_defaults(profile)
        for value in defaults:
            if value not in studies: studies.append(value)
        if not studies: studies = ["STDZ", "Clinical", "Validity"]
        if not defaults: defaults = studies
        self.study_picker.configure(values=["", *studies])
        if self.study_type.get() not in defaults: self.study_type.set(defaults[0])
        if not self.sample_size.get().strip() and profile.default_target_n: self.sample_size.set(str(profile.default_target_n))
        self._render_profile_rules(profile)
        self._populate_report_table_selector()
        self._render_profile_resources()

    def _render_profile_rules(self, profile) -> None:
        """Show the profile's sample rules and optimization tolerances in context."""
        if not hasattr(self, "profile_rules_label"): return
        rules = [f"{rule['label']} {'=' if rule['operator'] == 'include' else '≠'} {', '.join(rule['values'])}" for rule in profile_sample_filters(profile)]
        optimization = describe_profile_optimization_rules(profile)
        parts = []
        if rules: parts.append("Profile sample rules: " + "; ".join(rules))
        if optimization: parts.append("Optimization tolerances: " + optimization)
        self.profile_rules_label.configure(text="\n".join(parts))

    def _populate_report_table_selector(self, selected: dict[str, list[str]] | None = None) -> None:
        if not hasattr(self, "report_tables_frame") or not self.profile_id.get():
            return
        for child in self.report_tables_frame.winfo_children(): child.destroy()
        specification = self._profile().required_demographic_tables()
        self.report_table_lists = {}; self.report_table_values = {key: list(specification.get(key, [])) for key in ("one_way", "two_way", "three_way")}
        ttk.Label(
            self.report_tables_frame,
            text="Choose the exact sheets to include. Ctrl+click selects or clears individual tables; selections are saved with the project.",
            foreground=COLORS["muted"], wraplength=930, justify="left",
        ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 6))
        labels = {"one_way": "1-way tables", "two_way": "2-way crosses", "three_way": "3-way crosses"}
        for column, kind in enumerate(("one_way", "two_way", "three_way")):
            panel = ttk.Frame(self.report_tables_frame); panel.grid(row=1, column=column, sticky="nsew", padx=(0 if column == 0 else 8, 0))
            values = self.report_table_values[kind]
            ttk.Label(panel, text=f"{labels[kind]} ({len(values)} available)", foreground=COLORS["navy"], font=("Segoe UI", 9, "bold")).pack(anchor="w")
            box = tk.Listbox(panel, selectmode="multiple", exportselection=False, height=min(6, max(3, len(values))), activestyle="dotbox")
            box.pack(fill="both", expand=True, pady=(3, 4))
            for value in values: box.insert("end", value)
            wanted = set((selected or {}).get(kind, values))
            for index, value in enumerate(values):
                if value in wanted: box.selection_set(index)
            box.bind("<<ListboxSelect>>", self._report_table_selection_changed)
            self.report_table_lists[kind] = box
            actions = ttk.Frame(panel); actions.pack(fill="x")
            ttk.Button(actions, text="Select all", style="Secondary.TButton", command=lambda value=kind: self._set_report_table_family(value, True)).pack(side="left")
            ttk.Button(actions, text="Clear", style="Secondary.TButton", command=lambda value=kind: self._set_report_table_family(value, False)).pack(side="left", padx=(5, 0))
            self.report_tables_frame.grid_columnconfigure(column, weight=1)
        if not any(self.report_table_values.values()):
            ttk.Label(self.report_tables_frame, text="No demographic tables are defined in this Test Profile. Add them in SCOPE Profile Manager.", foreground=COLORS["warning"]).grid(row=1, column=0, columnspan=3, sticky="w")
        self._update_report_table_selector_visibility()

    def _update_report_table_selector_visibility(self) -> None:
        if not hasattr(self, "report_tables_frame"): return
        available = any(self.report_table_values.values())
        if self.workflow.get() == "sample_optimization" and available:
            if not self.report_tables_frame.winfo_manager(): self.report_tables_frame.pack(fill="x", pady=(6, 2))
        else:
            self.report_tables_frame.pack_forget()

    def _selected_report_tables(self) -> dict[str, list[str]]:
        return {
            kind: [self.report_table_values.get(kind, [])[index] for index in box.curselection()]
            for kind, box in self.report_table_lists.items()
        }

    def _set_report_table_family(self, kind: str, selected: bool) -> None:
        box = self.report_table_lists.get(kind)
        if not box: return
        box.selection_set(0, "end") if selected else box.selection_clear(0, "end")
        self._report_table_selection_changed()

    def _report_table_selection_changed(self, _event=None) -> None:
        if not self._suspend_dirty: self._invalidate("Demographic table selections changed. Validate the project again.")

    def _render_profile_resources(self) -> None:
        if not hasattr(self, "profile_resources_frame"): return
        for child in self.profile_resources_frame.winfo_children(): child.destroy()
        if not self.profile_id.get():
            ttk.Label(self.profile_resources_frame, text="Choose a Test Profile to resolve its approved reusable resources.", foreground=COLORS["muted"]).pack(anchor="w")
            return
        profile = self._profile(); statuses = profile.resource_status()
        if not statuses:
            ttk.Label(self.profile_resources_frame, text="This validated profile uses its bundled specification metadata and built-in report templates.", foreground=COLORS["muted"]).pack(anchor="w")
            return
        ttk.Label(self.profile_resources_frame, text="Automatically loaded from the approved Test Profile. Project participant data are never stored here.", foreground=COLORS["muted"]).pack(anchor="w", pady=(0, 5))
        grid = ttk.Frame(self.profile_resources_frame); grid.pack(fill="x")
        for index, (label, path, required) in enumerate(statuses):
            exists = bool(path and path.is_file()); symbol = "✓" if exists else "✕"; color = COLORS["success"] if exists else COLORS["error"] if required else COLORS["warning"]
            tk.Label(grid, text=f"{symbol}  {label}" + ("" if exists else " missing"), bg=COLORS["background"], fg=color, anchor="w", font=("Segoe UI", 9, "bold" if not exists else "normal")).grid(row=index // 2, column=index % 2, sticky="w", padx=(0, 30), pady=2)

    def _toggle_age_override(self) -> None:
        if self.override_age_defaults.get():
            if not self.age_frame.winfo_manager(): self.age_frame.pack(fill="x", pady=(5, 0))
        else:
            self.age_frame.pack_forget()
        self._invalidate("Age-band settings changed. Validate the project again.")

    def _age_text_changed(self, _event=None) -> None:
        if not self.age_text.edit_modified(): return
        self.age_text.edit_modified(False)
        if not self._suspend_dirty: self._invalidate("Age-band settings changed. Validate the project again.")

    def _render_workflow_inputs(self) -> None:
        for child in self.workflow_inputs.winfo_children(): child.destroy()
        self.input_status_labels = {}
        workflow = self.workflow.get(); accent = COLORS.get(workflow, COLORS["navy"])
        self._update_report_table_selector_visibility()
        if hasattr(self, "scope_frame"):
            if workflow in {"targeted_swap", "targeted_adjustment"}:
                self.scope_frame.pack_forget()
            elif not self.scope_frame.winfo_manager():
                self.scope_frame.pack(fill="x", pady=4, before=self.workflow_inputs)
        if workflow == "sample_optimization":
            self.workflow_inputs.configure(text="Workflow inputs")
            ttk.Label(self.workflow_inputs, text="Provide a candidate TEID list with the demographic columns used by the target workbook. A Research Plan is not required for target-workbook CP-SAT profiles.", foreground=accent, wraplength=800, justify="left").pack(anchor="w", pady=(0, 5))
            if not self.profile_id.get():
                ttk.Label(self.workflow_inputs, text="Choose a Test Profile to see its assessment-specific input requirements.", foreground=COLORS["muted"]).pack(anchor="w", pady=5)
            else:
                requirements = self._profile().requirements_for(workflow)
                required = [item for item in requirements if item.required]
                additional = [item for item in requirements if not item.required and item.role != "current_sample"]
                existing = next((item for item in requirements if item.role == "current_sample"), None)
                ttk.Label(self.workflow_inputs, text="REQUIRED INPUTS", foreground=COLORS["navy"], font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(4, 2))
                for requirement in required:
                    if requirement.role == "master_dataset":
                        self._candidate_dataset_editor(self.workflow_inputs)
                    else:
                        self._input_requirement_row(self.workflow_inputs, requirement)
                ttk.Label(self.workflow_inputs, text="OPTIONAL INPUTS", foreground=COLORS["navy"], font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(10, 2))
                if existing:
                    existing_row = ttk.Frame(self.workflow_inputs); existing_row.pack(fill="x", pady=(1, 3))
                    ttk.Checkbutton(existing_row, text="Use an existing selected sample (re-optimization only)", variable=self.use_existing_sample, command=self._toggle_existing_sample).pack(anchor="w")
                    if self.use_existing_sample.get(): self._input_requirement_row(self.workflow_inputs, existing)
                for requirement in additional:
                    if self.show_additional_sources or self._input_variable(requirement.role).get().strip(): self._input_requirement_row(self.workflow_inputs, requirement)
                if additional:
                    ttk.Button(self.workflow_inputs, text="Hide additional sources" if self.show_additional_sources else "Show additional profile-specific sources", style="Secondary.TButton", command=self._toggle_additional_sources).pack(anchor="w", pady=(8, 2))
                if not requirements:
                    ttk.Label(self.workflow_inputs, text="This profile has not declared Sample Optimization inputs.", foreground=COLORS["error"]).pack(anchor="w", pady=5)
                self._update_input_statuses()
            self.seeding_check.configure(state="normal")
        elif workflow == "recruitment_planning":
            self.workflow_inputs.configure(text="Research Plan / targets")
            ttk.Label(self.workflow_inputs, text="SCOPE will use the approved recruitment specification bundled with the selected Test Profile.", foreground=accent).pack(anchor="w", pady=(0, 5))
            self._path_row(self.workflow_inputs, "Project-specific target amendment (optional)", self.recruitment_plan_file, file=True)
            self._path_row(self.workflow_inputs, "Current Recruitment / Progress Report (optional)", self.recruitment_actuals_file, file=True)
            self.seeding_check.configure(state="disabled")

        elif workflow in {"targeted_swap", "targeted_adjustment"}:
            self.workflow_inputs.configure(text="TEID membership and participant data")
            ttk.Label(self.workflow_inputs, text="MEMBERSHIP LISTS", foreground=COLORS["navy"], font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(0, 2))
            ttk.Label(
                self.workflow_inputs,
                text="The current selected list must describe membership before these swaps and match the Original dataset below. The add-candidate/refreshed source must match the Updated dataset. Overlap is expected; SCOPE keeps current TEIDs selected and excludes them from the effective add pool.",
                foreground=COLORS["muted"], wraplength=800, justify="left",
            ).pack(anchor="w", pady=(0, 5))
            self._path_row(self.workflow_inputs, "Current selected TEID list (before swaps)", self.swap_sample, file=True)
            self._path_row(self.workflow_inputs, "Add-candidate / refreshed membership source", self.swap_available, file=True)
            ttk.Label(self.workflow_inputs, text="PARTICIPANT DATA FOR THOSE LISTS", foreground=COLORS["navy"], font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(9, 2))
            ttk.Label(
                self.workflow_inputs,
                text="Use these when the membership files are TEID-only. The original dataset supplies attributes for selected TEIDs; the updated dataset supplies attributes for available candidates. A shared case dataset from Step 1 remains a fallback.",
                foreground=COLORS["muted"], wraplength=800, justify="left",
            ).pack(anchor="w", pady=(0, 5))
            self._path_row(self.workflow_inputs, "Original dataset matching the current selected list (optional)", self.swap_sample_data, file=True)
            self._path_row(self.workflow_inputs, "Updated dataset matching the candidate/refreshed source (optional)", self.swap_available_data, file=True)
            lock_fields = default_lock_fields(self._profile())
            existing = {name: variable.get() for name, variable in self.lock_vars.items()}; self.lock_vars = {name: tk.BooleanVar(value=existing.get(name, True)) for name in lock_fields}
            for variable in self.lock_vars.values(): variable.trace_add("write", self._mark_project_dirty)
            lock_row = ttk.Frame(self.workflow_inputs); lock_row.pack(fill="x", pady=(7, 1))
            ttk.Label(lock_row, text="Required locks:", foreground=COLORS["navy"], font=("Segoe UI", 9, "bold")).pack(side="left")
            for name, variable in self.lock_vars.items():
                ttk.Checkbutton(lock_row, text=name.replace("_", " ").title(), variable=variable).pack(side="left", padx=(10, 0))
            section_name = "OSEP / clinical case adjustments" if workflow == "targeted_adjustment" else "Case swaps"
            ttk.Label(self.workflow_inputs, text=section_name, foreground=COLORS["navy"], font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(10, 3))
            entry = ttk.Frame(self.workflow_inputs); entry.pack(fill="x", pady=(4, 2))
            for label, variable, width in (("Add TEID", self.swap_add, 15), ("Remove TEID", self.swap_remove, 15), ("Reason", self.swap_reason, 25)):
                ttk.Label(entry, text=label + ":").pack(side="left")
                field = ttk.Entry(entry, textvariable=variable, width=width); field.pack(side="left", padx=(4, 10)); field.bind("<Return>", lambda _event: self._add_swap())
            actions = ttk.Frame(self.workflow_inputs); actions.pack(fill="x", pady=(2, 4))
            ttk.Button(actions, text="Add Swap Selections", style="Secondary.TButton", command=self._add_swap).pack(side="left", padx=(0, 6))
            ttk.Button(actions, text="Clear Swap Selections", style="Secondary.TButton", command=self._clear_swaps).pack(side="left")
            ttk.Label(actions, textvariable=self.swap_selection_summary, foreground=COLORS["success"]).pack(side="left", padx=(12, 0))
            self.swap_tree = ttk.Treeview(self.workflow_inputs, columns=("add", "remove", "reason"), show="headings", height=4)
            for column, title, width in (("add", "Add TEID", 170), ("remove", "Remove TEID", 170), ("reason", "Reason", 390)):
                self.swap_tree.heading(column, text=title); self.swap_tree.column(column, width=width)
            self.swap_tree.pack(fill="x", pady=4)
            for row in self.swap_rows: self.swap_tree.insert("", "end", values=(row["add"], row["remove"], row.get("reason", "")))
            self._update_swap_selection_summary()
            self.seeding_check.configure(state="disabled")

    def _candidate_dataset_editor(self, parent) -> None:
        ttk.Label(parent, text="Candidate case data / TEID list", font=("Segoe UI", 9, "bold")).pack(anchor="w")
        ttk.Label(parent, text="The file may be XLSX, SAS7BDAT, XPT, or CSV. It must contain TEID plus the demographic fields used in the target workbook; a TEID-only file cannot be optimized.", foreground=COLORS["muted"], wraplength=800, justify="left").pack(anchor="w", pady=(0, 4))
        source_types = list(dict.fromkeys([*(self._profile().study_types or []), "STDZ", "Demo", "NB395", "Clinical", "Pilot", "Other"]))
        for index, source in enumerate(self.candidate_datasets):
            row = ttk.Frame(parent); row.pack(fill="x", pady=2)
            ttk.Label(row, text=f"{index + 1}.", width=3).pack(side="left")
            ttk.Label(row, text=Path(str(source.get("path", ""))).name or "(missing)", width=30).pack(side="left")
            type_var = tk.StringVar(value=str(source.get("source_type") or "STDZ"))
            role_var = tk.StringVar(value=str(source.get("source_role") or ("Primary" if index == 0 else "Supplemental")))
            type_box = ttk.Combobox(row, textvariable=type_var, values=source_types, width=13); type_box.pack(side="left", padx=3)
            role_box = ttk.Combobox(row, textvariable=role_var, values=("Primary", "Supplemental", "Seeding only"), width=14); role_box.pack(side="left", padx=3)
            def update(_event=None, item=source, type_value=type_var, role_value=role_var):
                item["source_type"] = type_value.get(); item["source_role"] = role_value.get(); self._candidate_sources_changed()
            type_box.bind("<<ComboboxSelected>>", update); role_box.bind("<<ComboboxSelected>>", update)
            ttk.Button(row, text="Remove", style="Secondary.TButton", command=lambda i=index: self._remove_candidate_dataset(i)).pack(side="right")
        ttk.Button(parent, text="+ Add Dataset", style="Secondary.TButton", command=self._add_candidate_dataset).pack(anchor="w", pady=(5, 2))
        if not self.candidate_datasets:
            tk.Label(parent, text="✕ Missing · Add at least one dataset", bg=COLORS["background"], fg=COLORS["error"], font=("Segoe UI", 8, "bold")).pack(anchor="w")

    def _add_candidate_dataset(self) -> None:
        values = filedialog.askopenfilenames(title="Add Master / Candidate Datasets", filetypes=[("Case data", "*.sas7bdat *.xpt *.xlsx *.xlsm *.csv"), ("All files", "*.*")])
        if not values: return
        existing = {str(Path(item.get("path", "")).resolve()).casefold() for item in self.candidate_datasets}
        for value in values:
            resolved = str(Path(value).resolve())
            if resolved.casefold() in existing: continue
            index = len(self.candidate_datasets)
            name = Path(value).stem.casefold()
            source_type = "Clinical" if "clin" in name else "NB395" if "395" in name else "Demo" if "demo" in name else "STDZ"
            role = "Primary" if index == 0 else "Supplemental"
            self.candidate_datasets.append({"path": resolved, "source_type": source_type, "source_role": role, "priority": index + 1})
            self._append_log(f"{role} dataset added: {Path(value).name}")
            existing.add(resolved.casefold())
        self._candidate_sources_changed()

    def _remove_candidate_dataset(self, index: int) -> None:
        if 0 <= index < len(self.candidate_datasets):
            removed = self.candidate_datasets.pop(index)
            for priority, item in enumerate(self.candidate_datasets, start=1): item["priority"] = priority
            self._append_log(f"Candidate dataset removed: {Path(str(removed.get('path', ''))).name}")
            self._candidate_sources_changed()

    def _candidate_sources_changed(self) -> None:
        self.dataset.set(str(self.candidate_datasets[0]["path"]) if self.candidate_datasets else "")
        self.targets_confirmed.set(False); self.canonical_target_rows = []
        self._render_workflow_inputs(); self._invalidate("Candidate sources changed. Validate the project again.")

    def _toggle_additional_sources(self) -> None:
        self.show_additional_sources = not self.show_additional_sources
        self._render_workflow_inputs()

    def _toggle_existing_sample(self) -> None:
        self.targets_confirmed.set(False)
        self._render_workflow_inputs()
        self._invalidate("Sample scenario changed. Validate the project again.")

    def _input_variable(self, role: str):
        return {
            "master_dataset": self.dataset, "current_sample": self.master_file,
            "scored_data": self.scored_dir, "target_specification": self.target_file,
            "initial_teid_source": self.fr_report_file, "field_research_report": self.fr_report_file, "case_review_data": self.case_swap_dir,
            "recruitment_targets": self.recruitment_plan_file, "recruitment_progress": self.recruitment_actuals_file,
            "swap_sample": self.swap_sample, "swap_available_pool": self.swap_available, "swap_case_universe": self.swap_universe,
            "swap_sample_data": self.swap_sample_data, "swap_available_data": self.swap_available_data,
        }[role]

    def _input_requirement_row(self, parent, requirement) -> None:
        row = ttk.Frame(parent); row.pack(fill="x", pady=3)
        badge = tk.Label(row, text="", bg=COLORS["background"], width=22, anchor="w", font=("Segoe UI", 8, "bold")); badge.pack(side="left")
        self.input_status_labels[requirement.role] = (badge, requirement)
        content = ttk.Frame(row); content.pack(side="left", fill="x", expand=True)
        ttk.Label(content, text=requirement.label, font=("Segoe UI", 9, "bold")).pack(anchor="w")
        ttk.Label(content, text=requirement.purpose, foreground=COLORS["muted"], wraplength=620, justify="left").pack(anchor="w")
        selector = ttk.Frame(content); selector.pack(fill="x", pady=(2, 0))
        variable = self._input_variable(requirement.role)
        entry = ttk.Entry(selector, textvariable=variable); entry.pack(side="left", fill="x", expand=True)
        entry.bind("<FocusOut>", lambda _event: self._input_path_changed(requirement.role))
        ttk.Button(selector, text="Select Folder…" if requirement.kind == "folder" else "Browse…", style="Secondary.TButton", command=lambda value=requirement: self._choose_profile_input(value)).pack(side="left", padx=(7, 0))

    def _choose_profile_input(self, requirement) -> None:
        if requirement.kind == "folder":
            value = filedialog.askdirectory(title=requirement.label)
        else:
            patterns = " ".join(f"*{extension}" for extension in requirement.accepted_types) or "*.*"
            value = filedialog.askopenfilename(title=requirement.label, filetypes=[("Accepted files", patterns), ("All files", "*.*")])
        if not value: return
        self._input_variable(requirement.role).set(value)
        self._input_path_changed(requirement.role)

    def _input_path_changed(self, role: str) -> None:
        if role == "master_dataset":
            value = self.dataset.get().strip()
            if self.candidate_datasets:
                self.candidate_datasets[0]["path"] = value
            elif value:
                self.candidate_datasets = [{"path": value, "source_type": self.study_type.get() or "STDZ", "source_role": "Primary", "priority": 1}]
            self.targets_confirmed.set(False); self.canonical_target_rows = []
        if role in {"target_specification", "field_research_report"}:
            self.targets_confirmed.set(False); self.canonical_target_rows = []
            self._refresh_target_review(show_errors=False)
        self._update_input_statuses(); self._invalidate("An input changed. Validate the project again.")

    def _update_input_statuses(self) -> None:
        for role, (label, requirement) in self.input_status_labels.items():
            value = self._input_variable(role).get().strip()
            ready = Path(value).is_dir() if requirement.kind == "folder" and value else Path(value).is_file() if value else False
            profile_default = resolved_profile_input(
                ProjectConfig(profile_id=self.profile_id.get(), project_dir=self.project_dir.get(), target_file=value),
                self._profile(), role,
            ) if role == "target_specification" else None
            if not ready and profile_default and profile_default.is_file():
                ready = True
            status = requirement.status.casefold()
            if ready:
                text, color = (f"✓ {requirement.status.title()} · Selected" if value else f"✓ {requirement.status.title()} · Profile default"), COLORS["success"]
            elif status == "required":
                text, color = "✕ Missing · Required", COLORS["error"]
            elif status == "conditional":
                text, color = "⚠ Conditional", COLORS["warning"]
            else:
                text, color = "○ Optional", COLORS["muted"]
            label.configure(text=text, fg=color)

    def _refresh_target_review(self, show_errors: bool = False) -> None:
        if self.workflow.get() != "sample_optimization" or not self.profile_id.get(): return
        try:
            config = self._config(save=False)
            targets = build_target_specification(config, self._profile())
            rows = canonical_target_records(targets)
        except Exception as exc:
            self.canonical_target_rows = []; self.targets_confirmed.set(False)
            self.target_review_summary.set(f"⚠ {exc}")
            self._populate_target_review()
            if show_errors: self._show_correction("Review Targets", exc)
            return
        if rows != self.canonical_target_rows:
            self.targets_confirmed.set(False)
        self.canonical_target_rows = rows
        self._populate_target_review()

    def _populate_target_review(self) -> None:
        if not hasattr(self, "target_review_tree"): return
        self.target_review_tree.delete(*self.target_review_tree.get_children())
        for row in self.canonical_target_rows:
            raw_target = row.get("target", ""); unit = str(row.get("unit", ""))
            target = f"{raw_target}{unit}" if unit == "%" else f"{raw_target} {unit}".strip()
            source = str(row.get("source", "")); tags = ()
            if row.get("override"):
                original = row.get("profile_value", "")
                source = f"⚠ Project Override {original}{unit} → {target}"
                tags = ("override",)
            self.target_review_tree.insert("", "end", values=(row.get("category", ""), row.get("group", ""), row.get("rule", ""), target, row.get("tolerance", ""), row.get("priority", ""), source), tags=tags)
        if self.canonical_target_rows:
            overrides = sum(bool(row.get("override")) for row in self.canonical_target_rows)
            derived = sum(str(row.get("source", "")).startswith("Derived from") for row in self.canonical_target_rows)
            interpretation = describe_target_interpretation(self._profile())
            summary = f"{len(self.canonical_target_rows)} effective target rule(s) built automatically"
            if derived: summary += f" · {derived} derived marginal target(s) from the crossed source"
            if overrides: summary += f" · {overrides} project override(s)"
            summary += "."
            if interpretation: summary += f"\nTarget interpretation: {interpretation}"
            self.target_review_summary.set(summary)
            if hasattr(self, "accept_targets_button"): self.accept_targets_button.configure(state="normal")
        else:
            if not self.target_review_summary.get().startswith("⚠"):
                self.target_review_summary.set("No effective targets could be built from the selected profile and project inputs.")
            self.targets_confirmed.set(False)
            if hasattr(self, "accept_targets_button"): self.accept_targets_button.configure(state="disabled")

    def _confirm_target_review(self) -> None:
        if self.targets_confirmed.get() and not self.canonical_target_rows:
            self.targets_confirmed.set(False); messagebox.showinfo("Review Targets", "Load the canonical targets before confirming them."); return
        message = "Canonical targets confirmed. Continue to Validate." if self.targets_confirmed.get() else "Target confirmation cleared."
        self._invalidate(message)

    def _edit_targets(self) -> None:
        selected = self.target_review_tree.selection()
        if not selected:
            messagebox.showinfo("Edit Targets", "Select one target row to create or update a project override."); return
        values = self.target_review_tree.item(selected[0], "values")
        row = next((item for item in self.canonical_target_rows if str(item.get("category", "")) == str(values[0]) and str(item.get("group", "")) == str(values[1]) and str(item.get("rule", "")) == str(values[2])), None)
        if row is None: return
        window = tk.Toplevel(self.root); window.title("Project Target Override"); window.geometry("470x260"); window.configure(bg=COLORS["background"]); window.transient(self.root); window.grab_set()
        tk.Label(window, text="Project Target Override", bg=COLORS["background"], fg=COLORS["navy"], font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=18, pady=(16, 4))
        tk.Label(window, text=f"{row.get('category')} · {row.get('group')} · {row.get('rule')}\nProfile value: {row.get('profile_value', row.get('target'))} {row.get('unit', '')}", bg=COLORS["background"], fg=COLORS["text"], justify="left").pack(anchor="w", padx=18)
        target_var = tk.StringVar(value=str(row.get("target", ""))); reason_var = tk.StringVar(value=str(row.get("override_reason", "")))
        form = ttk.Frame(window); form.pack(fill="x", padx=18, pady=12)
        ttk.Label(form, text="Effective target:", width=18).grid(row=0, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=target_var).grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Reason (optional):", width=18).grid(row=1, column=0, sticky="w", pady=4); ttk.Entry(form, textvariable=reason_var).grid(row=1, column=1, sticky="ew", pady=4); form.grid_columnconfigure(1, weight=1)
        def save_override() -> None:
            value = target_var.get().strip()
            if not value: messagebox.showerror("Project Target Override", "Enter an effective target value.", parent=window); return
            try: parsed = float(value)
            except ValueError: parsed = value
            identity = (str(row.get("category", "")), str(row.get("group", "")), str(row.get("rule", "")))
            self.target_overrides = [item for item in self.target_overrides if (str(item.get("category", "")), str(item.get("group", "")), str(item.get("rule", ""))) != identity]
            self.target_overrides.append({"category": identity[0], "group": identity[1], "rule": identity[2], "target": parsed, "reason": reason_var.get().strip()})
            window.destroy(); self._refresh_target_review(show_errors=True); self._append_log(f"Project target override created: {identity[0]} / {identity[1]} = {parsed}")
        actions = ttk.Frame(window); actions.pack(fill="x", padx=18)
        ttk.Button(actions, text="Cancel", style="Secondary.TButton", command=window.destroy).pack(side="right", padx=(7, 0)); ttk.Button(actions, text="Save Override", style="Primary.TButton", command=save_override).pack(side="right")

    def _accept_targets_and_run(self) -> None:
        if not self.canonical_target_rows:
            messagebox.showinfo("Review Targets", "No effective targets are available to accept."); return
        self.targets_confirmed.set(True)
        self.status.set("Targets accepted. Starting Sample Optimization…")
        self._append_log("Effective targets accepted. Starting Sample Optimization.")
        self._show_step(3); self._run_selected()

    def _update_run_copy(self) -> None:
        workflow = self.workflow.get()
        if not workflow: return
        name, _audience, description = WORKFLOWS[workflow]
        heading = {"sample_optimization": "RUNNING CP-SAT", "recruitment_planning": "BUILDING SAMPLE PLAN", "targeted_swap": "APPLYING TARGETED CASE SWAPS", "targeted_adjustment": "APPLYING OSEP / CLINICAL ADJUSTMENTS"}[workflow]
        self.run_heading.configure(text=heading); self.run_blurb.configure(text=description)
        self.run_button.configure(text="GENERATE SELECTED OUTPUTS" if APP_MODE != "sample_planning" else "RUN SAMPLE PLANNING")
        self._reset_stages()

    def _stage_names(self) -> list[str]:
        return {
            "sample_optimization": ["Preparing candidate pool", "Applying eligibility rules", "Checking demographic targets", "Searching feasible solutions", "Profiling optimized sample", "Creating outputs"],
            "recruitment_planning": ["Reading Research Plan / targets", "Building target structure", "Reconciling integer target Ns", "Profiling recruitment plan", "Creating reports"],
            "targeted_swap": ["Loading sample", "Validating requested swaps", "Applying swaps", "Recalculating stratification", "Updating outputs"],
            "targeted_adjustment": ["Loading sample", "Validating group changes", "Applying adjustments", "Recalculating OSEP / clinical tables", "Updating outputs"],
        }.get(self.workflow.get(), [])

    def _reset_stages(self) -> None:
        names = self._stage_names()
        for index, label in enumerate(self.stage_labels): label.configure(text=f"○  {names[index]}" if index < len(names) else "", fg=COLORS["muted"])

    def _update_stages(self, current: int, complete: bool = False) -> None:
        names = self._stage_names()
        for index, label in enumerate(self.stage_labels):
            if index >= len(names): label.configure(text=""); continue
            if complete or index < current: label.configure(text=f"✓  {names[index]}", fg=COLORS["success"])
            elif index == current: label.configure(text=f"→  {names[index]}", fg=COLORS.get(self.workflow.get(), COLORS["navy"]))
            else: label.configure(text=f"○  {names[index]}", fg=COLORS["muted"])

    def _path_row(self, parent, label, variable, file=False, folder=False) -> None:
        row = ttk.Frame(parent); row.pack(fill="x", pady=2); ttk.Label(row, text=label + ":", width=38).pack(side="left"); ttk.Entry(row, textvariable=variable).pack(side="left", fill="x", expand=True)
        def choose() -> None:
            value = filedialog.askopenfilename() if file else filedialog.askdirectory() if folder else ""
            if value: variable.set(value); self._invalidate("An input changed. Validate the project again.")
        ttk.Button(row, text="Browse…", style="Secondary.TButton", command=choose).pack(side="left", padx=(7, 0))

    def _parse_age_bands(self) -> tuple[list[dict], dict[str, int]]:
        bands, targets = [], {}
        for group, line in enumerate((value.strip() for value in self.age_text.get("1.0", "end").splitlines()), start=1):
            if not line: continue
            match = re.fullmatch(r"(\d+):(\d+)\s*[-–]\s*(\d+):(\d+)(?:\s*=\s*(\d+))?", line)
            if not match: raise ValueError(f"Age band '{line}' must look like 5:0-5:5 or 5:0-5:5 = 100.")
            y1, m1, y2, m2, target = (int(value) if value is not None else None for value in match.groups())
            if m1 > 11 or m2 > 11 or y2 * 12 + m2 < y1 * 12 + m1: raise ValueError(f"Age band '{line}' is not a valid increasing month range.")
            bands.append({"group": group, "label": f"{y1}:{m1}-{y2}:{m2}", "min_months": y1 * 12 + m1, "max_months": y2 * 12 + m2})
            if target is not None: targets[str(group)] = target
        return bands, targets

    def _config(self, save: bool = False) -> ProjectConfig:
        if not self.workflow.get(): raise ValueError("Choose a workflow first.")
        if not self.profile_id.get(): raise ValueError("Choose a Test Profile first.")
        raw_n = self.sample_size.get().strip()
        try: sample_n = int(raw_n) if raw_n else None
        except ValueError as exc: raise ValueError("Target N must be a whole number.") from exc
        if sample_n is not None and sample_n < 1: raise ValueError("Target N must be greater than zero.")
        age_bands, target_by_age = self._parse_age_bands() if self.override_age_defaults.get() else ([], {}); profile = self._profile()
        targeted = self.workflow.get() in {"targeted_swap", "targeted_adjustment"}
        if targeted:
            self._commit_pending_swap()
        locks = [name for name, variable in self.lock_vars.items() if variable.get()] if targeted else ["vform", "agegrp"]
        membership = self.swap_sample.get().strip()
        available = self.swap_available.get().strip() or membership
        universe = self.dataset.get().strip() or self.swap_universe.get().strip()
        report_tables = self._selected_report_tables()
        config = ProjectConfig(profile_id=profile.id, profile_version=profile.version, project_file=str(self.current_project_file or ""), run_history=list(self.run_history), project_dir=self.project_dir.get().strip(), project_name=self.project_name.get().strip(), workflow=self.workflow.get(), output_dir=self.output_dir.get().strip(), output_selection=self.output_selection.get(), dataset=self.dataset.get().strip(), candidate_datasets=[dict(item) for item in self.candidate_datasets], master_file=self.master_file.get().strip() if self.use_existing_sample.get() else "", use_existing_sample=self.use_existing_sample.get(), fr_teid_file=self.fr_report_file.get().strip(), scored_dir=self.scored_dir.get().strip(), case_swap_dir=self.case_swap_dir.get().strip(), target_file=self.target_file.get().strip(), recruitment_plan_file=self.recruitment_plan_file.get().strip(), recruitment_actuals_file=self.recruitment_actuals_file.get().strip(), swap_sample_file=membership, swap_available_file=available, swap_case_universe_file=universe, swap_sample_data_file=self.swap_sample_data.get().strip(), swap_available_data_file=self.swap_available_data.get().strip(), targeted_swaps=list(self.swap_rows), forms=[name for name, variable in self.form_vars.items() if variable.get()], language=self.language.get(), study_type=self.study_type.get(), demographics=[name for name, variable in self.demographic_vars.items() if variable.get()], demographic_tables_1way=report_tables.get("one_way", []), demographic_tables_2way=report_tables.get("two_way", []), demographic_tables_3way=report_tables.get("three_way", []), demographic_tables_configured=self.workflow.get() == "sample_optimization" and bool(self.report_table_lists), include_census=self.include_census.get(), include_osep=self.include_osep.get(), include_clinical=self.include_clinical.get(), clinical_seeding=self.clinical_seeding.get(), sample_size=sample_n, age_bands=age_bands, target_n_by_agegrp=target_by_age, canonical_targets=list(self.canonical_target_rows), target_overrides=list(self.target_overrides), target_review_confirmed=self.targets_confirmed.get(), locks=locks, exclusions=list(profile.eligibility_rules or []))
        if save:
            destination = self.current_project_file or config.path
            config.save(destination); self.current_project_file = destination; self._record_recent(destination)
        return config

    def _validate(self) -> None:
        self._show_status_bar()
        try: config = self._config(); checks = validate(config)
        except Exception as exc: self._show_correction("Validation", exc); return
        self.validation_tree.delete(*self.validation_tree.get_children())
        for name, passed, detail in checks:
            review = passed and (name in {"Available pool", "Exclusions"} or str(detail).startswith(("Conditional", "Optional"))); status = "⚠ Review" if review else "✓ Passed" if passed else "✕ Missing"
            tag = "review" if review else "pass" if passed else "fail"; self.validation_tree.insert("", "end", values=(status, name, detail), tags=(tag,))
        self.validation_tree.tag_configure("pass", foreground=COLORS["success"]); self.validation_tree.tag_configure("review", foreground=COLORS["warning"]); self.validation_tree.tag_configure("fail", foreground=COLORS["error"])
        failures = [name for name, passed, _detail in checks if not passed and name != "READY TO RUN SCOPE"]
        if failures:
            self.validation_passed = False; self.completed.discard(2); self.completed.discard(3); self.completed.discard(4); self.unlocked_step = 2; self.run_button.configure(state="disabled")
            self.continue_run_button.configure(state="disabled"); self.run_readiness.configure(text="Correct the blocking items on Validate before running.", fg=COLORS["error"]); self.status.set("Something needs attention. Check Validation for details."); self._set_companion("attention_party", f"{len(failures)} blocking item(s) need correction.")
            self.target_review_frame.pack_forget()
        else:
            self.validation_passed = True; self.completed.add(2); self.unlocked_step = max(self.unlocked_step, 3); self.run_button.configure(state="normal")
            if self.workflow.get() == "sample_optimization":
                self._refresh_target_review(show_errors=True)
                if not self.target_review_frame.winfo_manager(): self.target_review_frame.pack(fill="x", pady=(0, 8), before=self.validation_tree)
                self.continue_run_button.pack_forget()
                self.status.set("Project validation passed. Review the effective targets.")
                self._append_log("Project validation passed. Ready to review targets.")
            else:
                if not self.continue_run_button.winfo_manager(): self.continue_run_button.pack(side="right")
                self.continue_run_button.configure(state="normal")
                self.status.set("Validation passed. Continue to Run.")
            self.run_readiness.configure(text="✓ Validation passed. This project is ready to run.", fg=COLORS["success"]); self._set_companion("bunny", "Bunny is checking demographic targets...")
        self._update_navigation()

    def _run_selected(self) -> None:
        try:
            config = self._config(); checks = validate(config); failures = [(name, detail) for name, passed, detail in checks if not passed]
            if failures: raise ValueError("Validation must pass first: " + "; ".join(f"{name} — {detail}" for name, detail in failures))
        except Exception as exc: self._show_correction("Run SCOPE", exc); self._show_step(3); return
        action = {"sample_optimization": "optimize", "recruitment_planning": "recruitment_plan", "targeted_swap": "targeted_swap", "targeted_adjustment": "targeted_adjustment"}[config.workflow]
        self._start(WORKFLOWS[config.workflow][0], lambda: run(action, config, self._log))

    def _start(self, label: str, operation) -> None:
        if self.busy: return
        self._show_status_bar()
        self.busy = True; self.run_started = time.monotonic(); self.progress_value = 4; self.progress["value"] = 4; self.run_screen_progress["value"] = 4; self.run_button.configure(state="disabled"); self._update_stages(0); self._update_navigation()
        self.status.set(f"Running {label}…"); self._set_companion("kitty", "Kitty is sorting candidate cases..."); self.elapsed.set("Elapsed 00:00"); self._tick_elapsed()
        def worker() -> None:
            try: result = operation()
            except Exception as exc: self.events.put(("error", (label, str(exc))))
            else: self.events.put(("result", result))
        threading.Thread(target=worker, daemon=True).start()

    def _tick_elapsed(self) -> None:
        if not self.busy: return
        seconds = int(time.monotonic() - self.run_started); self.elapsed.set(f"Elapsed {seconds // 60:02d}:{seconds % 60:02d}"); self.root.after(1000, self._tick_elapsed)

    def _log(self, text: str) -> None: self.events.put(("progress", str(text)))

    def _apply_progress(self, text: str) -> None:
        self._append_technical(text)
        lowered = text.casefold().strip()
        if not lowered.startswith(("traceback", "file ", "warning:", "debug:")) and "opt_engine_" not in lowered and not re.search(r"[a-z]:\\", lowered):
            self._append_log(text)
        match = re.search(r"(?:^|\[|\s)(\d{1,3})\s*%", text)
        if match: value = max(self.progress_value, min(96, int(match.group(1))))
        else: value = min(94, self.progress_value + 12)
        self.progress_value = value; self.progress["value"] = value; self.run_screen_progress["value"] = value; names = self._stage_names(); current = min(len(names) - 1, int(value / max(1, 100 / len(names)))) if names else 0; self._update_stages(current)
        if value < 28: animal, companion = "kitty", "Kitty is sorting candidate cases..."
        elif value < 52: animal, companion = "bunny", "Bunny is checking demographic targets..."
        elif value < 82: animal, companion = "puppy", "Puppy is searching the feasible region..."
        else: animal, companion = "guinea_pig", "Guinea pig is inspecting the outputs..."
        self._set_companion(animal, companion); self.status.set(text[:150] if text.strip() else "SCOPE is running…")

    def _finish_success(self, result) -> None:
        self.busy = False; self.last_result = result; self.progress_value = 100; self.progress["value"] = 100; self.run_screen_progress["value"] = 100; self._update_stages(0, complete=True)
        run_reference = getattr(result, "metrics", {}).get("run_reference")
        if run_reference and run_reference not in self.run_history: self.run_history.append(run_reference)
        self.completed.add(3); self.completed.add(4); self.unlocked_step = 4; self.run_button.configure(state="normal"); self.status.set("Complete! Everything is ready for review."); self._set_companion("success_party", ""); self._show_results(result); self._update_navigation(); self.root.after(3000, self._collapse_success_companion)

    def _collapse_success_companion(self) -> None:
        if not self.busy and self.last_result is not None:
            self._set_companion("guinea_pig", "")
            self.status.set("Complete! Everything is ready for review.")

    def _celebrate(self, frame: int) -> None:
        self._set_companion("success_party", "")

    def _finish_error(self, label: str, detail: str) -> None:
        self.busy = False; self.run_button.configure(state="normal" if self.validation_passed else "disabled"); self._set_companion("attention_party", "Something needs attention. Check Validation for details.")
        short = detail.strip() or "SCOPE could not complete the run"; self.status.set("Something needs attention. Check Validation for details."); self._append_log(f"{label}: {short}")
        messagebox.showerror(label, f"{short}\n\nCheck Validation for details."); self._show_step(3)

    def _show_results(self, result) -> None:
        self.preview_cards.pack_forget()
        for child in self.result_metrics.winfo_children(): child.destroy()
        for child in self.results_buttons.winfo_children(): child.destroy()
        workflow, outputs, metrics = self.workflow.get(), result.outputs, getattr(result, "metrics", {}); config = self._config()
        self.result_summary.set(f"{result.message}\nRun folder: {result.run_dir}")
        if workflow == "recruitment_planning":
            self._metric("Target N", metrics.get("target_n", config.sample_size or "From plan")); self._metric("Age bands", metrics.get("age_bands", "See report")); self._metric("Demographic groups", metrics.get("demographic_target_groups", "See report")); self._metric("OSEP / Clinical", f"{metrics.get('osep_groups', 0)} / {metrics.get('clinical_groups', 0)}")
            self.result_details.configure(text="Recruitment target and profile tables are internally reconciled. If current recruited cases were supplied, Remaining Recruitment Need reflects the updated counts.")
            items = [("Open Field Research Report", outputs.get("recruitment_report")), ("Open Recruitment Targets", outputs.get("target_tables"))]
        elif workflow == "sample_optimization":
            final_n = metrics.get("final_sample_n") or self._workbook_n(outputs.get("optimized_sample")); self._metric("Final sample N", final_n or config.sample_size or "See output"); self._metric("Cases added", metrics.get("cases_added", "See report")); self._metric("Cases replaced", metrics.get("cases_replaced", "See report")); self._metric("Hard constraints", metrics.get("hard_constraints", "Passed"), COLORS["success"])
            self.result_details.configure(text="The optimized sample completed successfully. Review the stratification report for the Before → After target error and alignment summary.")
            items = [("Open Demographic Tables", outputs.get("demographic_tables") or outputs.get("stratification_report")), ("Open OSEP Tables", outputs.get("osep_tables") or outputs.get("osep_clinical_tables")), ("Open Clinical Tables", outputs.get("clinical_tables")), ("Open Final TEID List", outputs.get("teid_list") or outputs.get("optimized_sample")), ("Open Available Pool", outputs.get("available_pool")), ("Open QA Summary", outputs.get("qa_summary"))]
        else:
            before_n = metrics.get("sample_n_before") or self._workbook_n(config.swap_sample_file); after_n = metrics.get("sample_n_after") or self._workbook_n(outputs.get("sample")); self._metric("Sample N", f"{before_n or '?'} → {after_n or '?'}"); self._metric("Constraint status", metrics.get("constraint_status", "Passed"), COLORS["success"]); self._metric("Optimizer", metrics.get("optimizer", "CP-SAT NOT RERUN"), COLORS["targeted_swap"])
            pairs = "\n".join(f"Removed {row['remove']}  →  Added {row['add']}  ·  {row.get('reason') or 'No reason entered'}" for row in config.targeted_swaps)
            self.result_details.configure(text=pairs or "No swap details available.")
            items = [("Open Selected Cases", outputs.get("selected_cases") or outputs.get("sample")), ("Open Available Pool", outputs.get("available_pool")), ("Open Swap Log", outputs.get("swap_log") or outputs.get("audit_log")), ("Open Updated Stratification", outputs.get("updated_stratification") or outputs.get("qa"))]
        for label, path in items:
            if not path or not Path(path).exists(): continue
            button = ttk.Button(self.results_buttons, text=label, style="Secondary.TButton", command=lambda value=path: self._open(value)); button.pack(side="left", padx=(0, 7), pady=3)
        ttk.Button(self.results_buttons, text="Open Run Folder", style="Primary.TButton", command=lambda: self._open(result.run_dir)).pack(side="left", pady=3)
        ttk.Button(self.results_buttons, text="New Run", style="Secondary.TButton", command=self._new_project).pack(side="left", padx=(7, 0), pady=3)
        self._show_step(4)

    def _metric(self, label: str, value, color: str | None = None) -> None:
        card = tk.Frame(self.result_metrics, bg=COLORS["surface"], highlightbackground=COLORS["border"], highlightthickness=1, width=205, height=72); card.pack(side="left", padx=(0, 8)); card.pack_propagate(False)
        tk.Label(card, text=label, bg=COLORS["surface"], fg=COLORS["muted"], font=("Segoe UI", 8)).pack(anchor="w", padx=10, pady=(9, 1)); tk.Label(card, text=str(value), bg=COLORS["surface"], fg=color or COLORS["navy"], font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=10)

    def _workbook_n(self, path) -> int | None:
        if not path or not Path(path).is_file() or Path(path).suffix.lower() not in {".xlsx", ".xls"}: return None
        try:
            with pd.ExcelFile(path) as book:
                preferred = next((name for name in ("final_all", "selected_teids", "final_teid_list") if name in book.sheet_names), book.sheet_names[0]); frame = pd.read_excel(book, sheet_name=preferred, dtype={"teid": str})
            return int(frame["teid"].nunique()) if "teid" in frame else len(frame)
        except Exception: return None

    def _add_swap(self) -> None:
        try:
            self._commit_pending_swap(require_entry=True)
        except ValueError as exc:
            messagebox.showerror("Targeted Case Swap", str(exc))

    def _commit_pending_swap(self, require_entry: bool = False) -> bool:
        row = stage_swap_selection(self.swap_add.get(), self.swap_remove.get(), self.swap_reason.get(), self.swap_rows)
        if row is None:
            if require_entry:
                raise ValueError("Enter both the Add TEID and Remove TEID for this swap selection.")
            return False
        self.swap_rows.append(row)
        if hasattr(self, "swap_tree") and self.swap_tree.winfo_exists():
            self.swap_tree.insert("", "end", values=(row["add"], row["remove"], row["reason"]))
        self.swap_add.set(""); self.swap_remove.set(""); self.swap_reason.set("")
        self._update_swap_selection_summary()
        self._invalidate("Swap selection added. Validate again.")
        return True

    def _update_swap_selection_summary(self) -> None:
        count = len(self.swap_rows)
        self.swap_selection_summary.set(f"{count} swap selection{'s' if count != 1 else ''} added")

    def _remove_swap(self) -> None:
        selected = self.swap_tree.selection()
        if not selected: return
        values = self.swap_tree.item(selected[0], "values"); self.swap_rows = [row for row in self.swap_rows if not (row["add"] == str(values[0]) and row["remove"] == str(values[1]))]; self.swap_tree.delete(selected[0]); self._update_swap_selection_summary(); self._invalidate("Swap list changed. Validate again.")

    def _clear_swaps(self) -> None:
        self.swap_rows = []
        self.swap_remove.set(""); self.swap_add.set(""); self.swap_reason.set("")
        if hasattr(self, "swap_tree"):
            self.swap_tree.delete(*self.swap_tree.get_children())
        self._update_swap_selection_summary()
        self._invalidate("Swap selections cleared. Add selections and review again.")

    def _save_project(self) -> None:
        try: config = self._config(save=False)
        except Exception as exc: self._show_correction("Save Project", exc); return
        destination = self.current_project_file
        if destination is None:
            selected = filedialog.asksaveasfilename(
                title="Save SCOPE Project", initialdir=config.project_dir,
                initialfile=config.path.name, defaultextension=".scope",
                filetypes=[("SCOPE project", "*.scope")],
            )
            if not selected: return
            destination = Path(selected)
        try: config.save(destination)
        except Exception as exc: self._show_correction("Save Project", exc); return
        self.current_project_file = Path(destination); self.project_dirty = False; self._record_recent(self.current_project_file); self._update_project_banner()
        self.status.set(f"Project configuration saved: {self.current_project_file}"); messagebox.showinfo("Save Project", f"SCOPE project configuration saved.\n\n{self.current_project_file}")

    def _open_project(self, path: str | Path | None = None) -> None:
        if not self._confirm_unsaved_changes(): return
        selected = str(path or filedialog.askopenfilename(title="Open SCOPE Project", filetypes=[("SCOPE project", "*.scope"), ("Legacy SCOPE project", "*.json")]))
        if not selected: return
        try:
            self._reset_activity_context()
            config = ProjectConfig.load(selected); self._suspend_dirty = True; self._apply_project(config)
            self.current_project_file = Path(selected); self.project_dirty = False; self._record_recent(Path(selected))
        except Exception as exc: self._show_correction("Open Project", exc); return
        finally: self._suspend_dirty = False
        missing = [f"{role}: {value}" for role, value in config.input_sources().items() if value and not Path(value).exists()]
        self._update_project_banner(); self.status.set(f"Project opened: {selected}" + (f" — {len(missing)} referenced file(s) need attention." if missing else "")); self._show_step(1)
        self._append_log(f"Project opened: {config.project_name or Path(selected).stem}")
        if missing: messagebox.showwarning("Open Project", "The project was restored without running. These referenced files were not found:\n\n" + "\n".join(missing))

    def _apply_project(self, config: ProjectConfig) -> None:
        if config.profile_id not in {profile.id for profile in self.profiles}: raise ValueError("This assessment does not yet have an available SCOPE Test Profile. Contact the SCOPE administrator.")
        self._select_profile(config.profile_id, navigate=False); self._choose_workflow(config.workflow, navigate=False)
        primary_data = config.dataset or (config.swap_case_universe_file if config.workflow in {"targeted_swap", "targeted_adjustment"} else "")
        separate_available = "" if config.swap_available_file == config.swap_sample_file else config.swap_available_file
        for variable, value in ((self.project_name, config.project_name), (self.project_dir, config.project_dir), (self.output_dir, config.output_dir), (self.output_selection, config.output_selection), (self.sample_size, "" if config.sample_size is None else str(config.sample_size)), (self.language, config.language), (self.study_type, config.study_type), (self.dataset, primary_data), (self.master_file, config.master_file), (self.fr_report_file, config.fr_teid_file), (self.scored_dir, config.scored_dir), (self.case_swap_dir, config.case_swap_dir), (self.target_file, config.target_file), (self.recruitment_plan_file, config.recruitment_plan_file), (self.recruitment_actuals_file, config.recruitment_actuals_file), (self.swap_sample, config.swap_sample_file), (self.swap_available, separate_available), (self.swap_universe, config.swap_case_universe_file), (self.swap_sample_data, config.swap_sample_data_file), (self.swap_available_data, config.swap_available_data_file)):
            variable.set(value)
        self.use_existing_sample.set(config.use_existing_sample or bool(config.master_file)); self.override_age_defaults.set(bool(config.age_bands)); self._toggle_age_override()
        self.include_census.set(config.include_census); self.include_osep.set(config.include_osep); self.include_clinical.set(config.include_clinical); self.clinical_seeding.set(config.clinical_seeding)
        for name, variable in self.form_vars.items(): variable.set(name in config.forms)
        for name, variable in self.demographic_vars.items(): variable.set(name in config.demographics)
        if config.demographic_tables_configured:
            self._populate_report_table_selector({
                "one_way": config.demographic_tables_1way,
                "two_way": config.demographic_tables_2way,
                "three_way": config.demographic_tables_3way,
            })
        self.canonical_target_rows = list(config.canonical_targets); self.targets_confirmed.set(config.target_review_confirmed)
        self.candidate_datasets = [dict(item) for item in config.candidate_datasets]
        if not self.candidate_datasets and config.dataset:
            self.candidate_datasets = [{"path": config.dataset, "source_type": config.study_type or "STDZ", "source_role": "Primary", "priority": 1}]
        self.target_overrides = [dict(item) for item in config.target_overrides]
        self.swap_rows = list(config.targeted_swaps); self._render_workflow_inputs()
        self.run_history = list(config.run_history)
        normalized_locks = normalize_lock_fields(self._profile(), config.locks)
        for name, variable in self.lock_vars.items(): variable.set(name in normalized_locks)
        self.age_text.delete("1.0", "end")
        for band in config.age_bands:
            label = band.get("label") or f"{int(band['min_months']) // 12}:{int(band['min_months']) % 12}-{int(band['max_months']) // 12}:{int(band['max_months']) % 12}"; target = config.target_n_by_agegrp.get(str(band.get("group")), "")
            self.age_text.insert("end", label + (f" = {target}" if target != "" else "") + "\n")
        self.age_text.edit_modified(False)
        if config.workflow == "sample_optimization": self._refresh_target_review(show_errors=False)
        self.completed = {0, 1}; self.unlocked_step = 2; self.validation_passed = False; self._update_navigation()

    def _recent_file(self) -> Path:
        base = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else ROOT
        return base / "scope_user" / "recent_projects.json"

    def _load_recent(self) -> list[str]:
        try: return [value for value in json.loads(self._recent_file().read_text(encoding="utf-8")) if Path(value).is_file()][:8]
        except Exception: return []

    def _record_recent(self, path: Path) -> None:
        recent = [str(path)] + [value for value in self._load_recent() if Path(value) != path]; target = self._recent_file(); target.parent.mkdir(parents=True, exist_ok=True); target.write_text(json.dumps(recent[:8], indent=2), encoding="utf-8"); self._refresh_recent_menu()

    def _refresh_recent_menu(self) -> None:
        if not hasattr(self, "recent_menu"): return
        self.recent_menu.delete(0, "end"); recent = self._load_recent()
        if not recent: self.recent_menu.add_command(label="No recent projects", state="disabled")
        for path in recent: self.recent_menu.add_command(label=Path(path).stem + " — " + path, command=lambda value=path: self._open_project(value))

    def _confirm_unsaved_changes(self) -> bool:
        if not self.project_dirty: return True
        decision = messagebox.askyesnocancel("Unsaved SCOPE Project", "Save changes to the current SCOPE project configuration first?")
        if decision is None: return False
        if decision:
            self._save_project()
            return not self.project_dirty
        return True

    def _track_project_changes(self) -> None:
        variables = (
            self.workflow, self.profile_id, self.project_name, self.project_dir, self.output_dir,
            self.sample_size, self.language, self.study_type, self.include_census, self.include_osep,
            self.include_clinical, self.clinical_seeding, self.dataset, self.master_file,
            self.scored_dir, self.case_swap_dir, self.target_file, self.fr_report_file,
            self.recruitment_plan_file, self.recruitment_actuals_file, self.swap_sample,
            self.swap_available, self.swap_universe, self.swap_sample_data, self.swap_available_data,
            self.output_selection, self.targets_confirmed, self.use_existing_sample, self.override_age_defaults,
        )
        for variable in variables: variable.trace_add("write", self._mark_project_dirty)
        for variable in (self.sample_size, self.study_type, self.language, self.include_census, self.include_osep, self.include_clinical, self.clinical_seeding):
            variable.trace_add("write", self._target_inputs_changed)

    def _target_inputs_changed(self, *_args) -> None:
        if self._suspend_dirty: return
        self._invalidate("Project inputs changed. Effective targets will be rebuilt automatically.")

    def _mark_project_dirty(self, *_args) -> None:
        if self._suspend_dirty: return
        self.project_dirty = True; self._update_project_banner()

    def _update_project_banner(self) -> None:
        name = self.project_name.get().strip() or "Untitled"
        if self.project_dirty:
            text = f"Project: {name}  •  Unsaved changes"
        elif self.current_project_file:
            text = f"Project: {name}\nSaved: {self.current_project_file}"
        else:
            text = f"Project: {name}  •  Not saved"
        self.project_banner.set(text)

    def _invalidate(self, message: str) -> None:
        self._mark_project_dirty()
        self.validation_passed = False; self.completed.discard(2); self.completed.discard(3); self.completed.discard(4); self.unlocked_step = min(self.unlocked_step, 2); self.run_button.configure(state="disabled"); self.continue_run_button.configure(state="disabled"); self.run_readiness.configure(text="Validate the project before running.", fg=COLORS["warning"]); self.status.set(message); self._update_navigation()
        if self.workflow.get() == "sample_optimization" and self.profile_id.get():
            self.targets_confirmed.set(False)
            if not self.target_refresh_pending:
                self.target_refresh_pending = True
                self.root.after_idle(self._automatic_target_refresh)

    def _automatic_target_refresh(self) -> None:
        self.target_refresh_pending = False
        self._refresh_target_review(show_errors=False)

    def _set_companion(self, name: str, text: str) -> None:
        if name != self.companion_name:
            self.companion_name = name
            self.companion_frame = 0
        self._render_companion()
        self.companion.set(text)

    def _render_companion(self) -> None:
        name = self.companion_name
        motion_name = f"{name}_motion"
        image_name = motion_name if self.companion_frame and motion_name in self.animal_images else name
        image = self.animal_images.get(image_name)
        self.animal_label.configure(image=image or "", text="" if image else name.replace("_", " ").title())
        self.animal_label.image = image

    def _animate_companion(self) -> None:
        """Give the small status helper a restrained two-frame motion."""
        motion_name = f"{self.companion_name}_motion"
        if self.statusbar.winfo_manager() and motion_name in self.animal_images:
            self.companion_frame = 1 - self.companion_frame
            self._render_companion()
        self.root.after(600, self._animate_companion)

    def _show_status_bar(self) -> None:
        if not self.activity.winfo_manager():
            self.activity.pack(fill="x", padx=(220, 18), pady=(0, 4))
        if not self.statusbar.winfo_manager():
            self.statusbar.pack(fill="x")

    def _toggle_activity(self) -> None:
        self.activity_expanded = not self.activity_expanded
        self.activity_toggle.configure(text=("▾ Activity" if self.activity_expanded else "▸ Activity"))
        if self.activity_expanded:
            self.activity_content.pack(fill="x")
        else:
            self.activity_content.pack_forget()

    def _open_technical_log(self) -> None:
        window = tk.Toplevel(self.root); window.title("SCOPE Technical Log"); window.geometry("900x520"); window.configure(bg=COLORS["background"])
        tk.Label(window, text="Technical / System Log", bg=COLORS["background"], fg=COLORS["navy"], font=("Segoe UI", 15, "bold")).pack(anchor="w", padx=16, pady=(14, 4))
        tk.Label(window, text="Profile discovery, startup diagnostics, resource paths, and engine details.", bg=COLORS["background"], fg=COLORS["muted"]).pack(anchor="w", padx=16, pady=(0, 8))
        text_widget = scrolledtext.ScrolledText(window, wrap="word", bg="#FBFAF7", fg=COLORS["text"], font=("Consolas", 9)); text_widget.pack(fill="both", expand=True, padx=16, pady=(0, 12))
        text_widget.insert("1.0", "\n".join(self.technical_messages) or "No technical messages recorded."); text_widget.configure(state="disabled")
        ttk.Button(window, text="Close", style="Primary.TButton", command=window.destroy).pack(anchor="e", padx=16, pady=(0, 14))

    def _reset_activity_context(self) -> None:
        self.user_messages = []
        if hasattr(self, "log"):
            self.log.configure(state="normal"); self.log.delete("1.0", "end"); self.log.configure(state="disabled")

    def _hide_status_bar(self) -> None:
        if self.statusbar.winfo_manager():
            self.statusbar.pack_forget()
        if self.activity.winfo_manager():
            self.activity.pack_forget()
        self.animal_label.configure(image="", text="")
        self.animal_label.image = None
        self.companion_name = ""
        self.companion_frame = 0
        self.companion.set("")
        self.progress_value = 0
        self.progress["value"] = 0
        self.run_screen_progress["value"] = 0
        self.elapsed.set("Elapsed 00:00")

    def _reset_project_scope(self) -> None:
        """Restore the visible project choices to the selected profile defaults."""
        self.sample_size.set("")
        profile = self._profile() if self.profile_id.get() else None
        languages = list(profile.languages or ["English", "Spanish"]) if profile else ["English", "Spanish"]
        defaults = profile_study_type_defaults(profile) if profile else []
        studies = list(profile.study_types or ["STDZ", "Clinical", "Validity"]) if profile else ["STDZ", "Clinical", "Validity"]
        for value in defaults:
            if value not in studies: studies.append(value)
        self.language.set(languages[0] if languages else "")
        self.study_type.set(defaults[0] if defaults else (studies[0] if studies else ""))
        for variable in self.form_vars.values():
            variable.set(True)
        for variable in self.demographic_vars.values():
            variable.set(True)
        for kind in self.report_table_lists:
            self._set_report_table_family(kind, True)
        self.include_census.set(True); self.include_osep.set(True); self.include_clinical.set(True); self.clinical_seeding.set(False)
        self.use_existing_sample.set(False); self.override_age_defaults.set(False); self.age_frame.pack_forget()
        self.age_text.delete("1.0", "end")
        self.canonical_target_rows = []; self.targets_confirmed.set(False); self._populate_target_review()
        self._invalidate("Project scope selections reset to the Test Profile defaults.")

    def _new_project(self) -> None:
        if self.busy:
            return
        if not self._confirm_unsaved_changes(): return
        self.profile_catalog = discover_profile_catalog(); self.profiles = self.profile_catalog.user_profiles; self.profile_errors = self.profile_catalog.errors
        self._suspend_dirty = True; self.current_project_file = None
        self.workflow.set(""); self.profile_id.set(""); self.project_name.set("")
        default_project_dir = Path.home() / "Documents" / "SCOPE Projects"
        try: default_project_dir.mkdir(parents=True, exist_ok=True)
        except OSError: default_project_dir = Path.cwd()
        self.project_dir.set(str(default_project_dir)); self.output_dir.set(str(default_project_dir / "scope_outputs")); self.sample_size.set("")
        self.output_selection.set("both")
        self.use_existing_sample.set(False); self.override_age_defaults.set(False); self.age_frame.pack_forget()
        for variable in (self.dataset, self.master_file, self.scored_dir, self.case_swap_dir, self.target_file, self.fr_report_file, self.recruitment_plan_file, self.recruitment_actuals_file, self.swap_sample, self.swap_available, self.swap_universe, self.swap_sample_data, self.swap_available_data):
            variable.set("")
        self.swap_rows = []; self.candidate_datasets = []; self.target_overrides = []; self.canonical_target_rows = []; self.targets_confirmed.set(False); self.last_result = None; self.run_history = []
        self.age_text.delete("1.0", "end")
        self.age_text.edit_modified(False)
        self.completed.clear(); self.unlocked_step = 0; self.validation_passed = False
        for card in self.workflow_cards.values(): card.set_selected(None)
        self._build_profile_cards(); self._refresh_profile_detail(); self._update_profile_continue_state(); self._hide_status_bar(); self._populate_target_review(); self._reset_stages()
        self.validation_tree.delete(*self.validation_tree.get_children())
        self.result_summary.set("No completed run yet."); self.status.set("Choose what SCOPE should do.")
        self._reset_activity_context()
        self.project_dirty = False; self._suspend_dirty = False; self._update_project_banner()
        self._show_step(0, force=True)

    def _show_correction(self, title: str, error: Exception) -> None:
        text = str(error).strip() or "SCOPE could not complete this step. Review the selected inputs and try again."; messagebox.showerror(title, text); self.status.set(f"{title} needs attention.")

    def _open(self, path) -> None:
        if path and Path(path).exists(): os.startfile(Path(path))
        else: messagebox.showinfo("Results", "That output was not produced by this run.")

    def _append_log(self, text: str) -> None:
        self.user_messages.append(str(text))
        self.log.configure(state="normal"); self.log.insert("end", text + "\n"); self.log.see("end"); self.log.configure(state="disabled")

    def _append_technical(self, text: str) -> None:
        self.technical_messages.append(str(text))

    def _log_profile_catalog(self) -> None:
        self._append_technical(f"SCOPE Test Profiles discovered: {len(self.profiles)}")
        for profile in self.profiles:
            self._append_technical(
                f"{profile.id}  status={profile.status}  "
                f"recruitment={profile.supports('recruitment_planning')}  "
                f"optimization={profile.supports('sample_optimization')}  "
                f"swap={profile.supports('targeted_swap')}"
            )
        for error in self.profile_errors:
            self._append_technical(f"PROFILE ERROR  {error.path}: {error.message}")

    def _drain(self) -> None:
        while not self.events.empty():
            kind, value = self.events.get_nowait()
            if kind == "progress": self._apply_progress(value)
            elif kind == "result": self._finish_success(value)
            elif kind == "error": self._finish_error(*value)
        self.root.after(100, self._drain)


if __name__ == "__main__":
    SCOPEApp()
