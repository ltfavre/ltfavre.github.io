from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from scope.configure import ProjectConfig
from scope.engine import apply_targeted_swaps, run
from scope.output import write_targeted_swap_outputs
from scope.profile_results import profile_sample
from scope.profiles import load_profile


ROOT = Path(__file__).resolve().parents[1]


def test_normal_gui_has_no_profile_authoring_and_uses_bottom_navigation():
    source = (ROOT / "gui" / "app.py").read_text(encoding="utf-8")
    assert "Map a New Assessment" not in source
    assert "save_user_profile" not in source
    assert "def _bottom_navigation" in source
    assert (ROOT / "tools" / "scope_profile_builder.py").is_file()


def test_scope_project_extension(tmp_path):
    config = ProjectConfig(profile_id="BASC4_NB", project_dir=str(tmp_path), project_name="BASC4NB MOOD")
    assert config.save().name == "BASC4NB_MOOD.scope"


def test_wisc_reporting_spec_generates_every_required_cross_tab():
    profile = load_profile("WISC6")
    sample = pd.DataFrame({
        "agegrp": [6, 6, 7], "sex": ["F", "M", "F"], "eth": ["A", "B", "A"],
        "edl": ["1", "2", "1"], "region": ["NE", "SO", "NE"],
    })
    tables = profile_sample(sample, profile.demographics, [], profile.required_demo_tables())
    cross_tabs = [name for name in tables if name.startswith("Cross_")]
    assert len(profile.required_demo_tables()) == 10
    assert len(cross_tabs) == 10


def test_targeted_swap_writes_distinct_deliverables(tmp_path):
    selected = pd.DataFrame({"teid": ["1"], "vform": ["TRS_C"], "agegrp": [3], "sex": ["F"], "eth": ["A"], "edl": ["1"], "region": ["NE"]})
    available = pd.DataFrame({"teid": ["2"], "vform": ["TRS_C"], "agegrp": [3], "sex": ["F"], "eth": ["B"], "edl": ["1"], "region": ["NE"]})
    selected_file, available_file = tmp_path / "selected.xlsx", tmp_path / "available.xlsx"
    selected.to_excel(selected_file, sheet_name="final_all", index=False)
    available.to_excel(available_file, sheet_name="available_pool", index=False)
    result = apply_targeted_swaps(selected_file, available_file, [{"remove": "1", "add": "2", "reason": "QA"}], tmp_path, load_profile("BASC4_NB"))
    for key in ("selected_cases", "available_pool", "swap_log", "updated_stratification"):
        assert result.outputs[key].is_file()
    log = pd.read_excel(result.outputs["swap_log"], sheet_name="Swap Log")
    assert {"Run ID", "Removed TEID", "Added TEID", "Reason", "Validation Status", "Date/Time"}.issubset(log.columns)
    summary = pd.read_excel(result.outputs["swap_log"], sheet_name="Run Summary")
    assert "CP-SAT NOT RERUN" in summary["Value"].astype(str).tolist()


def test_targeted_swap_teid_outputs_do_not_duplicate_wide_participant_data(tmp_path):
    rows = 40
    wide_columns = {f"item_{index}": list(range(rows)) for index in range(1000)}
    selected = pd.DataFrame({"teid": [str(index) for index in range(rows)], **wide_columns})
    available = pd.DataFrame({"teid": [str(index + rows) for index in range(rows)], **wide_columns})
    outputs = write_targeted_swap_outputs(
        tmp_path, selected, available, pd.DataFrame([{"status": "passed"}]), {}, rows,
        output_selection="teid_lists",
    )

    selected_output = pd.read_excel(outputs["selected_cases"], sheet_name="final_all", dtype={"teid": str})
    available_output = pd.read_excel(outputs["available_pool"], sheet_name="Available Pool", dtype={"teid": str})
    assert selected_output.columns.tolist() == ["teid"]
    assert available_output.columns.tolist() == ["teid"]
    assert selected_output.teid.tolist() == selected.teid.tolist()
    assert available_output.teid.tolist() == available.teid.tolist()


def test_recruitment_report_has_diff_and_negative_values_are_red(tmp_path):
    targets = tmp_path / "targets.xlsx"
    actuals = tmp_path / "actuals.xlsx"
    pd.DataFrame([
        {"Dimension": "Sex", "Category": "Female", "Target_N": 2},
        {"Dimension": "Sex", "Category": "Male", "Target_N": 2},
    ]).to_excel(targets, sheet_name="Recruitment Targets", index=False)
    pd.DataFrame([
        {"Dimension": "Sex", "Category": "Female", "Actual_N": 1},
        {"Dimension": "Sex", "Category": "Male", "Actual_N": 3},
    ]).to_excel(actuals, index=False)
    config = ProjectConfig(
        profile_id="BASC4_NB", project_dir=str(tmp_path), output_dir=str(tmp_path),
        workflow="recruitment_planning", sample_size=4,
        recruitment_plan_file=str(targets), recruitment_actuals_file=str(actuals),
    )
    result = run("recruitment_plan", config)
    frame = pd.read_excel(result.outputs["recruitment_report"], sheet_name="Field Research Report")
    assert "Diff" in frame.columns
    assert frame.loc[frame["Category"].eq("Female"), "Diff"].iloc[0] < 0
    workbook = load_workbook(result.outputs["recruitment_report"])
    sheet = workbook["Field Research Report"]
    diff_column = next(cell.column for cell in sheet[1] if cell.value == "Diff")
    negative_cell = next(sheet.cell(row, diff_column) for row in range(2, sheet.max_row + 1) if isinstance(sheet.cell(row, diff_column).value, (int, float)) and sheet.cell(row, diff_column).value < 0)
    assert negative_cell.font.color and negative_cell.font.color.rgb.endswith("B84444")
