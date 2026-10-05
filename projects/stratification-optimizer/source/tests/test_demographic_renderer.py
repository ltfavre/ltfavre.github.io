from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from scope.demographic_renderer import render_demographic_workbook
from scope.profiles import TestProfile as ScopeTestProfile, load_profile


TEMPLATE = Path(__file__).resolve().parents[1] / "profiles" / "_shared" / "templates" / "SCOPE_Demographic_Table_Templates.xlsx"


def _profile() -> ScopeTestProfile:
    variables = {
        "GroupA": {
            "canonical": "group_a", "field": "a", "display": "Primary Group",
            "levels": [{"value": "a1", "label": "Alpha"}, {"value": "a2", "label": "Beta"}],
        },
        "GroupB": {
            "canonical": "group_b", "field": "b", "display": "Secondary Group",
            "levels": [{"value": "b1", "label": "One"}, {"value": "b2", "label": "Two"}],
        },
        "GroupC": {
            "canonical": "group_c", "field": "c", "display": "Third Group",
            "levels": [{"value": "c1", "label": "First"}, {"value": "c2", "label": "Second"}],
        },
    }
    return ScopeTestProfile(
        id="TEST", display_name="Renderer Test", version="1", source_config="", forms=[],
        demographics=[], target_flags=[], reporting_variables=variables,
        reporting={"required_tables": {
            "demographics": ["GroupA", "GroupB"],
            "demo_cross_tabs": ["GroupA × GroupB"],
            "demo_three_way": ["GroupA × GroupB × GroupC"],
        }},
        resources={"demographic_template": str(TEMPLATE)},
    )


def _sample() -> pd.DataFrame:
    return pd.DataFrame([
        {"a": "a1", "b": "b1", "c": "c1"},
        {"a": "a1", "b": "b1", "c": "c2"},
        {"a": "a1", "b": "b2", "c": "c1"},
        {"a": "a1", "b": "b2", "c": "c2"},
        {"a": "a2", "b": "b1", "c": "c1"},
        {"a": "a2", "b": "b1", "c": "c1"},
        {"a": "a2", "b": "b1", "c": "c2"},
        {"a": "a2", "b": "b2", "c": "c1"},
    ])


def _targets() -> list[dict]:
    return [
        {"category": "GroupA", "group": "a1", "target": 75, "unit": "%"},
        {"category": "GroupA", "group": "a2", "target": 25, "unit": "%"},
        {"category": "GroupB", "group": "b1", "target": 75, "unit": "%"},
        {"category": "GroupB", "group": "b2", "target": 25, "unit": "%"},
        {"category": "GroupC", "group": "c1", "target": 75, "unit": "%"},
        {"category": "GroupC", "group": "Form A / age group 1 / c2", "target": 25, "unit": "%"},
        {"category": "GroupC", "group": "Form B / age group 2 / c2", "target": 25, "unit": "%"},
    ]


def test_generic_renderers_preserve_profile_order_calculations_and_template(tmp_path):
    before = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
    output = render_demographic_workbook(tmp_path / "demographics.xlsx", _sample(), _profile(), _targets())
    assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == before

    workbook = load_workbook(output)
    assert workbook.sheetnames[:2] == ["Primary Group", "Secondary Group"]
    assert workbook.sheetnames[2].startswith("Primary Group × Secondary")
    assert workbook.sheetnames[3].startswith("Primary Group × Secondary") and workbook.sheetnames[3].endswith("(2)")
    assert not {"README", "TEMPLATE_1WAY", "TEMPLATE_2WAY", "TEMPLATE_3WAY"} & set(workbook.sheetnames)

    one_way = workbook["Primary Group"]
    assert [one_way.cell(row, 1).value for row in (4, 5)] == ["Alpha", "Beta"]
    assert one_way["D4"].value == 4
    assert one_way["E4"].value == 0.5
    assert one_way.max_column == 5
    assert all(cell.value != "Diff" for row in one_way.iter_rows() for cell in row)

    two_way = workbook[workbook.sheetnames[2]]
    assert [two_way.cell(3, column).value for column in (3, 4)] == ["One", "Two"]
    assert [two_way.cell(row, 2).value for row in range(4, 7)] == ["Target %", "Sample %", "Sample N"]
    assert two_way["A4"].value == "Alpha"
    assert two_way["A7"].value == "Beta"
    assert two_way["C5"].value == 0.5
    assert two_way["C6"].value == 2
    assert all(cell.value != "Diff" for row in two_way.iter_rows() for cell in row)

    three_way = workbook[workbook.sheetnames[3]]
    section_rows = [row for row in range(1, three_way.max_row + 1) if str(three_way.cell(row, 1).value or "").startswith("Primary Group:")]
    assert [three_way.cell(row, 1).value for row in section_rows] == ["Primary Group: Alpha", "Primary Group: Beta"]
    for section_row in section_rows:
        assert [three_way.cell(section_row + 1, column).value for column in (3, 4)] == ["First", "Second"]
        assert [three_way.cell(section_row + row, 2).value for row in range(2, 5)] == ["Target %", "Sample %", "Sample N"]
    assert three_way["C6"].value == 0.5
    assert three_way["C7"].value == 1
    assert all(cell.value != "Diff" for row in three_way.iter_rows() for cell in row)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert "{{" not in str(cell.value or "")
                assert "PLACEHOLDER" not in str(cell.value or "").upper()


def test_wisc_profile_generates_every_required_demographic_worksheet(tmp_path):
    profile = load_profile("WISC6")
    definitions = profile.reporting_variables or {}
    row = {}
    for specification in definitions.values():
        level = specification["levels"][0]
        row[specification["field"]] = level.get("value") if isinstance(level, dict) else level
    output = render_demographic_workbook(tmp_path / "wisc.xlsx", pd.DataFrame([row]), profile, [])
    workbook = load_workbook(output, read_only=True)
    required = profile.required_demographic_tables()
    assert len(workbook.sheetnames) == sum(len(values) for values in required.values()) == 16
    assert len(set(workbook.sheetnames)) == 16
    assert all(not name.startswith("TEMPLATE_") for name in workbook.sheetnames)
    assert all("{" not in name for name in workbook.sheetnames)


def test_results_preview_exposes_profile_driven_template_structures():
    profile = load_profile("BASC4_NB")
    previews = profile.output_previews("sample_optimization")
    template_previews = {item.get("template_sheet"): item for item in previews if item.get("kind") == "demographic_template"}
    assert {"TEMPLATE_1WAY", "TEMPLATE_2WAY"}.issubset(template_previews)
    for item in template_previews.values():
        assert Path(str(item["template"])).is_file()
        assert item["tables"]


def test_renderer_honors_project_table_selections(tmp_path):
    selected = {
        "one_way": ["GroupA"],
        "two_way": ["GroupA × GroupB"],
        "three_way": [],
    }
    output = render_demographic_workbook(tmp_path / "selected.xlsx", _sample(), _profile(), _targets(), selected_tables=selected)
    workbook = load_workbook(output, read_only=True)
    assert len(workbook.sheetnames) == 2
    assert workbook.sheetnames[0] == "Primary Group"
    assert workbook.sheetnames[1].startswith("Primary Group × Secondary")


def test_fully_crossed_counts_roll_up_to_conditional_template_targets(tmp_path):
    import json

    targets = []
    for group_a, group_b, count in (("a1", "b1", 30), ("a1", "b2", 10), ("a2", "b1", 20), ("a2", "b2", 40)):
        targets.append({
            "category": "Group A × Group B", "group": f"group_a={group_a} | group_b={group_b}",
            "target": count, "unit": "N", "criteria": json.dumps({"group_a": group_a, "group_b": group_b}),
        })
    selected = {"one_way": ["GroupA"], "two_way": ["GroupA × GroupB"], "three_way": []}
    output = render_demographic_workbook(tmp_path / "crossed.xlsx", _sample(), _profile(), targets, selected_tables=selected)
    workbook = load_workbook(output, data_only=True)
    one_way = workbook[workbook.sheetnames[0]]
    assert one_way["B4"].value == 40
    assert one_way["C4"].value == 0.4
    two_way = workbook[workbook.sheetnames[1]]
    assert two_way["C4"].value == 0.75
    assert two_way["D4"].value == 0.25
