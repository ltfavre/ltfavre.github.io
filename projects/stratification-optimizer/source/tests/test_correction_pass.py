from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from scope.canonical import build_target_specification, prepare_canonical_inputs
from scope.configure import ProjectConfig
from scope.output import write_recruitment_workbook
from scope.profiles import load_profile
from scope.recruitment import (
    build_recruitment_plan,
    profile_age_bands,
    targets_from_approved_template,
    validate_actual_age_mapping,
    validate_profile_age_targets,
)


def _cases(path: Path, rows: list[tuple[str, str]]) -> Path:
    frame = pd.DataFrame([
        {
            "teid": teid, "vform": "TRS_C", "agey": 8, "agegrp": 3,
            "sex": "F", "eth": "WH", "edl": "3", "region": "MW",
            "testlang": "English", "study": source_type,
            "tmp_ADHD": 0, "tmp_EBD": 0, "tmp_ASD": 0, "tmp_DD": 0,
            "tmp_SLD": 0, "tmp_SLI": 0, "tmp_clin": 0,
        }
        for teid, source_type in rows
    ])
    frame.to_excel(path, index=False)
    return path


def test_multiple_candidate_sources_preserve_provenance_and_resolve_duplicates(tmp_path: Path) -> None:
    primary = _cases(tmp_path / "stdz.xlsx", [("1", "STDZ"), ("2", "STDZ")])
    supplemental = _cases(tmp_path / "nb395.xlsx", [("2", "NB395"), ("3", "NB395")])
    config = ProjectConfig(
        profile_id="BASC4_NB", project_dir=str(tmp_path), workflow="sample_optimization",
        candidate_datasets=[
            {"path": str(primary), "source_type": "STDZ", "source_role": "Primary", "priority": 1},
            {"path": str(supplemental), "source_type": "NB395", "source_role": "Supplemental", "priority": 2},
        ],
        forms=["TRS_C"], sample_size=2,
    )
    prepared = prepare_canonical_inputs(config, load_profile("BASC4_NB"))
    assert prepared.case_universe.SCOPE_ID.tolist() == ["1", "2", "3"]
    assert {"SCOPE_SOURCE_FILE", "SCOPE_SOURCE_TYPE", "SCOPE_SOURCE_ROLE", "SCOPE_SOURCE_PRIORITY"}.issubset(prepared.case_universe)
    duplicate = prepared.case_universe.loc[prepared.case_universe.SCOPE_ID.eq("2")].iloc[0]
    assert duplicate.SCOPE_SOURCE_FILE == "stdz.xlsx"
    assert duplicate.SCOPE_SOURCE_ROLE == "Primary"
    assert set(prepared.current_sample.SCOPE_ID) == {"1", "2"}
    assert set(prepared.available_pool.SCOPE_ID) == {"3"}
    assert any("Resolved 1 duplicate" in note for note in prepared.notes)
    loaded = ProjectConfig.load(config.save())
    assert loaded.candidate_datasets == config.candidate_datasets


def test_target_edit_is_project_override_and_does_not_change_profile(tmp_path: Path) -> None:
    profile = load_profile("BASC4_NB")
    baseline = ProjectConfig(profile_id=profile.id, project_dir=str(tmp_path), workflow="sample_optimization")
    original = build_target_specification(baseline, profile)
    first = original.iloc[0]
    config = ProjectConfig(
        profile_id=profile.id, project_dir=str(tmp_path), workflow="sample_optimization",
        target_overrides=[{
            "category": first.category, "group": first.group, "rule": first.rule,
            "target": 20, "reason": "Project-approved adjustment",
        }],
    )
    effective = build_target_specification(config, profile)
    edited = effective.loc[
        effective.category.eq(first.category) & effective.group.eq(first.group) & effective.rule.eq(first.rule)
    ].iloc[0]
    assert edited.target == 20
    assert edited.profile_value == first.target
    assert edited.source == "Project Override"
    assert edited.override is True or edited.override == 1
    assert build_target_specification(baseline, profile).iloc[0].target == first.target
    assert ProjectConfig.load(config.save()).target_overrides == config.target_overrides


@pytest.mark.parametrize(
    ("profile_id", "study_type", "filename", "expected_first", "expected_last"),
    [
        ("WISC6", "WISC6_STDZ_P", "WISC6_STDZ_Field_Research_Report.xlsx", "6:0-6:11", "16:0-16:11"),
        ("WISC6", "Clinical", "WISC6_Clinical_Field_Research_Report.xlsx", "6:0-6:11", "16:0-16:11"),
        ("WISC6", "OSEP", "WISC6_OSEP_Field_Research_Report.xlsx", "6:0-6:11", "16:0-16:11"),
        ("CELF6", "Standardization", "CELF6_STDZ_Field_Research_Report.xlsx", "5:0-5:5", "17:0-21:11"),
        ("CELF6", "Clinical", "CELF6_Clinical_Field_Research_Report.xlsx", "5:0-8:11", "9:0-21:11"),
    ],
)
def test_profile_study_type_selects_approved_template_and_exact_age_bands(
    profile_id: str, study_type: str, filename: str, expected_first: str, expected_last: str,
) -> None:
    profile = load_profile(profile_id)
    template = Path(profile.report_template_for("recruitment_planning", study_type))
    assert template.name == filename
    bands = profile_age_bands(profile, study_type)
    assert bands[0] == expected_first and bands[-1] == expected_last
    targets, total_n = targets_from_approved_template(template, profile, study_type, 300)
    age = targets.loc[targets.Dimension.eq("Age")]
    assert age.Category.tolist() == bands
    assert age.Target_N.sum() == total_n == 300
    assert all(part.Target_N.sum() == 300 for _name, part in targets.groupby("Dimension"))


@pytest.mark.parametrize(("profile_id", "study_type"), [("WISC6", "WISC6_STDZ_P"), ("CELF6", "Standardization")])
def test_native_recruitment_renderer_preserves_workbook_structure_and_populates_targets(
    tmp_path: Path, profile_id: str, study_type: str,
) -> None:
    profile = load_profile(profile_id)
    template = Path(profile.report_template_for("recruitment_planning", study_type))
    target_frame, total_n = targets_from_approved_template(template, profile, study_type, 100)
    outputs = build_recruitment_plan(target_frame, total_n, max_seconds=10)
    destination = write_recruitment_workbook(tmp_path / f"{profile_id}.xlsx", outputs, str(template))
    source_book = load_workbook(template, read_only=False, data_only=False)
    result_book = load_workbook(destination, read_only=False, data_only=False)
    assert result_book.sheetnames == source_book.sheetnames
    assert list(result_book["Summary"].merged_cells.ranges) == list(source_book["Summary"].merged_cells.ranges)
    source_formulas = [cell.coordinate for row in source_book["Summary"].iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    result_formulas = [cell.coordinate for row in result_book["Summary"].iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert result_formulas == source_formulas
    aurora = result_book["aurora_data"]
    header_row = next(row for row in range(1, 20) if str(aurora.cell(row, 1).value).casefold() == "ageband")
    total_column = next(column for column in range(1, aurora.max_column + 1) if str(aurora.cell(header_row, column).value).casefold() == "total")
    first_target_row = next(row for row in range(header_row + 1, aurora.max_row + 1) if str(aurora.cell(row, 3).value).casefold() == "target n")
    assert aurora.cell(first_target_row, total_column).value == int(target_frame.loc[target_frame.Dimension.eq("Age")].iloc[0].Target_N)


def test_unmapped_recruitment_age_blocks_run(tmp_path: Path) -> None:
    progress = tmp_path / "progress.csv"
    pd.DataFrame({"agey": [6, 22], "sex": ["F", "M"]}).to_csv(progress, index=False)
    with pytest.raises(ValueError, match="do not map to exactly one approved age band"):
        validate_actual_age_mapping(progress, load_profile("WISC6"), "WISC6_STDZ_P")


def test_project_amendment_must_keep_approved_age_bands(tmp_path: Path) -> None:
    amendment = tmp_path / "amendment.csv"
    pd.DataFrame([
        {"Dimension": "Age", "Category": "6:0-6:11", "Target_N": 100},
        {"Dimension": "Age", "Category": "7:0-7:11", "Target_N": 100},
    ]).to_csv(amendment, index=False)
    with pytest.raises(ValueError, match="exactly match the approved profile order"):
        validate_profile_age_targets(amendment, 200, load_profile("WISC6"), "WISC6_STDZ_P")
