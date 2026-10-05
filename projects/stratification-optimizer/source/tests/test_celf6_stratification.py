from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from scope.celf6_stratification import render_celf6_stratification_workbook
from scope.profiles import load_profile


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _template_roster(template: Path) -> list[str]:
    workbook = load_workbook(template, read_only=True, data_only=True)
    sheet = workbook["TEID_List"]
    return [
        str(sheet.cell(row, 1).value).removesuffix(".0")
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, 1).value is not None
    ]


def test_celf6_targeted_swap_output_preserves_reviewed_workbook(tmp_path: Path) -> None:
    profile = load_profile("CELF6")
    template = Path(profile.report_template_for("targeted_swap"))
    roster = _template_roster(template)
    sample = pd.DataFrame({
        "teid": roster,
        "agey": 5,
        "agem": 0,
        "form": "form_1",
        "ped": 1,
        "eth": "AS",
        "region": "MW",
        "sex": "F",
    })
    available = pd.DataFrame({
        "teid": ["900001", "900002"],
        "agey": [5, 5],
        "agem": [0, 0],
        "form": ["form_1", "form_1"],
        "ped": [1, 1],
        "eth": ["AS", "AS"],
        "region": ["MW", "MW"],
        "sex": ["F", "F"],
    })
    before_hash = _sha256(template)
    output = render_celf6_stratification_workbook(
        tmp_path / "Updated_Stratification.xlsx", sample, available, template
    )

    source = load_workbook(template, data_only=False)
    result = load_workbook(output, data_only=False)
    assert _sha256(template) == before_hash
    assert result.sheetnames == source.sheetnames == [
        "byage", "form_1", "form_2", "edl X sex", "edl X reg",
        "eth X sex", "eth X edl", "eth X reg", "reg X sex",
        "osep", "clinical_sources", "TEID_List",
    ]
    for sheet_name in result.sheetnames[:-1]:
        expected, actual = source[sheet_name], result[sheet_name]
        assert expected.merged_cells.ranges == actual.merged_cells.ranges
        assert len(expected.conditional_formatting) == len(actual.conditional_formatting)
        for row in range(1, expected.max_row + 1):
            for column in range(1, expected.max_column + 1):
                expected_value = expected.cell(row, column).value
                actual_value = actual.cell(row, column).value
                if isinstance(expected_value, (int, float)) and isinstance(actual_value, (int, float)):
                    assert actual_value == pytest.approx(expected_value, abs=1e-12)
                else:
                    assert actual_value == expected_value
                assert actual.cell(row, column).style_id == expected.cell(row, column).style_id
    assert result["TEID_List"]["A1"].value == "Selected TEIDs (N=1500)"
    assert result["TEID_List"]["B1"].value == "Available Pool TEIDs (N=2)"
    assert result["TEID_List"]["B2"].value == "900001"
    assert result["TEID_List"]["B3"].value == "900002"
