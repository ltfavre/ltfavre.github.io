from pathlib import Path
from datetime import datetime
import math
import sys
import json
import os

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# =====================================================
# CONFIG
# =====================================================
# This script is now configuration-driven.
# Preferred use: edit the Excel config workbook instead of editing this script.
# Optional override:
#   PowerShell: $env:STRAT_OPTIMIZER_CONFIG = "C:\path\to\100_optimizer_config.xlsx"

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent
DEFAULT_CONFIG_WORKBOOK = PROJECT_ROOT / "inputs" / "configs" / "stratification" / "100_optimizer_config.xlsx"
CONFIG_WORKBOOK = Path(
    sys.argv[1]
    if len(sys.argv) > 1
    else os.environ.get("STRAT_OPTIMIZER_CONFIG", DEFAULT_CONFIG_WORKBOOK)
)

DEFAULT_CONFIG = {
    "project_dir": str(PROJECT_ROOT),
    "target_workbook": r"inputs\data_raw\reference_tables\adhd_demo.xlsx",
    "master_sample_file": r"outputs\qa_tables\ADHD_final_samples_after_exact_sex_balance_from_90_latest.xlsx",
    "qa_dir": r"outputs\qa_tables",
    "run_label": "nightmare_sample_optimizer_configurable",
    "output_workbook_name": "nightmare_optimized_output.xlsx",
    "max_candidate_adds_per_remove": 40,
    "max_solver_seconds": 900,
    "num_workers": 8,
    "scale": 1000,
    "force_exact_sex_balance": "Y",
    "sex_balance_column": "sex",
    "sex_balance_female_value": "F",
    "sex_balance_male_value": "M",
    "sex_balance_odd_n_action": "skip",
}

DEFAULT_CASE_SWAP_FILES = [
    "ADHD_PRS_A_AGEGRP5_case_swap.xlsx",
    "ADHD_PRS_A_AGEGRP6_case_swap.xlsx",
    "ADHD_PRS_C_AGEGRP3_case_swap.xlsx",
    "ADHD_PRS_C_AGEGRP4_case_swap.xlsx",
    "ADHD_SRP_A_AGEGRP5_case_swap.xlsx",
    "ADHD_SRP_A_AGEGRP6_case_swap.xlsx",
    "ADHD_SRP_C_AGEGRP4_case_swap.xlsx",
    "ADHD_TRS_A_AGEGRP5_case_swap.xlsx",
    "ADHD_TRS_A_AGEGRP6_case_swap.xlsx",
    "ADHD_TRS_C_AGEGRP3_case_swap.xlsx",
    "ADHD_TRS_C_AGEGRP4_case_swap.xlsx",
]

DEFAULT_KEY_COLS = ["teid", "vform", "agegrp"]
DEFAULT_DEMO_COLS = ["sex", "eth", "edl", "region"]
DEFAULT_OSEP_COMPONENT_FLAG_COLS = [
    "tmp_ADHD", "tmp_EBD", "tmp_ASD", "tmp_DD", "tmp_SLD", "tmp_SLI", "tmp_OTHER"
]
DEFAULT_HARD_LOCK_COLS = ["vform", "agegrp", "sex", "tmp_ADHD"]
DEFAULT_AGE_SOURCE_COLS = ["agey", "basc4_agey", "o_basc4_agey", "adjagey"]
DEFAULT_WEIGHTS = {
    "tmp_ADHD": 1_000_000,
    "tmp_clin": 200_000,
    "tmp_EBD": 40_000,
    "tmp_ASD": 30_000,
    "demo": 4_000,
    "swap": 10,
    "sex_balance_prefilter": 500_000,
}
DEFAULT_OSEP_TARGETS = {
    "tmp_ADHD": {"target_pct": 5.0},
    "tmp_EBD": {"target_pct": 0.70},
    "tmp_ASD": {"target_pct": 2.00},
    "tmp_clin": {"min_pct": 11.0, "max_pct": 13.0},
}
DEFAULT_VFORM_ORDER = ["TRS_C", "TRS_A", "PRS_C", "PRS_A", "SRP_C", "SRP_A"]
DEFAULT_BYAGE_AGES = list(range(6, 19))
DEFAULT_DEMO_TARGET_SHEETS = {"trs": "TRS", "prs": "PRS", "srp": "SRP"}


def _is_enabled(value):
    if pd.isna(value):
        return True
    return str(value).strip().upper() not in {"N", "NO", "FALSE", "0", "OFF"}


def _clean_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return None
    return value


def _read_sheet(path, sheet_name):
    try:
        return pd.read_excel(path, sheet_name=sheet_name)
    except ValueError:
        return pd.DataFrame()


def _to_number_if_possible(value):
    value = _clean_value(value)
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip()
    try:
        num = float(text)
        return int(num) if num.is_integer() else num
    except ValueError:
        return value


def _resolve_path(value, project_dir):
    value = _clean_value(value)
    if value is None:
        return None
    text = str(value).replace("{PROJECT_DIR}", str(project_dir))
    path = Path(text)
    return path if path.is_absolute() else project_dir / path


def _load_json_overrides():
    """Load optional pipeline JSON config from STRAT_CONFIG for cross-stage consistency."""
    config_path = os.environ.get("STRAT_CONFIG")
    if not config_path:
        return {}

    path = Path(config_path)
    if not path.exists():
        print(f"STRAT_CONFIG provided but not found: {path}")
        return {}

    with open(path, "r", encoding="utf-8") as f:
        raw = json.load(f)

    mapped = {}
    key_map = {
        "project_dir": "project_dir",
        "demo_osep_template_file": "target_workbook",
        "sex_balance_output_file": "master_sample_file",
        "master_input_file": "master_sample_file",
        "qa_output_dir": "qa_dir",
        "run_label": "run_label",
        "optimizer_output_workbook": "output_workbook_path",
    }

    for src_key, dst_key in key_map.items():
        value = raw.get(src_key)
        if value is not None and str(value).strip() != "":
            mapped[dst_key] = value

    swap_groups = raw.get("swap_groups", [])
    if swap_groups:
        project = str(raw.get("project", "ADHD")).strip().upper()
        mapped["case_swap_files"] = [
            f"{project}_{g['vform']}_AGEGRP{int(g['agegrp'])}_case_swap.xlsx"
            for g in swap_groups
            if isinstance(g, dict) and g.get("vform") is not None and g.get("agegrp") is not None
        ]

    return mapped


def _load_excel_config(path):
    """Load all optimizer settings from the Excel config workbook."""
    config = {}
    settings = DEFAULT_CONFIG.copy()

    if not path.exists():
        print(f"Config workbook not found; using script defaults: {path}")
        config_path_used = None
    else:
        print(f"Reading optimizer config workbook: {path}")
        config_path_used = path
        settings_df = _read_sheet(path, "settings")
        if not settings_df.empty:
            for _, row in settings_df.iterrows():
                if "setting" not in row or "value" not in row:
                    continue
                setting = _clean_value(row.get("setting"))
                if setting is None:
                    continue
                settings[str(setting).strip()] = _to_number_if_possible(row.get("value"))

    # JSON overrides (when STRAT_CONFIG is set) take precedence over workbook/defaults.
    json_overrides = _load_json_overrides()
    for key, value in json_overrides.items():
        if key in {"case_swap_files", "output_workbook_path"}:
            continue
        settings[key] = _to_number_if_possible(value)

    project_dir = Path(str(settings["project_dir"]))
    target_workbook = _resolve_path(settings["target_workbook"], project_dir)
    master_sample_file = _resolve_path(settings["master_sample_file"], project_dir)
    qa_dir = _resolve_path(settings["qa_dir"], project_dir)

    case_swap_files = DEFAULT_CASE_SWAP_FILES.copy()
    columns = {
        "KEY_COLS": DEFAULT_KEY_COLS.copy(),
        "DEMO_COLS": DEFAULT_DEMO_COLS.copy(),
        "OSEP_COMPONENT_FLAG_COLS": DEFAULT_OSEP_COMPONENT_FLAG_COLS.copy(),
        "HARD_LOCK_COLS": DEFAULT_HARD_LOCK_COLS.copy(),
        "AGE_SOURCE_COLS": DEFAULT_AGE_SOURCE_COLS.copy(),
    }
    weights = DEFAULT_WEIGHTS.copy()
    osep_targets = DEFAULT_OSEP_TARGETS.copy()
    vform_order = DEFAULT_VFORM_ORDER.copy()
    byage_ages = DEFAULT_BYAGE_AGES.copy()
    demo_target_sheets = DEFAULT_DEMO_TARGET_SHEETS.copy()

    if "case_swap_files" in json_overrides and json_overrides["case_swap_files"]:
        case_swap_files = [str(x).strip() for x in json_overrides["case_swap_files"] if str(x).strip()]

    if config_path_used is not None:
        df = _read_sheet(path, "case_swap_files")
        if not df.empty and "file_name" in df.columns:
            case_swap_files = [
                str(r["file_name"]).strip()
                for _, r in df.iterrows()
                if _clean_value(r.get("file_name")) is not None and _is_enabled(r.get("enabled", True))
            ]

        df = _read_sheet(path, "columns")
        if not df.empty and {"list_name", "column_name"}.issubset(df.columns):
            out = {}
            for list_name, grp in df.groupby("list_name"):
                vals = [
                    str(r["column_name"]).strip()
                    for _, r in grp.iterrows()
                    if _clean_value(r.get("column_name")) is not None and _is_enabled(r.get("enabled", True))
                ]
                if vals:
                    out[str(list_name).strip()] = vals
            columns.update(out)

        df = _read_sheet(path, "weights")
        if not df.empty and {"objective", "weight"}.issubset(df.columns):
            weights = {
                str(r["objective"]).strip(): int(float(r["weight"]))
                for _, r in df.iterrows()
                if _clean_value(r.get("objective")) is not None
                and _clean_value(r.get("weight")) is not None
                and _is_enabled(r.get("enabled", True))
            }

        df = _read_sheet(path, "osep_targets")
        if not df.empty and {"flag", "mode"}.issubset(df.columns):
            parsed = {}
            for _, r in df.iterrows():
                if not _is_enabled(r.get("enabled", True)):
                    continue
                flag = _clean_value(r.get("flag"))
                mode = str(_clean_value(r.get("mode")) or "target").strip().lower()
                if flag is None:
                    continue
                flag = str(flag).strip()
                if mode == "range":
                    min_pct = _clean_value(r.get("min_pct"))
                    max_pct = _clean_value(r.get("max_pct"))
                    if min_pct is not None and max_pct is not None:
                        parsed[flag] = {"min_pct": float(min_pct), "max_pct": float(max_pct)}
                else:
                    target_pct = _clean_value(r.get("target_pct"))
                    if target_pct is not None:
                        parsed[flag] = {"target_pct": float(target_pct)}
            if parsed:
                osep_targets = parsed

        df = _read_sheet(path, "vform_order")
        if not df.empty and "vform" in df.columns:
            sort_col = "sort_order" if "sort_order" in df.columns else None
            if sort_col:
                df = df.sort_values(sort_col)
            vform_order = [
                str(r["vform"]).strip().upper()
                for _, r in df.iterrows()
                if _clean_value(r.get("vform")) is not None and _is_enabled(r.get("enabled", True))
            ] or vform_order

        df = _read_sheet(path, "byage_ages")
        if not df.empty and "age" in df.columns:
            byage_ages = [
                int(float(r["age"]))
                for _, r in df.iterrows()
                if _clean_value(r.get("age")) is not None and _is_enabled(r.get("enabled", True))
            ] or byage_ages

        df = _read_sheet(path, "demo_target_sheets")
        if not df.empty and {"sheet_name", "family_prefix"}.issubset(df.columns):
            demo_target_sheets = {
                str(r["sheet_name"]).strip(): str(r["family_prefix"]).strip().upper()
                for _, r in df.iterrows()
                if _clean_value(r.get("sheet_name")) is not None
                and _clean_value(r.get("family_prefix")) is not None
                and _is_enabled(r.get("enabled", True))
            } or demo_target_sheets

    config.update({
        "CONFIG_PATH_USED": config_path_used,
        "PROJECT_DIR": project_dir,
        "TARGET_WORKBOOK": target_workbook,
        "MASTER_SAMPLE_FILE": master_sample_file,
        "QA_DIR": qa_dir,
        "RUN_LABEL": str(settings.get("run_label") or DEFAULT_CONFIG["run_label"]),
        "OUTPUT_WORKBOOK_NAME": str(settings.get("output_workbook_name") or DEFAULT_CONFIG["output_workbook_name"]),
        "MAX_CANDIDATE_ADDS_PER_REMOVE": int(settings.get("max_candidate_adds_per_remove") or DEFAULT_CONFIG["max_candidate_adds_per_remove"]),
        "MAX_SOLVER_SECONDS": int(settings.get("max_solver_seconds") or DEFAULT_CONFIG["max_solver_seconds"]),
        "NUM_WORKERS": int(settings.get("num_workers") or DEFAULT_CONFIG["num_workers"]),
        "SCALE": int(settings.get("scale") or DEFAULT_CONFIG["scale"]),
        "FORCE_EXACT_SEX_BALANCE": _is_enabled(settings.get("force_exact_sex_balance", DEFAULT_CONFIG["force_exact_sex_balance"])),
        "SEX_BALANCE_COLUMN": str(settings.get("sex_balance_column") or DEFAULT_CONFIG["sex_balance_column"]).strip(),
        "SEX_BALANCE_FEMALE_VALUE": str(settings.get("sex_balance_female_value") or DEFAULT_CONFIG["sex_balance_female_value"]).strip().upper(),
        "SEX_BALANCE_MALE_VALUE": str(settings.get("sex_balance_male_value") or DEFAULT_CONFIG["sex_balance_male_value"]).strip().upper(),
        "SEX_BALANCE_ODD_N_ACTION": str(settings.get("sex_balance_odd_n_action") or DEFAULT_CONFIG["sex_balance_odd_n_action"]).strip().lower(),
        "CASE_SWAP_FILES": case_swap_files,
        "KEY_COLS": columns["KEY_COLS"],
        "DEMO_COLS": columns["DEMO_COLS"],
        "OSEP_COMPONENT_FLAG_COLS": columns["OSEP_COMPONENT_FLAG_COLS"],
        "HARD_LOCK_COLS": columns["HARD_LOCK_COLS"],
        "AGE_SOURCE_COLS": columns["AGE_SOURCE_COLS"],
        "WEIGHTS": weights,
        "OSEP_TARGETS": osep_targets,
        "VFORM_ORDER": vform_order,
        "BYAGE_AGES": byage_ages,
        "DEMO_TARGET_SHEETS": demo_target_sheets,
    })

    output_workbook_path = json_overrides.get("output_workbook_path")
    if output_workbook_path:
        config["OUTPUT_WORKBOOK_PATH"] = _resolve_path(output_workbook_path, project_dir)

    return config


_config = _load_excel_config(CONFIG_WORKBOOK)
PROJECT_DIR = _config["PROJECT_DIR"]
TARGET_WORKBOOK = _config["TARGET_WORKBOOK"]
MASTER_SAMPLE_FILE = _config["MASTER_SAMPLE_FILE"]
QA_DIR = _config["QA_DIR"]
RUN_LABEL = _config["RUN_LABEL"]
MAX_CANDIDATE_ADDS_PER_REMOVE = _config["MAX_CANDIDATE_ADDS_PER_REMOVE"]
MAX_SOLVER_SECONDS = _config["MAX_SOLVER_SECONDS"]
NUM_WORKERS = _config["NUM_WORKERS"]
SCALE = _config["SCALE"]
FORCE_EXACT_SEX_BALANCE = _config["FORCE_EXACT_SEX_BALANCE"]
SEX_BALANCE_COLUMN = _config["SEX_BALANCE_COLUMN"]
SEX_BALANCE_FEMALE_VALUE = _config["SEX_BALANCE_FEMALE_VALUE"]
SEX_BALANCE_MALE_VALUE = _config["SEX_BALANCE_MALE_VALUE"]
SEX_BALANCE_ODD_N_ACTION = _config["SEX_BALANCE_ODD_N_ACTION"]
CASE_SWAP_FILES = _config["CASE_SWAP_FILES"]
KEY_COLS = _config["KEY_COLS"]
DEMO_COLS = _config["DEMO_COLS"]
OSEP_COMPONENT_FLAG_COLS = _config["OSEP_COMPONENT_FLAG_COLS"]
OSEP_FLAGS = OSEP_COMPONENT_FLAG_COLS + ["tmp_clin"]
HARD_LOCK_COLS = _config["HARD_LOCK_COLS"]
SWAP_LOCK_COLS = [
    c for c in HARD_LOCK_COLS
    if not (FORCE_EXACT_SEX_BALANCE and c == SEX_BALANCE_COLUMN)
]
AGE_SOURCE_COLS = _config["AGE_SOURCE_COLS"]
WEIGHTS = _config["WEIGHTS"]
OSEP_TARGETS = _config["OSEP_TARGETS"]
VFORM_ORDER = _config["VFORM_ORDER"]
BYAGE_AGES = _config["BYAGE_AGES"]
DEMO_TARGET_SHEETS = _config["DEMO_TARGET_SHEETS"]

RUN_ID = datetime.now().strftime("%Y%m%d_%H%M%S")
RUN_DIR = PROJECT_DIR / "outputs" / "stratification_runs" / f"run_{RUN_ID}_{RUN_LABEL}"
RUN_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_WORKBOOK = _config.get("OUTPUT_WORKBOOK_PATH", RUN_DIR / _config["OUTPUT_WORKBOOK_NAME"])
OUTPUT_WORKBOOK.parent.mkdir(parents=True, exist_ok=True)

print("Optimizer configuration loaded:")
print(f"  project_dir: {PROJECT_DIR}")
print(f"  target_workbook: {TARGET_WORKBOOK}")
print(f"  master_sample_file: {MASTER_SAMPLE_FILE}")
print(f"  qa_dir: {QA_DIR}")
print(f"  hard_locks: {HARD_LOCK_COLS}")
print(f"  swap_generation_locks: {SWAP_LOCK_COLS}")
print(f"  force_exact_sex_balance: {FORCE_EXACT_SEX_BALANCE}")
print(f"  osep_targets: {OSEP_TARGETS}")
print(f"  weights: {WEIGHTS}")

# =====================================================
# BASIC HELPERS
# =====================================================
def require_columns(df, cols, label):
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise ValueError(f"{label} is missing required columns: {missing}")


def normalize_teid(series):
    text = series.astype(str).str.strip()
    numeric = pd.to_numeric(text, errors="coerce")
    mask = numeric.notna()
    text.loc[mask] = numeric.loc[mask].astype("Int64").astype(str)
    return text.replace({"NAN": pd.NA, "NONE": pd.NA, "<NA>": pd.NA})


def normalize_sample(df):
    df = df.copy()
    require_columns(df, KEY_COLS, "sample")
    df["teid"] = normalize_teid(df["teid"])
    df["vform"] = df["vform"].astype(str).str.strip().str.upper()
    df["agegrp"] = pd.to_numeric(df["agegrp"], errors="raise").astype(int)
    age_source = None
    for col in AGE_SOURCE_COLS:
        if col not in df.columns:
            continue
        values = pd.to_numeric(df[col], errors="coerce")
        age_source = values if age_source is None else age_source.fillna(values)
    if age_source is not None:
        df["agey"] = age_source
    for col in DEMO_COLS:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip().str.upper()
    for flag in OSEP_COMPONENT_FLAG_COLS:
        flag_aliases = [flag, flag.lower(), flag.upper()]
        source_col = next((col for col in flag_aliases if col in df.columns), None)
        if source_col is None:
            df[flag] = 0
        elif source_col != flag:
            df[flag] = df[source_col]
        df[flag] = pd.to_numeric(df[flag], errors="coerce").fillna(0).astype(int)
    df["tmp_clin"] = df[OSEP_COMPONENT_FLAG_COLS].astype(bool).any(axis=1).astype(int)
    return df


def flag_int(row, flag):
    return int(pd.to_numeric(pd.Series([row.get(flag, 0)]), errors="coerce").fillna(0).iloc[0])


def clean_category_value(value):
    if pd.isna(value):
        return ""
    return str(value).strip().upper()


def is_sex_value(value, target):
    return int(clean_category_value(value) == str(target).strip().upper())


def sex_count(df, value):
    if SEX_BALANCE_COLUMN not in df.columns:
        return 0
    return int(df[SEX_BALANCE_COLUMN].apply(lambda x: clean_category_value(x) == str(value).strip().upper()).sum())


def build_sex_balance_qa(df):
    rows = []
    if SEX_BALANCE_COLUMN not in df.columns:
        return pd.DataFrame(rows)
    for (vform, agegrp), grp in df.groupby(["vform", "agegrp"], dropna=False):
        total_n = len(grp)
        female_n = sex_count(grp, SEX_BALANCE_FEMALE_VALUE)
        male_n = sex_count(grp, SEX_BALANCE_MALE_VALUE)
        target_each = total_n // 2 if total_n % 2 == 0 else None
        rows.append({
            "vform": vform,
            "agegrp": agegrp,
            "total_n": total_n,
            "female_value": SEX_BALANCE_FEMALE_VALUE,
            "female_n": female_n,
            "female_pct": round(female_n / total_n * 100, 3) if total_n else 0,
            "male_value": SEX_BALANCE_MALE_VALUE,
            "male_n": male_n,
            "male_pct": round(male_n / total_n * 100, 3) if total_n else 0,
            "target_each": target_each,
            "exact_50_50": (female_n == male_n) if total_n % 2 == 0 else False,
            "odd_n_impossible": total_n % 2 != 0,
        })
    return pd.DataFrame(rows).sort_values(["vform", "agegrp"])


def pct_to_count(n, pct):
    return int(round(float(pct) * n / 100.0 * SCALE))


def range_error_scaled(count, n, min_pct, max_pct):
    lo = float(min_pct) * n / 100.0 * SCALE
    hi = float(max_pct) * n / 100.0 * SCALE
    val = count * SCALE
    if val < lo:
        return lo - val
    if val > hi:
        return val - hi
    return 0.0


def target_error_scaled(count, n, target_pct):
    return abs(count * SCALE - pct_to_count(n, target_pct))


# =====================================================
# READ TARGET WORKBOOK
# =====================================================
def parse_demo_targets(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    targets = {}
    sheet_for_prefix = DEMO_TARGET_SHEETS

    def parse_demo_header(header):
        header = str(header).strip()
        if header.startswith("edl_"):
            return "edl", header.split("_", 1)[1], header
        if header.startswith("eth_"):
            suffix = header.split("_", 1)[1].lstrip("0123456789")
            value = "BL" if suffix == "AA" else suffix
            return "eth", value, header
        if header.startswith("region_"):
            value = header.split("_", 1)[1].lstrip("0123456789")
            return "region", value, header
        return None

    for sheet_name, prefix in sheet_for_prefix.items():
        ws = wb[sheet_name]
        headers = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]
        current_vform = None
        for r in range(6, ws.max_row + 1):
            raw_vform = ws.cell(r, 1).value
            if raw_vform:
                current_vform = str(raw_vform).strip().upper()
            source = ws.cell(r, 4).value
            agegrp = ws.cell(r, 2).value
            if current_vform is None or source != "US population" or agegrp is None:
                continue
            key = (current_vform, int(agegrp))
            targets[key] = {}
            for c, header in enumerate(headers, start=1):
                if not header or c <= 4:
                    continue
                value = ws.cell(r, c).value
                if value is None:
                    continue
                header = str(header).strip()
                parsed = parse_demo_header(header)
                if parsed:
                    targets[key][parsed] = float(value)
    return targets


def parse_osep_targets(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["osep"]
    out = {}
    for r in range(3, ws.max_row + 1):
        label = ws.cell(r, 1).value
        target = ws.cell(r, 8).value
        if not label:
            continue
        label = str(label).strip()
        flag = label.replace("pct_", "tmp_", 1)
        if flag == "tmp_clin":
            out[flag] = {"min_pct": 11.0, "max_pct": 13.0}
        elif target not in [None, ""]:
            out[flag] = {"target_pct": float(target)}
    out.update(OSEP_TARGETS)
    return out


# =====================================================
# READ DATA
# =====================================================
def read_inputs():
    master = pd.read_excel(MASTER_SAMPLE_FILE, sheet_name="final_all")
    master = normalize_sample(master)
    require_columns(master, KEY_COLS + DEMO_COLS + OSEP_COMPONENT_FLAG_COLS, "master")

    pools = []
    read_log = [{"file": str(MASTER_SAMPLE_FILE), "sheet": "final_all", "n": len(master)}]
    for fname in CASE_SWAP_FILES:
        path = QA_DIR / fname
        if not path.exists():
            continue
        xls = pd.ExcelFile(path)
        if "replacement_pool" not in xls.sheet_names:
            continue
        parts = fname.replace(".xlsx", "").split("_")
        vform = f"{parts[1]}_{parts[2]}"
        agegrp = int(parts[3].replace("AGEGRP", ""))
        pool = pd.read_excel(path, sheet_name="replacement_pool").copy()
        pool["vform"] = vform
        pool["agegrp"] = agegrp
        pool["source_file"] = fname
        pools.append(pool)
        read_log.append({"file": str(path), "sheet": "replacement_pool", "n": len(pool)})

    if pools:
        available = normalize_sample(pd.concat(pools, ignore_index=True))
        available = available.loc[~available["teid"].isin(set(master["teid"]))].copy()
    else:
        available = pd.DataFrame(columns=master.columns)
    return master, available, pd.DataFrame(read_log)


# =====================================================
# LOSS MODEL
# =====================================================
def build_state(master, demo_targets):
    state = {}
    for vform, grp in master.groupby("vform"):
        state[(vform, None)] = {
            "n": len(grp),
            "flags": {
                flag: int(pd.to_numeric(grp[flag], errors="coerce").fillna(0).sum())
                for flag in OSEP_FLAGS
            },
        }
    for (vform, agegrp), grp in master.groupby(["vform", "agegrp"]):
        state[(vform, agegrp)] = {
            "n": len(grp),
            "demo": {},
            "sex": {
                SEX_BALANCE_FEMALE_VALUE: sex_count(grp, SEX_BALANCE_FEMALE_VALUE),
                SEX_BALANCE_MALE_VALUE: sex_count(grp, SEX_BALANCE_MALE_VALUE),
            },
        }
        for (target_vform, target_agegrp), target_map in demo_targets.items():
            if target_vform != vform or target_agegrp != agegrp:
                continue
            for (col, value, _header), _target_pct in target_map.items():
                state[(vform, agegrp)]["demo"][(col, value)] = int(
                    (grp[col].astype(str).str.upper() == str(value).upper()).sum()
                )
    return state


def local_loss_delta(vform, agegrp, rem, add, state, demo_targets, osep_targets):
    pre = 0.0
    post = 0.0
    form_state = state[(vform, None)]
    n_form = form_state["n"]
    for flag, rule in osep_targets.items():
        if flag not in form_state["flags"]:
            continue
        base = form_state["flags"][flag]
        new = base - flag_int(rem, flag) + flag_int(add, flag)
        if "target_pct" in rule:
            pre += target_error_scaled(base, n_form, rule["target_pct"]) * WEIGHTS.get(flag, 1)
            post += target_error_scaled(new, n_form, rule["target_pct"]) * WEIGHTS.get(flag, 1)
        else:
            pre += range_error_scaled(base, n_form, rule["min_pct"], rule["max_pct"]) * WEIGHTS.get(flag, 1)
            post += range_error_scaled(new, n_form, rule["min_pct"], rule["max_pct"]) * WEIGHTS.get(flag, 1)

    cell_targets = demo_targets.get((vform, agegrp), {})
    cell_state = state[(vform, agegrp)]
    n_cell = cell_state["n"]
    for (col, value, _header), target_pct in cell_targets.items():
        base = cell_state["demo"].get((col, value), 0)
        new = base - int(str(rem.get(col)).upper() == str(value).upper()) + int(str(add.get(col)).upper() == str(value).upper())
        pre += target_error_scaled(base, n_cell, target_pct) * WEIGHTS["demo"]
        post += target_error_scaled(new, n_cell, target_pct) * WEIGHTS["demo"]
    return pre - post


# =====================================================
# EXACT SEX BALANCE HELPERS
# =====================================================
def sex_balance_quick_delta(vform, agegrp, rem, add, state):
    """Positive value means this swap moves the vform+agegrp cell closer to exact 50/50."""
    if not FORCE_EXACT_SEX_BALANCE:
        return 0
    if SEX_BALANCE_COLUMN not in rem.index or SEX_BALANCE_COLUMN not in add.index:
        return 0
    cell_state = state.get((vform, agegrp))
    if not cell_state:
        return 0
    n = int(cell_state["n"])
    if n % 2 != 0:
        return 0
    target_each = n // 2
    base_f = int(cell_state["sex"].get(SEX_BALANCE_FEMALE_VALUE, 0))
    new_f = base_f - is_sex_value(rem.get(SEX_BALANCE_COLUMN), SEX_BALANCE_FEMALE_VALUE) + is_sex_value(add.get(SEX_BALANCE_COLUMN), SEX_BALANCE_FEMALE_VALUE)
    pre_dev = abs(base_f - target_each)
    post_dev = abs(new_f - target_each)
    return (pre_dev - post_dev) * int(WEIGHTS.get("sex_balance_prefilter", 500000))


def add_exact_sex_balance_constraints(model, y, swaps, state):
    """Add Script-64a-style exact 50/50 sex constraints inside the CP-SAT model."""
    if not FORCE_EXACT_SEX_BALANCE:
        return []
    messages = []
    if SEX_BALANCE_COLUMN not in swaps.columns and f"remove_{SEX_BALANCE_COLUMN}" not in swaps.columns:
        messages.append("Skipped exact sex balance constraints because sex columns were not found in swaps.")
        return messages

    for (vform, agegrp), cell_state in sorted(state.items()):
        if agegrp is None:
            continue
        n = int(cell_state["n"])
        if n % 2 != 0:
            msg = f"Skipped exact sex balance for {vform} agegrp {agegrp}: odd N={n}."
            messages.append(msg)
            if SEX_BALANCE_ODD_N_ACTION == "error":
                raise ValueError(msg)
            continue
        target_each = n // 2
        cell_swaps = swaps.loc[(swaps["vform"] == vform) & (swaps["agegrp"] == agegrp)]
        female_expr = int(cell_state["sex"].get(SEX_BALANCE_FEMALE_VALUE, 0))
        male_expr = int(cell_state["sex"].get(SEX_BALANCE_MALE_VALUE, 0))
        for r in cell_swaps.itertuples():
            sid = int(r.swap_id)
            rem_sex = getattr(r, f"remove_{SEX_BALANCE_COLUMN}")
            add_sex = getattr(r, f"add_{SEX_BALANCE_COLUMN}")
            female_delta = is_sex_value(add_sex, SEX_BALANCE_FEMALE_VALUE) - is_sex_value(rem_sex, SEX_BALANCE_FEMALE_VALUE)
            male_delta = is_sex_value(add_sex, SEX_BALANCE_MALE_VALUE) - is_sex_value(rem_sex, SEX_BALANCE_MALE_VALUE)
            if female_delta:
                female_expr += female_delta * y[sid]
            if male_delta:
                male_expr += male_delta * y[sid]
        model.Add(female_expr == target_each)
        model.Add(male_expr == target_each)
        messages.append(f"Added exact sex balance for {vform} agegrp {agegrp}: {SEX_BALANCE_FEMALE_VALUE}={target_each}, {SEX_BALANCE_MALE_VALUE}={target_each}.")
    return messages


# =====================================================
# OPTIMIZE
# =====================================================
def optimize(master, available, demo_targets, osep_targets):
    try:
        from ortools.sat.python import cp_model
    except ImportError as exc:
        raise ImportError("Install OR-Tools in your environment: python -m pip install ortools") from exc

    state = build_state(master, demo_targets)
    swap_rows = []
    for lock_vals, used_group in master.groupby(SWAP_LOCK_COLS, dropna=False):
        if not isinstance(lock_vals, tuple):
            lock_vals = (lock_vals,)
        mask = pd.Series(True, index=available.index)
        for col, val in zip(SWAP_LOCK_COLS, lock_vals):
            mask &= available[col].astype(str).eq(str(val))
        avail_group = available.loc[mask].copy()
        if avail_group.empty:
            continue
        for _, rem in used_group.iterrows():
            scored = []
            for _, add in avail_group.iterrows():
                base_delta = local_loss_delta(rem["vform"], rem["agegrp"], rem, add, state, demo_targets, osep_targets)
                base_delta += sex_balance_quick_delta(rem["vform"], rem["agegrp"], rem, add, state)
                scored.append((base_delta, add))
            scored = sorted(scored, key=lambda x: x[0], reverse=True)[:MAX_CANDIDATE_ADDS_PER_REMOVE]
            for quick_delta, add in scored:
                row = {
                    "swap_id": len(swap_rows),
                    "vform": rem["vform"],
                    "agegrp": int(rem["agegrp"]),
                    "sex": rem["sex"],
                    "remove_teid": rem["teid"],
                    "add_teid": add["teid"],
                    "quick_loss_reduction": quick_delta,
                }
                for flag in OSEP_FLAGS:
                    row[f"delta_{flag}"] = flag_int(add, flag) - flag_int(rem, flag)
                    row[f"remove_{flag}"] = flag_int(rem, flag)
                    row[f"add_{flag}"] = flag_int(add, flag)
                for col in DEMO_COLS:
                    row[f"remove_{col}"] = rem.get(col)
                    row[f"add_{col}"] = add.get(col)
                if SEX_BALANCE_COLUMN not in DEMO_COLS:
                    row[f"remove_{SEX_BALANCE_COLUMN}"] = rem.get(SEX_BALANCE_COLUMN)
                    row[f"add_{SEX_BALANCE_COLUMN}"] = add.get(SEX_BALANCE_COLUMN)
                row["sex_balance_delta"] = sex_balance_quick_delta(rem["vform"], rem["agegrp"], rem, add, state)
                swap_rows.append(row)

    swaps = pd.DataFrame(swap_rows)
    if swaps.empty:
        raise ValueError("No feasible hard-locked swaps were generated.")

    model = cp_model.CpModel()
    y = {int(r.swap_id): model.NewBoolVar(f"swap_{int(r.swap_id)}") for r in swaps.itertuples()}

    for teid, grp in swaps.groupby("remove_teid"):
        model.Add(sum(y[int(s)] for s in grp["swap_id"]) <= 1)
    for teid, grp in swaps.groupby("add_teid"):
        model.Add(sum(y[int(s)] for s in grp["swap_id"]) <= 1)

    sex_constraint_messages = add_exact_sex_balance_constraints(model, y, swaps, state)

    objective_terms = []

    def post_flag_expr(vform, flag):
        expr = int(state[(vform, None)]["flags"].get(flag, 0))
        for r in swaps.loc[swaps["vform"] == vform].itertuples():
            d = int(getattr(r, f"delta_{flag}"))
            if d:
                expr += d * y[int(r.swap_id)]
        return expr

    def add_abs_dev(expr, target, ub, name):
        dev = model.NewIntVar(0, int(ub), name)
        target = int(target)
        model.Add(dev >= expr - target)
        model.Add(dev >= target - expr)
        return dev

    for vform in VFORM_ORDER:
        if (vform, None) not in state:
            continue
        n = state[(vform, None)]["n"]
        for flag, rule in osep_targets.items():
            expr = post_flag_expr(vform, flag)
            if "target_pct" in rule:
                dev = add_abs_dev(expr * SCALE, pct_to_count(n, rule["target_pct"]), n * SCALE, f"dev_{vform}_{flag}")
            else:
                lo = math.ceil(float(rule["min_pct"]) * n / 100.0 * SCALE)
                hi = math.floor(float(rule["max_pct"]) * n / 100.0 * SCALE)
                under = model.NewIntVar(0, n * SCALE, f"under_{vform}_{flag}")
                over = model.NewIntVar(0, n * SCALE, f"over_{vform}_{flag}")
                model.Add(expr * SCALE + under >= lo)
                model.Add(expr * SCALE - over <= hi)
                dev = model.NewIntVar(0, n * SCALE, f"dev_{vform}_{flag}")
                model.Add(dev == under + over)
            objective_terms.append(int(WEIGHTS.get(flag, 1)) * dev)

    for (vform, agegrp), target_map in demo_targets.items():
        if (vform, agegrp) not in state:
            continue
        n = state[(vform, agegrp)]["n"]
        cell_swaps = swaps.loc[(swaps["vform"] == vform) & (swaps["agegrp"] == agegrp)]
        for (col, value, _header), target_pct in target_map.items():
            expr = int(state[(vform, agegrp)]["demo"].get((col, value), 0))
            for r in cell_swaps.itertuples():
                d = int(str(getattr(r, f"add_{col}")).upper() == str(value).upper()) - int(
                    str(getattr(r, f"remove_{col}")).upper() == str(value).upper()
                )
                if d:
                    expr += d * y[int(r.swap_id)]
            dev = add_abs_dev(expr * SCALE, pct_to_count(n, target_pct), n * SCALE, f"dev_{vform}_{agegrp}_{col}_{value}")
            objective_terms.append(WEIGHTS["demo"] * dev)

    objective_terms.append(WEIGHTS["swap"] * sum(y.values()))
    model.Minimize(sum(objective_terms))

    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = float(MAX_SOLVER_SECONDS)
    solver.parameters.num_search_workers = int(NUM_WORKERS)
    status = solver.Solve(model)
    if status not in [cp_model.OPTIMAL, cp_model.FEASIBLE]:
        raise ValueError(f"Solver failed: {solver.StatusName(status)}")

    selected = swaps.loc[[solver.Value(y[int(s)]) == 1 for s in swaps["swap_id"]]].copy()
    remove_teids = set(selected["remove_teid"])
    add_teids = set(selected["add_teid"])
    final_sample = pd.concat(
        [
            master.loc[~master["teid"].isin(remove_teids)].copy(),
            available.loc[available["teid"].isin(add_teids)].copy(),
        ],
        ignore_index=True,
        sort=False,
    )
    final_sample = normalize_sample(final_sample)
    selected["solver_status"] = solver.StatusName(status)
    selected["solver_objective"] = solver.ObjectiveValue()
    if sex_constraint_messages:
        selected["sex_balance_constraints"] = " | ".join(sex_constraint_messages[:20])
    return final_sample, selected, swaps


# =====================================================
# OUTPUT TABLES
# =====================================================
def pct(series):
    return float(pd.to_numeric(series, errors="coerce").fillna(0).mean() * 100)


def build_byage(final_sample):
    rows = []
    families = sorted({str(v).split("_")[0] for v in VFORM_ORDER})
    for age in BYAGE_AGES:
        row = {"Age": age}
        for vform in VFORM_ORDER:
            if "agey" in final_sample.columns:
                age_series = pd.to_numeric(final_sample["agey"], errors="coerce").round().astype("Int64")
                row[vform] = int(((final_sample["vform"] == vform) & (age_series == age)).sum())
            else:
                row[vform] = ""
        for family in families:
            row[f"{family}_total"] = sum(row.get(v, 0) for v in VFORM_ORDER if str(v).startswith(family))
        rows.append(row)
    total = {"Age": ""}
    for col in VFORM_ORDER + [f"{family}_total" for family in families]:
        total[col] = sum(r[col] for r in rows if isinstance(r.get(col), int))
    total["Grand_total"] = len(final_sample)
    rows.append(total)
    return pd.DataFrame(rows)

def build_demo_sheet(final_sample, demo_targets, family):
    vforms = [v for v in VFORM_ORDER if v.startswith(family)]
    rows = []
    for vform in vforms:
        vf_rows = [k for k in demo_targets if k[0] == vform]
        for v_idx, (target_vform, agegrp) in enumerate(sorted(vf_rows)):
            target_map = demo_targets[(target_vform, agegrp)]
            grp = final_sample.loc[(final_sample["vform"] == vform) & (final_sample["agegrp"] == agegrp)]
            target_row = {"Vform": vform if v_idx == 0 else "", "Age Group": agegrp, "Age Band": "", "Source": "US population"}
            actual_row = {"Vform": "", "Age Group": agegrp, "Age Band": "", "Source": ""}
            for (col, value, header), target_pct in target_map.items():
                target_row[header] = target_pct
                actual_row[header] = pct(grp[col].astype(str).str.upper().eq(str(value).upper()).astype(int)) if len(grp) else 0
            actual_row["female"] = int((grp["sex"] == "F").sum())
            actual_row["male"] = int((grp["sex"] == "M").sum())
            rows.extend([target_row, actual_row])
    return pd.DataFrame(rows)


def build_osep(final_sample):
    display = []
    for flag in OSEP_FLAGS:
        label = flag.replace("tmp_", "pct_", 1)
        rule = OSEP_TARGETS.get(flag, {})
        if "target_pct" in rule:
            target = f"{float(rule['target_pct']):.2f}"
        elif "min_pct" in rule and "max_pct" in rule:
            target = f"{float(rule['min_pct']):.2f}-{float(rule['max_pct']):.2f}"
        else:
            target = ""
        display.append((label, flag, target))
    display.append(("total_n", None, ""))

    rows = []
    for label, flag, target in display:
        row = {"Group": label}
        for vform in VFORM_ORDER:
            grp = final_sample.loc[final_sample["vform"] == vform]
            row[vform] = int(len(grp)) if flag is None else round(pct(grp[flag]), 2)
        row["Census Target"] = target
        rows.append(row)
    return pd.DataFrame(rows)

def write_formatted_workbook(final_sample, swaps, all_swaps, read_log, demo_targets):
    wb = Workbook()
    wb.remove(wb.active)
    title_fill = PatternFill("solid", fgColor="002060")
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    label_fill = PatternFill("solid", fgColor="D9D2E9")
    total_fill = PatternFill("solid", fgColor="EDEDED")
    border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    def write_df(sheet_name, title, df):
        ws = wb.create_sheet(sheet_name)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(df.columns)))
        ws.cell(1, 1).value = title
        ws.cell(1, 1).fill = title_fill
        ws.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=12, name="Courier New")
        ws.cell(1, 1).alignment = Alignment(horizontal="center")
        for c, col in enumerate(df.columns, start=1):
            cell = ws.cell(2, c, col)
            cell.fill = header_fill
            cell.font = Font(bold=True)
        for r, row in enumerate(df.itertuples(index=False), start=3):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
                if isinstance(value, float):
                    cell.number_format = "0.00"
            if str(ws.cell(r, 1).value) == "total_n":
                for c in range(1, len(df.columns) + 1):
                    ws.cell(r, c).fill = total_fill
                    ws.cell(r, c).font = Font(bold=True)
            else:
                ws.cell(r, 1).fill = label_fill
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = border
        for col_idx in range(1, ws.max_column + 1):
            ws.column_dimensions[ws.cell(2, col_idx).column_letter].width = 14
        return ws

    write_df("byage", "GROUP SIZES: AFTER OPTIMIZED SWAPS", build_byage(final_sample))
    write_df("trs", "DEMOGRAPHIC VARIABLES: AFTER OPTIMIZED SWAPS", build_demo_sheet(final_sample, demo_targets, "TRS"))
    write_df("prs", "DEMOGRAPHIC VARIABLES: AFTER OPTIMIZED SWAPS", build_demo_sheet(final_sample, demo_targets, "PRS"))
    write_df("srp", "DEMOGRAPHIC VARIABLES: AFTER OPTIMIZED SWAPS", build_demo_sheet(final_sample, demo_targets, "SRP"))
    write_df("osep", "OSEP: AFTER OPTIMIZED SWAPS", build_osep(final_sample))
    write_df("sex_balance", "SEX BALANCE: AFTER OPTIMIZED SWAPS", build_sex_balance_qa(final_sample))

    for name, df in [
        ("swap_pairs", swaps),
        ("final_teids", final_sample[[c for c in ["teid", "vform", "agegrp", "agey", "sex", "tmp_ADHD", "tmp_clin"] if c in final_sample.columns]].sort_values([c for c in ["vform", "agegrp", "agey", "teid"] if c in final_sample.columns])),
        ("all_candidate_swaps", all_swaps.head(200000)),
        ("read_log", read_log),
    ]:
        ws = wb.create_sheet(name[:31])
        for c, col in enumerate(df.columns, start=1):
            ws.cell(1, c, col).font = Font(bold=True)
        for r, row in enumerate(df.itertuples(index=False), start=2):
            for c, value in enumerate(row, start=1):
                ws.cell(r, c, value)
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[ws.cell(1, c).column_letter].width = 16

    # Save a copy of the active configuration inside the output for auditability.
    config_rows = []
    for k, v in {
        "config_workbook": str(CONFIG_WORKBOOK),
        "project_dir": str(PROJECT_DIR),
        "target_workbook": str(TARGET_WORKBOOK),
        "master_sample_file": str(MASTER_SAMPLE_FILE),
        "qa_dir": str(QA_DIR),
        "run_label": RUN_LABEL,
        "max_candidate_adds_per_remove": MAX_CANDIDATE_ADDS_PER_REMOVE,
        "max_solver_seconds": MAX_SOLVER_SECONDS,
        "num_workers": NUM_WORKERS,
        "scale": SCALE,
        "force_exact_sex_balance": FORCE_EXACT_SEX_BALANCE,
        "sex_balance_column": SEX_BALANCE_COLUMN,
        "sex_balance_female_value": SEX_BALANCE_FEMALE_VALUE,
        "sex_balance_male_value": SEX_BALANCE_MALE_VALUE,
        "sex_balance_odd_n_action": SEX_BALANCE_ODD_N_ACTION,
        "hard_lock_cols": ", ".join(HARD_LOCK_COLS),
        "swap_generation_locks": ", ".join(SWAP_LOCK_COLS),
        "vform_order": ", ".join(VFORM_ORDER),
        "byage_ages": ", ".join(map(str, BYAGE_AGES)),
    }.items():
        config_rows.append({"setting": k, "value": v})
    for k, v in WEIGHTS.items():
        config_rows.append({"setting": f"weight:{k}", "value": v})
    for k, v in OSEP_TARGETS.items():
        config_rows.append({"setting": f"osep_target:{k}", "value": str(v)})
    ws = wb.create_sheet("config_used")
    config_df = pd.DataFrame(config_rows)
    for c, col in enumerate(config_df.columns, start=1):
        ws.cell(1, c, col).font = Font(bold=True)
    for r, row in enumerate(config_df.itertuples(index=False), start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    for c in range(1, ws.max_column + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 35

    wb.save(OUTPUT_WORKBOOK)


def main():
    print("Reading targets from target workbook and config...")
    demo_targets = parse_demo_targets(TARGET_WORKBOOK)
    osep_targets = parse_osep_targets(TARGET_WORKBOOK)

    print("Reading sample and replacement pools...")
    master, available, read_log = read_inputs()
    print(f"Master N={len(master):,}; available N={len(available):,}")

    print("Optimizing sample...")
    final_sample, selected_swaps, all_swaps = optimize(master, available, demo_targets, osep_targets)
    if len(final_sample) != len(master):
        raise ValueError(f"Final sample size changed: {len(master)} -> {len(final_sample)}")
    if final_sample["teid"].duplicated().any():
        raise ValueError("Final sample has duplicate TEIDs.")
    if FORCE_EXACT_SEX_BALANCE:
        sex_qa = build_sex_balance_qa(final_sample)
        failed = sex_qa.loc[(sex_qa["odd_n_impossible"] == False) & (sex_qa["exact_50_50"] == False)]
        if not failed.empty:
            raise ValueError("Exact sex balance failed for one or more even-N cells:\n" + failed.to_string(index=False))

    print("Writing output workbook...")
    write_formatted_workbook(final_sample, selected_swaps, all_swaps, read_log, demo_targets)
    print(f"Done: {OUTPUT_WORKBOOK}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"FAILED: {exc}", file=sys.stderr)
        raise
