from pathlib import Path
from datetime import datetime
import math
import sys

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# =====================================================
# CONFIG
# =====================================================
# This script is now configuration-driven.
# Preferred use: edit the Excel config workbook instead of editing this script.
# Optional override:
#   PowerShell: $env:STRAT_OPTIMIZER_CONFIG = "C:\path\to\100_optimizer_config.xlsx"

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_CONFIG_WORKBOOK = SCRIPT_DIR / "100_optimizer_config.xlsx"
CONFIG_WORKBOOK = Path(
    sys.argv[1]
    if len(sys.argv) > 1
    else __import__("os").environ.get("STRAT_OPTIMIZER_CONFIG", DEFAULT_CONFIG_WORKBOOK)
)

DEFAULT_CONFIG = {
    "project_dir": r"C:\Users\UFAVRLA\PythonProjects\Fake_BASC4NB",
    "target_workbook": r"C:\Users\UFAVRLA\OneDrive - Pearson\Desktop\Stratification\stratification nightmare.xlsx",
    "master_sample_file": r"outputs\qa_tables\ADHD_final_samples_after_exact_sex_balance_from_90_latest.xlsx",
    "qa_dir": r"outputs\qa_tables",
    "run_label": "nightmare_sample_optimizer_configurable",
    "output_workbook_name": "nightmare_optimized_output.xlsx",
    "max_candidate_adds_per_remove": 40,
    "max_solver_seconds": 900,
    "num_workers": 8,
    "scale": 1000,
    "force_exact_sex_balance": "Y",
    "sex_balance_column": "sex",
    "sex_balance_female_value": "F",
    "sex_balance_male_value": "M",
    "sex_balance_odd_n_action": "skip",
}

DEFAULT_CASE_SWAP_FILES = [
    "ADHD_PRS_A_AGEGRP5_case_swap.xlsx",
    "ADHD_PRS_A_AGEGRP6_case_swap.xlsx",
    "ADHD_PRS_C_AGEGRP3_case_swap.xlsx",
    "ADHD_PRS_C_AGEGRP4_case_swap.xlsx",
    "ADHD_SRP_A_AGEGRP5_case_swap.xlsx",
    "ADHD_SRP_A_AGEGRP6_case_swap.xlsx",
    "ADHD_SRP_C_AGEGRP4_case_swap.xlsx",
    "ADHD_TRS_A_AGEGRP5_case_swap.xlsx",
    "ADHD_TRS_A_AGEGRP6_case_swap.xlsx",
    "ADHD_TRS_C_AGEGRP3_case_swap.xlsx",
    "ADHD_TRS_C_AGEGRP4_case_swap.xlsx",
]

DEFAULT_KEY_COLS = ["teid", "vform", "agegrp"]
DEFAULT_DEMO_COLS = ["sex", "eth", "edl", "region"]
DEFAULT_OSEP_COMPONENT_FLAG_COLS = [
    "tmp_ADHD", "tmp_EBD", "tmp_ASD", "tmp_DD", "tmp_SLD", "tmp_SLI", "tmp_OTHER"
]
DEFAULT_HARD_LOCK_COLS = ["vform", "agegrp", "sex", "tmp_ADHD"]
DEFAULT_AGE_SOURCE_COLS = ["agey", "basc4_agey", "o_basc4_agey", "adjagey"]
DEFAULT_WEIGHTS = {
    "tmp_ADHD": 1_000_000,
    "tmp_clin": 200_000,
    "tmp_EBD": 40_000,
    "tmp_ASD": 30_000,
    "demo": 4_000,
    "swap": 10,
    "sex_balance_prefilter": 500_000,
}
DEFAULT_OSEP_TARGETS = {
    "tmp_ADHD": {"target_pct": 5.0},
    "tmp_EBD": {"target_pct": 0.70},
    "tmp_ASD": {"target_pct": 2.00},
    "tmp_clin": {"min_pct": 11.0, "max_pct": 13.0},
}
DEFAULT_VFORM_ORDER = ["TRS_C", "TRS_A", "PRS_C", "PRS_A", "SRP_C", "SRP_A"]
DEFAULT_BYAGE_AGES = list(range(6, 19))
DEFAULT_DEMO_TARGET_SHEETS = {"trs": "TRS", "prs": "PRS", "srp": "SRP"}


def _is_enabled(value):
    if pd.isna(value):
        return True
    return str(value).strip().upper() not in {"N", "NO", "FALSE", "0", "OFF"}


def _clean_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return None
    return value


def _read_sheet(path, sheet_name):
    try:
        return pd.read_excel(path, sheet_name=sheet_name)
    except ValueError:
        return pd.DataFrame()


def _to_number_if_possible(value):
    value = _clean_value(value)
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip()
    try:
        num = float(text)
        return int(num) if num.is_integer() else num
    except ValueError:
        return value


def _resolve_path(value, project_dir):
    value = _clean_value(value)
    if value is None:
        return None
    text = str(value).replace("{PROJECT_DIR}", str(project_dir))
    path = Path(text)
    return path if path.is_absolute() else project_dir / path



def _table_from_guided_sheet(path, sheet_name):
    """Read a user-facing config sheet whose real header may start after title/help rows."""
    df_raw = pd.read_excel(path, sheet_name=sheet_name, header=None)
    header_idx = None
    for idx, row in df_raw.iterrows():
        values = [str(v).strip() for v in row.tolist() if not pd.isna(v)]
        if sheet_name == "1_Project_Files" and {"Section", "Setting", "Value"}.issubset(set(values)):
            header_idx = idx
            break
        if sheet_name == "2_Data_Map" and {"Variable Category", "Variable Name: In Python", "Variable Name:  In Raw Data"}.issubset(set(values)):
            header_idx = idx
            break
        if sheet_name == "3_Pools_Constraints" and {"Section", "Item / Constraint", "Enabled?", "File or Column"}.issubset(set(values)):
            header_idx = idx
            break
        if sheet_name == "4_Targets_Weights" and {"Section", "Target Type", "Column Name", "Weight"}.issubset(set(values)):
            header_idx = idx
            break
    if header_idx is None:
        return pd.DataFrame()
    headers = [str(v).strip() if not pd.isna(v) else f"Unnamed_{i}" for i, v in enumerate(df_raw.iloc[header_idx].tolist())]
    df = df_raw.iloc[header_idx + 1:].copy()
    df.columns = headers
    df = df.dropna(how="all")
    return df


def _sheet_exists(path, sheet_name):
    try:
        return sheet_name in pd.ExcelFile(path).sheet_names
    except Exception:
        return False


def _parse_guided_config(path, settings):
    """Parse the 5-sheet Stratification_Optimizer_Config.xlsx format."""
    config = {}
    df = _table_from_guided_sheet(path, "1_Project_Files")
    if not df.empty:
        for _, row in df.iterrows():
            setting = _clean_value(row.get("Setting"))
            if setting is None:
                continue
            settings[str(setting).strip()] = _to_number_if_possible(row.get("Value"))

    project_dir = Path(str(settings.get("project_dir", DEFAULT_CONFIG["project_dir"])))
    target_workbook = _resolve_path(settings.get("target_workbook", DEFAULT_CONFIG["target_workbook"]), project_dir)
    master_sample_file = _resolve_path(settings.get("master_sample_file", DEFAULT_CONFIG["master_sample_file"]), project_dir)
    qa_dir = _resolve_path(settings.get("replacement_pool_dir", settings.get("qa_dir", DEFAULT_CONFIG["qa_dir"])), project_dir)
    output_dir = _resolve_path(settings.get("output_dir", r"outputs\stratification_runs"), project_dir)
    master_sample_sheet = str(settings.get("master_sample_sheet", "final_all"))
    target_osep_sheet = str(settings.get("target_osep_sheet", "osep"))

    key_rename_map = {}
    demo_cols = []
    osep_flags = []
    demo_target_sheets = DEFAULT_DEMO_TARGET_SHEETS.copy()
    age_source_cols = DEFAULT_AGE_SOURCE_COLS.copy()

    df = _table_from_guided_sheet(path, "2_Data_Map")
    if not df.empty:
        for _, row in df.iterrows():
            if not _is_enabled(row.get("Enabled?", True)):
                continue
            category = str(_clean_value(row.get("Variable Category")) or "").strip().lower()
            py_name = str(_clean_value(row.get("Variable Name: In Python")) or "").strip()
            raw_name = _clean_value(row.get("Variable Name:  In Raw Data"))
            if raw_name is None:
                continue
            raw_name = str(raw_name).strip()

            if "teid" in py_name.lower() or category == "case id":
                key_rename_map[raw_name] = "teid"
            elif "vform" in py_name.lower() or category == "form":
                key_rename_map[raw_name] = "vform"
            elif "agegrp" in py_name.lower() or category == "age group":
                key_rename_map[raw_name] = "agegrp"
            elif "agey" in py_name.lower() or category == "age year":
                key_rename_map[raw_name] = "agey"
                if raw_name not in age_source_cols:
                    age_source_cols = [raw_name] + age_source_cols
            elif category == "demographic":
                demo_cols.append(raw_name)
            elif category == "clinical flag":
                osep_flags.append(raw_name)
            elif category == "target workbook" and "demo target sheet" in py_name.lower():
                text = raw_name.replace("→", "->")
                if "->" in text:
                    sheet, family = [x.strip() for x in text.split("->", 1)]
                    if sheet and family:
                        demo_target_sheets[sheet] = family.upper()

    if not demo_cols:
        demo_cols = DEFAULT_DEMO_COLS.copy()
    if not osep_flags:
        osep_flags = DEFAULT_OSEP_COMPONENT_FLAG_COLS.copy()

    case_swap_specs = []
    hard_lock_cols = []
    force_exact_sex_balance = False
    sex_balance_column = "sex"
    sex_balance_female_value = "F"
    sex_balance_male_value = "M"

    df = _table_from_guided_sheet(path, "3_Pools_Constraints")
    if not df.empty:
        for _, row in df.iterrows():
            if not _is_enabled(row.get("Enabled?", True)):
                continue
            section = str(_clean_value(row.get("Section")) or "").strip().lower()
            item = str(_clean_value(row.get("Item / Constraint")) or "").strip()
            file_or_col = _clean_value(row.get("File or Column"))
            mode = str(_clean_value(row.get("Mode")) or "").strip().lower()
            if section == "replacement pool":
                fname = str(file_or_col or item).strip()
                if fname:
                    case_swap_specs.append({
                        "file_name": fname,
                        "sheet_name": "replacement_pool",
                        "vform": _clean_value(row.get("Vform")),
                        "agegrp": _clean_value(row.get("Age Group")),
                    })
            elif section == "hard constraint":
                if item == "force_exact_sex_balance":
                    force_exact_sex_balance = True
                    if file_or_col is not None:
                        sex_balance_column = str(file_or_col).strip()
                elif item == "preserve_sex_when_not_balancing":
                    if file_or_col is not None:
                        hard_lock_cols.append(str(file_or_col).strip())
                elif mode == "same_value_swap" and file_or_col is not None:
                    hard_lock_cols.append(str(file_or_col).strip())

    for required_lock in ["vform", "agegrp"]:
        if required_lock not in hard_lock_cols:
            hard_lock_cols.insert(0, required_lock)

    if force_exact_sex_balance:
        hard_lock_cols = [c for c in hard_lock_cols if c != sex_balance_column]

    if not case_swap_specs:
        case_swap_specs = [{"file_name": f, "sheet_name": "replacement_pool", "vform": None, "agegrp": None} for f in DEFAULT_CASE_SWAP_FILES]
    if not hard_lock_cols:
        hard_lock_cols = ["vform", "agegrp", "tmp_ADHD"]

    weights = {}
    osep_targets = {}
    df = _table_from_guided_sheet(path, "4_Targets_Weights")
    if not df.empty:
        for _, row in df.iterrows():
            if not _is_enabled(row.get("Enabled?", True)):
                continue
            section = str(_clean_value(row.get("Section")) or "").strip().lower()
            col = _clean_value(row.get("Column Name"))
            mode = str(_clean_value(row.get("Mode")) or "").strip().lower()
            weight = _clean_value(row.get("Weight"))
            target_pct = _clean_value(row.get("Target %"))
            min_pct = _clean_value(row.get("Min %"))
            max_pct = _clean_value(row.get("Max %"))

            if col is not None and weight is not None:
                if section == "demographic target":
                    weights["demo"] = int(float(weight))
                elif section == "swap control":
                    weights["swap"] = int(float(weight))
                else:
                    weights[str(col).strip()] = int(float(weight))

            if section == "clinical target" and col is not None:
                flag = str(col).strip()
                if mode == "range" and min_pct is not None and max_pct is not None:
                    osep_targets[flag] = {"min_pct": float(min_pct), "max_pct": float(max_pct)}
                elif mode == "target" and target_pct is not None:
                    osep_targets[flag] = {"target_pct": float(target_pct)}

    weights = {**DEFAULT_WEIGHTS, **weights}
    if not osep_targets:
        osep_targets = DEFAULT_OSEP_TARGETS.copy()

    vform_order = []
    for spec in case_swap_specs:
        vf = spec.get("vform")
        if vf is not None and not pd.isna(vf) and str(vf).strip() and str(vf).strip().upper() not in vform_order:
            vform_order.append(str(vf).strip().upper())
    if not vform_order:
        vform_order = DEFAULT_VFORM_ORDER.copy()

    return {
        "CONFIG_PATH_USED": path,
        "PROJECT_DIR": project_dir,
        "TARGET_WORKBOOK": target_workbook,
        "MASTER_SAMPLE_FILE": master_sample_file,
        "MASTER_SAMPLE_SHEET": master_sample_sheet,
        "TARGET_OSEP_SHEET": target_osep_sheet,
        "QA_DIR": qa_dir,
        "OUTPUT_DIR_BASE": output_dir,
        "RUN_LABEL": str(settings.get("run_label") or DEFAULT_CONFIG["run_label"]),
        "OUTPUT_WORKBOOK_NAME": str(settings.get("output_workbook_name") or DEFAULT_CONFIG["output_workbook_name"]),
        "MAX_CANDIDATE_ADDS_PER_REMOVE": int(settings.get("max_candidate_adds_per_remove") or DEFAULT_CONFIG["max_candidate_adds_per_remove"]),
        "MAX_SOLVER_SECONDS": int(settings.get("max_solver_seconds") or DEFAULT_CONFIG["max_solver_seconds"]),
        "NUM_WORKERS": int(settings.get("num_workers") or DEFAULT_CONFIG["num_workers"]),
        "SCALE": int(settings.get("scale") or DEFAULT_CONFIG["scale"]),
        "FORCE_EXACT_SEX_BALANCE": force_exact_sex_balance,
        "SEX_BALANCE_COLUMN": sex_balance_column,
        "SEX_BALANCE_FEMALE_VALUE": sex_balance_female_value,
        "SEX_BALANCE_MALE_VALUE": sex_balance_male_value,
        "SEX_BALANCE_ODD_N_ACTION": str(settings.get("sex_balance_odd_n_action") or DEFAULT_CONFIG["sex_balance_odd_n_action"]).strip().lower(),
        "CASE_SWAP_FILES": [s["file_name"] for s in case_swap_specs],
        "CASE_SWAP_SPECS": case_swap_specs,
        "KEY_COLS": ["teid", "vform", "agegrp"],
        "RAW_TO_CANONICAL_RENAME": key_rename_map,
        "DEMO_COLS": demo_cols,
        "OSEP_COMPONENT_FLAG_COLS": osep_flags,
        "HARD_LOCK_COLS": hard_lock_cols,
        "AGE_SOURCE_COLS": age_source_cols,
        "WEIGHTS": weights,
        "OSEP_TARGETS": osep_targets,
        "VFORM_ORDER": vform_order,
        "BYAGE_AGES": DEFAULT_BYAGE_AGES.copy(),
        "DEMO_TARGET_SHEETS": demo_target_sheets,
    }


def _load_excel_config(path):
    """Load optimizer settings from either the legacy machine-readable config or the 5-sheet guided workbook."""
    config = {}
    settings = DEFAULT_CONFIG.copy()

    if not path.exists():
        print(f"Config workbook not found; using script defaults: {path}")
        config_path_used = None
    else:
        print(f"Reading optimizer config workbook: {path}")
        config_path_used = path

        if _sheet_exists(path, "1_Project_Files") and _sheet_exists(path, "2_Data_Map"):
            return _parse_guided_config(path, settings)

        settings_df = _read_sheet(path, "settings")
        if not settings_df.empty:
            for _, row in settings_df.iterrows():
                if "setting" not in row or "value" not in row:
                    continue
                setting = _clean_value(row.get("setting"))
                if setting is None:
                    continue
                settings[str(setting).strip()] = _to_number_if_possible(row.get("value"))

    project_dir = Path(str(settings["project_dir"]))
    target_workbook = _resolve_path(settings["target_workbook"], project_dir)
    master_sample_file = _resolve_path(settings["master_sample_file"], project_dir)
    qa_dir = _resolve_path(settings["qa_dir"], project_dir)

    case_swap_files = DEFAULT_CASE_SWAP_FILES.copy()
    columns = {
        "KEY_COLS": DEFAULT_KEY_COLS.copy(),
        "DEMO_COLS": DEFAULT_DEMO_COLS.copy(),
        "OSEP_COMPONENT_FLAG_COLS": DEFAULT_OSEP_COMPONENT_FLAG_COLS.copy(),
        "HARD_LOCK_COLS": DEFAULT_HARD_LOCK_COLS.copy(),
        "AGE_SOURCE_COLS": DEFAULT_AGE_SOURCE_COLS.copy(),
    }
    weights = DEFAULT_WEIGHTS.copy()
    osep_targets = DEFAULT_OSEP_TARGETS.copy()
    vform_order = DEFAULT_VFORM_ORDER.copy()
    byage_ages = DEFAULT_BYAGE_AGES.copy()
    demo_target_sheets = DEFAULT_DEMO_TARGET_SHEETS.copy()

    if config_path_used is not None:
        df = _read_sheet(path, "case_swap_files")
        if not df.empty and "file_name" in df.columns:
            case_swap_files = [
                str(r["file_name"]).strip()
                for _, r in df.iterrows()
                if _clean_value(r.get("file_name")) is not None and _is_enabled(r.get("enabled", True))
            ]

        df = _read_sheet(path, "columns")
        if not df.empty and {"list_name", "column_name"}.issubset(df.columns):
            out = {}
            for list_name, grp in df.groupby("list_name"):
                vals = [
                    str(r["column_name"]).strip()
                    for _, r in grp.iterrows()
                    if _clean_value(r.get("column_name")) is not None and _is_enabled(r.get("enabled", True))
                ]
                if vals:
                    out[str(list_name).strip()] = vals
            columns.update(out)

        df = _read_sheet(path, "weights")
        if not df.empty and {"objective", "weight"}.issubset(df.columns):
            weights = {
                str(r["objective"]).strip(): int(float(r["weight"]))
                for _, r in df.iterrows()
                if _clean_value(r.get("objective")) is not None
                and _clean_value(r.get("weight")) is not None
                and _is_enabled(r.get("enabled", True))
            }

        df = _read_sheet(path, "osep_targets")
        if not df.empty and {"flag", "mode"}.issubset(df.columns):
            parsed = {}
            for _, r in df.iterrows():
                if not _is_enabled(r.get("enabled", True)):
                    continue
                flag = _clean_value(r.get("flag"))
                mode = str(_clean_value(r.get("mode")) or "target").strip().lower()
                if flag is None:
                    continue
                flag = str(flag).strip()
                if mode == "range":
                    min_pct = _clean_value(r.get("min_pct"))
                    max_pct = _clean_value(r.get("max_pct"))
                    if min_pct is not None and max_pct is not None:
                        parsed[flag] = {"min_pct": float(min_pct), "max_pct": float(max_pct)}
                else:
                    target_pct = _clean_value(r.get("target_pct"))
                    if target_pct is not None:
                        parsed[flag] = {"target_pct": float(target_pct)}
            if parsed:
                osep_targets = parsed

        df = _read_sheet(path, "vform_order")
        if not df.empty and "vform" in df.columns:
            sort_col = "sort_order" if "sort_order" in df.columns else None
            if sort_col:
                df = df.sort_values(sort_col)
            vform_order = [
                str(r["vform"]).strip().upper()
                for _, r in df.iterrows()
                if _clean_value(r.get("vform")) is not None and _is_enabled(r.get("enabled", True))
            ] or vform_order

        df = _read_sheet(path, "byage_ages")
        if not df.empty and "age" in df.columns:
            byage_ages = [
                int(float(r["age"]))
                for _, r in df.iterrows()
                if _clean_value(r.get("age")) is not None and _is_enabled(r.get("enabled", True))
            ] or byage_ages

        df = _read_sheet(path, "demo_target_sheets")
        if not df.empty and {"sheet_name", "family_prefix"}.issubset(df.columns):
            demo_target_sheets = {
                str(r["sheet_name"]).strip(): str(r["family_prefix"]).strip().upper()
                for _, r in df.iterrows()
                if _clean_value(r.get("sheet_name")) is not None
                and _clean_value(r.get("family_prefix")) is not None
                and _is_enabled(r.get("enabled", True))
            } or demo_target_sheets

    config.update({
        "CONFIG_PATH_USED": config_path_used,
        "PROJECT_DIR": project_dir,
        "TARGET_WORKBOOK": target_workbook,
        "MASTER_SAMPLE_FILE": master_sample_file,
        "MASTER_SAMPLE_SHEET": "final_all",
        "TARGET_OSEP_SHEET": "osep",
        "QA_DIR": qa_dir,
        "OUTPUT_DIR_BASE": project_dir / "outputs" / "stratification_runs",
        "RUN_LABEL": str(settings.get("run_label") or DEFAULT_CONFIG["run_label"]),
        "OUTPUT_WORKBOOK_NAME": str(settings.get("output_workbook_name") or DEFAULT_CONFIG["output_workbook_name"]),
        "MAX_CANDIDATE_ADDS_PER_REMOVE": int(settings.get("max_candidate_adds_per_remove") or DEFAULT_CONFIG["max_candidate_adds_per_remove"]),
        "MAX_SOLVER_SECONDS": int(settings.get("max_solver_seconds") or DEFAULT_CONFIG["max_solver_seconds"]),
        "NUM_WORKERS": int(settings.get("num_workers") or DEFAULT_CONFIG["num_workers"]),
        "SCALE": int(settings.get("scale") or DEFAULT_CONFIG["scale"]),
        "FORCE_EXACT_SEX_BALANCE": _is_enabled(settings.get("force_exact_sex_balance", DEFAULT_CONFIG["force_exact_sex_balance"])),
        "SEX_BALANCE_COLUMN": str(settings.get("sex_balance_column") or DEFAULT_CONFIG["sex_balance_column"]).strip(),
        "SEX_BALANCE_FEMALE_VALUE": str(settings.get("sex_balance_female_value") or DEFAULT_CONFIG["sex_balance_female_value"]).strip().upper(),
        "SEX_BALANCE_MALE_VALUE": str(settings.get("sex_balance_male_value") or DEFAULT_CONFIG["sex_balance_male_value"]).strip().upper(),
        "SEX_BALANCE_ODD_N_ACTION": str(settings.get("sex_balance_odd_n_action") or DEFAULT_CONFIG["sex_balance_odd_n_action"]).strip().lower(),
        "CASE_SWAP_FILES": case_swap_files,
        "CASE_SWAP_SPECS": [{"file_name": f, "sheet_name": "replacement_pool", "vform": None, "agegrp": None} for f in case_swap_files],
        "KEY_COLS": columns["KEY_COLS"],
        "RAW_TO_CANONICAL_RENAME": {},
        "DEMO_COLS": columns["DEMO_COLS"],
        "OSEP_COMPONENT_FLAG_COLS": columns["OSEP_COMPONENT_FLAG_COLS"],
        "HARD_LOCK_COLS": columns["HARD_LOCK_COLS"],
        "AGE_SOURCE_COLS": columns["AGE_SOURCE_COLS"],
        "WEIGHTS": weights,
        "OSEP_TARGETS": osep_targets,
        "VFORM_ORDER": vform_order,
        "BYAGE_AGES": byage_ages,
        "DEMO_TARGET_SHEETS": demo_target_sheets,
    })
    return config


_config = _load_excel_config(CONFIG_WORKBOOK)
PROJECT_DIR = _config["PROJECT_DIR"]
TARGET_WORKBOOK = _config["TARGET_WORKBOOK"]
MASTER_SAMPLE_FILE = _config["MASTER_SAMPLE_FILE"]
MASTER_SAMPLE_SHEET = _config["MASTER_SAMPLE_SHEET"]
TARGET_OSEP_SHEET = _config["TARGET_OSEP_SHEET"]
QA_DIR = _config["QA_DIR"]
OUTPUT_DIR_BASE = _config["OUTPUT_DIR_BASE"]
RUN_LABEL = _config["RUN_LABEL"]
MAX_CANDIDATE_ADDS_PER_REMOVE = _config["MAX_CANDIDATE_ADDS_PER_REMOVE"]
MAX_SOLVER_SECONDS = _config["MAX_SOLVER_SECONDS"]
NUM_WORKERS = _config["NUM_WORKERS"]
SCALE = _config["SCALE"]
FORCE_EXACT_SEX_BALANCE = _config["FORCE_EXACT_SEX_BALANCE"]
SEX_BALANCE_COLUMN = _config["SEX_BALANCE_COLUMN"]
SEX_BALANCE_FEMALE_VALUE = _config["SEX_BALANCE_FEMALE_VALUE"]
SEX_BALANCE_MALE_VALUE = _config["SEX_BALANCE_MALE_VALUE"]
SEX_BALANCE_ODD_N_ACTION = _config["SEX_BALANCE_ODD_N_ACTION"]
CASE_SWAP_FILES = _config["CASE_SWAP_FILES"]
CASE_SWAP_SPECS = _config["CASE_SWAP_SPECS"]
RAW_TO_CANONICAL_RENAME = _config["RAW_TO_CANONICAL_RENAME"]
KEY_COLS = _config["KEY_COLS"]
DEMO_COLS = _config["DEMO_COLS"]
OSEP_COMPONENT_FLAG_COLS = _config["OSEP_COMPONENT_FLAG_COLS"]
OSEP_FLAGS = OSEP_COMPONENT_FLAG_COLS + ["tmp_clin"]
HARD_LOCK_COLS = _config["HARD_LOCK_COLS"]
SWAP_LOCK_COLS = [
    c for c in HARD_LOCK_COLS
    if not (FORCE_EXACT_SEX_BALANCE and c == SEX_BALANCE_COLUMN)
]
AGE_SOURCE_COLS = _config["AGE_SOURCE_COLS"]
WEIGHTS = _config["WEIGHTS"]
OSEP_TARGETS = _config["OSEP_TARGETS"]
VFORM_ORDER = _config["VFORM_ORDER"]
BYAGE_AGES = _config["BYAGE_AGES"]
DEMO_TARGET_SHEETS = _config["DEMO_TARGET_SHEETS"]

RUN_ID = datetime.now().strftime("%Y%m%d_%H%M%S")
RUN_DIR = OUTPUT_DIR_BASE / f"run_{RUN_ID}_{RUN_LABEL}"
RUN_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_WORKBOOK = RUN_DIR / _config["OUTPUT_WORKBOOK_NAME"]

print("Optimizer configuration loaded:")
print(f"  project_dir: {PROJECT_DIR}")
print(f"  target_workbook: {TARGET_WORKBOOK}")
print(f"  master_sample_file: {MASTER_SAMPLE_FILE}")
print(f"  master_sample_sheet: {MASTER_SAMPLE_SHEET}")
print(f"  qa_dir: {QA_DIR}")
print(f"  output_dir_base: {OUTPUT_DIR_BASE}")
print(f"  hard_locks: {HARD_LOCK_COLS}")
print(f"  swap_generation_locks: {SWAP_LOCK_COLS}")
print(f"  force_exact_sex_balance: {FORCE_EXACT_SEX_BALANCE}")
print(f"  osep_targets: {OSEP_TARGETS}")
print(f"  weights: {WEIGHTS}")

# =====================================================
# BASIC HELPERS
# =====================================================
def require_columns(df, cols, label):
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise ValueError(f"{label} is missing required columns: {missing}")


def normalize_teid(series):
    text = series.astype(str).str.strip()
    numeric = pd.to_numeric(text, errors="coerce")
    mask = numeric.notna()
    text.loc[mask] = numeric.loc[mask].astype("Int64").astype(str)
    return text.replace({"NAN": pd.NA, "NONE": pd.NA, "<NA>": pd.NA})


def normalize_sample(df):
    df = df.copy()
    rename_map = {
        raw: internal
        for raw, internal in RAW_TO_CANONICAL_RENAME.items()
        if raw in df.columns and raw != internal
    }
    if rename_map:
        df = df.rename(columns=rename_map)

    require_columns(df, KEY_COLS, "sample")
    df["teid"] = normalize_teid(df["teid"])
    df["vform"] = df["vform"].astype(str).str.strip().str.upper()
    df["agegrp"] = pd.to_numeric(df["agegrp"], errors="raise").astype(int)
    age_source = None
    for col in AGE_SOURCE_COLS:
        if col not in df.columns:
            continue
        values = pd.to_numeric(df[col], errors="coerce")
        age_source = values if age_source is None else age_source.fillna(values)
    if age_source is not None:
        df["agey"] = age_source
    for col in DEMO_COLS:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip().str.upper()
    for flag in OSEP_COMPONENT_FLAG_COLS:
        flag_aliases = [flag, flag.lower(), flag.upper()]
        source_col = next((col for col in flag_aliases if col in df.columns), None)
        if source_col is None:
            df[flag] = 0
        elif source_col != flag:
            df[flag] = df[source_col]
        df[flag] = pd.to_numeric(df[flag], errors="coerce").fillna(0).astype(int)
    df["tmp_clin"] = df[OSEP_COMPONENT_FLAG_COLS].astype(bool).any(axis=1).astype(int)
    return df


def flag_int(row, flag):
    return int(pd.to_numeric(pd.Series([row.get(flag, 0)]), errors="coerce").fillna(0).iloc[0])


def clean_category_value(value):
    if pd.isna(value):
        return ""
    return str(value).strip().upper()


def is_sex_value(value, target):
    return int(clean_category_value(value) == str(target).strip().upper())


def sex_count(df, value):
    if SEX_BALANCE_COLUMN not in df.columns:
        return 0
    return int(df[SEX_BALANCE_COLUMN].apply(lambda x: clean_category_value(x) == str(value).strip().upper()).sum())


def build_sex_balance_qa(df):
    rows = []
    if SEX_BALANCE_COLUMN not in df.columns:
        return pd.DataFrame(rows)
    for (vform, agegrp), grp in df.groupby(["vform", "agegrp"], dropna=False):
        total_n = len(grp)
        female_n = sex_count(grp, SEX_BALANCE_FEMALE_VALUE)
        male_n = sex_count(grp, SEX_BALANCE_MALE_VALUE)
        target_each = total_n // 2 if total_n % 2 == 0 else None
        rows.append({
            "vform": vform,
            "agegrp": agegrp,
            "total_n": total_n,
            "female_value": SEX_BALANCE_FEMALE_VALUE,
            "female_n": female_n,
            "female_pct": round(female_n / total_n * 100, 3) if total_n else 0,
            "male_value": SEX_BALANCE_MALE_VALUE,
            "male_n": male_n,
            "male_pct": round(male_n / total_n * 100, 3) if total_n else 0,
            "target_each": target_each,
            "exact_50_50": (female_n == male_n) if total_n % 2 == 0 else False,
            "odd_n_impossible": total_n % 2 != 0,
        })
    return pd.DataFrame(rows).sort_values(["vform", "agegrp"])


def pct_to_count(n, pct):
    return int(round(float(pct) * n / 100.0 * SCALE))


def range_error_scaled(count, n, min_pct, max_pct):
    lo = float(min_pct) * n / 100.0 * SCALE
    hi = float(max_pct) * n / 100.0 * SCALE
    val = count * SCALE
    if val < lo:
        return lo - val
    if val > hi:
        return val - hi
    return 0.0


def target_error_scaled(count, n, target_pct):
    return abs(count * SCALE - pct_to_count(n, target_pct))


# =====================================================
# READ TARGET WORKBOOK
# =====================================================
def parse_demo_targets(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    targets = {}
    sheet_for_prefix = DEMO_TARGET_SHEETS

    def parse_demo_header(header):
        header = str(header).strip()
        for demo_col in DEMO_COLS:
            prefix = f"{demo_col}_"
            if header.startswith(prefix):
                value = header.split("_", 1)[1].lstrip("0123456789")
                if demo_col == "eth" and value == "AA":
                    value = "BL"
                return demo_col, value, header
        if header.startswith("edl_"):
            return "edl", header.split("_", 1)[1], header
        if header.startswith("eth_"):
            suffix = header.split("_", 1)[1].lstrip("0123456789")
            value = "BL" if suffix == "AA" else suffix
            return "eth", value, header
        if header.startswith("region_"):
            value = header.split("_", 1)[1].lstrip("0123456789")
            return "region", value, header
        return None

    for sheet_name, prefix in sheet_for_prefix.items():
        ws = wb[sheet_name]
        headers = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]
        current_vform = None
        for r in range(6, ws.max_row + 1):
            raw_vform = ws.cell(r, 1).value
            if raw_vform:
                current_vform = str(raw_vform).strip().upper()
            source = ws.cell(r, 4).value
            agegrp = ws.cell(r, 2).value
            if current_vform is None or source != "US population" or agegrp is None:
                continue
            key = (current_vform, int(agegrp))
            targets[key] = {}
            for c, header in enumerate(headers, start=1):
                if not header or c <= 4:
                    continue
                value = ws.cell(r, c).value
                if value is None:
                    continue
                header = str(header).strip()
                parsed = parse_demo_header(header)
                if parsed:
                    targets[key][parsed] = float(value)
    return targets


def parse_osep_targets(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[TARGET_OSEP_SHEET]
    out = {}
    for r in range(3, ws.max_row + 1):
        label = ws.cell(r, 1).value
        target = ws.cell(r, 8).value
        if not label:
            continue
        label = str(label).strip()
        flag = label.replace("pct_", "tmp_", 1)
        if flag == "tmp_clin":
            out[flag] = {"min_pct": 11.0, "max_pct": 13.0}
        elif target not in [None, ""]:
            out[flag] = {"target_pct": float(target)}
    out.update(OSEP_TARGETS)
    return out


# =====================================================
# READ DATA
# =====================================================
def _infer_vform_agegrp_from_filename(fname):
    stem = Path(fname).stem.upper()
    parts = stem.split("_")
    vform = None
    agegrp = None

    for i in range(len(parts) - 1):
        candidate = f"{parts[i]}_{parts[i + 1]}"
        if candidate in VFORM_ORDER or parts[i] in {"TRS", "PRS", "SRP"}:
            vform = candidate
            break

    match = __import__("re").search(r"AGEGRP[_-]?(\d+)", stem)
    if match:
        agegrp = int(match.group(1))

    return vform, agegrp


def _read_excel_sheet_if_present(path, preferred_sheet):
    xls = pd.ExcelFile(path)
    if preferred_sheet in xls.sheet_names:
        return pd.read_excel(path, sheet_name=preferred_sheet), preferred_sheet
    if xls.sheet_names:
        return pd.read_excel(path, sheet_name=xls.sheet_names[0]), xls.sheet_names[0]
    raise ValueError(f"No sheets found in {path}")


def read_inputs():
    read_log = []
    master_parts = []

    if MASTER_SAMPLE_FILE is not None and Path(MASTER_SAMPLE_FILE).exists():
        master_raw, used_sheet = _read_excel_sheet_if_present(MASTER_SAMPLE_FILE, MASTER_SAMPLE_SHEET)
        master_parts.append(master_raw)
        read_log.append({"file": str(MASTER_SAMPLE_FILE), "sheet": used_sheet, "n": len(master_raw), "role": "master_sample"})
    else:
        print("Master sample file not found or not provided. Falling back to final_sample sheets in replacement-pool workbooks.")

    pools = []
    for spec in CASE_SWAP_SPECS:
        fname = spec.get("file_name") if isinstance(spec, dict) else str(spec)
        sheet_name = spec.get("sheet_name", "replacement_pool") if isinstance(spec, dict) else "replacement_pool"
        path = QA_DIR / str(fname)
        if not path.exists():
            read_log.append({"file": str(path), "sheet": sheet_name, "n": 0, "role": "missing_pool_file"})
            continue

        xls = pd.ExcelFile(path)

        if not master_parts and "final_sample" in xls.sheet_names:
            final_part = pd.read_excel(path, sheet_name="final_sample").copy()
            vf = spec.get("vform") if isinstance(spec, dict) else None
            ag = spec.get("agegrp") if isinstance(spec, dict) else None
            if vf is None or pd.isna(vf) or ag is None or pd.isna(ag):
                vf2, ag2 = _infer_vform_agegrp_from_filename(fname)
                vf = vf if vf is not None and not pd.isna(vf) else vf2
                ag = ag if ag is not None and not pd.isna(ag) else ag2
            if vf is not None and "vform" not in final_part.columns:
                final_part["vform"] = str(vf).upper()
            if ag is not None and "agegrp" not in final_part.columns:
                final_part["agegrp"] = int(float(ag))
            master_parts.append(final_part)
            read_log.append({"file": str(path), "sheet": "final_sample", "n": len(final_part), "role": "master_part_from_case_swap"})

        if sheet_name not in xls.sheet_names:
            if "replacement_pool" in xls.sheet_names:
                sheet_name = "replacement_pool"
            else:
                read_log.append({"file": str(path), "sheet": sheet_name, "n": 0, "role": "pool_sheet_missing"})
                continue

        pool = pd.read_excel(path, sheet_name=sheet_name).copy()
        vf = spec.get("vform") if isinstance(spec, dict) else None
        ag = spec.get("agegrp") if isinstance(spec, dict) else None
        if vf is None or pd.isna(vf) or ag is None or pd.isna(ag):
            vf2, ag2 = _infer_vform_agegrp_from_filename(fname)
            vf = vf if vf is not None and not pd.isna(vf) else vf2
            ag = ag if ag is not None and not pd.isna(ag) else ag2
        if vf is not None and "vform" not in pool.columns:
            pool["vform"] = str(vf).upper()
        if ag is not None and "agegrp" not in pool.columns:
            pool["agegrp"] = int(float(ag))
        pool["source_file"] = fname
        pools.append(pool)
        read_log.append({"file": str(path), "sheet": sheet_name, "n": len(pool), "role": "replacement_pool"})

    if not master_parts:
        raise ValueError(
            "No master sample could be read. Provide master_sample_file/master_sample_sheet "
            "or case-swap files with final_sample sheets."
        )

    master = normalize_sample(pd.concat(master_parts, ignore_index=True))
    require_columns(master, KEY_COLS + [c for c in DEMO_COLS + OSEP_COMPONENT_FLAG_COLS if c in master.columns], "master")

    if pools:
        available = normalize_sample(pd.concat(pools, ignore_index=True))
        available = available.loc[~available["teid"].isin(set(master["teid"]))].copy()
    else:
        available = pd.DataFrame(columns=master.columns)

    return master, available, pd.DataFrame(read_log)


# =====================================================
# LOSS MODEL
# =====================================================
def build_state(master, demo_targets):
    state = {}
    for vform, grp in master.groupby("vform"):
        state[(vform, None)] = {
            "n": len(grp),
            "flags": {
                flag: int(pd.to_numeric(grp[flag], errors="coerce").fillna(0).sum())
                for flag in OSEP_FLAGS
            },
        }
    for (vform, agegrp), grp in master.groupby(["vform", "agegrp"]):
        state[(vform, agegrp)] = {
            "n": len(grp),
            "demo": {},
            "sex": {
                SEX_BALANCE_FEMALE_VALUE: sex_count(grp, SEX_BALANCE_FEMALE_VALUE),
                SEX_BALANCE_MALE_VALUE: sex_count(grp, SEX_BALANCE_MALE_VALUE),
            },
        }
        for (target_vform, target_agegrp), target_map in demo_targets.items():
            if target_vform != vform or target_agegrp != agegrp:
                continue
            for (col, value, _header), _target_pct in target_map.items():
                state[(vform, agegrp)]["demo"][(col, value)] = int(
                    (grp[col].astype(str).str.upper() == str(value).upper()).sum()
                )
    return state


def local_loss_delta(vform, agegrp, rem, add, state, demo_targets, osep_targets):
    pre = 0.0
    post = 0.0
    form_state = state[(vform, None)]
    n_form = form_state["n"]
    for flag, rule in osep_targets.items():
        if flag not in form_state["flags"]:
            continue
        base = form_state["flags"][flag]
        new = base - flag_int(rem, flag) + flag_int(add, flag)
        if "target_pct" in rule:
            pre += target_error_scaled(base, n_form, rule["target_pct"]) * WEIGHTS.get(flag, 1)
            post += target_error_scaled(new, n_form, rule["target_pct"]) * WEIGHTS.get(flag, 1)
        else:
            pre += range_error_scaled(base, n_form, rule["min_pct"], rule["max_pct"]) * WEIGHTS.get(flag, 1)
            post += range_error_scaled(new, n_form, rule["min_pct"], rule["max_pct"]) * WEIGHTS.get(flag, 1)

    cell_targets = demo_targets.get((vform, agegrp), {})
    cell_state = state[(vform, agegrp)]
    n_cell = cell_state["n"]
    for (col, value, _header), target_pct in cell_targets.items():
        base = cell_state["demo"].get((col, value), 0)
        new = base - int(str(rem.get(col)).upper() == str(value).upper()) + int(str(add.get(col)).upper() == str(value).upper())
        pre += target_error_scaled(base, n_cell, target_pct) * WEIGHTS["demo"]
        post += target_error_scaled(new, n_cell, target_pct) * WEIGHTS["demo"]
    return pre - post


# =====================================================
# EXACT SEX BALANCE HELPERS
# =====================================================
def sex_balance_quick_delta(vform, agegrp, rem, add, state):
    """Positive value means this swap moves the vform+agegrp cell closer to exact 50/50."""
    if not FORCE_EXACT_SEX_BALANCE:
        return 0
    if SEX_BALANCE_COLUMN not in rem.index or SEX_BALANCE_COLUMN not in add.index:
        return 0
    cell_state = state.get((vform, agegrp))
    if not cell_state:
        return 0
    n = int(cell_state["n"])
    if n % 2 != 0:
        return 0
    target_each = n // 2
    base_f = int(cell_state["sex"].get(SEX_BALANCE_FEMALE_VALUE, 0))
    new_f = base_f - is_sex_value(rem.get(SEX_BALANCE_COLUMN), SEX_BALANCE_FEMALE_VALUE) + is_sex_value(add.get(SEX_BALANCE_COLUMN), SEX_BALANCE_FEMALE_VALUE)
    pre_dev = abs(base_f - target_each)
    post_dev = abs(new_f - target_each)
    return (pre_dev - post_dev) * int(WEIGHTS.get("sex_balance_prefilter", 500000))


def add_exact_sex_balance_constraints(model, y, swaps, state):
    """Add Script-64a-style exact 50/50 sex constraints inside the CP-SAT model."""
    if not FORCE_EXACT_SEX_BALANCE:
        return []
    messages = []
    if SEX_BALANCE_COLUMN not in swaps.columns and f"remove_{SEX_BALANCE_COLUMN}" not in swaps.columns:
        messages.append("Skipped exact sex balance constraints because sex columns were not found in swaps.")
        return messages

    for (vform, agegrp), cell_state in sorted(state.items()):
        if agegrp is None:
            continue
        n = int(cell_state["n"])
        if n % 2 != 0:
            msg = f"Skipped exact sex balance for {vform} agegrp {agegrp}: odd N={n}."
            messages.append(msg)
            if SEX_BALANCE_ODD_N_ACTION == "error":
                raise ValueError(msg)
            continue
        target_each = n // 2
        cell_swaps = swaps.loc[(swaps["vform"] == vform) & (swaps["agegrp"] == agegrp)]
        female_expr = int(cell_state["sex"].get(SEX_BALANCE_FEMALE_VALUE, 0))
        male_expr = int(cell_state["sex"].get(SEX_BALANCE_MALE_VALUE, 0))
        for r in cell_swaps.itertuples():
            sid = int(r.swap_id)
            rem_sex = getattr(r, f"remove_{SEX_BALANCE_COLUMN}")
            add_sex = getattr(r, f"add_{SEX_BALANCE_COLUMN}")
            female_delta = is_sex_value(add_sex, SEX_BALANCE_FEMALE_VALUE) - is_sex_value(rem_sex, SEX_BALANCE_FEMALE_VALUE)
            male_delta = is_sex_value(add_sex, SEX_BALANCE_MALE_VALUE) - is_sex_value(rem_sex, SEX_BALANCE_MALE_VALUE)
            if female_delta:
                female_expr += female_delta * y[sid]
            if male_delta:
                male_expr += male_delta * y[sid]
        model.Add(female_expr == target_each)
        model.Add(male_expr == target_each)
        messages.append(f"Added exact sex balance for {vform} agegrp {agegrp}: {SEX_BALANCE_FEMALE_VALUE}={target_each}, {SEX_BALANCE_MALE_VALUE}={target_each}.")
    return messages


# =====================================================
# OPTIMIZE
# =====================================================
def optimize(master, available, demo_targets, osep_targets):
    try:
        from ortools.sat.python import cp_model
    except ImportError as exc:
        raise ImportError("Install OR-Tools in your environment: python -m pip install ortools") from exc

    state = build_state(master, demo_targets)
    swap_rows = []
    for lock_vals, used_group in master.groupby(SWAP_LOCK_COLS, dropna=False):
        if not isinstance(lock_vals, tuple):
            lock_vals = (lock_vals,)
        mask = pd.Series(True, index=available.index)
        for col, val in zip(SWAP_LOCK_COLS, lock_vals):
            mask &= available[col].astype(str).eq(str(val))
        avail_group = available.loc[mask].copy()
        if avail_group.empty:
            continue
        for _, rem in used_group.iterrows():
            scored = []
            for _, add in avail_group.iterrows():
                base_delta = local_loss_delta(rem["vform"], rem["agegrp"], rem, add, state, demo_targets, osep_targets)
                base_delta += sex_balance_quick_delta(rem["vform"], rem["agegrp"], rem, add, state)
                scored.append((base_delta, add))
            scored = sorted(scored, key=lambda x: x[0], reverse=True)[:MAX_CANDIDATE_ADDS_PER_REMOVE]
            for quick_delta, add in scored:
                row = {
                    "swap_id": len(swap_rows),
                    "vform": rem["vform"],
                    "agegrp": int(rem["agegrp"]),
                    "sex": rem["sex"],
                    "remove_teid": rem["teid"],
                    "add_teid": add["teid"],
                    "quick_loss_reduction": quick_delta,
                }
                for flag in OSEP_FLAGS:
                    row[f"delta_{flag}"] = flag_int(add, flag) - flag_int(rem, flag)
                    row[f"remove_{flag}"] = flag_int(rem, flag)
                    row[f"add_{flag}"] = flag_int(add, flag)
                for col in DEMO_COLS:
                    row[f"remove_{col}"] = rem.get(col)
                    row[f"add_{col}"] = add.get(col)
                if SEX_BALANCE_COLUMN not in DEMO_COLS:
                    row[f"remove_{SEX_BALANCE_COLUMN}"] = rem.get(SEX_BALANCE_COLUMN)
                    row[f"add_{SEX_BALANCE_COLUMN}"] = add.get(SEX_BALANCE_COLUMN)
                row["sex_balance_delta"] = sex_balance_quick_delta(rem["vform"], rem["agegrp"], rem, add, state)
                swap_rows.append(row)

    swaps = pd.DataFrame(swap_rows)
    if swaps.empty:
        raise ValueError("No feasible hard-locked swaps were generated.")

    model = cp_model.CpModel()
    y = {int(r.swap_id): model.NewBoolVar(f"swap_{int(r.swap_id)}") for r in swaps.itertuples()}

    for teid, grp in swaps.groupby("remove_teid"):
        model.Add(sum(y[int(s)] for s in grp["swap_id"]) <= 1)
    for teid, grp in swaps.groupby("add_teid"):
        model.Add(sum(y[int(s)] for s in grp["swap_id"]) <= 1)

    sex_constraint_messages = add_exact_sex_balance_constraints(model, y, swaps, state)

    objective_terms = []

    def post_flag_expr(vform, flag):
        expr = int(state[(vform, None)]["flags"].get(flag, 0))
        for r in swaps.loc[swaps["vform"] == vform].itertuples():
            d = int(getattr(r, f"delta_{flag}"))
            if d:
                expr += d * y[int(r.swap_id)]
        return expr

    def add_abs_dev(expr, target, ub, name):
        dev = model.NewIntVar(0, int(ub), name)
        target = int(target)
        model.Add(dev >= expr - target)
        model.Add(dev >= target - expr)
        return dev

    for vform in VFORM_ORDER:
        if (vform, None) not in state:
            continue
        n = state[(vform, None)]["n"]
        for flag, rule in osep_targets.items():
            expr = post_flag_expr(vform, flag)
            if "target_pct" in rule:
                dev = add_abs_dev(expr * SCALE, pct_to_count(n, rule["target_pct"]), n * SCALE, f"dev_{vform}_{flag}")
            else:
                lo = math.ceil(float(rule["min_pct"]) * n / 100.0 * SCALE)
                hi = math.floor(float(rule["max_pct"]) * n / 100.0 * SCALE)
                under = model.NewIntVar(0, n * SCALE, f"under_{vform}_{flag}")
                over = model.NewIntVar(0, n * SCALE, f"over_{vform}_{flag}")
                model.Add(expr * SCALE + under >= lo)
                model.Add(expr * SCALE - over <= hi)
                dev = model.NewIntVar(0, n * SCALE, f"dev_{vform}_{flag}")
                model.Add(dev == under + over)
            objective_terms.append(int(WEIGHTS.get(flag, 1)) * dev)

    for (vform, agegrp), target_map in demo_targets.items():
        if (vform, agegrp) not in state:
            continue
        n = state[(vform, agegrp)]["n"]
        cell_swaps = swaps.loc[(swaps["vform"] == vform) & (swaps["agegrp"] == agegrp)]
        for (col, value, _header), target_pct in target_map.items():
            expr = int(state[(vform, agegrp)]["demo"].get((col, value), 0))
            for r in cell_swaps.itertuples():
                d = int(str(getattr(r, f"add_{col}")).upper() == str(value).upper()) - int(
                    str(getattr(r, f"remove_{col}")).upper() == str(value).upper()
                )
                if d:
                    expr += d * y[int(r.swap_id)]
            dev = add_abs_dev(expr * SCALE, pct_to_count(n, target_pct), n * SCALE, f"dev_{vform}_{agegrp}_{col}_{value}")
            objective_terms.append(WEIGHTS["demo"] * dev)

    objective_terms.append(WEIGHTS["swap"] * sum(y.values()))
    model.Minimize(sum(objective_terms))

    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = float(MAX_SOLVER_SECONDS)
    solver.parameters.num_search_workers = int(NUM_WORKERS)
    status = solver.Solve(model)
    if status not in [cp_model.OPTIMAL, cp_model.FEASIBLE]:
        raise ValueError(f"Solver failed: {solver.StatusName(status)}")

    selected = swaps.loc[[solver.Value(y[int(s)]) == 1 for s in swaps["swap_id"]]].copy()
    remove_teids = set(selected["remove_teid"])
    add_teids = set(selected["add_teid"])
    final_sample = pd.concat(
        [
            master.loc[~master["teid"].isin(remove_teids)].copy(),
            available.loc[available["teid"].isin(add_teids)].copy(),
        ],
        ignore_index=True,
        sort=False,
    )
    final_sample = normalize_sample(final_sample)
    selected["solver_status"] = solver.StatusName(status)
    selected["solver_objective"] = solver.ObjectiveValue()
    if sex_constraint_messages:
        selected["sex_balance_constraints"] = " | ".join(sex_constraint_messages[:20])
    return final_sample, selected, swaps


# =====================================================
# OUTPUT TABLES
# =====================================================
def pct(series):
    return float(pd.to_numeric(series, errors="coerce").fillna(0).mean() * 100)


def build_byage(final_sample):
    rows = []
    families = sorted({str(v).split("_")[0] for v in VFORM_ORDER})
    for age in BYAGE_AGES:
        row = {"Age": age}
        for vform in VFORM_ORDER:
            if "agey" in final_sample.columns:
                age_series = pd.to_numeric(final_sample["agey"], errors="coerce").round().astype("Int64")
                row[vform] = int(((final_sample["vform"] == vform) & (age_series == age)).sum())
            else:
                row[vform] = ""
        for family in families:
            row[f"{family}_total"] = sum(row.get(v, 0) for v in VFORM_ORDER if str(v).startswith(family))
        rows.append(row)
    total = {"Age": ""}
    for col in VFORM_ORDER + [f"{family}_total" for family in families]:
        total[col] = sum(r[col] for r in rows if isinstance(r.get(col), int))
    total["Grand_total"] = len(final_sample)
    rows.append(total)
    return pd.DataFrame(rows)

def build_demo_sheet(final_sample, demo_targets, family):
    vforms = [v for v in VFORM_ORDER if v.startswith(family)]
    rows = []
    for vform in vforms:
        vf_rows = [k for k in demo_targets if k[0] == vform]
        for v_idx, (target_vform, agegrp) in enumerate(sorted(vf_rows)):
            target_map = demo_targets[(target_vform, agegrp)]
            grp = final_sample.loc[(final_sample["vform"] == vform) & (final_sample["agegrp"] == agegrp)]
            target_row = {"Vform": vform if v_idx == 0 else "", "Age Group": agegrp, "Age Band": "", "Source": "US population"}
            actual_row = {"Vform": "", "Age Group": agegrp, "Age Band": "", "Source": ""}
            for (col, value, header), target_pct in target_map.items():
                target_row[header] = target_pct
                actual_row[header] = pct(grp[col].astype(str).str.upper().eq(str(value).upper()).astype(int)) if len(grp) else 0
            actual_row["female"] = int((grp["sex"] == "F").sum())
            actual_row["male"] = int((grp["sex"] == "M").sum())
            rows.extend([target_row, actual_row])
    return pd.DataFrame(rows)


def build_osep(final_sample):
    display = []
    for flag in OSEP_FLAGS:
        label = flag.replace("tmp_", "pct_", 1)
        rule = OSEP_TARGETS.get(flag, {})
        if "target_pct" in rule:
            target = f"{float(rule['target_pct']):.2f}"
        elif "min_pct" in rule and "max_pct" in rule:
            target = f"{float(rule['min_pct']):.2f}-{float(rule['max_pct']):.2f}"
        else:
            target = ""
        display.append((label, flag, target))
    display.append(("total_n", None, ""))

    rows = []
    for label, flag, target in display:
        row = {"Group": label}
        for vform in VFORM_ORDER:
            grp = final_sample.loc[final_sample["vform"] == vform]
            row[vform] = int(len(grp)) if flag is None else round(pct(grp[flag]), 2)
        row["Census Target"] = target
        rows.append(row)
    return pd.DataFrame(rows)

def write_formatted_workbook(final_sample, swaps, all_swaps, read_log, demo_targets):
    wb = Workbook()
    wb.remove(wb.active)
    title_fill = PatternFill("solid", fgColor="002060")
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    label_fill = PatternFill("solid", fgColor="D9D2E9")
    total_fill = PatternFill("solid", fgColor="EDEDED")
    border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    def write_df(sheet_name, title, df):
        ws = wb.create_sheet(sheet_name)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(df.columns)))
        ws.cell(1, 1).value = title
        ws.cell(1, 1).fill = title_fill
        ws.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=12, name="Courier New")
        ws.cell(1, 1).alignment = Alignment(horizontal="center")
        for c, col in enumerate(df.columns, start=1):
            cell = ws.cell(2, c, col)
            cell.fill = header_fill
            cell.font = Font(bold=True)
        for r, row in enumerate(df.itertuples(index=False), start=3):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
                if isinstance(value, float):
                    cell.number_format = "0.00"
            if str(ws.cell(r, 1).value) == "total_n":
                for c in range(1, len(df.columns) + 1):
                    ws.cell(r, c).fill = total_fill
                    ws.cell(r, c).font = Font(bold=True)
            else:
                ws.cell(r, 1).fill = label_fill
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = border
        for col_idx in range(1, ws.max_column + 1):
            ws.column_dimensions[ws.cell(2, col_idx).column_letter].width = 14
        return ws

    write_df("byage", "GROUP SIZES: AFTER OPTIMIZED SWAPS", build_byage(final_sample))
    write_df("trs", "DEMOGRAPHIC VARIABLES: AFTER OPTIMIZED SWAPS", build_demo_sheet(final_sample, demo_targets, "TRS"))
    write_df("prs", "DEMOGRAPHIC VARIABLES: AFTER OPTIMIZED SWAPS", build_demo_sheet(final_sample, demo_targets, "PRS"))
    write_df("srp", "DEMOGRAPHIC VARIABLES: AFTER OPTIMIZED SWAPS", build_demo_sheet(final_sample, demo_targets, "SRP"))
    write_df("osep", "OSEP: AFTER OPTIMIZED SWAPS", build_osep(final_sample))
    write_df("sex_balance", "SEX BALANCE: AFTER OPTIMIZED SWAPS", build_sex_balance_qa(final_sample))

    for name, df in [
        ("swap_pairs", swaps),
        ("final_teids", final_sample[[c for c in ["teid", "vform", "agegrp", "agey", "sex", "tmp_ADHD", "tmp_clin"] if c in final_sample.columns]].sort_values([c for c in ["vform", "agegrp", "agey", "teid"] if c in final_sample.columns])),
        ("all_candidate_swaps", all_swaps.head(200000)),
        ("read_log", read_log),
    ]:
        ws = wb.create_sheet(name[:31])
        for c, col in enumerate(df.columns, start=1):
            ws.cell(1, c, col).font = Font(bold=True)
        for r, row in enumerate(df.itertuples(index=False), start=2):
            for c, value in enumerate(row, start=1):
                ws.cell(r, c, value)
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[ws.cell(1, c).column_letter].width = 16

    # Save a copy of the active configuration inside the output for auditability.
    config_rows = []
    for k, v in {
        "config_workbook": str(CONFIG_WORKBOOK),
        "project_dir": str(PROJECT_DIR),
        "target_workbook": str(TARGET_WORKBOOK),
        "master_sample_file": str(MASTER_SAMPLE_FILE),
        "qa_dir": str(QA_DIR),
        "run_label": RUN_LABEL,
        "max_candidate_adds_per_remove": MAX_CANDIDATE_ADDS_PER_REMOVE,
        "max_solver_seconds": MAX_SOLVER_SECONDS,
        "num_workers": NUM_WORKERS,
        "scale": SCALE,
        "force_exact_sex_balance": FORCE_EXACT_SEX_BALANCE,
        "sex_balance_column": SEX_BALANCE_COLUMN,
        "sex_balance_female_value": SEX_BALANCE_FEMALE_VALUE,
        "sex_balance_male_value": SEX_BALANCE_MALE_VALUE,
        "sex_balance_odd_n_action": SEX_BALANCE_ODD_N_ACTION,
        "hard_lock_cols": ", ".join(HARD_LOCK_COLS),
        "swap_generation_locks": ", ".join(SWAP_LOCK_COLS),
        "vform_order": ", ".join(VFORM_ORDER),
        "byage_ages": ", ".join(map(str, BYAGE_AGES)),
    }.items():
        config_rows.append({"setting": k, "value": v})
    for k, v in WEIGHTS.items():
        config_rows.append({"setting": f"weight:{k}", "value": v})
    for k, v in OSEP_TARGETS.items():
        config_rows.append({"setting": f"osep_target:{k}", "value": str(v)})
    ws = wb.create_sheet("config_used")
    config_df = pd.DataFrame(config_rows)
    for c, col in enumerate(config_df.columns, start=1):
        ws.cell(1, c, col).font = Font(bold=True)
    for r, row in enumerate(config_df.itertuples(index=False), start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    for c in range(1, ws.max_column + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 35

    wb.save(OUTPUT_WORKBOOK)


def main():
    print("Reading targets from target workbook and config...")
    demo_targets = parse_demo_targets(TARGET_WORKBOOK)
    osep_targets = parse_osep_targets(TARGET_WORKBOOK)

    print("Reading sample and replacement pools...")
    master, available, read_log = read_inputs()
    print(f"Master N={len(master):,}; available N={len(available):,}")

    print("Optimizing sample...")
    final_sample, selected_swaps, all_swaps = optimize(master, available, demo_targets, osep_targets)
    if len(final_sample) != len(master):
        raise ValueError(f"Final sample size changed: {len(master)} -> {len(final_sample)}")
    if final_sample["teid"].duplicated().any():
        raise ValueError("Final sample has duplicate TEIDs.")
    if FORCE_EXACT_SEX_BALANCE:
        sex_qa = build_sex_balance_qa(final_sample)
        failed = sex_qa.loc[(sex_qa["odd_n_impossible"] == False) & (sex_qa["exact_50_50"] == False)]
        if not failed.empty:
            raise ValueError("Exact sex balance failed for one or more even-N cells:\n" + failed.to_string(index=False))

    print("Writing output workbook...")
    write_formatted_workbook(final_sample, selected_swaps, all_swaps, read_log, demo_targets)
    print(f"Done: {OUTPUT_WORKBOOK}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"FAILED: {exc}", file=sys.stderr)
        raise
